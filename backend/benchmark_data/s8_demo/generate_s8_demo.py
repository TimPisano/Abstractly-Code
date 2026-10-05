"""
Generates the Oak Hollow Apartments fake affordable-housing demo: a
40-unit Project-Based Section 8 (PBRA) property with a rent roll, 10
complete tenant files (one multi-page PDF each), and deliberately planted
compliance errors, plus expected_findings.json and README.md describing
each error and the rule it breaks.

Every name, address, employer, case number, and dollar figure is invented.
SSNs all use the never-issued 000 area number (000-00-00NN). The income
limits, contract rents, and utility allowances are made up for the demo
and are NOT published HUD figures.

Everything is computed bottom-up from the household dictionaries below
(never hand-typed totals), so the PDFs, the rent roll,
expected_findings.json, and README.md cannot drift out of sync.
Run from backend/: venv/bin/python benchmark_data/s8_demo/generate_s8_demo.py
"""
import csv
from xml.sax.saxutils import escape
import json
import os
import random
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfReader
from reportlab import rl_config
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

# Byte-reproducible PDFs: no creation timestamps or random document IDs.
rl_config.invariant = 1

OUT_DIR = os.path.dirname(os.path.abspath(__file__))
FILES_DIR = os.path.join(OUT_DIR, "tenant_files")
RENT_ROLL_DIR = os.path.join(OUT_DIR, "rent_roll")

PROPERTY_NAME = "Oak Hollow Apartments"
PROPERTY_ADDRESS = "1180 Oak Hollow Drive, Riverton, OH 45099"
OWNER_ENTITY = "Oak Hollow Housing Partners, LP"
MANAGEMENT_CO = "Brightwater Residential Management, LLC"
AGENT_NAME = "Jordan Avery"
AGENT_TITLE = "Compliance Specialist"
PROGRAM = "Project-Based Section 8 (PBRA)"
HAP_CONTRACT = "DEMO-HAP-0040"
# The date the file review "happens". Every planted error is defined
# relative to this fixed date, so the demo never goes stale.
REVIEW_DATE = date(2026, 10, 1)
RENT_ROLL_GENERATED = date(2026, 10, 2)
RNG_SEED = "oak-hollow-s8-demo-2026"

FOOTER = ("FICTIONAL DEMO DATA - Oak Hollow Apartments is not a real property and every person "
          "in this file is invented. Forms are simplified models, not official HUD forms.")

# Demo calculation rules (pre-HOTMA, see README "Rules the demo applies").
DEPENDENT_DEDUCTION = 480
ELDERLY_DISABLED_DEDUCTION = 400
MEDICAL_THRESHOLD = 0.03
MINIMUM_RENT = 25
VERIFICATION_VALID_DAYS = 120
EIV_VALID_DAYS = 120

# Fictional demo limits -- NOT published HUD figures.
LOW_INCOME_LIMIT = {1: 52000, 2: 59400, 3: 66850, 4: 74250, 5: 80200, 6: 86150}
UNIT_TYPES = {
    "1BR": {"beds": 1, "sqft": 640, "contract_rent": 1050, "utility_allowance": 85},
    "2BR": {"beds": 2, "sqft": 860, "contract_rent": 1240, "utility_allowance": 110},
    "3BR": {"beds": 3, "sqft": 1080, "contract_rent": 1480, "utility_allowance": 135},
}
# Per floor: 01-04, 06, 09 = 2BR; 05, 07, 10 = 1BR; 08 = 3BR. 4 floors = 40 units.
FLOOR_LAYOUT = {1: "2BR", 2: "2BR", 3: "2BR", 4: "2BR", 5: "1BR",
                6: "2BR", 7: "1BR", 8: "3BR", 9: "2BR", 10: "1BR"}
VACANT_UNITS = {"402", "410"}


def d(s):
    return date.fromisoformat(s)


def fmt_date(x):
    return x.strftime("%m/%d/%Y") if x else ""


def money(x):
    return f"${x:,.0f}" if float(x).is_integer() else f"${x:,.2f}"


def age_on(dob, on):
    return on.year - dob.year - ((on.month, on.day) < (dob.month, dob.day))


def round_dollar(x):
    # HUD rounds TTP to the nearest dollar; half rounds up (no banker's rounding).
    return int(x + 0.5)


def months_between(start, end):
    """Monthly HAP payments from start (1st of month) through the month before end."""
    return (end.year - start.year) * 12 + (end.month - start.month)


# --------------------------------------------------------------------------
# The 10 tenant files. `on_cert: False` marks an item that exists in the
# file but was left off the certification (a planted error); `claimed`
# marks which deductions the certification actually took.
# --------------------------------------------------------------------------
JAYLEN = {"name": "Jaylen Brooks", "rel": "Dependent", "dob": "2019-02-11",
          "ssn": "000-00-0021", "student": "No"}

HOUSEHOLDS = [
    {   # CLEAN control file -- nothing planted.
        "unit": "102", "slug": "delgado",
        "members": [
            {"name": "Marisol Delgado", "rel": "Head", "dob": "1992-03-14", "ssn": "000-00-0001", "student": "No"},
            {"name": "Mateo Delgado", "rel": "Dependent", "dob": "2017-05-02", "ssn": "000-00-0002", "student": "No"},
            {"name": "Lucia Delgado", "rel": "Dependent", "dob": "2020-01-27", "ssn": "000-00-0003", "student": "No"},
        ],
        "move_in": "2021-06-01", "application": "2021-04-12 10:42 AM",
        "cert": {"type": "Annual Recertification (AR)", "effective": "2026-06-01", "signed": "2026-05-12"},
        "history": [("Move-in (MI)", "2021-06-01", 402), ("AR", "2025-06-01", 511)],
        "income": [
            {"member": "Marisol Delgado", "kind": "Employment", "source": "Riverton Family Dental",
             "detail": "$17.50/hr x 32 hrs/wk x 52 wks", "annual": 29120,
             "sent": "2026-03-20", "received": "2026-03-27"},
        ],
        "assets": [{"member": "Marisol Delgado", "desc": "Checking - Riverton Community Credit Union", "value": 640, "income": 0}],
        "childcare": {"annual": 2400, "provider": "Little Acorns Learning Center", "received": "2026-03-30"},
        "eiv": {"pulled": "2026-03-18", "rows": [("Marisol Delgado", "Wages - Riverton Family Dental", "Q4 2025", 7280)],
                "note": "EIV wages match third-party verification. No discrepancies. - J. Avery 03/18/2026"},
        "notices": {"120": "2026-02-01", "90": None, "60": None, "30": None,
                    "response": "Tenant returned packet 02/20/2026; interview 03/20/2026"},
        "consents_signed": "2026-03-20", "lease_start": "2021-06-01",
    },
    {   # E1: spouse never signed the certification.
        "unit": "105", "slug": "whitfield",
        "members": [
            {"name": "Harold Whitfield", "rel": "Head", "dob": "1955-02-08", "ssn": "000-00-0004", "student": "No"},
            {"name": "Doreen Whitfield", "rel": "Spouse", "dob": "1958-09-30", "ssn": "000-00-0005", "student": "No"},
        ],
        "elderly": True,
        "move_in": "2018-04-01", "application": "2018-01-09 2:15 PM",
        "cert": {"type": "Annual Recertification (AR)", "effective": "2026-04-01", "signed": "2026-03-10"},
        "history": [("Move-in (MI)", "2018-04-01", 488), ("AR", "2025-04-01", 551)],
        "income": [
            {"member": "Harold Whitfield", "kind": "Social Security", "source": "Social Security Administration",
             "detail": "$1,200.00/mo x 12", "annual": 14400, "sent": "2026-01-15", "received": "2026-01-22"},
            {"member": "Harold Whitfield", "kind": "Pension", "source": "Ohio Valley Machinists Pension Fund",
             "detail": "$400.00/mo x 12", "annual": 4800, "sent": "2026-01-15", "received": "2026-01-29"},
            {"member": "Doreen Whitfield", "kind": "Social Security", "source": "Social Security Administration",
             "detail": "$800.00/mo x 12", "annual": 9600, "sent": "2026-01-15", "received": "2026-01-22"},
        ],
        "assets": [{"member": "Harold Whitfield", "desc": "Savings - Riverton Community Credit Union", "value": 3900, "income": 4}],
        "medical": 3100,
        "eiv": {"pulled": "2026-01-14", "rows": [("Harold Whitfield", "SS benefit", "Dec 2025", 1200),
                                                  ("Doreen Whitfield", "SS benefit", "Dec 2025", 800)],
                "note": "Benefit amounts match SSA letters. No discrepancies. - J. Avery 01/14/2026"},
        "notices": {"120": "2025-12-02", "90": None, "60": None, "30": None,
                    "response": "Tenant returned packet 12/18/2025; interview 01/15/2026"},
        "consents_signed": "2026-01-15", "lease_start": "2018-04-01",
        "cert_unsigned": ["Doreen Whitfield"],
    },
    {   # E2: adult son never signed consent forms. E3: stale employment verification.
        "unit": "108", "slug": "okafor",
        "members": [
            {"name": "Grace Okafor", "rel": "Head", "dob": "1981-07-19", "ssn": "000-00-0006", "student": "No"},
            {"name": "Emeka Okafor", "rel": "Other adult", "dob": "2007-11-14", "ssn": "000-00-0007", "student": "No"},
            {"name": "Adaeze Okafor", "rel": "Dependent", "dob": "2014-04-03", "ssn": "000-00-0008", "student": "No"},
        ],
        "move_in": "2019-07-01", "application": "2019-05-02 9:05 AM",
        "cert": {"type": "Annual Recertification (AR)", "effective": "2026-07-01", "signed": "2026-06-24"},
        "history": [("Move-in (MI)", "2019-07-01", 702), ("AR", "2025-07-01", 841)],
        "income": [
            {"member": "Grace Okafor", "kind": "Employment", "source": "Riverton Senior Living",
             "detail": "$19.25/hr x 36 hrs/wk x 52 wks", "annual": 36036,
             "sent": "2026-01-12", "received": "2026-01-20"},
            {"member": "Emeka Okafor", "kind": "Employment", "source": "Quickstop Grocery #14",
             "detail": "$14.00/hr x 20 hrs/wk x 52 wks", "annual": 14560,
             "sent": "2026-05-26", "received": "2026-06-02"},
        ],
        "assets": [{"member": "Grace Okafor", "desc": "Checking - First Riverton Bank", "value": 1150, "income": 0}],
        "eiv": {"pulled": "2026-05-20", "rows": [("Grace Okafor", "Wages - Riverton Senior Living", "Q1 2026", 9009),
                                                  ("Emeka Okafor", "Wages - Quickstop Grocery #14", "Q1 2026", 3640)],
                "note": "EIV wages consistent with verified income. No discrepancies. - J. Avery 05/20/2026"},
        "notices": {"120": "2026-03-03", "90": "2026-04-02", "60": None, "30": None,
                    "response": "Tenant returned packet 04/20/2026; interview 05/26/2026"},
        "consents_signed": "2026-05-26", "lease_start": "2019-07-01",
        "consent_unsigned": ["Emeka Okafor"],
    },
    {   # E4: EIV shows a second employer that never made it onto the cert.
        "unit": "203", "slug": "petrakis",
        "members": [
            {"name": "Elena Petrakis", "rel": "Head", "dob": "1988-12-01", "ssn": "000-00-0009", "student": "No"},
            {"name": "Nikos Petrakis", "rel": "Dependent", "dob": "2016-08-22", "ssn": "000-00-0010", "student": "No"},
        ],
        "move_in": "2022-07-01", "application": "2022-05-16 11:30 AM",
        "cert": {"type": "Annual Recertification (AR)", "effective": "2026-07-01", "signed": "2026-06-15"},
        "history": [("Move-in (MI)", "2022-07-01", 361), ("AR", "2025-07-01", 372)],
        "income": [
            {"member": "Elena Petrakis", "kind": "Employment", "source": "Riverton Public Library",
             "detail": "$15.50/hr x 25 hrs/wk x 52 wks", "annual": 20150,
             "sent": "2026-05-12", "received": "2026-05-19"},
        ],
        # Shows up only in EIV; never verified, never on the cert.
        "unreported_income": {"member": "Elena Petrakis", "source": "Lakeside Logistics, Inc.",
                              "annual": 12480, "basis": "2 quarters x $3,120 in EIV, annualized"},
        "assets": [{"member": "Elena Petrakis", "desc": "Checking - First Riverton Bank", "value": 420, "income": 0}],
        "eiv": {"pulled": "2026-05-10", "rows": [("Elena Petrakis", "Wages - Riverton Public Library", "Q1 2026", 5038),
                                                  ("Elena Petrakis", "Wages - Lakeside Logistics, Inc.", "Q4 2025", 3120),
                                                  ("Elena Petrakis", "Wages - Lakeside Logistics, Inc.", "Q1 2026", 3120)],
                "note": None},
        "notices": {"120": "2026-03-03", "90": None, "60": None, "30": None,
                    "response": "Tenant returned packet 03/30/2026; interview 05/12/2026"},
        "consents_signed": "2026-05-12", "lease_start": "2022-07-01",
    },
    {   # E5: dependent and child-care deductions verified but not taken.
        "unit": "206", "slug": "hollis",
        "members": [
            {"name": "Andre Hollis", "rel": "Head", "dob": "1995-04-17", "ssn": "000-00-0011", "student": "No"},
            {"name": "Kayla Hollis", "rel": "Dependent", "dob": "2018-03-09", "ssn": "000-00-0012", "student": "No"},
            {"name": "Marcus Hollis", "rel": "Dependent", "dob": "2021-06-30", "ssn": "000-00-0013", "student": "No"},
        ],
        "move_in": "2023-06-01", "application": "2023-04-03 3:48 PM",
        "cert": {"type": "Annual Recertification (AR)", "effective": "2026-06-01", "signed": "2026-05-14"},
        "history": [("Move-in (MI)", "2023-06-01", 788), ("AR", "2025-06-01", 819)],
        "income": [
            {"member": "Andre Hollis", "kind": "Employment", "source": "Northgate Distribution Center",
             "detail": "$21.00/hr x 40 hrs/wk x 52 wks", "annual": 43680,
             "sent": "2026-03-24", "received": "2026-03-31"},
        ],
        "assets": [{"member": "Andre Hollis", "desc": "Checking - Riverton Community Credit Union", "value": 980, "income": 0}],
        "childcare": {"annual": 3600, "provider": "Sunrise Kids Daycare", "received": "2026-04-02"},
        "cert_omits_dependents": True, "cert_omits_childcare": True,
        "eiv": {"pulled": "2026-03-20", "rows": [("Andre Hollis", "Wages - Northgate Distribution Center", "Q4 2025", 10920)],
                "note": "EIV wages match third-party verification. No discrepancies. - J. Avery 03/20/2026"},
        "notices": {"120": "2026-02-01", "90": None, "60": None, "30": None,
                    "response": "Tenant returned packet 02/26/2026; interview 03/24/2026"},
        "consents_signed": "2026-03-24", "lease_start": "2023-06-01",
    },
    {   # E6: annual recert due 08/01/2026 never done; reminders never sent.
        "unit": "207", "slug": "ramirez_cole",
        "members": [
            {"name": "Sofia Ramirez-Cole", "rel": "Head", "dob": "1974-10-05", "ssn": "000-00-0014",
             "student": "No", "disabled": True},
        ],
        "disabled_family": True,
        "move_in": "2020-08-01", "application": "2020-06-08 1:20 PM",
        "cert": {"type": "Annual Recertification (AR)", "effective": "2025-08-01", "signed": "2025-07-16"},
        "history": [("Move-in (MI)", "2020-08-01", 214), ("AR", "2024-08-01", 228)],
        "income": [
            {"member": "Sofia Ramirez-Cole", "kind": "SSDI", "source": "Social Security Administration",
             "detail": "$1,100.00/mo x 12", "annual": 13200, "sent": "2025-05-28", "received": "2025-06-04"},
        ],
        "assets": [{"member": "Sofia Ramirez-Cole", "desc": "Checking - First Riverton Bank", "value": 310, "income": 0}],
        "eiv": {"pulled": "2025-05-20", "rows": [("Sofia Ramirez-Cole", "SSDI benefit", "Apr 2025", 1100)],
                "note": "Benefit matches SSA letter. No discrepancies. - J. Avery 05/20/2025"},
        # This is the 2026 cycle: first notice went out, nothing after it.
        "notices": {"120": "2026-04-03", "90": None, "60": None, "30": None,
                    "response": "No response recorded. No follow-up recorded."},
        "next_ar": "2026-08-01", "recert_missed": True,
        "consents_signed": "2025-05-28", "lease_start": "2020-08-01",
    },
    {   # E7: adult daughter's income left off the move-in cert -> household was over income.
        "unit": "301", "slug": "lindqvist",
        "members": [
            {"name": "Karin Lindqvist", "rel": "Head", "dob": "1978-06-12", "ssn": "000-00-0015", "student": "No"},
            {"name": "Freya Lindqvist", "rel": "Other adult", "dob": "2004-02-20", "ssn": "000-00-0016", "student": "No"},
        ],
        "move_in": "2026-03-01", "application": "2026-01-06 10:10 AM",
        "cert": {"type": "Move-in (MI)", "effective": "2026-03-01", "signed": "2026-02-24"},
        "history": [],
        "income": [
            {"member": "Karin Lindqvist", "kind": "Employment", "source": "Riverton School District",
             "detail": "$48,500 annual salary", "annual": 48500, "sent": "2026-01-12", "received": "2026-01-21"},
            {"member": "Freya Lindqvist", "kind": "Employment", "source": "Harbor Coffee Roasters",
             "detail": "$15.00/hr x 17.2 hrs/wk x 52 wks (rounded)", "annual": 13400,
             "sent": "2026-01-12", "received": "2026-01-23", "on_cert": False},
        ],
        "assets": [{"member": "Karin Lindqvist", "desc": "Savings - First Riverton Bank", "value": 2200, "income": 2}],
        "eiv": {"pulled": "2026-01-15", "rows": [("Karin Lindqvist", "EIV Existing Tenant Search", "-", 0)],
                "note": "Existing Tenant Search: no current HUD assistance found for any member. - J. Avery 01/15/2026"},
        "notices": None,
        "consents_signed": "2026-01-12", "lease_start": "2026-03-01",
    },
    {   # E8 (with 309): joint-custody child counted in both households.
        # E9: rent roll never updated after the 05/01 interim decrease.
        "unit": "304", "slug": "brooks_t",
        "members": [
            {"name": "Tanya Brooks", "rel": "Head", "dob": "1997-01-23", "ssn": "000-00-0017", "student": "No"},
            dict(JAYLEN),
            {"name": "Amira Brooks", "rel": "Dependent", "dob": "2023-05-15", "ssn": "000-00-0018", "student": "No"},
        ],
        "move_in": "2022-01-01", "application": "2021-11-01 4:02 PM",
        "cert": {"type": "Interim Recertification (IR) - income decrease", "effective": "2026-05-01", "signed": "2026-04-22"},
        "history": [("Move-in (MI)", "2022-01-01", 540), ("AR", "2026-01-01", 615)],
        "income": [
            {"member": "Tanya Brooks", "kind": "Employment", "source": "Riverton Pediatrics",
             "detail": "$18.00/hr x 24 hrs/wk x 52 wks (hours cut from 32 on 04/06/2026)", "annual": 22464,
             "sent": "2026-04-08", "received": "2026-04-14"},
        ],
        "assets": [{"member": "Tanya Brooks", "desc": "Checking - Riverton Community Credit Union", "value": 275, "income": 0}],
        "eiv": {"pulled": "2026-04-09", "rows": [("Tanya Brooks", "Wages - Riverton Pediatrics", "Q4 2025", 7488)],
                "note": "Q4 2025 wages reflect prior 32-hr schedule; reduction verified by employer 04/14/2026. - J. Avery 04/15/2026"},
        "notices": None,
        "consents_signed": "2026-04-08", "lease_start": "2022-01-01",
        "rent_roll_tenant_rent": 615,  # pre-interim amount, never updated
    },
    {   # E10: full-time student, no exemption, no parental-income documentation.
        "unit": "307", "slug": "ostrander",
        "members": [
            {"name": "Tyler Ostrander", "rel": "Head", "dob": "2006-03-28", "ssn": "000-00-0019",
             "student": "Full-time - Riverton Community College"},
        ],
        "move_in": "2025-09-01", "application": "2025-07-14 12:55 PM",
        "cert": {"type": "Annual Recertification (AR)", "effective": "2026-09-01", "signed": "2026-08-18"},
        "history": [("Move-in (MI)", "2025-09-01", 145)],
        "income": [
            {"member": "Tyler Ostrander", "kind": "Employment", "source": "Riverton Cinemas",
             "detail": "$14.00/hr x 13.5 hrs/wk x 52 wks (rounded)", "annual": 9800,
             "sent": "2026-06-22", "received": "2026-06-29"},
        ],
        "assets": [{"member": "Tyler Ostrander", "desc": "Checking - First Riverton Bank", "value": 190, "income": 0}],
        "eiv": {"pulled": "2026-06-20", "rows": [("Tyler Ostrander", "Wages - Riverton Cinemas", "Q1 2026", 2450)],
                "note": "EIV wages match third-party verification. No discrepancies. - J. Avery 06/20/2026"},
        "notices": {"120": "2026-05-04", "90": None, "60": None, "30": None,
                    "response": "Tenant returned packet 05/29/2026; interview 06/22/2026"},
        "student_detail": {"Tyler Ostrander": {"enrollment": "Full-time, Fall 2026 (15 credit hours)", "age_under_24": "Yes",
                                               "married": "No", "veteran": "No", "disabled": "No",
                                               "dependent_child": "No", "living_with_parents": "No",
                                               "independent_student_docs": "None in file",
                                               "parent_income_cert": "None in file"}},
        "consents_signed": "2026-06-22", "lease_start": "2025-09-01",
    },
    {   # E8 (with 304): this household also counts Jaylen; its own custody order says he lives mostly with his mother.
        "unit": "309", "slug": "brooks_d",
        "members": [
            {"name": "Darnell Brooks", "rel": "Head", "dob": "1993-08-04", "ssn": "000-00-0020", "student": "No"},
            dict(JAYLEN),
        ],
        "move_in": "2024-05-01", "application": "2024-03-11 9:40 AM",
        "cert": {"type": "Annual Recertification (AR)", "effective": "2026-05-01", "signed": "2026-04-16"},
        "history": [("Move-in (MI)", "2024-05-01", 872), ("AR", "2025-05-01", 896)],
        "income": [
            {"member": "Darnell Brooks", "kind": "Employment", "source": "Riverton Transit Authority",
             "detail": "$20.00/hr x 40 hrs/wk x 52 wks", "annual": 41600,
             "sent": "2026-02-23", "received": "2026-03-02"},
        ],
        "assets": [{"member": "Darnell Brooks", "desc": "Checking - First Riverton Bank", "value": 1320, "income": 0}],
        "eiv": {"pulled": "2026-02-18", "rows": [("Darnell Brooks", "Wages - Riverton Transit Authority", "Q4 2025", 10400)],
                "note": "EIV wages match third-party verification. No discrepancies. - J. Avery 02/18/2026"},
        "notices": {"120": "2026-01-01", "90": None, "60": None, "30": None,
                    "response": "Tenant returned packet 01/26/2026; interview 02/23/2026"},
        "not_a_member": ["Jaylen Brooks"],
        "custody": {"case": "DEMO-DR-2024-0117 (fictional)", "child": "Jaylen Brooks",
                    "terms": ("Mother (Tanya Brooks) has residential custody Sunday evening through Friday "
                              "morning. Father (Darnell Brooks) has parenting time Friday evening through "
                              "Sunday evening and alternating holidays."),
                    "overnights_with_father": "2 of 7 nights (~29%)"},
        "consents_signed": "2026-02-23", "lease_start": "2024-05-01",
    },
]


def unit_type(unit):
    return FLOOR_LAYOUT[int(unit[1:])]


def next_ar(h):
    """Next annual recertification date. An interim doesn't move the anniversary."""
    if h.get("next_ar"):
        return d(h["next_ar"])
    anchor = d(h["history"][-1][1]) if "Interim" in h["cert"]["type"] else d(h["cert"]["effective"])
    while anchor <= d(h["cert"]["effective"]):
        anchor = anchor.replace(year=anchor.year + 1)
    return anchor


def is_dependent(m, on):
    return m["rel"] == "Dependent" and (age_on(d(m["dob"]), on) < 18 or m["student"].startswith("Full-time"))


def calculate(h, as_certified):
    """Rent calculation. as_certified=True reproduces what the file's
    certification shows (including planted mistakes); False is correct."""
    eff = d(h["cert"]["effective"])
    ut = UNIT_TYPES[unit_type(h["unit"])]
    incomes = [i for i in h["income"] if i.get("on_cert", True) or not as_certified]
    annual = sum(i["annual"] for i in incomes) + sum(a["income"] for a in h["assets"])
    if not as_certified and h.get("unreported_income"):
        annual += h["unreported_income"]["annual"]
    members = [m for m in h["members"] if as_certified or m["name"] not in h.get("not_a_member", [])]
    deps = [m for m in members if is_dependent(m, eff)]
    dep_ded = 0 if (as_certified and h.get("cert_omits_dependents")) else DEPENDENT_DEDUCTION * len(deps)
    cc = h.get("childcare", {}).get("annual", 0)
    cc_ded = 0 if (as_certified and h.get("cert_omits_childcare")) else cc
    ed_ded = ELDERLY_DISABLED_DEDUCTION if (h.get("elderly") or h.get("disabled_family")) else 0
    med_ded = max(0, h.get("medical", 0) - round(annual * MEDICAL_THRESHOLD)) if ed_ded else 0
    adjusted = annual - dep_ded - cc_ded - ed_ded - med_ded
    ttp = max(round_dollar(adjusted * 0.30 / 12), round_dollar(annual * 0.10 / 12), MINIMUM_RENT)
    tenant_rent = max(ttp - ut["utility_allowance"], 0)
    return {
        "annual": annual, "dependents": len(deps), "dep_ded": dep_ded, "cc_ded": cc_ded,
        "ed_ded": ed_ded, "med_ded": med_ded, "adjusted": adjusted, "ttp": ttp,
        "utility_allowance": ut["utility_allowance"], "contract_rent": ut["contract_rent"],
        "tenant_rent": tenant_rent, "utility_reimbursement": max(ut["utility_allowance"] - ttp, 0),
        "hap": ut["contract_rent"] - tenant_rent, "household_size": len(members),
    }


# --------------------------------------------------------------------------
# PDF rendering
# --------------------------------------------------------------------------
styles = getSampleStyleSheet()
H1 = ParagraphStyle("h1", parent=styles["Heading1"], fontSize=15, spaceAfter=4)
H2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=11, spaceBefore=6, spaceAfter=3)
BODY = ParagraphStyle("body", parent=styles["BodyText"], fontSize=9, leading=11.5)
SMALL = ParagraphStyle("small", parent=BODY, fontSize=8, leading=10, textColor=colors.HexColor("#555555"))
SIG = ParagraphStyle("sig", parent=BODY, fontName="Helvetica-Oblique")


def table(rows, widths, header=True):
    t = Table([[c if isinstance(c, Paragraph) else Paragraph(escape(str(c)), BODY) for c in r] for r in rows],
              colWidths=widths)
    style = [("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
             ("VALIGN", (0, 0), (-1, -1), "TOP"),
             ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]
    if header:
        style.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef4")))
    t.setStyle(TableStyle(style))
    return t


def kv(pairs):
    return table([[k, v] for k, v in pairs], [2.3 * inch, 4.7 * inch], header=False)


def sig(name, signed):
    """A signature line: '/s/ Name' and a date if signed, blank if not."""
    return [Paragraph(f"/s/ {name}" if signed else "", SIG), fmt_date(signed) if signed else ""]


def page_header(h, title, subtitle=None):
    out = [Paragraph(escape(title), H1),
           Paragraph(f"{PROPERTY_NAME} &middot; Unit {h['unit']} &middot; {PROGRAM} &middot; HAP Contract {HAP_CONTRACT}", SMALL)]
    if subtitle:
        out.append(Paragraph(escape(subtitle), SMALL))
    out.append(Spacer(1, 8))
    return out


def adults(h, on):
    return [m for m in h["members"] if age_on(d(m["dob"]), on) >= 18]


def build_docs(h, calc):
    """Return [(doc_key, title, flowables)], one page each."""
    eff = d(h["cert"]["effective"])
    signed = d(h["cert"]["signed"])
    head = h["members"][0]["name"]
    ut = unit_type(h["unit"])
    docs = []

    # Rental application + supplement
    app_rows = [["Name", "Relationship", "Date of birth", "Employer / income source"]]
    for m in h["members"]:
        srcs = ", ".join(i["source"] for i in h["income"] if i["member"] == m["name"]) or "-"
        app_rows.append([m["name"], m["rel"], fmt_date(d(m["dob"])), srcs])
    docs.append(("application", "Rental Application & Supplement", page_header(h, "Rental Application") + [
        kv([("Application received", fmt_date(d(h["application"][:10])) + h["application"][10:]),
            ("Unit size requested", ut), ("Household size", str(len(h["members"]))),
            ("Current address", "Fictional - withheld in demo"), ("Phone", f"(555) 555-01{h['unit'][1:]}")]),
        Paragraph("Household members", H2), table(app_rows, [1.8 * inch, 1.1 * inch, 1.1 * inch, 3.0 * inch]),
        Paragraph("Supplement to Application (modeled on HUD-92006)", H2),
        Paragraph("Applicant was offered the option to name an additional contact person or organization. "
                  "Response: declined to provide additional contact.", BODY),
        Spacer(1, 10), table([["Applicant signature", "Date"], sig(head, d(h["application"][:10]))],
                             [4.5 * inch, 2.5 * inch]),
    ]))

    # Consent forms (9887 / 9887-A)
    consent_date = d(h["consents_signed"])
    rows = [["Adult household member", "Age", "HUD-9887 signature", "Date", "HUD-9887-A signature", "Date"]]
    for m in adults(h, eff):
        ok = m["name"] not in h.get("consent_unsigned", [])
        s = sig(m["name"], consent_date if ok else None)
        rows.append([m["name"], str(age_on(d(m["dob"]), eff)), s[0], s[1], sig(m["name"], consent_date if ok else None)[0], s[1]])
    docs.append(("consents", "Consent Forms (9887 / 9887-A)", page_header(
        h, "Authorization for Release of Information & Applicant/Tenant Certification",
        "Models HUD-9887 (Notice and Consent for the Release of Information) and HUD-9887-A. "
        f"Signature table for the {h['cert']['type']} effective {fmt_date(eff)}.") + [
        Paragraph("Each household member age 18 or older signs both forms.", BODY), Spacer(1, 6),
        table(rows, [1.6 * inch, 0.45 * inch, 1.45 * inch, 0.85 * inch, 1.45 * inch, 0.85 * inch]),
    ]))

    # Household composition, SSN, citizenship
    rows = [["Member", "Relationship", "DOB", "Age", "SSN", "SSN document in file", "Citizenship"]]
    for m in h["members"]:
        rows.append([m["name"], m["rel"], fmt_date(d(m["dob"])), str(age_on(d(m["dob"]), eff)), m["ssn"],
                     "SS card copy", "U.S. citizen (declaration signed)"])
    docs.append(("household", "Household Composition, SSN & Citizenship", page_header(
        h, "Household Composition, SSN & Citizenship Documentation") + [
        table(rows, [1.35 * inch, 0.85 * inch, 0.8 * inch, 0.4 * inch, 0.95 * inch, 1.15 * inch, 1.5 * inch]),
        Spacer(1, 6), Paragraph("SSNs in this demo use the never-issued 000 area number.", SMALL),
    ]))

    # Student status
    rows = [["Member", "Age", "Student status"]]
    for m in h["members"]:
        rows.append([m["name"], str(age_on(d(m["dob"]), eff)), m["student"]])
    flow = page_header(h, "Student Status Certification") + [
        table(rows, [2.6 * inch, 0.6 * inch, 3.8 * inch])]
    for name, det in h.get("student_detail", {}).items():
        flow += [Paragraph(f"Student eligibility questionnaire - {name}", H2), kv([
            ("Enrollment", det["enrollment"]), ("Under age 24", det["age_under_24"]),
            ("Married", det["married"]), ("Veteran", det["veteran"]), ("Person with disabilities", det["disabled"]),
            ("Has a dependent child", det["dependent_child"]), ("Lives with parents", det["living_with_parents"]),
            ("Independent-student documentation", det["independent_student_docs"]),
            ("Parents' income certification", det["parent_income_cert"])])]
    flow += [Spacer(1, 8), table([["Head of household signature", "Date"], sig(head, signed)], [4.5 * inch, 2.5 * inch])]
    docs.append(("student", "Student Status Certification", flow))

    # Owner certification (50059 model)
    inc_rows = [["Member", "Income type", "Source", "Annual amount"]]
    for i in h["income"]:
        if i.get("on_cert", True):
            inc_rows.append([i["member"], i["kind"], i["source"], money(i["annual"])])
    for a in h["assets"]:
        if a["income"]:
            inc_rows.append([a["member"], "Asset income", a["desc"], money(a["income"])])
    sig_rows = [["Signer", "Signature", "Date"]]
    for m in adults(h, eff):
        ok = m["name"] not in h.get("cert_unsigned", [])
        sig_rows.append([f"{m['name']} ({m['rel']})"] + sig(m["name"], signed if ok else None))
    sig_rows.append([f"{AGENT_NAME}, {AGENT_TITLE} (owner/agent)"] + sig(AGENT_NAME, signed))
    calc_pairs = [("Total annual income", money(calc["annual"])),
                  (f"Dependent deduction ({calc['dependents'] if calc['dep_ded'] else 0} x ${DEPENDENT_DEDUCTION})", money(calc["dep_ded"])),
                  ("Child care expense deduction", money(calc["cc_ded"]))]
    if calc["ed_ded"]:
        calc_pairs += [("Elderly/disabled family deduction", money(calc["ed_ded"])),
                       ("Medical expense deduction (over 3% of income)", money(calc["med_ded"]))]
    calc_pairs += [("Adjusted annual income", money(calc["adjusted"])),
                   ("Total Tenant Payment (TTP)", money(calc["ttp"])),
                   ("Utility allowance", money(calc["utility_allowance"])),
                   ("Tenant rent", money(calc["tenant_rent"])),
                   ("Contract rent", money(calc["contract_rent"])),
                   ("Assistance payment (HAP)", money(calc["hap"]))]
    docs.append(("certification", "Owner Certification (50059 model)", page_header(
        h, "Owner's Certification of Compliance - Household Income & Rent",
        "Simplified model of form HUD-50059. Not an official HUD form.") + [
        kv([("Certification type", h["cert"]["type"]), ("Effective date", fmt_date(eff)),
            ("Next annual recertification", fmt_date(next_ar(h))),
            ("Head of household", head), ("Household size", str(calc["household_size"])), ("Unit", f"{h['unit']} ({ut})")]),
        Paragraph("Household members", H2),
        table([["Member", "Relationship", "Age", "SSN"]] + [[m["name"], m["rel"], str(age_on(d(m["dob"]), eff)), m["ssn"]] for m in h["members"]],
              [2.4 * inch, 1.6 * inch, 0.8 * inch, 2.2 * inch]),
        Paragraph("Income", H2), table(inc_rows, [1.7 * inch, 1.3 * inch, 2.6 * inch, 1.4 * inch]),
        Paragraph("Calculation", H2), kv(calc_pairs),
        Paragraph("Signatures - all adult household members and the owner/agent", H2),
        table(sig_rows, [3.2 * inch, 2.4 * inch, 1.4 * inch]),
    ]))

    # Third-party verifications, one page each
    for n, i in enumerate(h["income"], 1):
        docs.append((f"verification_{n}", f"Verification - {i['kind']} ({i['member']})", page_header(
            h, f"Third-Party Verification - {i['kind']}") + [
            kv([("Household member", i["member"]), ("Source", i["source"]), ("Income type", i["kind"]),
                ("Amount verified", i["detail"]), ("Annualized", money(i["annual"])),
                ("Request sent by owner", fmt_date(d(i["sent"]))), ("Received by owner", fmt_date(d(i["received"]))),
                ("Verification method", "Written third-party verification" if i["kind"] == "Employment" else "Benefit/award letter")]),
            Spacer(1, 10), table([["Verifier signature", "Date"],
                                  [Paragraph(f"/s/ Payroll / benefits office, {i['source']}", SIG), fmt_date(d(i["received"]))]],
                                 [4.5 * inch, 2.5 * inch]),
        ]))
    if h.get("childcare"):
        c = h["childcare"]
        docs.append(("verification_childcare", "Verification - Child Care Expense", page_header(
            h, "Third-Party Verification - Child Care Expense") + [
            kv([("Provider", c["provider"]), ("Children in care", ", ".join(m["name"] for m in h["members"] if m["rel"] == "Dependent")),
                ("Annual out-of-pocket cost", money(c["annual"])), ("Reason", "Enables head of household to work"),
                ("Received by owner", fmt_date(d(c["received"])))]),
        ]))
    docs.append(("verification_assets", "Verification - Assets", page_header(h, "Asset Verification / Self-Certification") + [
        table([["Member", "Asset", "Cash value", "Annual income"]] +
              [[a["member"], a["desc"], money(a["value"]), money(a["income"])] for a in h["assets"]],
              [1.7 * inch, 2.9 * inch, 1.2 * inch, 1.2 * inch]),
        Spacer(1, 6), Paragraph("Net family assets under $5,000: household self-certified. "
                                "Disposal-of-assets declaration: no assets disposed of for less than fair market value in the past 2 years.", BODY),
    ]))
    if h.get("medical"):
        docs.append(("verification_medical", "Verification - Medical Expenses", page_header(
            h, "Third-Party Verification - Unreimbursed Medical Expenses") + [
            kv([("Household", head), ("Unreimbursed medical expenses (12 months)", money(h["medical"])),
                ("Components", "Medicare Part B & D premiums, prescription copays"),
                ("Received by owner", fmt_date(d(h["income"][0]["received"])))]),
        ]))
    if h.get("custody"):
        c = h["custody"]
        docs.append(("custody", "Custody Arrangement", page_header(
            h, "Custody Arrangement - Summary of Parenting-Time Order") + [
            kv([("Case number", c["case"]), ("Child", c["child"]), ("Terms", c["terms"]),
                ("Overnights with this household", c["overnights_with_father"]),
                ("Copy of order received", fmt_date(d(h["income"][0]["received"])))]),
        ]))

    # EIV
    e = h["eiv"]
    docs.append(("eiv", "EIV Income Report & Discrepancy Resolution", page_header(
        h, "EIV Income Report & Discrepancy Resolution",
        "Summary of the Enterprise Income Verification report printed for this certification.") + [
        kv([("Report pulled", fmt_date(d(e["pulled"]))), ("For certification effective", fmt_date(eff))]),
        Paragraph("Income reported to HUD", H2),
        table([["Member", "Source", "Period", "Amount"]] + [[r[0], r[1], r[2], money(r[3]) if r[3] else "-"] for r in e["rows"]],
              [1.7 * inch, 2.9 * inch, 1.1 * inch, 1.3 * inch]),
        Paragraph("Discrepancy review and written resolution", H2),
        Paragraph(escape(e["note"] or ""), BODY), Spacer(1, 30 if not e["note"] else 4),
    ]))

    # Recert notices
    if h["notices"]:
        n = h["notices"]
        due = next_ar(h) if h.get("recert_missed") else d(h["cert"]["effective"])
        rows = [["Notice", "Due by", "Date sent"]]
        for k, label in (("120", "Initial recertification notice (120 days)"), ("90", "First reminder (90 days)"),
                         ("60", "Second reminder (60 days)"), ("30", "Final reminder / notice of rent change (30 days)")):
            due_by = date.fromordinal(due.toordinal() - int(k))
            rows.append([label, fmt_date(due_by), fmt_date(d(n[k])) if n[k] else ""])
        docs.append(("notices", "Recertification Notice Log", page_header(
            h, "Annual Recertification Notice Log", f"Recertification due {fmt_date(due)}.") + [
            table(rows, [3.6 * inch, 1.6 * inch, 1.8 * inch]), Spacer(1, 6),
            kv([("Tenant response", n["response"])]),
            Spacer(1, 4), Paragraph("Reminders at 90/60/30 days are required only if the tenant has not responded.", SMALL),
        ]))

    # Lease & rent
    docs.append(("lease", "Lease & Rent Summary", page_header(h, "Lease & Rent Summary") + [
        kv([("Lease", "HUD Model Lease for Subsidized Programs (fictional copy)"),
            ("Original lease start", fmt_date(d(h["lease_start"]))), ("Term", "12 months, renews month-to-month"),
            ("Contract rent", money(calc["contract_rent"])),
            ("Tenant rent (per current rent change notice)", money(calc["tenant_rent"])),
            ("Assistance payment (HAP)", money(calc["hap"])),
            ("Rent change effective", fmt_date(eff)),
            ("House rules, pet & lead-paint addenda", "Signed at move-in"),
            ("Resident rights & responsibilities brochure", "Acknowledged " + fmt_date(d(h["consents_signed"])))]),
    ]))

    docs.append(("race_ethnicity", "Race & Ethnicity Data", page_header(h, "Race and Ethnic Data Reporting Form") + [
        Paragraph("Form offered to every household member. Completion is voluntary.", BODY), Spacer(1, 6),
        table([["Member", "Form status"]] + [[m["name"], "On file - " + fmt_date(d(h["consents_signed"]))] for m in h["members"]],
              [3.5 * inch, 3.5 * inch]),
    ]))
    return docs


def render_file(h, calc):
    docs = build_docs(h, calc)
    pages = {"cover": 1}
    for n, (key, _title, _flow) in enumerate(docs, 2):
        pages[key] = n
    total = len(docs) + 1

    index_rows = [["Document", "Page"]] + [[title, str(pages[key])] for key, title, _ in docs]
    hist = [["Certification", "Effective", "Tenant rent"]] + \
           [[t, fmt_date(d(e)), money(r)] for t, e, r in h["history"]] + \
           [[h["cert"]["type"] + " (current)", fmt_date(d(h["cert"]["effective"])), money(calc["tenant_rent"])]]
    story = page_header(h, f"Tenant File - Unit {h['unit']} - {h['members'][0]['name']}") + [
        kv([("Property", PROPERTY_NAME), ("Address", PROPERTY_ADDRESS), ("Owner", OWNER_ENTITY),
            ("Management agent", MANAGEMENT_CO), ("Move-in date", fmt_date(d(h["move_in"])))]),
        Paragraph("File index", H2), table(index_rows, [5.8 * inch, 1.2 * inch]),
        Paragraph("Certification history", H2), table(hist, [3.8 * inch, 1.6 * inch, 1.6 * inch]),
    ]
    for _key, _title, flow in docs:
        story += [PageBreak()] + flow

    fname = f"unit_{h['unit']}_{h['slug']}.pdf"

    def on_page(canv, doc):
        canv.saveState()
        canv.setFont("Helvetica", 7)
        canv.setFillColor(colors.HexColor("#666666"))
        canv.drawString(0.75 * inch, 0.5 * inch, FOOTER)
        canv.drawRightString(7.75 * inch, 0.35 * inch, f"{fname} - Page {doc.page} of {total}")
        canv.restoreState()

    path = os.path.join(FILES_DIR, fname)
    SimpleDocTemplate(path, pagesize=letter, leftMargin=0.75 * inch, rightMargin=0.75 * inch,
                      topMargin=0.7 * inch, bottomMargin=0.8 * inch,
                      title=f"{PROPERTY_NAME} Unit {h['unit']} tenant file (fictional)",
                      author=MANAGEMENT_CO).build(story, onFirstPage=on_page, onLaterPages=on_page)
    # Page citations only mean something if every document fit on one page.
    actual = len(PdfReader(path).pages)
    assert actual == total, f"{fname}: expected {total} pages, got {actual} (a document overflowed)"
    return fname, pages


# --------------------------------------------------------------------------
# Findings: each one computed from the same data the PDFs were drawn from.
# --------------------------------------------------------------------------
def build_findings(by_unit, calcs, cites):
    def cite(unit, *keys):
        f, pages = cites[unit]
        return [{"file": f"tenant_files/{f}", "page": pages[k], "document": k} for k in keys]

    def impact(monthly, start, basis, potential=False):
        months = months_between(start, REVIEW_DATE)
        return {"monthly": monthly, "months": months, "total": monthly * months,
                "through": fmt_date(date(REVIEW_DATE.year, REVIEW_DATE.month, 1)),
                "kind": "potential" if potential else "confirmed", "basis": basis}

    F = []
    h = by_unit["105"]; c = calcs["105"]["cert"]
    F.append({"id": "E1", "units": ["105"], "severity": "high", "category": "Missing signature",
              "title": "Spouse did not sign the annual certification",
              "detail": f"Doreen Whitfield (Spouse, age {age_on(d(h['members'][1]['dob']), d(h['cert']['effective']))}) has a blank signature and date on the {fmt_date(d(h['cert']['effective']))} certification. Only the head and the agent signed.",
              "rule": "Every adult household member (head, spouse, co-head, and other adults) must sign the certification. Without all signatures it is not a valid certification, and assistance paid on it can be recovered.",
              "citation": "HUD Handbook 4350.3, Ch. 7 (recertification) and the HUD-50059 signature requirements",
              "evidence": cite("105", "certification"),
              "dollar_impact": impact(c["hap"], d(h["cert"]["effective"]), "HAP paid on an invalid certification is at risk of recovery", True)})

    F.append({"id": "E2", "units": ["108"], "severity": "medium", "category": "Missing consent form",
              "title": "Adult son never signed HUD-9887 / 9887-A",
              "detail": "Emeka Okafor turned 18 on 11/14/2025 and is listed as an adult on the 07/01/2026 certification, but his 9887 and 9887-A signature lines are blank.",
              "rule": "Every household member age 18 or older must sign the consent forms (HUD-9887 and 9887-A), including members who turn 18 between certifications. Without consent the owner cannot verify that member's income.",
              "citation": "24 CFR 5.230; HUD Handbook 4350.3, Ch. 5",
              "evidence": cite("108", "consents", "household"), "dollar_impact": None})

    h = by_unit["108"]
    rec = d(h["income"][0]["received"]); eff = d(h["cert"]["effective"])
    F.append({"id": "E3", "units": ["108"], "severity": "medium", "category": "Stale verification",
              "title": "Head's employment verification was too old when used",
              "detail": f"Grace Okafor's employment verification was received {fmt_date(rec)}; the certification it supports was signed {fmt_date(d(h['cert']['signed']))} and took effect {fmt_date(eff)} - {(eff - rec).days} days after receipt.",
              "rule": f"Third-party verifications are valid for {VERIFICATION_VALID_DAYS} days from the date the owner receives them. Older verifications must be redone before the certification is completed.",
              "citation": "HUD Handbook 4350.3, Ch. 5 (verification); see docs/research/section8.md §2",
              "evidence": cite("108", "verification_1", "certification"), "dollar_impact": None})

    h = by_unit["203"]; cert, true = calcs["203"]["cert"], calcs["203"]["true"]
    diff = true["tenant_rent"] - cert["tenant_rent"]
    F.append({"id": "E4", "units": ["203"], "severity": "high", "category": "Unresolved EIV discrepancy",
              "title": "EIV shows a second employer that is not on the certification",
              "detail": f"The EIV report pulled {fmt_date(d(h['eiv']['pulled']))} shows wages from Lakeside Logistics, Inc. ($3,120 in each of Q4 2025 and Q1 2026). The certification lists only Riverton Public Library income, and the discrepancy-resolution section is blank.",
              "rule": "Owners must review EIV income reports and resolve every income discrepancy with the tenant, documenting the resolution in writing in the file. EIV discrepancies are the most common MOR finding.",
              "citation": "24 CFR 5.233; HUD Handbook 4350.3, Ch. 9 (EIV)",
              "evidence": cite("203", "eiv", "certification"),
              "dollar_impact": impact(diff, d(h["cert"]["effective"]),
                                      f"If the Lakeside wages are current (~{money(h['unreported_income']['annual'])}/yr), tenant rent is understated by {money(diff)}/mo ({money(cert['tenant_rent'])} certified vs {money(true['tenant_rent'])}) and HAP is overpaid by the same amount", True)})

    h = by_unit["206"]; cert, true = calcs["206"]["cert"], calcs["206"]["true"]
    diff = cert["tenant_rent"] - true["tenant_rent"]
    F.append({"id": "E5", "units": ["206"], "severity": "high", "category": "Rent calculation error",
              "title": "Dependent and child-care deductions verified but not taken",
              "detail": f"The file verifies two dependent children and {money(h['childcare']['annual'])}/yr of child care, but the certification shows $0 for both deductions. Adjusted income is {money(cert['adjusted'])} on the cert; it should be {money(true['adjusted'])}.",
              "rule": f"Adjusted income must subtract ${DEPENDENT_DEDUCTION} per dependent and reasonable child-care expenses that enable a family member to work. TTP is the greater of 30% of monthly adjusted income, 10% of monthly gross income, or the minimum rent.",
              "citation": "24 CFR 5.611 (adjusted income), 5.628 (TTP); HUD Handbook 4350.3, Ch. 5",
              "evidence": cite("206", "certification", "household", "verification_childcare"),
              "dollar_impact": impact(diff, d(h["cert"]["effective"]),
                                      f"Tenant overcharged {money(diff)}/mo ({money(cert['tenant_rent'])} charged vs {money(true['tenant_rent'])} correct); owed back to the tenant")})

    h = by_unit["207"]; c = calcs["207"]["cert"]; due = d(h["next_ar"])
    F.append({"id": "E6", "units": ["207"], "severity": "high", "category": "Late recertification",
              "title": "Annual recertification is 2 months overdue and reminders were never sent",
              "detail": f"The annual recertification was due {fmt_date(due)}. The 120-day notice went out {fmt_date(d(h['notices']['120']))}; the tenant never responded and no 90-, 60-, or 30-day reminder was sent. As of {fmt_date(REVIEW_DATE)} the household is still on the {fmt_date(d(h['cert']['effective']))} certification.",
              "rule": "Owners must send the initial notice 120 days before the recertification date and reminders at 90 and 60 days (plus a 30-day notice) if the tenant has not responded, and complete the recertification by the anniversary date. Because the reminders were not sent, the owner cannot raise the tenant to market rent and is out of compliance.",
              "citation": "HUD Handbook 4350.3, Ch. 7 (annual recertification notices and deadlines)",
              "evidence": cite("207", "notices", "certification"),
              "dollar_impact": impact(c["hap"], due, "HAP paid after the anniversary date on an expired certification is at risk", True)})

    h = by_unit["301"]; cert, true = calcs["301"]["cert"], calcs["301"]["true"]
    limit = LOW_INCOME_LIMIT[len(h["members"])]
    F.append({"id": "E7", "units": ["301"], "severity": "critical", "category": "Over income at admission",
              "title": "Adult daughter's income left off the move-in certification; household was over the income limit",
              "detail": f"The application and a third-party verification in the file show Freya Lindqvist employed at Harbor Coffee Roasters ({money(h['income'][1]['annual'])}/yr), but the move-in certification counts only Karin's {money(h['income'][0]['annual'])}. True household income is {money(true['annual'])}, over the {money(limit)} low-income limit for 2 people. The certified {money(cert['annual'])} made the household look eligible.",
              "rule": "Annual income includes all adult members' income. At admission, household income may not exceed the applicable low-income limit for the household size.",
              "citation": "24 CFR 5.609 (annual income), 5.653 (income eligibility at admission); HUD Handbook 4350.3, Ch. 3",
              "evidence": cite("301", "application", "verification_2", "certification"),
              "dollar_impact": impact(cert["hap"], d(h["cert"]["effective"]), "Household was not eligible; all HAP since move-in is at risk", True)})

    h304, h309 = by_unit["304"], by_unit["309"]; c309, t309 = calcs["309"]["cert"], calcs["309"]["true"]
    F.append({"id": "E8", "units": ["304", "309"], "severity": "high", "category": "Duplicate household member",
              "title": "Same child counted in two assisted households",
              "detail": f"Jaylen Brooks (SSN {JAYLEN['ssn']}, DOB {fmt_date(d(JAYLEN['dob']))}) is a dependent on both the Unit 304 certification (mother) and the Unit 309 certification (father). Unit 309's own custody summary says he spends 2 of 7 nights there, so he belongs to Unit 304 only.",
              "rule": "A person may be a member of only one assisted household. For joint custody, the child is counted in the household where they live more than 50% of the time, and only that household takes the dependent deduction.",
              "citation": "24 CFR 5.216 (SSN disclosure lets duplicates be caught); HUD Handbook 4350.3, Ch. 3 (household composition, joint custody)",
              "evidence": cite("304", "household", "certification") + cite("309", "household", "certification", "custody"),
              "dollar_impact": impact(t309["tenant_rent"] - c309["tenant_rent"], d(h309["cert"]["effective"]),
                                      f"Unit 309 took a ${DEPENDENT_DEDUCTION} dependent deduction it isn't entitled to: tenant rent {money(c309['tenant_rent'])} should be {money(t309['tenant_rent'])}, so HAP is overpaid by the difference. Unit 309 is also certified at household size 2 instead of 1.")})
    h = by_unit["304"]; c = calcs["304"]["cert"]
    diff = h["rent_roll_tenant_rent"] - c["tenant_rent"]
    F.append({"id": "E9", "units": ["304"], "severity": "high", "category": "Rent roll vs certification",
              "title": "Rent roll still charges the pre-interim rent",
              "detail": f"The 05/01/2026 interim recertification lowered tenant rent to {money(c['tenant_rent'])}, but the rent roll still charges {money(h['rent_roll_tenant_rent'])} (the 01/01/2026 amount). On the rent roll, tenant rent plus HAP ({money(c['hap'])}) adds up to {money(h['rent_roll_tenant_rent'] + c['hap'])}, which is {money(diff)} more than the {money(c['contract_rent'])} contract rent.",
              "rule": "Tenant rent must equal the amount on the current effective certification. When an interim decrease is processed, the new rent applies from its effective date.",
              "citation": "HUD Handbook 4350.3, Ch. 7 (interim recertifications)",
              "evidence": cite("304", "certification", "lease") + [{"file": "rent_roll/oak_hollow_rent_roll_40unit.xlsx", "row": "Unit 304"}],
              "dollar_impact": impact(diff, d(h["cert"]["effective"]), f"Tenant overcharged {money(diff)}/mo; owed back to the tenant")})

    h = by_unit["307"]; c = calcs["307"]["cert"]
    F.append({"id": "E10", "units": ["307"], "severity": "critical", "category": "Ineligible student",
              "title": "Full-time student household with no exemption or parental-income documentation",
              "detail": "Tyler Ostrander (age 20) is a full-time student living alone. He is unmarried, not a veteran, not disabled, and has no dependent child. The file has no independent-student documentation and no certification of his parents' income.",
              "rule": "A student under 24 who is not a veteran, not married, has no dependent child, and is not a person with disabilities is ineligible for Section 8 assistance unless the student is independent of their parents or the parents are also income-eligible - and that must be documented.",
              "citation": "24 CFR 5.612; HUD Handbook 4350.3, Ch. 3 (student eligibility)",
              "evidence": cite("307", "student", "household", "certification"),
              "dollar_impact": impact(c["hap"], d(h["move_in"]), "Household may never have been eligible; HAP since move-in is at risk (amount uses the current HAP for every month)", True)})
    return F


# --------------------------------------------------------------------------
# Rent roll: 40 units, the 10 file units drawn from their certifications.
# --------------------------------------------------------------------------
FIRST = ["Avery", "Bianca", "Carlos", "Dana", "Elijah", "Fatima", "Gavin", "Hana", "Isaac", "Jasmine",
         "Kofi", "Lena", "Malik", "Nadia", "Owen", "Priya", "Quinn", "Rosa", "Samuel", "Tessa",
         "Uriel", "Vivian", "Wesley", "Ximena", "Yusuf", "Zoe", "Bernard", "Celeste", "Dmitri", "Esther"]
LAST = ["Abernathy", "Barrow", "Castellanos", "Dunmore", "Ellison", "Fairbanks", "Galloway", "Hartigan",
        "Iverson", "Jablonski", "Kettering", "Larkspur", "Montague", "Northcott", "Oyelaran", "Pembrook",
        "Quintero", "Rasmussen", "Sorensen", "Thackeray", "Underhill", "Valdivia", "Winslow", "Yarborough",
        "Zelinski", "Ashcombe", "Blackwood", "Corrigan", "Delacroix", "Esposito"]


def rent_roll_rows(by_unit, calcs):
    rng = random.Random(RNG_SEED)
    names = [f"{l}, {f}" for f, l in zip(FIRST, LAST)]
    rng.shuffle(names)
    rows = []
    for floor in range(1, 5):
        for n in range(1, 11):
            unit = f"{floor}{n:02d}"
            ut = FLOOR_LAYOUT[n]
            t = UNIT_TYPES[ut]
            base = {"Unit": unit, "Type": ut, "Sq Ft": t["sqft"], "Contract Rent": t["contract_rent"],
                    "Utility Allowance": t["utility_allowance"]}
            if unit in VACANT_UNITS:
                rows.append({**base, "Status": "Vacant", "Head of Household": "VACANT", "HH Size": "",
                             "Move-In": "", "Cert Type": "", "Cert Effective": "", "Next AR Due": "",
                             "Tenant Rent": 0, "HAP": 0, "Utility Reimb.": 0, "Balance": 0})
                continue
            if unit in by_unit:
                h, c = by_unit[unit], calcs[unit]["cert"]
                eff = d(h["cert"]["effective"])
                ctype = h["cert"]["type"].split("(")[1].split(")")[0]
                tenant_rent = h.get("rent_roll_tenant_rent", c["tenant_rent"])
                last, first = h["members"][0]["name"].split(" ", 1)[1], h["members"][0]["name"].split(" ")[0]
                rows.append({**base, "Status": "Occupied", "Head of Household": f"{last}, {first}",
                             "HH Size": c["household_size"], "Move-In": d(h["move_in"]), "Cert Type": ctype,
                             "Cert Effective": eff, "Next AR Due": next_ar(h), "Tenant Rent": tenant_rent,
                             "HAP": c["hap"], "Utility Reimb.": c["utility_reimbursement"], "Balance": 0})
                continue
            # Filler household: consistent numbers, nothing planted.
            size = min(t["beds"] * 2, rng.choice([1, 2, 2, 3, 3, 4]))
            annual = rng.randrange(9000, 38000, 10)
            deps = max(0, size - rng.choice([1, 1, 2]))
            adjusted = annual - DEPENDENT_DEDUCTION * deps
            ttp = max(round_dollar(adjusted * 0.3 / 12), round_dollar(annual * 0.1 / 12), MINIMUM_RENT)
            tr = max(ttp - t["utility_allowance"], 0)
            # Certified sometime Nov 2025 - Sep 2026, so no filler recert is overdue.
            months_after_oct_2025 = rng.randrange(1, 12)
            eff = date(2025 + (9 + months_after_oct_2025) // 12, (9 + months_after_oct_2025) % 12 + 1, 1)
            rows.append({**base, "Status": "Occupied", "Head of Household": names.pop(), "HH Size": size,
                         "Move-In": date(rng.randrange(2014, 2025), eff.month, 1), "Cert Type": "AR",
                         "Cert Effective": eff, "Next AR Due": eff.replace(year=eff.year + 1),
                         "Tenant Rent": tr, "HAP": t["contract_rent"] - tr,
                         "Utility Reimb.": max(t["utility_allowance"] - ttp, 0), "Balance": 0})
    return rows


RR_COLS = ["Unit", "Type", "Sq Ft", "Status", "Head of Household", "HH Size", "Move-In", "Cert Type",
           "Cert Effective", "Next AR Due", "Contract Rent", "Utility Allowance", "Tenant Rent", "HAP",
           "Utility Reimb.", "Balance"]


def write_rent_roll(rows):
    wb = Workbook()
    wb.properties.creator = MANAGEMENT_CO
    wb.properties.created = wb.properties.modified = datetime(2026, 10, 2, 8, 0, 0)
    ws = wb.active
    ws.title = "Rent Roll"
    ws.append([PROPERTY_NAME])
    ws.append([PROPERTY_ADDRESS])
    ws.append([f"Affordable Rent Roll - {PROGRAM} - HAP Contract {HAP_CONTRACT}"])
    ws.append([f"As of {fmt_date(REVIEW_DATE)} - Generated {fmt_date(RENT_ROLL_GENERATED)} - FICTIONAL DEMO DATA"])
    ws.append([])
    ws.append(RR_COLS)
    for c in ws[1]:
        c.font = Font(bold=True, size=13)
    for c in ws[6]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8EEF4")
        c.alignment = Alignment(wrap_text=True)
    money_cols = {"Contract Rent", "Utility Allowance", "Tenant Rent", "HAP", "Utility Reimb.", "Balance"}
    for r in rows:
        ws.append([r[c] for c in RR_COLS])
        for i, col in enumerate(RR_COLS, 1):
            cell = ws.cell(row=ws.max_row, column=i)
            if col in money_cols:
                cell.number_format = '"$"#,##0'
            elif isinstance(r[col], date):
                cell.number_format = "mm/dd/yyyy"
    occ = [r for r in rows if r["Status"] == "Occupied"]
    ws.append([])
    ws.append(["Totals", "", "", f"{len(occ)}/{len(rows)} occupied", "", "", "", "", "", "",
               sum(r["Contract Rent"] for r in occ), "", sum(r["Tenant Rent"] for r in occ),
               sum(r["HAP"] for r in occ), sum(r["Utility Reimb."] for r in occ), 0])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
        if c.column >= 11:
            c.number_format = '"$"#,##0'
    widths = [7, 6, 7, 10, 24, 8, 11, 9, 12, 12, 10, 10, 10, 9, 10, 9]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B7"
    wb.save(os.path.join(RENT_ROLL_DIR, "oak_hollow_rent_roll_40unit.xlsx"))

    with open(os.path.join(RENT_ROLL_DIR, "oak_hollow_rent_roll_40unit.csv"), "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(RR_COLS)
        for r in rows:
            w.writerow([fmt_date(r[c]) if isinstance(r[c], date) else r[c] for c in RR_COLS])


# --------------------------------------------------------------------------
# README
# --------------------------------------------------------------------------
def write_readme(findings, cites, rows):
    def ev(e):
        return ", ".join(f"`{x['file'].split('/')[-1]}` p.{x['page']}" if "page" in x else f"`{x['file'].split('/')[-1]}` ({x['row']})"
                         for x in e)

    def dol(di):
        if not di:
            return "None directly - compliance finding"
        return (f"{money(di['total'])} {di['kind']} ({money(di['monthly'])}/mo x {di['months']} mo). {di['basis']}.")

    lines = []
    for f in findings:
        lines += [f"### {f['id']} - Unit {' & '.join(f['units'])}: {f['title']}", "",
                  f"- **Severity:** {f['severity']} · **Category:** {f['category']}",
                  f"- **What's wrong:** {f['detail']}",
                  f"- **Rule it breaks:** {f['rule']}",
                  f"- **Citation:** {f['citation']}",
                  f"- **Where to look:** {ev(f['evidence'])}",
                  f"- **Dollar impact:** {dol(f['dollar_impact'])}", ""]
    confirmed = sum(f["dollar_impact"]["total"] for f in findings if f["dollar_impact"] and f["dollar_impact"]["kind"] == "confirmed")
    potential = sum(f["dollar_impact"]["total"] for f in findings if f["dollar_impact"] and f["dollar_impact"]["kind"] == "potential")
    file_rows = "\n".join(
        f"| {u} | `tenant_files/{cites[u][0]}` | {len(cites[u][1])} | {', '.join(f['id'] for f in findings if u in f['units']) or '**clean** (control)'} |"
        for u in sorted(cites))
    occ = [r for r in rows if r["Status"] == "Occupied"]
    text = f"""# Demo Property: Oak Hollow Apartments (Section 8)

A fully fictional 40-unit **Project-Based Section 8 (PBRA)** property with
10 complete tenant files and **{len(findings)} planted compliance errors**, for
demoing affordable-housing file review. Nothing here is real: every person,
employer, address, case number, and dollar figure was invented. SSNs all use
the never-issued `000` area number (`000-00-0001` to `000-00-0021`).
Forms are simplified models of the HUD forms, not the official forms.

- **Property:** {PROPERTY_NAME}, {PROPERTY_ADDRESS}
- **Owner:** {OWNER_ENTITY} · **Agent:** {MANAGEMENT_CO}
- **Program:** {PROGRAM}, HAP contract `{HAP_CONTRACT}` (fictional)
- **Units:** 40 (12 one-bed / 24 two-bed / 4 three-bed), {len(occ)} occupied, 2 vacant (402, 410)
- **Review date:** {fmt_date(REVIEW_DATE)} - every "late" or "stale" error and every
  dollar total is measured to this fixed date, so the demo never drifts.

## Files

```
tenant_files/unit_XXX_<name>.pdf               10 tenant files, one multi-page PDF each
rent_roll/oak_hollow_rent_roll_40unit.xlsx     40-unit affordable rent roll (also .csv)
expected_findings.json                         machine-readable answer key (pages + dollars)
generate_s8_demo.py                            generator - the source of truth for everything here
```

Each tenant file has: a cover page (file index + certification history),
rental application & supplement, 9887/9887-A consent signatures, household
composition / SSN / citizenship, student status, the owner certification
(50059 model) with the full rent calculation and signatures, one page per
third-party verification (income, child care, assets, medical), EIV income
report with the written discrepancy resolution, recertification notice log,
lease & rent summary, and race/ethnicity form status. **Every document is
exactly one page**, so page citations are stable (the generator asserts this).

| Unit | File | Pages | Errors |
|---|---|---|---|
{file_rows}

The other 30 occupied units appear only on the rent roll, with internally
consistent numbers and nothing planted.

## Planted errors

Totals as of {fmt_date(REVIEW_DATE)}: **{money(confirmed)} confirmed** (tenant
overcharges owed back, plus HAP overpaid on the duplicate dependent), and **{money(potential)} potential** subsidy at risk
(HAP that HUD could recover, or understated rent if the EIV wages hold up).

{chr(10).join(lines)}
## Rules the demo applies

The calculations use the **pre-HOTMA** rules (dependent deduction $480,
elderly/disabled family deduction $400, medical expenses over 3% of income,
TTP = greater of 30% of monthly adjusted income, 10% of monthly gross
income, or $25 minimum rent; tenant rent = TTP minus utility allowance). Per
`docs/research/section8.md` §5, HUD extended full multifamily HOTMA
compliance to 01/01/2027, so a 2026 file plausibly still uses these. Income
limits, contract rents, and utility allowances are **made up** (2-person low
income limit {money(LOW_INCOME_LIMIT[2])}; 1BR/2BR/3BR contract rents
{money(UNIT_TYPES['1BR']['contract_rent'])}/{money(UNIT_TYPES['2BR']['contract_rent'])}/{money(UNIT_TYPES['3BR']['contract_rent'])}).

⚠️ The rule summaries and citations are paraphrased from public sources
(HUD Handbook 4350.3, 24 CFR Part 5, and `docs/research/section8.md`). They
have **not been checked by a compliance practitioner**. Confirm the
citations before showing this to anyone who does this work for a living.

## Regenerating

From `backend/`: `venv/bin/python benchmark_data/s8_demo/generate_s8_demo.py`.
Output is deterministic (fixed dates, seeded RNG, invariant PDFs): a re-run
with no code changes leaves the PDFs, CSV, JSON, and README byte-identical.
The `.xlsx` differs only because openpyxl stamps the save time into it.
"""
    with open(os.path.join(OUT_DIR, "README.md"), "w") as f:
        f.write(text)


def main():
    os.makedirs(FILES_DIR, exist_ok=True)
    os.makedirs(RENT_ROLL_DIR, exist_ok=True)
    by_unit = {h["unit"]: h for h in HOUSEHOLDS}
    assert len(by_unit) == 10
    calcs = {u: {"cert": calculate(h, True), "true": calculate(h, False)} for u, h in by_unit.items()}
    cites = {u: render_file(h, calcs[u]["cert"]) for u, h in by_unit.items()}
    findings = build_findings(by_unit, calcs, cites)
    rows = rent_roll_rows(by_unit, calcs)
    write_rent_roll(rows)

    # Sanity checks on the planted errors.
    assert calcs["102"]["cert"] == calcs["102"]["true"], "control file must be clean"
    assert calcs["301"]["true"]["annual"] > LOW_INCOME_LIMIT[2] >= calcs["301"]["cert"]["annual"]
    assert all(m["ssn"].startswith("000-00-") for h in HOUSEHOLDS for m in h["members"])

    with open(os.path.join(OUT_DIR, "expected_findings.json"), "w") as f:
        json.dump({"property": PROPERTY_NAME, "program": PROGRAM, "review_date": REVIEW_DATE.isoformat(),
                   "fictional": True, "findings": findings,
                   "clean_units": [u for u in cites if not any(u in f["units"] for f in findings)]},
                  f, indent=2)
        f.write("\n")
    write_readme(findings, cites, rows)
    for f in findings:
        di = f["dollar_impact"]
        print(f"{f['id']:>3} units {','.join(f['units']):<8} {f['title'][:60]:<60} "
              f"{(money(di['total']) + ' ' + di['kind']) if di else '-'}")


if __name__ == "__main__":
    main()
