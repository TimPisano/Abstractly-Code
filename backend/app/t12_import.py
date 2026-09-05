"""
T12 (trailing 12-month) operating statement import: parses a property's
operating statement well enough to extract ONE specific number --
actual, collected annual rental income -- for cross-checking against a
rent roll's own annualized total (see compute_t12_reconciliation in
portfolio.py). Deliberately does NOT attempt to parse the rest of a
full operating statement (vacancy loss, expense categories, NOI): the
Platform feature this exists for is specifically "cross-checks the
rent roll against the T12," not "ingest and store a full P&L" -- a
much bigger, different feature nobody asked for.

Same "no fixed format" philosophy as rent_roll_import.py: every
operator builds their own T12 template, with a real-world quirk on top
-- unlike a rent roll (one row per tenant, columns are fields), a T12
is one row per LINE ITEM (Total Rental Income, Vacancy Loss, Operating
Expenses, ...), with columns typically being months (Jan-Dec) plus
sometimes an explicit Total/Annual column. This module matches the
LABEL in each row's first cell against a list of known line-item
aliases (see _RENTAL_INCOME_LABEL_ALIASES), the same alias-matching
spirit as the rent roll importer applied to a different axis (rows
instead of columns).

The most important design decision here isn't parsing mechanics, it's
which number counts as "rent": a T12 routinely has BOTH a "Gross
Potential Rent" line (the theoretical, vacancy-inclusive maximum at
100% occupancy / market rate) and an actual "Rental Income" /"Rent
Revenue" line (what was really collected). Only the latter is the
right comparison against a rent roll's own actual, currently-occupied
rent -- comparing against gross potential would make a perfectly
healthy, fully-consistent portfolio look "short" by definition, since
potential > actual accounts for vacancy on its own. See
_POTENTIAL_INCOME_WORDS.
"""

import csv
import io
import re
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook

from app.normalize import normalize_header as _normalize_header, parse_currency


class T12ImportError(Exception):
    """
    Raised when the file can't be usefully parsed AT ALL -- wrong file
    type, completely empty, no recognizable way to compute an annual
    total for any row (no Total/Annual column AND fewer than all 12
    month columns present), or no row whose label is recognizable as
    actual (not potential/market) rental income. Without a real number
    to extract, there's nothing to cross-check against a rent roll.
    """


# Header aliases for the column holding each line item's annual figure.
_TOTAL_COLUMN_ALIASES = [
    "total", "annual total", "annual", "ytd", "ytd total",
    "12 month total", "trailing 12", "t12", "ttm",
]

# Twelve months' worth of header aliases, each key a canonical month id.
_MONTH_COLUMN_ALIASES: Dict[str, List[str]] = {
    "jan": ["jan", "january"], "feb": ["feb", "february"], "mar": ["mar", "march"],
    "apr": ["apr", "april"], "may": ["may"], "jun": ["jun", "june"],
    "jul": ["jul", "july"], "aug": ["aug", "august"], "sep": ["sep", "sept", "september"],
    "oct": ["oct", "october"], "nov": ["nov", "november"], "dec": ["dec", "december"],
}

# Row-label aliases for the ACTUAL (collected) rental income line item
# -- "Gross" here contrasts with "Net" (before vs. after operating
# expenses are subtracted), standard operating-statement terminology,
# NOT "gross potential" (before vs. after vacancy) -- a "Gross Rental
# Income" line is legitimately an actual-income alias, not a potential-
# income one. See _POTENTIAL_INCOME_WORDS for what IS excluded.
_RENTAL_INCOME_LABEL_ALIASES = [
    "total rental income", "rental income", "gross rental income",
    "base rental income", "total rent income", "rent revenue",
    "scheduled rent income", "total rent revenue",
]

# Label words that mean a rent-like line item is the THEORETICAL,
# vacancy-inclusive maximum -- not what was actually collected. A label
# containing "rent" or "income" together with any of these is never
# treated as the actual rental income line, even if no better line
# exists: see module docstring for why comparing against this would be
# systematically wrong, not just imprecise.
_POTENTIAL_INCOME_WORDS = {"potential", "market", "proforma", "projected", "asking", "scheduled"}
# "scheduled" is ambiguous in the wild (some operators use "Scheduled
# Rent" to mean the actual contracted rent, matching rent_roll_import.
# py's own "Scheduled Rent" alias for rent_amount there) -- but on a
# T12 specifically, "Scheduled Gross Income" is common terminology for
# the potential/asking figure before vacancy loss is deducted, a
# meaningfully different context from a rent roll's per-tenant column.
# Only "scheduled rent income" (an unambiguous actual-income phrasing)
# is allow-listed above as its own exact alias despite this; a bare
# "Scheduled Income" or "Scheduled Gross Income" is correctly excluded.


_HEADER_SCAN_WINDOW = 20


def _is_potential_income_label(normalized_label: str) -> bool:
    words = set(normalized_label.split())
    return bool(("rent" in words or "income" in words) and (words & _POTENTIAL_INCOME_WORDS))


def _match_t12_columns(headers: List[Any]) -> Dict[str, Any]:
    """Maps 'total' -> column index (if a Total/Annual column exists) and 'months' -> {month_id: column index} for whichever month columns exist."""
    normalized = [_normalize_header(h) for h in headers]
    total_norms = {_normalize_header(a) for a in _TOTAL_COLUMN_ALIASES}
    month_norms = {month: {_normalize_header(a) for a in aliases} for month, aliases in _MONTH_COLUMN_ALIASES.items()}

    mapping: Dict[str, Any] = {"months": {}}
    for i, h in enumerate(normalized):
        if h in total_norms and "total" not in mapping:
            mapping["total"] = i
            continue
        for month, norms in month_norms.items():
            if h in norms and month not in mapping["months"]:
                mapping["months"][month] = i
                break
    return mapping


def _find_t12_header_row(all_rows: List[List[Any]]) -> int:
    """
    Same reasoning as rent_roll_import.py's _find_header_row -- a T12
    export routinely opens with decorative rows (property name,
    "Operating Statement", a date range) before the real column-header
    row. Considered found once a row has either a real Total/Annual
    column, or at least 2 recognizable month columns (a T12 has no
    single-column analog to a rent roll's "must have both tenant AND
    rent" signal, so this is the closest bounded, low-false-positive
    equivalent -- a decorative title row essentially never contains 2+
    literal month names).
    """
    for i, row in enumerate(all_rows[:_HEADER_SCAN_WINDOW]):
        mapping = _match_t12_columns(row)
        if "total" in mapping or len(mapping["months"]) >= 2:
            return i
    return 0


def _parse_t12_currency(value: Any) -> Optional[float]:
    """Same tolerant-of-a-bare-number behavior as rent_roll_import.py's _parse_import_currency -- T12 spreadsheet cells routinely have no literal "$"."""
    if isinstance(value, (int, float)):
        return float(value)
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    strict = parse_currency(text)
    if strict is not None:
        return strict
    match = re.search(r"[\d,]+(?:\.\d+)?", text)
    if not match:
        return None
    try:
        return float(match.group(0).replace(",", ""))
    except ValueError:
        return None


def _row_annual_total(row: List[Any], column_mapping: Dict[str, Any]) -> Optional[float]:
    """A row's annual figure: the Total/Annual column if the file has one, otherwise the sum of all 12 month columns if the file has every one of them. None if neither is available."""
    if "total" in column_mapping:
        idx = column_mapping["total"]
        return _parse_t12_currency(row[idx]) if idx < len(row) else None
    months = column_mapping.get("months", {})
    if len(months) < 12:
        return None
    total = 0.0
    for idx in months.values():
        if idx < len(row):
            value = _parse_t12_currency(row[idx])
            if value is not None:
                total += value
    return total


def parse_t12_rows(
    headers: List[Any],
    rows: List[List[Any]],
    filename: str,
    header_row_offset: int = 0,
    row_numbers: Optional[List[int]] = None,
) -> Dict[str, Any]:
    """
    Core parsing logic, independent of file format -- see
    parse_csv_t12/parse_xlsx_t12 below. Scans every data row's first
    cell (the line-item label) for the first one recognizable as ACTUAL
    rental income (never a potential/market one, see
    _POTENTIAL_INCOME_WORDS), and computes that row's annual total from
    either a Total/Annual column or a sum of all 12 month columns.

    `row_numbers`, when given, is the TRUE 1-indexed original file row
    number for each entry in `rows` -- same precision fix as
    rent_roll_import.py's parse_rent_roll_rows (see that module's
    docstring for the full reasoning): header_row_offset's arithmetic
    alone silently drifts wrong by however many fully-blank rows were
    dropped before parsing, which matters once a decorative header
    block (routinely containing a blank spacer row) is in play.

    Returns `{annual_rental_income, source: {row, file, quote},
    column_mapping}`.

    Raises T12ImportError (fatal) if the file has no way to compute an
    annual figure at all (no Total column AND fewer than 12 month
    columns), or no row whose label is recognizable as actual rental
    income within the file.
    """
    column_mapping = _match_t12_columns(headers)

    if "total" not in column_mapping and len(column_mapping.get("months", {})) < 12:
        raise T12ImportError(
            "Couldn't find a column to compute an annual total from in this file's "
            "header row -- expected either a \"Total\"/\"Annual\" column, or all 12 "
            "month columns (Jan through Dec) to sum. Check that the file actually has "
            "a real column-header row within the first "
            f"{_HEADER_SCAN_WINDOW} rows (a title/date block before it is fine and is "
            "skipped automatically)."
        )

    for row_index, row in enumerate(rows):
        if not row:
            continue
        label_raw = row[0]
        label_str = str(label_raw).strip() if label_raw is not None else ""
        if not label_str:
            continue
        normalized_label = _normalize_header(label_str)
        alias_norms = {_normalize_header(a) for a in _RENTAL_INCOME_LABEL_ALIASES}

        # An EXACT match to an explicit allow-listed alias always wins,
        # even over the potential-income denylist below -- this is what
        # makes "Scheduled Rent Income" (allow-listed, unambiguous)
        # correctly override "scheduled" otherwise being a denylisted
        # word (needed to exclude the different, ambiguous "Scheduled
        # Gross Income" phrasing). Only a non-exact, substring-based
        # match is subject to the denylist check.
        if normalized_label in alias_norms:
            is_match = True
        elif _is_potential_income_label(normalized_label):
            is_match = False  # e.g. "Gross Potential Rent" -- never treated as actual income
        else:
            is_match = False
            for alias in _RENTAL_INCOME_LABEL_ALIASES:
                if re.search(rf"\b{re.escape(_normalize_header(alias))}\b", normalized_label):
                    is_match = True
                    break
        if not is_match:
            continue

        value = _row_annual_total(row, column_mapping)
        if value is None:
            continue  # this row matched the label but had no usable numbers -- keep scanning for a better row

        if row_numbers is not None:
            row_num = row_numbers[row_index]
        else:
            row_num = row_index + header_row_offset + 2  # +1 for 1-indexing, +1 because the header row itself precedes the data
        return {
            "annual_rental_income": value,
            "source": {"row": row_num, "file": filename, "quote": label_str},
            "column_mapping": column_mapping,
        }

    raise T12ImportError(
        "Couldn't find a row labeled as actual rental income (e.g. \"Total Rental "
        "Income\", \"Rental Income\", \"Gross Rental Income\") anywhere in this file "
        "-- a \"Gross Potential Rent\" or \"Market Rent\" line alone doesn't count, "
        "since that's the theoretical maximum before vacancy, not what was actually "
        "collected."
    )


def parse_csv_t12(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """Reads a CSV T12's bytes and parses it via parse_t12_rows. Raises T12ImportError for a genuinely empty file."""
    text = file_bytes.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    numbered_rows = [(i, row) for i, row in enumerate(reader, start=1) if any(cell.strip() for cell in row)]

    if not numbered_rows:
        raise T12ImportError("This CSV file is empty -- nothing to import.")

    row_contents = [row for _, row in numbered_rows]
    header_idx = _find_t12_header_row(row_contents)
    headers = row_contents[header_idx]
    data_entries = numbered_rows[header_idx + 1:]
    data_rows = [row for _, row in data_entries]
    data_row_numbers = [n for n, _ in data_entries]
    return parse_t12_rows(
        headers, data_rows, filename,
        header_row_offset=header_idx, row_numbers=data_row_numbers,
    )


def parse_xlsx_t12(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """Reads an .xlsx T12's bytes and parses it via parse_t12_rows. Uses the first (active) worksheet. Raises T12ImportError for a genuinely empty or unreadable file."""
    try:
        workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise T12ImportError(f"Couldn't read this file as an Excel workbook: {exc}")

    sheet = workbook.active
    numbered_rows = [
        (i, list(row)) for i, row in enumerate(sheet.iter_rows(values_only=True), start=1)
        if any(cell is not None and str(cell).strip() for cell in row)
    ]

    if not numbered_rows:
        raise T12ImportError("This Excel file is empty -- nothing to import.")

    row_contents = [row for _, row in numbered_rows]
    header_idx = _find_t12_header_row(row_contents)
    headers = row_contents[header_idx]
    data_entries = numbered_rows[header_idx + 1:]
    data_rows = [row for _, row in data_entries]
    data_row_numbers = [n for n, _ in data_entries]
    return parse_t12_rows(
        list(headers), data_rows, filename,
        header_row_offset=header_idx, row_numbers=data_row_numbers,
    )
