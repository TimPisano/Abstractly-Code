"""
Parse HUD's bulk annual files (CSV or XLSX) into store-ready rows.

HUD publishes one file per dataset per year, one row per county (or New
England town), keyed by a 10-digit FIPS code. Column names drift a
little year to year (`fmr_0` vs `fmr0`, `median2024` vs `median2025`,
`lim50_24p1` vs `lim50_25p1`), so headers are matched by pattern,
case-insensitively, rather than by exact name.

The parsers are strict on purpose: a missing column, a non-numeric
limit, a duplicate area, or a file whose year-stamped headers disagree
with the year being loaded all raise HudDataParseError with the row
number. Loading a wrong number silently is worse than loading nothing.
"""
from __future__ import annotations

import csv
import io
import re
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

from .models import (
    DATASET_FMR,
    DATASET_INCOME_LIMITS,
    DATASET_MTSP,
    DATASETS,
    AreaInfo,
    ParsedDataset,
)

HOUSEHOLD_SIZES = range(1, 9)
BEDROOM_SIZES = range(0, 5)


class HudDataParseError(ValueError):
    pass


_FIPS_NAMES = ("fips", "fips2020", "fips2010", "fips2000", "fips_code")
_HUD_AREA_CODE_NAMES = ("hud_area_code", "cbsasub", "hud_area")
_HUD_AREA_NAME_NAMES = ("hud_area_name", "areaname", "area_name")
_STATE_NAMES = ("stusps", "state_alpha", "state_abbr", "state")
_COUNTY_NAMES = ("county_name", "countyname", "cntyname", "county_town_name")

_MEDIAN_RE = re.compile(r"^median(\d{4})?$")
_IL_RES = {
    "extremely_low": re.compile(r"^eli_([1-8])$"),
    "very_low": re.compile(r"^l50_([1-8])$"),
    "low": re.compile(r"^l80_([1-8])$"),
}
_FMR_RE = re.compile(r"^fmr_?([0-4])$")
_MTSP_RE = re.compile(r"^lim(50|60)_(\d{2})?p([1-8])$")


def parse_file(dataset: str, content: bytes, filename: str, year: int) -> ParsedDataset:
    """Parse one HUD file for `dataset` ('il', 'fmr', 'mtsp') and `year`."""
    if dataset not in DATASETS:
        raise HudDataParseError(f"Unknown dataset {dataset!r}; expected one of {', '.join(DATASETS)}")
    _check_filename_year(filename, year)
    header, rows = _read_table(content, filename)
    columns = _Columns(header, year)
    parse_row = {
        DATASET_INCOME_LIMITS: columns.income_limits_row,
        DATASET_FMR: columns.fmr_row,
        DATASET_MTSP: columns.mtsp_row,
    }[dataset]
    columns.check_required(dataset)

    parsed = ParsedDataset(dataset=dataset)
    seen = set()
    for line_no, row in rows:
        if not any(_cell_text(v) for v in row):
            continue
        area = columns.area(row, line_no)
        if area.area_code in seen:
            raise HudDataParseError(f"Row {line_no}: duplicate area {area.area_code}")
        seen.add(area.area_code)
        parsed.rows.append((area, parse_row(row, line_no)))

    if not parsed.rows:
        raise HudDataParseError(f"{filename}: no data rows found")
    return parsed


_FILENAME_YEAR_RE = re.compile(r"fy[_-]?(\d{4}|\d{2})(?!\d)", re.IGNORECASE)


def _check_filename_year(filename: str, year: int) -> None:
    """
    FMR files and the income-limit columns (l50_N, ELI_N, l80_N) carry no
    year, so the header check can't catch an FY2025 file loaded as 2026.
    HUD's file names do carry it (Section8-FY25.xlsx, FY26_FMRs.xlsx), so
    refuse a name that names a different fiscal year.
    """
    m = _FILENAME_YEAR_RE.search(filename)
    if not m:
        return
    stated = int(m.group(1))
    if (stated if len(m.group(1)) == 4 else 2000 + stated) != year:
        raise HudDataParseError(
            f"File name {filename!r} says FY{m.group(1)}, but it is being loaded as {year}"
        )


# ---------------------------------------------------------------- reading

def _read_table(content: bytes, filename: str) -> Tuple[List[str], List[Tuple[int, list]]]:
    name = filename.lower()
    if name.endswith(".xlsx"):
        table = _read_xlsx(content)
    elif name.endswith(".csv"):
        table = _read_csv(content)
    else:
        raise HudDataParseError(f"{filename}: unsupported file type (expected .csv or .xlsx)")

    # Header = first row that names a FIPS column; HUD files sometimes
    # carry a title row above it.
    for idx, (line_no, row) in enumerate(table):
        normalized = [_normalize_header(v) for v in row]
        if any(h in _FIPS_NAMES for h in normalized):
            return normalized, table[idx + 1:]
    raise HudDataParseError(f"{filename}: no header row with a FIPS column ({', '.join(_FIPS_NAMES)})")


def _read_csv(content: bytes) -> List[Tuple[int, list]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    return [(i, row) for i, row in enumerate(csv.reader(io.StringIO(text)), start=1)]


def _read_xlsx(content: bytes) -> List[Tuple[int, list]]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several unrelated types
        raise HudDataParseError(f"Could not open workbook: {exc}") from exc
    try:
        ws = wb.worksheets[0]
        return [(i, list(row)) for i, row in enumerate(ws.iter_rows(values_only=True), start=1)]
    finally:
        wb.close()


def _normalize_header(value) -> str:
    return re.sub(r"\s+", "_", _cell_text(value).lower())


def _cell_text(value) -> str:
    return "" if value is None else str(value).strip()


# ---------------------------------------------------------------- columns

class _Columns:
    """Resolves header positions once, then extracts values per row."""

    def __init__(self, header: Sequence[str], year: int):
        self.header = list(header)
        self.year = year
        self.fips = self._first(_FIPS_NAMES)
        self.hud_area_code = self._first(_HUD_AREA_CODE_NAMES)
        self.hud_area_name = self._first(_HUD_AREA_NAME_NAMES)
        self.state = self._first(_STATE_NAMES)
        self.county = self._first(_COUNTY_NAMES)
        self.median = self._median()
        self.il = {key: self._indexed(rx) for key, rx in _IL_RES.items()}
        self.fmr = self._indexed(_FMR_RE)
        self.mtsp = {"50": {}, "60": {}}
        for pos, name in enumerate(self.header):
            m = _MTSP_RE.match(name)
            if m:
                pct, yy, size = m.groups()
                self._check_year_suffix(yy, name)
                self.mtsp[pct][int(size)] = pos

    def _first(self, names: Iterable[str]) -> Optional[int]:
        for name in names:
            if name in self.header:
                return self.header.index(name)
        return None

    def _indexed(self, rx) -> Dict[int, int]:
        out = {}
        for pos, name in enumerate(self.header):
            m = rx.match(name)
            if m:
                out[int(m.group(1))] = pos
        return out

    def _median(self) -> Optional[int]:
        for pos, name in enumerate(self.header):
            m = _MEDIAN_RE.match(name)
            if m:
                if m.group(1) and int(m.group(1)) != self.year:
                    raise HudDataParseError(
                        f"Column {name!r} says this file is for {m.group(1)}, "
                        f"but it is being loaded as {self.year}"
                    )
                return pos
        return None

    def _check_year_suffix(self, yy: Optional[str], name: str) -> None:
        if yy and int(yy) != self.year % 100:
            raise HudDataParseError(
                f"Column {name!r} says this file is for 20{yy}, "
                f"but it is being loaded as {self.year}"
            )

    def check_required(self, dataset: str) -> None:
        missing = []
        if dataset == DATASET_INCOME_LIMITS:
            for key, prefix in (("extremely_low", "ELI_"), ("very_low", "l50_"), ("low", "l80_")):
                missing += [f"{prefix}{n}" for n in HOUSEHOLD_SIZES if n not in self.il[key]]
        elif dataset == DATASET_FMR:
            missing += [f"fmr_{n}" for n in BEDROOM_SIZES if n not in self.fmr]
        else:
            for pct in ("50", "60"):
                missing += [f"lim{pct}_p{n}" for n in HOUSEHOLD_SIZES if n not in self.mtsp[pct]]
        if missing:
            raise HudDataParseError(f"Missing columns for {dataset}: {', '.join(missing)}")

    # -- per-row extraction

    def area(self, row: list, line_no: int) -> AreaInfo:
        return AreaInfo(
            area_code=normalize_area_code(self._cell(row, self.fips), line_no),
            hud_area_code=self._text(row, self.hud_area_code),
            hud_area_name=self._text(row, self.hud_area_name),
            state=self._state(row),
            county_name=self._text(row, self.county),
        )

    def income_limits_row(self, row: list, line_no: int) -> dict:
        payload = {"median_income": self._optional_amount(row, self.median, line_no)}
        for key, positions in self.il.items():
            payload[key] = self._by_household_size(row, positions, line_no)
        return payload

    def fmr_row(self, row: list, line_no: int) -> dict:
        return {"by_bedrooms": {n: self._amount(row, self.fmr[n], line_no) for n in BEDROOM_SIZES}}

    def mtsp_row(self, row: list, line_no: int) -> dict:
        return {
            "median_income": self._optional_amount(row, self.median, line_no),
            "limit_50": self._by_household_size(row, self.mtsp["50"], line_no),
            "limit_60": self._by_household_size(row, self.mtsp["60"], line_no),
        }

    def _by_household_size(self, row: list, positions: Dict[int, int], line_no: int) -> Dict[int, int]:
        # HUD income limits never fall as a household grows; a dip means
        # shifted or mislabeled columns.
        values = {n: self._amount(row, positions[n], line_no) for n in HOUSEHOLD_SIZES}
        for n in range(2, 9):
            if values[n] < values[n - 1]:
                raise HudDataParseError(
                    f"Row {line_no}: {self.header[positions[n]]!r} ({values[n]}) is lower than "
                    f"{self.header[positions[n - 1]]!r} ({values[n - 1]})"
                )
        return values

    def _cell(self, row: list, pos: Optional[int]):
        if pos is None or pos >= len(row):
            return None
        return row[pos]

    def _text(self, row: list, pos: Optional[int]) -> str:
        return _cell_text(self._cell(row, pos))

    def _state(self, row: list) -> str:
        # In the income-limit files `State` is the numeric state FIPS, not
        # the postal code; only keep a two-letter value.
        value = self._text(row, self.state).upper()
        return value if re.fullmatch(r"[A-Z]{2}", value) else ""

    def _amount(self, row: list, pos: int, line_no: int) -> int:
        value = parse_amount(self._cell(row, pos))
        if value is None:
            raise HudDataParseError(f"Row {line_no}: column {self.header[pos]!r} is blank")
        return _checked(value, self.header[pos], line_no)

    def _optional_amount(self, row: list, pos: Optional[int], line_no: int) -> Optional[int]:
        if pos is None:
            return None
        value = parse_amount(self._cell(row, pos))
        return None if value is None else _checked(value, self.header[pos], line_no)


def _checked(value, column: str, line_no: int) -> int:
    if isinstance(value, str):
        raise HudDataParseError(f"Row {line_no}: column {column!r} is not a number: {value!r}")
    if value <= 0:
        raise HudDataParseError(f"Row {line_no}: column {column!r} must be positive, got {value}")
    return value


def parse_amount(value):
    """
    Whole-dollar amount from a cell. Returns an int, None for blank, or
    the original text when it isn't a number (the caller reports it with
    the row and column).
    """
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        return int(round(value))
    text = str(value).strip().replace("$", "").replace(",", "")
    if not text:
        return None
    try:
        return int(round(float(text)))
    except (ValueError, OverflowError):  # "abc", "inf"; "nan" is ValueError too
        return str(value)


def normalize_area_code(value, line_no: Optional[int] = None) -> str:
    """
    HUD keys areas by a 10-digit FIPS: 2-digit state, 3-digit county,
    5-digit county subdivision (99999 when the whole county is one area,
    i.e. everywhere outside New England). Spreadsheets drop leading
    zeros, so Alabama's 0100199999 can arrive as 100199999 or 100199999.0.
    A 5-digit county FIPS is widened to the whole-county code.
    """
    where = f"Row {line_no}: " if line_no is not None else ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = _cell_text(value)
    if text.endswith(".0"):
        text = text[:-2]
    if not text.isdigit():
        raise HudDataParseError(f"{where}area code {text!r} is not a FIPS code")
    if len(text) <= 5:
        return text.zfill(5) + "99999"
    if len(text) <= 10:
        return text.zfill(10)
    raise HudDataParseError(f"{where}area code {text!r} is longer than 10 digits")
