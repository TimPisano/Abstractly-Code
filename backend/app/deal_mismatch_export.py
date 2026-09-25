"""
Deal Mismatch Report renderers: PDF and Excel, both reading from the one
data structure deal_mismatch.build_deal_mismatch_report_data returns, so
the two formats can never disagree about what they're reporting for the
same request -- same convention investment_memo.py's two renderers
already follow off build_investment_memo_data.

PDF: reportlab platypus (SimpleDocTemplate/Table/Paragraph), reusing
summary_memo.py's shared styles/footer/header machinery rather than
re-declaring fonts/colors/margins a third time in this codebase.

Excel: openpyxl, one sheet, styled the same way rent_roll_export.py's
workbook is (bold frozen header, currency number format on the impact
column) so a Deal Mismatch Report workbook looks like it came from the
same product as every other Abstractly export.
"""
import io
from datetime import date
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, Spacer, Table, TableStyle

from .summary_memo import (
    _build_pdf,
    _header_flowables,
    _styles,
    _INK,
    _MUTED,
    _RULE,
    _SEVERITY_COLORS,
)

DISCREPANCY_TYPE_LABELS = {
    "rent_mismatch": "Rent Mismatch",
    "expired_but_occupied": "Expired but Occupied",
    "unit_no_lease": "Unit on Rent Roll, No Lease",
    "lease_no_unit": "Lease on File, Not on Rent Roll",
    "concession_missing": "Concession Missing from Rent Roll",
    "dates_mismatch": "Lease Dates Mismatch",
    "tenant_mismatch": "Tenant Name Mismatch",
    "t12_income_gap": "Rent Roll vs. T12 Income",
    "t12_occupancy_mismatch": "Rent Roll vs. T12 Occupancy",
    "t12_concession_gap": "T12 Concessions Reported",
    "t12_bad_debt_trend": "T12 Bad Debt Trend",
}

_SEVERITY_ORDER = {"high": 0, "medium": 1, "low": 2}


def _fmt_money(value: Optional[float]) -> str:
    if value is None:
        return "—"
    sign = "-" if value < 0 else ""
    return f"{sign}${abs(value):,.2f}"


def _fmt_source(source: Optional[Dict[str, Any]]) -> str:
    if not source or not source.get("page"):
        return "—"
    return f"Page {source['page']}"


def _summary_flowables(data: Dict[str, Any]) -> List[Any]:
    scope = data.get("property_address") or "whole portfolio"
    overstatement = data.get("annual_income_overstatement")

    if overstatement is None:
        headline = "Not enough matched, dollar-quantified discrepancies to estimate a net income impact."
    elif overstatement > 0:
        headline = f"Rent roll <b>overstates</b> annual income by <b>{_fmt_money(overstatement)}</b>."
    elif overstatement < 0:
        headline = f"Rent roll <b>understates</b> annual income by <b>{_fmt_money(abs(overstatement))}</b>."
    else:
        headline = "Rent roll and lease-documented income agree overall (net impact: $0)."

    return [
        Paragraph(f"Scope: {scope}", _styles["body"]),
        Paragraph(f"Units checked: {data.get('total_units_checked', 0)}", _styles["body"]),
        Paragraph(f"Discrepancies found: {data.get('total_discrepancies', 0)}", _styles["body"]),
        Spacer(1, 4),
        Paragraph(headline, _styles["section"]),
    ]


def _discrepancies_table(rows: List[Dict[str, Any]]) -> Table:
    header = [
        Paragraph("<b>UNIT</b>", _styles["cell_label"]),
        Paragraph("<b>TYPE</b>", _styles["cell_label"]),
        Paragraph("<b>RENT ROLL</b>", _styles["cell_label"]),
        Paragraph("<b>LEASE</b>", _styles["cell_label"]),
        Paragraph("<b>SOURCE</b>", _styles["cell_label"]),
        Paragraph("<b>SEVERITY</b>", _styles["cell_label"]),
        Paragraph("<b>ANNUAL IMPACT</b>", _styles["cell_label"]),
    ]
    table_rows = [header]
    ordered = sorted(rows, key=lambda r: (_SEVERITY_ORDER.get(r.get("severity"), 99), r.get("unit") or ""))
    for row in ordered:
        severity = row.get("severity") or "low"
        color = _SEVERITY_COLORS.get(severity, _MUTED)
        table_rows.append([
            Paragraph(row.get("unit") or "—", _styles["cell_value"]),
            Paragraph(DISCREPANCY_TYPE_LABELS.get(row["discrepancy_type"], row["discrepancy_type"]), _styles["cell_value"]),
            Paragraph(str(row.get("rent_roll_value") or "—"), _styles["cell_value"]),
            Paragraph(str(row.get("lease_value") or "—"), _styles["cell_value"]),
            Paragraph(_fmt_source(row.get("source")), _styles["cell_value"]),
            Paragraph(f'<font color="{color.hexval()}"><b>{severity.upper()}</b></font>', _styles["cell_value"]),
            Paragraph(_fmt_money(row.get("annual_dollar_impact")), _styles["cell_value"]),
        ])

    table = Table(
        table_rows,
        colWidths=[1.3 * inch, 1.15 * inch, 0.95 * inch, 0.95 * inch, 0.55 * inch, 0.65 * inch, 0.85 * inch],
        repeatRows=1,
    )
    table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LINEBELOW", (0, 0), (-1, -1), 0.5, _RULE),
        ("LINEBELOW", (0, 0), (-1, 0), 1, _INK),
    ]))
    return table


def generate_deal_mismatch_report_pdf(data: Dict[str, Any]) -> bytes:
    scope = data.get("property_address") or "Whole Portfolio"
    generated_date = date.fromisoformat(data["generated_date"]) if data.get("generated_date") else date.today()
    flowables = _header_flowables(
        "Deal Mismatch Report",
        [f"Rent roll vs. lease documents &mdash; {scope}"],
        generated_date,
    )
    flowables.append(Paragraph("Summary", _styles["section"]))
    flowables.extend(_summary_flowables(data))

    flowables.append(Paragraph("Discrepancies", _styles["section"]))
    rows = data.get("discrepancies") or []
    if not rows:
        flowables.append(Paragraph("No discrepancies found.", _styles["muted"]))
    else:
        flowables.append(_discrepancies_table(rows))

    # T12 section
    t12_rows = data.get("rent_roll_vs_actual_collections")
    if t12_rows:
        flowables.append(Spacer(1, 12))
        flowables.append(Paragraph("Rent Roll vs. Actual Collections (T12 Cross-Check)", _styles["section"]))
        t12_overstatement = data.get("estimated_income_overstatement_from_t12")
        if t12_overstatement is not None:
            if t12_overstatement > 0:
                headline = f"Rent roll <b>overstates</b> actual collections by <b>{_fmt_money(t12_overstatement)}</b>."
            else:
                headline = f"Rent roll <b>understates</b> actual collections by <b>{_fmt_money(abs(t12_overstatement))}</b>."
        else:
            headline = "No dollar-quantified T12 findings."
        flowables.append(Paragraph(headline, _styles["body"]))
        flowables.append(Spacer(1, 8))
        flowables.append(_discrepancies_table(t12_rows))

    return _build_pdf(flowables)


_EXCEL_HEADER = [
    "Unit", "Discrepancy Type", "Field", "Rent Roll Value", "Lease Value",
    "Source Page", "Severity", "Monthly Impact", "Annual Impact", "Income Direction",
]
_EXCEL_COLUMN_WIDTHS = [26, 26, 16, 22, 22, 12, 10, 16, 16, 16]
_CURRENCY_FORMAT = "$#,##0.00"


def generate_deal_mismatch_report_excel(data: Dict[str, Any]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Deal Mismatch Report"

    scope = data.get("property_address") or "Whole Portfolio"
    header_font = Font(bold=True)
    title_font = Font(bold=True, size=14)

    sheet.cell(row=1, column=1, value="Deal Mismatch Report").font = title_font
    sheet.cell(row=2, column=1, value=f"Scope: {scope}")
    sheet.cell(row=3, column=1, value=f"Units checked: {data.get('total_units_checked', 0)}")
    sheet.cell(row=4, column=1, value=f"Discrepancies found: {data.get('total_discrepancies', 0)}")

    overstatement = data.get("annual_income_overstatement")
    if overstatement is None:
        summary_text = "Not enough matched, dollar-quantified discrepancies to estimate a net income impact."
    elif overstatement > 0:
        summary_text = f"Rent roll overstates annual income by {_fmt_money(overstatement)}."
    elif overstatement < 0:
        summary_text = f"Rent roll understates annual income by {_fmt_money(abs(overstatement))}."
    else:
        summary_text = "Rent roll and lease-documented income agree overall (net impact: $0)."
    sheet.cell(row=5, column=1, value=summary_text).font = Font(bold=True)

    header_row = 7
    for index, column in enumerate(_EXCEL_HEADER, start=1):
        cell = sheet.cell(row=header_row, column=index, value=column)
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.fill = PatternFill(start_color="F0EDE5", end_color="F0EDE5", fill_type="solid")
        sheet.column_dimensions[get_column_letter(index)].width = _EXCEL_COLUMN_WIDTHS[index - 1]

    rows = data.get("discrepancies") or []
    ordered = sorted(rows, key=lambda r: (_SEVERITY_ORDER.get(r.get("severity"), 99), r.get("unit") or ""))
    current_row = header_row + 1
    for row in ordered:
        source = row.get("source") or {}
        values = [
            row.get("unit"),
            DISCREPANCY_TYPE_LABELS.get(row["discrepancy_type"], row["discrepancy_type"]),
            row.get("field"),
            row.get("rent_roll_value"),
            row.get("lease_value"),
            source.get("page"),
            (row.get("severity") or "").upper(),
            row.get("monthly_dollar_impact"),
            row.get("annual_dollar_impact"),
            row.get("income_direction"),
        ]
        for col_offset, value in enumerate(values, start=1):
            cell = sheet.cell(row=current_row, column=col_offset, value=value)
            if col_offset in (8, 9) and value is not None:
                cell.number_format = _CURRENCY_FORMAT
        current_row += 1

    # T12 section
    t12_rows = data.get("rent_roll_vs_actual_collections")
    if t12_rows:
        # Add blank row for spacing
        current_row += 1

        # Add T12 header
        t12_header_row = current_row
        sheet.cell(row=t12_header_row, column=1, value="Rent Roll vs. Actual Collections (T12)").font = Font(bold=True)
        current_row += 1

        # T12 overstatement summary
        t12_overstatement = data.get("estimated_income_overstatement_from_t12")
        if t12_overstatement is not None:
            if t12_overstatement > 0:
                summary_text = f"Rent roll overstates actual collections by {_fmt_money(t12_overstatement)}."
            else:
                summary_text = f"Rent roll understates actual collections by {_fmt_money(abs(t12_overstatement))}."
        else:
            summary_text = "No dollar-quantified T12 findings."
        sheet.cell(row=current_row, column=1, value=summary_text)
        current_row += 1

        # T12 header row
        t12_header_start = current_row
        for index, column in enumerate(_EXCEL_HEADER, start=1):
            cell = sheet.cell(row=current_row, column=index, value=column)
            cell.font = header_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.fill = PatternFill(start_color="E8F4F8", end_color="E8F4F8", fill_type="solid")
        current_row += 1

        # T12 data rows
        t12_ordered = sorted(t12_rows, key=lambda r: (_SEVERITY_ORDER.get(r.get("severity"), 99), r.get("unit") or ""))
        for row in t12_ordered:
            source = row.get("source") or {}
            values = [
                row.get("unit"),
                DISCREPANCY_TYPE_LABELS.get(row["discrepancy_type"], row["discrepancy_type"]),
                row.get("field"),
                row.get("rent_roll_value"),
                row.get("lease_value"),
                source.get("page"),
                (row.get("severity") or "").upper(),
                row.get("monthly_dollar_impact"),
                row.get("annual_dollar_impact"),
                row.get("income_direction"),
            ]
            for col_offset, value in enumerate(values, start=1):
                cell = sheet.cell(row=current_row, column=col_offset, value=value)
                if col_offset in (8, 9) and value is not None:
                    cell.number_format = _CURRENCY_FORMAT
            current_row += 1

    sheet.freeze_panes = f"A{header_row + 1}"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
