"""
Tests for app/file_types.py: upload type is decided by content, not by the
filename. Every sample is real bytes produced by the library that writes
that format (reportlab, python-docx, openpyxl, Pillow, pillow-heif) or an
existing multi-format fixture (legacy .doc / .xls), never a hand-typed
header -- except for the deliberately-wrong files (executable, archives).
"""
import _test_env  # noqa: F401  -- first: temp DB, no .env, no real network (see _test_env.py)
import io
import os
import shutil
import sys
import tempfile
import zipfile

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from app.file_types import UnsupportedFileType, detect_file_type  # noqa: E402

TESTS_DIR = os.path.dirname(os.path.abspath(__file__))


def _write(tmpdir, name, data):
    path = os.path.join(tmpdir, name)
    with open(path, "wb") as f:
        f.write(data)
    return path


def _pdf_bytes():
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(72, 720, "Lease between Example Landlord LLC and Example Tenant Inc.")
    c.save()
    return buf.getvalue()


def _docx_bytes():
    from docx import Document
    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph("Lease between Example Landlord LLC and Example Tenant Inc.")
    doc.save(buf)
    return buf.getvalue()


def _xlsx_bytes():
    import openpyxl
    buf = io.BytesIO()
    wb = openpyxl.Workbook()
    wb.active.append(["Unit", "Tenant", "Rent"])
    wb.save(buf)
    return buf.getvalue()


def _image_bytes(fmt):
    from PIL import Image
    if fmt == "HEIF":
        import pillow_heif
        pillow_heif.register_heif_opener()
    buf = io.BytesIO()
    Image.new("RGB", (64, 64), "white").save(buf, fmt)
    return buf.getvalue()


def _pe_bytes():
    """The start of a real Windows executable: DOS header whose e_lfanew (bytes 60-63) points at "PE\\0\\0"."""
    header = bytearray(b"MZ" + b"\x90\x00" + bytes(60))
    header[60:64] = (64).to_bytes(4, "little")
    return bytes(header[:64]) + b"PE\x00\x00" + os.urandom(512)


def _zip_bytes(member_names):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for name in member_names:
            z.writestr(name, "x")
    return buf.getvalue()


def _detect(tmpdir, filename, data, kind="lease"):
    path = _write(tmpdir, "upload.bin", data)  # stored name never carries the user's extension
    return detect_file_type(path, filename, kind)


def _expect_rejection(tmpdir, filename, data, must_contain, kind="lease"):
    try:
        detected = _detect(tmpdir, filename, data, kind)
    except UnsupportedFileType as e:
        message = str(e)
        assert must_contain.lower() in message.lower(), f"{filename}: {message!r} lacks {must_contain!r}"
        assert "Please upload" in message or "upload" in message.lower(), message
        return message
    raise AssertionError(f"{filename} should have been rejected, detected {detected!r}")


def test_every_supported_format_is_recognized_from_its_bytes():
    tmpdir = tempfile.mkdtemp()
    try:
        with open(os.path.join(TESTS_DIR, "multiformat_lease.doc"), "rb") as f:
            doc = f.read()
        with open(os.path.join(TESTS_DIR, "multiformat_lease.xls"), "rb") as f:
            xls = f.read()
        cases = [
            ("lease.pdf", _pdf_bytes(), "pdf"),
            ("lease.docx", _docx_bytes(), "docx"),
            ("lease.doc", doc, "doc"),
            ("roll.xlsx", _xlsx_bytes(), "xlsx"),
            ("roll.xls", xls, "xls"),
            ("roll.csv", b"Unit,Tenant,Rent\n101,Acme,1200\n", "csv"),
            ("roll.tsv", b"Unit\tTenant\tRent\n101\tAcme\t1200\n", "tsv"),
            ("lease.txt", b"This Lease is made between Example Landlord LLC and Example Tenant Inc.\n", "txt"),
            ("scan.jpg", _image_bytes("JPEG"), "jpg"),
            ("scan.jpeg", _image_bytes("JPEG"), "jpg"),
            ("scan.png", _image_bytes("PNG"), "png"),
            ("scan.tif", _image_bytes("TIFF"), "tiff"),
            ("photo.heic", _image_bytes("HEIF"), "heic"),
        ]
        for filename, data, expected in cases:
            detected = _detect(tmpdir, filename, data)
            assert detected.extension == expected, f"{filename}: got {detected.extension}, expected {expected}"
            assert not detected.mismatched_name, f"{filename} wrongly flagged as mislabeled"
            assert detected.canonical_filename(filename) == filename
    finally:
        shutil.rmtree(tmpdir)
    print("✓ test_every_supported_format_is_recognized_from_its_bytes: PASS")


def test_mislabeled_files_are_processed_as_what_they_really_are():
    """The repro's scenario F: a PNG named .pdf used to fail as 'corrupted PDF'; a PDF named .txt was imported as garbage text."""
    tmpdir = tempfile.mkdtemp()
    try:
        png_as_pdf = _detect(tmpdir, "lease.pdf", _image_bytes("PNG"))
        assert png_as_pdf.extension == "png" and png_as_pdf.mismatched_name
        assert png_as_pdf.canonical_filename("lease.pdf") == "lease.png"

        pdf_as_txt = _detect(tmpdir, "lease.txt", _pdf_bytes())
        assert pdf_as_txt.extension == "pdf" and pdf_as_txt.canonical_filename("lease.txt") == "lease.pdf"

        docx_without_extension = _detect(tmpdir, "Lease Final", _docx_bytes())
        assert docx_without_extension.extension == "docx"
        assert docx_without_extension.canonical_filename("Lease Final") == "Lease Final.docx"

        # Text flavors are told apart by name or, failing that, by shape.
        assert _detect(tmpdir, "roll.txt", b"a,b,c\n1,2,3\n").extension == "txt"
        assert _detect(tmpdir, "export", b"a,b,c\n1,2,3\n4,5,6\n").extension == "csv"
        assert _detect(tmpdir, "export", b"a\tb\tc\n1\t2\t3\n").extension == "tsv"
    finally:
        shutil.rmtree(tmpdir)
    print("✓ test_mislabeled_files_are_processed_as_what_they_really_are: PASS")


def test_unsupported_and_dangerous_content_is_rejected_with_a_friendly_message():
    tmpdir = tempfile.mkdtemp()
    try:
        _expect_rejection(tmpdir, "tool.pdf", _pe_bytes(), "a program")
        _expect_rejection(tmpdir, "deck.pptx", _zip_bytes(["[Content_Types].xml", "ppt/presentation.xml"]), "PowerPoint")
        _expect_rejection(tmpdir, "docs.zip", _zip_bytes(["a.pdf", "b.pdf"]), "ZIP archive")
        _expect_rejection(tmpdir, "lease.odt", _zip_bytes(["mimetype", "content.xml"]), "OpenDocument")
        _expect_rejection(tmpdir, "scan.gif", _image_bytes("GIF"), "GIF")
        _expect_rejection(tmpdir, "lease.rtf", b"{\\rtf1\\ansi Lease}", "Rich Text")
        _expect_rejection(tmpdir, "lease.pdf", b"<!DOCTYPE html><html><body>Lease</body></html>", "web page")
        _expect_rejection(tmpdir, "backup.tar.gz", b"\x1f\x8b\x08" + os.urandom(512), "compressed archive")
        _expect_rejection(tmpdir, "mystery.dat", bytes(range(256)) * 8, "couldn't recognize")
        # Recognized but not an accepted lease format: names what it is.
        _expect_rejection(tmpdir, "scan.webp", _image_bytes("WEBP"), "WebP")
        # The same WebP is fine as a rent roll (that importer already read WebP).
        assert _detect(tmpdir, "roll.webp", _image_bytes("WEBP"), kind="rent_roll").extension == "webp"
        # .doc is a lease format but the rent-roll importer has no .doc reader.
        with open(os.path.join(TESTS_DIR, "multiformat_lease.doc"), "rb") as f:
            _expect_rejection(tmpdir, "roll.doc", f.read(), "Save it as .docx", kind="rent_roll")
    finally:
        shutil.rmtree(tmpdir)
    print("✓ test_unsupported_and_dangerous_content_is_rejected_with_a_friendly_message: PASS")


def test_empty_and_damaged_files_get_specific_messages():
    tmpdir = tempfile.mkdtemp()
    try:
        _expect_rejection(tmpdir, "lease.pdf", b"", "empty")
        _expect_rejection(tmpdir, "lease.docx", b"PK\x03\x04" + os.urandom(64), "damaged")
        _expect_rejection(tmpdir, "lease.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + os.urandom(64), "damaged")
    finally:
        shutil.rmtree(tmpdir)
    print("✓ test_empty_and_damaged_files_get_specific_messages: PASS")


def test_password_protected_office_file_is_named_as_such():
    """
    An encrypted .docx/.xlsx is an OLE2 container holding EncryptionInfo +
    EncryptedPackage streams, not a ZIP. No encrypted fixture exists in the
    repo (making one needs msoffcrypto-tool, not a dependency), so olefile's
    directory listing is stubbed to what such a file contains.
    """
    from unittest import mock
    tmpdir = tempfile.mkdtemp()
    try:
        fake = mock.MagicMock()
        fake.__enter__.return_value.listdir.return_value = [["EncryptionInfo"], ["EncryptedPackage"]]
        with mock.patch("app.file_types.olefile.OleFileIO", return_value=fake):
            _expect_rejection(tmpdir, "rent roll.xlsx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(504),
                              "password-protected")
    finally:
        shutil.rmtree(tmpdir)
    print("✓ test_password_protected_office_file_is_named_as_such: PASS")


def test_magic_checks_dont_misread_text_that_happens_to_start_like_a_binary():
    """Review finding: a CSV starting "BM..." was detected as a BMP image; text starting "MZ" as a program; any text mentioning %PDF- as a PDF."""
    tmpdir = tempfile.mkdtemp()
    try:
        assert _detect(tmpdir, "roll.csv", b"BMR,Unit,Rent\nA,101,1200\n", kind="rent_roll").extension == "csv"
        assert _detect(tmpdir, "notes.txt", b"MZ Holdings LLC lease summary for the 2026 renewal.\n" * 3).extension == "txt"
        assert _detect(tmpdir, "notes.txt", b"Scanner settings: save as %PDF-1.7 compatible.\nThanks!\n").extension == "txt"
        # The real formats are still recognized.
        assert _detect(tmpdir, "scan.bmp", _image_bytes("BMP"), kind="rent_roll").extension == "bmp"
        assert _detect(tmpdir, "lease.pdf", b"\n\n" + _pdf_bytes()).extension == "pdf", "leading junk before %PDF- is tolerated"
    finally:
        shutil.rmtree(tmpdir)
    print("✓ test_magic_checks_dont_misread_text_that_happens_to_start_like_a_binary: PASS")


def test_zip_bombs_are_refused_from_the_directory_alone():
    """Security review: a small .xlsx/.docx can inflate to gigabytes in openpyxl/python-docx. Refused before any parser runs."""
    from app import file_types
    tmpdir = tempfile.mkdtemp()
    try:
        bomb = os.path.join(tmpdir, "bomb.xlsx")
        with zipfile.ZipFile(bomb, "w", compression=zipfile.ZIP_DEFLATED) as z:
            z.writestr("[Content_Types].xml", "<Types/>")
            with z.open("xl/workbook.xml", "w") as member:
                chunk = b"\x00" * (1024 * 1024)
                for _ in range(file_types.MAX_ZIP_UNCOMPRESSED_BYTES // len(chunk) + 1):
                    member.write(chunk)
        assert os.path.getsize(bomb) < 2 * 1024 * 1024, "the bomb itself is small"
        try:
            detect_file_type(bomb, "bomb.xlsx", "rent_roll")
        except UnsupportedFileType as e:
            assert "expands to far more data" in str(e), e
        else:
            raise AssertionError("zip bomb accepted")

        many = os.path.join(tmpdir, "many.docx")
        with zipfile.ZipFile(many, "w") as z:
            z.writestr("word/document.xml", "<w:document/>")
            for i in range(file_types.MAX_ZIP_ENTRIES):
                z.writestr(f"x/{i}", "")
        try:
            detect_file_type(many, "many.docx", "lease")
        except UnsupportedFileType as e:
            assert "expands to far more data" in str(e), e
        else:
            raise AssertionError("20,000-entry zip accepted")
    finally:
        shutil.rmtree(tmpdir)
    print("✓ test_zip_bombs_are_refused_from_the_directory_alone: PASS")


if __name__ == "__main__":
    test_every_supported_format_is_recognized_from_its_bytes()
    test_mislabeled_files_are_processed_as_what_they_really_are()
    test_unsupported_and_dangerous_content_is_rejected_with_a_friendly_message()
    test_empty_and_damaged_files_get_specific_messages()
    test_password_protected_office_file_is_named_as_such()
    test_magic_checks_dont_misread_text_that_happens_to_start_like_a_binary()
    test_zip_bombs_are_refused_from_the_directory_alone()
    print("\nAll file-type tests passed!")
