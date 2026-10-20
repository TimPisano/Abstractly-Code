"""
Generates the large-file fixtures under tests/fixtures/large/ (gitignored --
together they are tens of MB, and they are fully reproducible from this
script, so they are never committed):

  large_text_lease_500p.pdf   500-page single lease with a real text layer.
                              Key terms are deliberately placed late in the
                              document (security deposit p.262, square
                              footage p.333, renewal p.431, exclusivity
                              p.478) so anything that reads only the first
                              part of a long lease visibly misses them.
  scanned_lease_300p.pdf      300-page image-only PDF (no text layer at all),
                              200 DPI grayscale JPEG pages -- the profile of
                              an office-scanner PDF.
  messy_rent_roll_5000.xlsx   5,000 unit rows across 50 buildings: a cover
                              "Summary" sheet first, the real "Rent Roll"
                              sheet second (title rows, a two-row header
                              with merged group titles, a heading row per
                              building, a subtotal row per building, a
                              blank row between buildings, a grand total),
                              and a trailing "Notes" sheet. Mixed cell
                              types: rents as numbers and as "$1,250.00"
                              strings, dates as real dates and as text.
  messy_rent_roll_5000.csv    The same "Rent Roll" sheet flattened to CSV.
  expected.json               The answer key the tests check against
                              (occupied/vacant counts, page numbers of the
                              planted lease terms, ...).

Every name, address and amount is fictional (CLAUDE.md: no PII in
fixtures).

Usage:
    python backend/tests/create_large_fixtures.py          # build what's missing
    python backend/tests/create_large_fixtures.py --force  # rebuild everything

Tests call ensure_large_fixtures(), which builds only missing files, so a
fresh worktree needs no manual step (the scanned PDF takes about a minute
the first time).
"""
import csv
import datetime as dt
import io
import json
import os
import random
import sys

FIXTURES_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures", "large")

TEXT_LEASE = "large_text_lease_500p.pdf"
SCANNED_LEASE = "scanned_lease_300p.pdf"
RENT_ROLL_XLSX = "messy_rent_roll_5000.xlsx"
RENT_ROLL_CSV = "messy_rent_roll_5000.csv"
EXPECTED = "expected.json"

TEXT_LEASE_PAGES = 500
SCANNED_LEASE_PAGES = 300
RENT_ROLL_BUILDINGS = 50
RENT_ROLL_UNITS_PER_BUILDING = 100

LANDLORD = "Harborline Commons LLC"
TENANT = "Northwind Outfitters, Inc."
PREMISES = "4100 Cedar Bluff Parkway, Suite 300, Denver, Colorado 80202"
MONTHLY_RENT = "$18,750.00"

# Planted terms: page -> sentence. Pages are 1-based, as a reader sees them.
PLANTED_TERMS = {
    262: ("security_deposit",
          "SECURITY DEPOSIT. Upon execution of this Lease, Tenant shall deposit with Landlord "
          "a security deposit in the amount of $56,250.00, to be held as security for the "
          "faithful performance of Tenant's obligations."),
    333: ("square_footage",
          "AREA OF THE PREMISES. The Premises contain approximately 12,500 rentable square feet."),
    431: ("renewal_options",
          "OPTION TO RENEW. Tenant shall have two (2) options to renew the Term for five (5) "
          "years each, exercisable by written notice given not less than nine (9) months before "
          "the expiration of the then-current Term, at ninety-five percent (95%) of fair market rent."),
    478: ("exclusivity_clause",
          "EXCLUSIVE USE. Landlord shall not lease any other space in the Shopping Center for the "
          "primary purpose of selling outdoor apparel and camping equipment."),
}

# Neutral filler: no dollar amounts, no dates, no party labels, no
# "LEASE AGREEMENT" heading -- so it can never be mistaken for a lease term
# or a new-lease boundary by the regex extractor.
_FILLER_SENTENCES = [
    "Each party shall perform its obligations under this Article in a good and workmanlike manner.",
    "The provisions of this Section shall be construed in harmony with the remaining provisions hereof.",
    "Any approval required under this Article shall not be unreasonably withheld, conditioned or delayed.",
    "The headings in this Article are for convenience only and do not define or limit its scope.",
    "Notices under this Section shall be given in the manner described in the Article titled Notices.",
    "The rights described in this Section are cumulative and in addition to any other rights at law.",
    "Nothing in this Section shall be deemed a waiver of any other provision of this instrument.",
    "The parties acknowledge that this Section was negotiated at arm's length by sophisticated parties.",
    "Where this Section refers to reasonable efforts, commercially reasonable efforts are intended.",
    "Compliance with applicable codes and ordinances is a continuing obligation under this Section.",
    "Any dispute concerning the interpretation of this Section shall be resolved under the governing law clause.",
    "The obligations in this Section survive only to the extent expressly stated elsewhere in this instrument.",
]


def _filler_paragraph(page_number, paragraph_index, sentences=7):
    rng = random.Random(page_number * 1000 + paragraph_index)
    picked = [rng.choice(_FILLER_SENTENCES) for _ in range(sentences)]
    return f"Section {page_number}.{paragraph_index + 1}. " + " ".join(picked)


def _text_lease_page_paragraphs(page_number):
    if page_number == 1:
        return [
            "COMMERCIAL LEASE",
            f"This Commercial Lease is entered into by and between {LANDLORD} (\"Landlord\") "
            f"and {TENANT} (\"Tenant\").",
            f"PREMISES. Landlord leases to Tenant the premises located at {PREMISES} (the \"Premises\").",
            "TERM. The Term of this Lease shall commence on March 1, 2026 and shall expire on "
            "February 28, 2036, unless sooner terminated as provided herein.",
            f"BASE RENT. Tenant shall pay Base Rent of {MONTHLY_RENT} per month, payable in advance "
            "on the first day of each calendar month.",
            _filler_paragraph(1, 0, sentences=4),
        ]
    if page_number == TEXT_LEASE_PAGES:
        return [
            "IN WITNESS WHEREOF, the parties have executed this instrument as of the date first written above.",
            f"LANDLORD: {LANDLORD}   By: ____________________   Name: Authorized Signatory",
            f"TENANT: {TENANT}   By: ____________________   Name: Authorized Signatory",
        ]
    paragraphs = [f"ARTICLE {page_number}. GENERAL PROVISIONS (CONTINUED)"]
    paragraphs += [_filler_paragraph(page_number, i) for i in range(4)]
    if page_number in PLANTED_TERMS:
        paragraphs.insert(2, f"Section {page_number}.0. " + PLANTED_TERMS[page_number][1])
    return paragraphs


def build_text_lease(path):
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate

    styles = getSampleStyleSheet()
    story = []
    for page_number in range(1, TEXT_LEASE_PAGES + 1):
        for text in _text_lease_page_paragraphs(page_number):
            story.append(Paragraph(text, styles["BodyText"]))
        if page_number < TEXT_LEASE_PAGES:
            story.append(PageBreak())
    SimpleDocTemplate(path, pagesize=letter, title="Fixture lease (fictional)").build(story)


def _wrap(text, width_chars):
    words, lines, line = text.split(), [], ""
    for word in words:
        if len(line) + len(word) + 1 > width_chars:
            lines.append(line)
            line = word
        else:
            line = f"{line} {word}".strip()
    if line:
        lines.append(line)
    return lines


def build_scanned_lease(path):
    """Image-only pages: each page is a picture of text drawn with Pillow, embedded as a JPEG. pypdf extracts 0 characters from every page."""
    from PIL import Image, ImageDraw, ImageFont
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    dpi = 200
    width_px, height_px = int(8.5 * dpi), int(11 * dpi)
    try:
        font = ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", 30)
    except Exception:
        font = ImageFont.load_default()

    pdf = canvas.Canvas(path, pagesize=letter)
    for page_number in range(1, SCANNED_LEASE_PAGES + 1):
        image = Image.new("L", (width_px, height_px), color=255)
        draw = ImageDraw.Draw(image)
        y = 150
        for paragraph in _text_lease_page_paragraphs(page_number)[:4]:
            for line in _wrap(paragraph, 80):
                draw.text((150, y), line, fill=0, font=font)
                y += 44
            y += 30
            if y > height_px - 200:
                break
        draw.text((width_px // 2 - 40, height_px - 120), f"- {page_number} -", fill=0, font=font)
        buf = io.BytesIO()
        image.save(buf, "JPEG", quality=70)
        buf.seek(0)
        pdf.drawImage(ImageReader(buf), 0, 0, width=letter[0], height=letter[1])
        pdf.showPage()
    pdf.save()


_TENANT_PREFIXES = ["Cedar", "Summit", "Bluebird", "Iron Gate", "Harbor", "Juniper", "Granite", "Lantern",
                    "Willow", "Copper", "Meadow", "Falcon", "Prairie", "Sterling", "Aspen", "Beacon"]
_TENANT_SUFFIXES = ["Labs LLC", "Dental Group", "Coffee Co.", "Fitness LLC", "Design Studio", "Logistics Inc.",
                    "Bakery", "Insurance Agency", "Yoga Collective", "Tax Services", "Pet Supply", "Books"]


def _rent_roll_rows():
    """Yields (row_kind, values) for the Rent Roll sheet, in sheet order. row_kind is one of title/header/building/unit/vacant/subtotal/blank/total."""
    rng = random.Random(5000)
    yield "title", ["Maple Ridge Portfolio (fictional) - Rent Roll"]
    yield "title", ["As of 10/01/2026", "", "", "", "", "", "", "Prepared by: Asset Management"]
    yield "blank", []
    yield "header", ["Unit", "Tenant", "Sq Ft", "Lease Dates", None, "Charges", None, "Status"]
    yield "header", [None, None, None, "Start", "End", "Monthly Rent", "CAM", None]
    grand_rent = 0.0
    for b in range(1, RENT_ROLL_BUILDINGS + 1):
        street = f"{1000 + b * 10} Juniper Row, Building {b}"
        yield "building", [f"Building {b} - {street}"]
        building_rent = 0.0
        for u in range(1, RENT_ROLL_UNITS_PER_BUILDING + 1):
            unit = f"{b}-{u:03d}"
            sqft = rng.choice([650, 800, 950, 1200, 1500, 2200])
            if rng.random() < 0.04:
                yield "vacant", [unit, "VACANT", sqft, None, None, None, None, "Vacant"]
                continue
            start = dt.date(2022, 1, 1) + dt.timedelta(days=rng.randrange(0, 1400))
            end = start + dt.timedelta(days=365 * rng.choice([1, 2, 3, 5]) - 1)
            rent = round(sqft * rng.uniform(1.4, 2.6), 2)
            cam = round(sqft * 0.25, 2)
            building_rent += rent
            name = f"{rng.choice(_TENANT_PREFIXES)} {rng.choice(_TENANT_SUFFIXES)} #{b}{u:03d}"
            # Messy cell types, as hand-maintained broker sheets really are.
            rent_cell = rent if u % 3 else f"${rent:,.2f}"
            start_cell = start if u % 4 else start.strftime("%m/%d/%Y")
            yield "unit", [unit, name, sqft, start_cell, end, rent_cell, cam, "Occupied"]
        grand_rent += building_rent
        yield "subtotal", [f"Subtotal Building {b}", None, None, None, None, round(building_rent, 2), None, None]
        yield "blank", []
    yield "total", ["Grand Total", None, None, None, None, round(grand_rent, 2), None, None]


def build_rent_rolls(xlsx_path, csv_path):
    import openpyxl

    rows = list(_rent_roll_rows())
    workbook = openpyxl.Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Maple Ridge Portfolio (fictional)"])
    summary.append(["Buildings", RENT_ROLL_BUILDINGS])
    summary.append(["Units", RENT_ROLL_BUILDINGS * RENT_ROLL_UNITS_PER_BUILDING])
    summary.append(["See the Rent Roll sheet for unit detail."])

    sheet = workbook.create_sheet("Rent Roll")
    header_rows = []
    for kind, values in rows:
        sheet.append(values)
        if kind == "header":
            header_rows.append(sheet.max_row)
    # Merged group titles over the two-row header, as broker templates do.
    top = header_rows[0]
    sheet.merge_cells(start_row=top, start_column=4, end_row=top, end_column=5)
    sheet.merge_cells(start_row=top, start_column=6, end_row=top, end_column=7)
    for col in (1, 2, 3, 8):
        sheet.merge_cells(start_row=top, start_column=col, end_row=top + 1, end_column=col)

    notes = workbook.create_sheet("Notes")
    notes.append(["Subtotals are per building. Vacant units have no rent."])
    workbook.save(xlsx_path)

    with open(csv_path, "w", newline="") as f:
        writer = csv.writer(f)
        for _, values in rows:
            writer.writerow(["" if v is None else (v.strftime("%m/%d/%Y") if isinstance(v, dt.date) else v) for v in values])

    counts = {}
    for kind, _ in rows:
        counts[kind] = counts.get(kind, 0) + 1
    return counts


def _expected(rent_roll_counts):
    return {
        "text_lease": {
            "pages": TEXT_LEASE_PAGES,
            "tenant": TENANT,
            "landlord": LANDLORD,
            "rent_amount": MONTHLY_RENT,
            "planted_terms": {name: page for page, (name, _) in PLANTED_TERMS.items()},
        },
        "scanned_lease": {"pages": SCANNED_LEASE_PAGES, "tenant": TENANT},
        "rent_roll": {
            "unit_rows": rent_roll_counts.get("unit", 0) + rent_roll_counts.get("vacant", 0),
            "occupied_rows": rent_roll_counts.get("unit", 0),
            "vacant_rows": rent_roll_counts.get("vacant", 0),
            "subtotal_rows": rent_roll_counts.get("subtotal", 0),
            "blank_rows": rent_roll_counts.get("blank", 0),
            "buildings": RENT_ROLL_BUILDINGS,
        },
    }


def path_of(name):
    return os.path.join(FIXTURES_DIR, name)


def ensure_large_fixtures(force=False, only=None):
    """Builds any missing fixture (or all, with force). `only` limits it to a set of filenames. Returns FIXTURES_DIR."""
    os.makedirs(FIXTURES_DIR, exist_ok=True)
    wanted = set(only or [TEXT_LEASE, SCANNED_LEASE, RENT_ROLL_XLSX, RENT_ROLL_CSV, EXPECTED])

    def needed(name):
        return name in wanted and (force or not os.path.exists(path_of(name)))

    if needed(TEXT_LEASE):
        build_text_lease(path_of(TEXT_LEASE))
    if needed(SCANNED_LEASE):
        build_scanned_lease(path_of(SCANNED_LEASE))
    if needed(RENT_ROLL_XLSX) or needed(RENT_ROLL_CSV) or needed(EXPECTED):
        counts = build_rent_rolls(path_of(RENT_ROLL_XLSX), path_of(RENT_ROLL_CSV))
        with open(path_of(EXPECTED), "w") as f:
            json.dump(_expected(counts), f, indent=2)
    return FIXTURES_DIR


def load_expected():
    ensure_large_fixtures(only=[EXPECTED, RENT_ROLL_XLSX, RENT_ROLL_CSV])
    with open(path_of(EXPECTED)) as f:
        return json.load(f)


if __name__ == "__main__":
    ensure_large_fixtures(force="--force" in sys.argv)
    for name in (TEXT_LEASE, SCANNED_LEASE, RENT_ROLL_XLSX, RENT_ROLL_CSV, EXPECTED):
        print(f"{name}: {os.path.getsize(path_of(name)) / (1024 * 1024):.1f} MB")
