"""
Tests for the investment memo export: build_investment_memo_data (the
shared data structure), generate_investment_memo_pdf/_excel, and the
POST /portfolio/investment-memo.pdf / .xlsx routes.

Uses Flask's in-process test_client() against an isolated temp SQLite
file, same pattern as test_discrepancies.py.
"""

import io
import os
import sys
import tempfile

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

import PyPDF2
from openpyxl import load_workbook

from app.api import app
from app import database
from app.portfolio import FIELD_NAMES, compute_rent_roll_reconciliation
from app.discrepancies import sync_rent_roll_reconciliation
from app.investment_memo import build_investment_memo_data, generate_investment_memo_pdf, generate_investment_memo_excel


def _fresh_temp_db():
    tmp = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
    tmp.close()
    database.configure(tmp.name)
    database.init_db()
    return tmp.name


def _authed_client():
    """
    A test_client() pre-authenticated as a logged-in analyst, via
    Flask's session_transaction() -- the standard way to test a
    session-gated route without driving an actual login POST through
    bcrypt for every test, same convention as test_admin_auth.py's
    _create_admin()/session pattern. Most routes now require at least
    a logged-in session (see app/auth.py's require_role()) since the
    RBAC audit -- analyst covers every route these tests exercise.
    """
    client = app.test_client()
    with client.session_transaction() as sess:
        sess["user_id"] = 1
        sess["email"] = "test-analyst@example.com"
        sess["name"] = "Test Analyst"
        sess["role"] = "analyst"
    return client


def _fields(**overrides):
    result = {}
    for name in FIELD_NAMES:
        if name in overrides:
            value = overrides[name]
            result[name] = {"value": value, "source": {"page": 1, "quote": f"...{value}..."}, "confidence": "high"}
        else:
            result[name] = {"value": None, "source": None, "confidence": None}
    return result


def _insert(filename="lease.pdf", **overrides):
    return database.insert_lease(filename, _fields(**overrides), display_name=overrides.get("tenant") or filename)


def _pdf_text(pdf_bytes):
    reader = PyPDF2.PdfReader(io.BytesIO(pdf_bytes))
    return "\n".join(page.extract_text() for page in reader.pages)


# ------------------------------------------------------------------
# build_investment_memo_data
# ------------------------------------------------------------------

def test_empty_portfolio_returns_honest_zero_state():
    db_path = _fresh_temp_db()
    try:
        data = build_investment_memo_data()
        assert data["scope"] == "portfolio"
        assert data["lease_count"] == 0
        assert data["leases"] == []
        assert data["discrepancy_summary"] == {"total": 0, "open": 0, "resolved": 0}
        assert data["t12_cross_check"]["available"] is False
    finally:
        os.unlink(db_path)
    print("✓ test_empty_portfolio_returns_honest_zero_state: PASS")


def test_property_scope_filters_by_building_ignoring_suite():
    db_path = _fresh_temp_db()
    try:
        _insert(tenant="Acme", rent_amount="$4,000.00", property_address="100 Elm St, Suite 200, Springfield, IL")
        _insert(tenant="Beta", rent_amount="$5,000.00", property_address="100 Elm St, Suite 300, Springfield, IL")
        _insert(tenant="Gamma", rent_amount="$6,000.00", property_address="999 Oak Ave, Springfield, IL")

        data = build_investment_memo_data(property_address="100 Elm St, Springfield, IL")
        assert data["scope"] == "property"
        assert data["lease_count"] == 2
        assert {l["extracted_fields"]["tenant"]["value"] for l in data["leases"]} == {"Acme", "Beta"}
    finally:
        os.unlink(db_path)
    print("✓ test_property_scope_filters_by_building_ignoring_suite: PASS")


def test_dedupes_rent_roll_duplicate_of_a_pdf_lease_for_financial_totals():
    """The real bug found during manual review: a rent-roll snapshot of a unit that also has a real PDF lease must not double-count that unit's rent in totals/WALT/rollover -- but must still appear in the full leases list for display."""
    db_path = _fresh_temp_db()
    try:
        _insert("lease.pdf", tenant="Acme", rent_amount="$6,250.00", property_address="1 Main St, Suite 100", square_footage="1,000 sq ft", lease_end_date="March 31, 2030")
        _insert("rentroll.csv", tenant="Acme", rent_amount="$6,000.00", property_address="1 Main St, Suite 100", lease_end_date="March 31, 2030")

        data = build_investment_memo_data(property_address="1 Main St")
        assert data["record_count"] == 2, "both records must be visible"
        assert data["lease_count"] == 1, "but only ONE counts toward financial totals"
        assert data["portfolio_metrics"]["total_monthly_rent"] == 6250.0, "the PDF lease, not the rent-roll snapshot, must win"
        assert data["rollover"]["rollover_schedule"]["total_rent"] == 6250.0
    finally:
        os.unlink(db_path)
    print("✓ test_dedupes_rent_roll_duplicate_of_a_pdf_lease_for_financial_totals: PASS")


def test_dedupe_prefers_pdf_lease_regardless_of_upload_order():
    db_path = _fresh_temp_db()
    try:
        _insert("rentroll.csv", tenant="Acme", rent_amount="$6,000.00", property_address="1 Main St")
        _insert("lease.pdf", tenant="Acme", rent_amount="$6,250.00", property_address="1 Main St")  # uploaded AFTER the rent roll row
        data = build_investment_memo_data(property_address="1 Main St")
        assert data["portfolio_metrics"]["total_monthly_rent"] == 6250.0, "PDF lease must win even if uploaded later"
    finally:
        os.unlink(db_path)
    print("✓ test_dedupe_prefers_pdf_lease_regardless_of_upload_order: PASS")


def test_dedupe_keeps_most_recent_when_both_are_the_same_source_type():
    db_path = _fresh_temp_db()
    try:
        _insert("rentroll_q1.csv", tenant="Acme", rent_amount="$5,000.00", property_address="1 Main St")
        _insert("rentroll_q2.csv", tenant="Acme", rent_amount="$5,500.00", property_address="1 Main St")
        data = build_investment_memo_data(property_address="1 Main St")
        assert data["portfolio_metrics"]["total_monthly_rent"] == 5500.0, "the more recent rent-roll snapshot must win when there's no PDF lease at all"
    finally:
        os.unlink(db_path)
    print("✓ test_dedupe_keeps_most_recent_when_both_are_the_same_source_type: PASS")


def test_discrepancies_and_resolutions_scoped_to_property():
    db_path = _fresh_temp_db()
    try:
        _insert("lease.pdf", tenant="Acme", rent_amount="$6,250.00", property_address="1 Main St, Suite 100", lease_end_date="March 31, 2030")
        _insert("rentroll.csv", tenant="Acme", rent_amount="$6,000.00", property_address="1 Main St, Suite 100", lease_end_date="March 31, 2030")
        _insert("other.pdf", tenant="Other Co", rent_amount="$3,000.00", property_address="999 Elsewhere Ave")  # unrelated property

        leases = database.get_all_effective_leases()
        recon = compute_rent_roll_reconciliation(leases)
        mismatches = sync_rent_roll_reconciliation(recon["mismatches"])
        assert mismatches, "fixture must produce a real mismatch"
        disc_id = mismatches[0]["discrepancy_id"]
        database.resolve_discrepancy(disc_id, "lease_document", "Confirmed via signed PDF.", "Jane Analyst")

        data = build_investment_memo_data(property_address="1 Main St")
        assert data["discrepancy_summary"] == {"total": 1, "open": 0, "resolved": 1}
        assert data["discrepancies"][0]["resolutions"][0]["resolved_by"] == "Jane Analyst"

        portfolio_data = build_investment_memo_data()
        assert portfolio_data["discrepancy_summary"]["total"] == 1, "still exactly 1 -- the unrelated property has no discrepancies of its own"
    finally:
        os.unlink(db_path)
    print("✓ test_discrepancies_and_resolutions_scoped_to_property: PASS")


def test_portfolio_wide_memo_excludes_discrepancies_for_deleted_leases():
    """
    Regression test for a real, confirmed bug: a portfolio-wide memo
    used to include EVERY discrepancy row ever recorded
    (database.list_discrepancies() with no filtering at all), including
    ones whose lease had since been deleted and is no longer part of
    the portfolio. Found live: after stress-testing with thousands of
    leases that were later deleted, a 3-lease portfolio's memo reported
    "19,230 flagged" -- directly contradicting its own "3 leases
    covered" header. Fixed in _relevant_discrepancies to only include a
    lease-scoped discrepancy if its lease_id/related_lease_id still
    points to a lease CURRENTLY in the portfolio.
    """
    db_path = _fresh_temp_db()
    try:
        current_id = _insert("current.pdf", tenant="Current Co", rent_amount="$5,000.00", property_address="1 Main St")
        doomed_id = _insert("doomed.pdf", tenant="Doomed Co", rent_amount="$3,000.00", property_address="2 Elm St")

        database.upsert_discrepancy(
            discrepancy_type="lease_risk_flag", natural_key="lease_risk:current:missing_clause:cam_charges:0",
            category="missing_clause", message="No CAM charges clause found in this lease",
            details={}, lease_id=current_id,
        )
        database.upsert_discrepancy(
            discrepancy_type="lease_risk_flag", natural_key="lease_risk:doomed:missing_clause:cam_charges:0",
            category="missing_clause", message="No CAM charges clause found in this lease",
            details={}, lease_id=doomed_id,
        )

        # Before deletion: both discrepancies are legitimately relevant.
        data = build_investment_memo_data()
        assert data["discrepancy_summary"]["total"] == 2, "both leases still exist -- both discrepancies should count"

        database.delete_lease(doomed_id)

        # After deletion: the orphaned discrepancy must disappear from
        # a portfolio-wide memo, even though the row itself is still
        # in the database (discrepancies are deliberately permanent
        # records -- see database.py's un-FK'd lease_id).
        data = build_investment_memo_data()
        assert data["discrepancy_summary"]["total"] == 1, f"expected only the current lease's discrepancy, got {data['discrepancy_summary']}"
        assert data["discrepancies"][0]["lease_id"] == current_id
        assert data["lease_count"] == 1, "the deleted lease itself must also be gone from the memo's lease list"
    finally:
        os.unlink(db_path)
    print("✓ test_portfolio_wide_memo_excludes_discrepancies_for_deleted_leases: PASS")


def test_portfolio_wide_memo_excludes_tenant_concentration_for_a_tenant_no_longer_in_the_portfolio():
    """Same bug class as the lease-scoped case above, for the no-lease-id discrepancy types (tenant_concentration/t12_reconciliation) -- these must also be excluded once the tenant/address they're about no longer appears anywhere in the current portfolio."""
    db_path = _fresh_temp_db()
    try:
        _insert("current.pdf", tenant="Current Co", rent_amount="$5,000.00", property_address="1 Main St")

        database.upsert_discrepancy(
            discrepancy_type="tenant_concentration", natural_key="tenant_concentration:stale co",
            category="tenant_concentration", message="Stale Co accounts for a large share of portfolio rent",
            details={"tenant": "Stale Co"},
        )
        database.upsert_discrepancy(
            discrepancy_type="tenant_concentration", natural_key="tenant_concentration:current co",
            category="tenant_concentration", message="Current Co accounts for a large share of portfolio rent",
            details={"tenant": "Current Co"},
        )

        data = build_investment_memo_data()
        assert data["discrepancy_summary"]["total"] == 1, f"expected only the Current Co discrepancy, got {data['discrepancy_summary']}"
        assert data["discrepancies"][0]["details"]["tenant"] == "Current Co"
    finally:
        os.unlink(db_path)
    print("✓ test_portfolio_wide_memo_excludes_tenant_concentration_for_a_tenant_no_longer_in_the_portfolio: PASS")


def test_t12_fresh_upload_takes_precedence_over_persisted():
    db_path = _fresh_temp_db()
    try:
        _insert("lease.pdf", tenant="Acme", rent_amount="$5,000.00", property_address="1 Main St")
        t12_parsed = {"annual_rental_income": 55000.0, "source": {"page": None, "quote": "Total Rental Income"}}
        data = build_investment_memo_data(property_address="1 Main St", t12_parsed=t12_parsed)
        assert data["t12_cross_check"]["available"] is True
        assert data["t12_cross_check"]["basis"] == "fresh_upload"
        assert data["t12_cross_check"]["rent_roll_annual_rent"] == 60000.0
    finally:
        os.unlink(db_path)
    print("✓ test_t12_fresh_upload_takes_precedence_over_persisted: PASS")


def test_t12_falls_back_to_persisted_discrepancy_when_no_fresh_upload():
    db_path = _fresh_temp_db()
    try:
        from app.discrepancies import sync_t12_reconciliation
        _insert("lease.pdf", tenant="Acme", rent_amount="$20,000.00", property_address="1 Main St")
        stale_result = {
            "property_address": "1 Main St", "flagged": True, "matched_lease_count": 1, "excluded_lease_count": 0,
            "rent_roll_annual_rent": 240000.0, "t12_annual_rental_income": 150000.0, "difference": 90000.0,
            "difference_pct": 37.5, "direction": "rent_roll_higher",
        }
        sync_t12_reconciliation(stale_result)

        data = build_investment_memo_data(property_address="1 Main St")
        assert data["t12_cross_check"]["available"] is True
        assert data["t12_cross_check"]["basis"] == "persisted"
        assert data["t12_cross_check"]["t12_annual_rental_income"] == 150000.0
    finally:
        os.unlink(db_path)
    print("✓ test_t12_falls_back_to_persisted_discrepancy_when_no_fresh_upload: PASS")


def test_t12_unavailable_for_portfolio_scope():
    db_path = _fresh_temp_db()
    try:
        data = build_investment_memo_data()  # no property_address
        assert data["t12_cross_check"]["available"] is False
        assert "single property" in data["t12_cross_check"]["reason"]
    finally:
        os.unlink(db_path)
    print("✓ test_t12_unavailable_for_portfolio_scope: PASS")


# ------------------------------------------------------------------
# PDF / Excel rendering
# ------------------------------------------------------------------

def test_pdf_renders_real_content_for_a_populated_property():
    db_path = _fresh_temp_db()
    try:
        _insert("lease.pdf", tenant="Acme Roasters", rent_amount="$6,250.00", property_address="1 Main St", lease_end_date="March 31, 2030", square_footage="2,000 sq ft")
        data = build_investment_memo_data(property_address="1 Main St")
        pdf_bytes = generate_investment_memo_pdf(data)
        assert len(pdf_bytes) > 1000
        text = _pdf_text(pdf_bytes)
        assert "Acme Roasters" in text
        assert "$6,250.00" in text
        assert "Rollover Risk Summary" in text
        assert "T12 Cross-Check Summary" in text
        assert "Flagged Discrepancies" in text
    finally:
        os.unlink(db_path)
    print("✓ test_pdf_renders_real_content_for_a_populated_property: PASS")


def test_pdf_renders_for_empty_portfolio_without_crashing():
    db_path = _fresh_temp_db()
    try:
        data = build_investment_memo_data()
        pdf_bytes = generate_investment_memo_pdf(data)
        assert len(pdf_bytes) > 500
        text = _pdf_text(pdf_bytes)
        assert "No leases in scope" in text
    finally:
        os.unlink(db_path)
    print("✓ test_pdf_renders_for_empty_portfolio_without_crashing: PASS")


def test_excel_has_all_five_sheets_with_correct_dedup():
    db_path = _fresh_temp_db()
    try:
        _insert("lease.pdf", tenant="Acme", rent_amount="$6,250.00", property_address="1 Main St", square_footage="1,000 sq ft", lease_end_date="March 31, 2030")
        _insert("rentroll.csv", tenant="Acme", rent_amount="$6,000.00", property_address="1 Main St", lease_end_date="March 31, 2030")

        data = build_investment_memo_data(property_address="1 Main St")
        excel_bytes = generate_investment_memo_excel(data)
        workbook = load_workbook(io.BytesIO(excel_bytes))

        assert workbook.sheetnames == ["Overview", "Key Lease Terms", "Discrepancies & Resolutions", "T12 Cross-Check", "Rollover Schedule"]

        overview = workbook["Overview"]
        rows = {row[0]: row[1] for row in overview.iter_rows(values_only=True) if row[0]}
        assert rows["Total Monthly Rent"] == 6250.0, "Excel overview must match the deduped PDF total, not double-count"

        key_terms = workbook["Key Lease Terms"]
        source_col = [row[1] for row in key_terms.iter_rows(min_row=2, values_only=True)]
        assert set(source_col) == {"Lease Document", "Rent Roll Import"}, "both records must still be LISTED, just not double-counted"
    finally:
        os.unlink(db_path)
    print("✓ test_excel_has_all_five_sheets_with_correct_dedup: PASS")


# ------------------------------------------------------------------
# API routes
# ------------------------------------------------------------------

def test_pdf_route_portfolio_wide():
    db_path = _fresh_temp_db()
    try:
        client = _authed_client()
        _insert("lease.pdf", tenant="Acme", rent_amount="$5,000.00", property_address="1 Main St")
        resp = client.post("/portfolio/investment-memo.pdf", data={})
        assert resp.status_code == 200
        assert resp.mimetype == "application/pdf"
        assert "attachment" in resp.headers.get("Content-Disposition", "")
        assert len(resp.data) > 500
    finally:
        os.unlink(db_path)
    print("✓ test_pdf_route_portfolio_wide: PASS")


def test_pdf_route_property_scoped():
    db_path = _fresh_temp_db()
    try:
        client = _authed_client()
        _insert("lease.pdf", tenant="Acme", rent_amount="$5,000.00", property_address="1 Main St")
        resp = client.post("/portfolio/investment-memo.pdf", data={"property_address": "1 Main St"})
        assert resp.status_code == 200
        text = _pdf_text(resp.data)
        assert "1 Main St" in text
        assert "Acme" in text
    finally:
        os.unlink(db_path)
    print("✓ test_pdf_route_property_scoped: PASS")


def test_excel_route():
    db_path = _fresh_temp_db()
    try:
        client = _authed_client()
        _insert("lease.pdf", tenant="Acme", rent_amount="$5,000.00", property_address="1 Main St")
        resp = client.post("/portfolio/investment-memo.xlsx", data={})
        assert resp.status_code == 200
        assert resp.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        workbook = load_workbook(io.BytesIO(resp.data))
        assert "Overview" in workbook.sheetnames
    finally:
        os.unlink(db_path)
    print("✓ test_excel_route: PASS")


def test_t12_file_without_property_address_returns_400():
    db_path = _fresh_temp_db()
    try:
        client = _authed_client()
        resp = client.post(
            "/portfolio/investment-memo.pdf",
            data={"t12_file": (io.BytesIO(b"Notes\nfake t12\n"), "t12.csv")},
            content_type="multipart/form-data",
        )
        assert resp.status_code == 400
        assert "property_address" in resp.get_json()["error"]
    finally:
        os.unlink(db_path)
    print("✓ test_t12_file_without_property_address_returns_400: PASS")


def test_invalid_t12_file_type_returns_400():
    db_path = _fresh_temp_db()
    try:
        client = _authed_client()
        resp = client.post(
            "/portfolio/investment-memo.pdf",
            data={"property_address": "1 Main St", "t12_file": (io.BytesIO(b"not a real t12"), "t12.txt")},
            content_type="multipart/form-data",
        )
        assert resp.status_code == 400
    finally:
        os.unlink(db_path)
    print("✓ test_invalid_t12_file_type_returns_400: PASS")


if __name__ == "__main__":
    test_empty_portfolio_returns_honest_zero_state()
    test_property_scope_filters_by_building_ignoring_suite()
    test_dedupes_rent_roll_duplicate_of_a_pdf_lease_for_financial_totals()
    test_dedupe_prefers_pdf_lease_regardless_of_upload_order()
    test_dedupe_keeps_most_recent_when_both_are_the_same_source_type()
    test_discrepancies_and_resolutions_scoped_to_property()
    test_portfolio_wide_memo_excludes_discrepancies_for_deleted_leases()
    test_portfolio_wide_memo_excludes_tenant_concentration_for_a_tenant_no_longer_in_the_portfolio()
    test_t12_fresh_upload_takes_precedence_over_persisted()
    test_t12_falls_back_to_persisted_discrepancy_when_no_fresh_upload()
    test_t12_unavailable_for_portfolio_scope()
    test_pdf_renders_real_content_for_a_populated_property()
    test_pdf_renders_for_empty_portfolio_without_crashing()
    test_excel_has_all_five_sheets_with_correct_dedup()
    test_pdf_route_portfolio_wide()
    test_pdf_route_property_scoped()
    test_excel_route()
    test_t12_file_without_property_address_returns_400()
    test_invalid_t12_file_type_returns_400()
    print("\nAll investment memo tests passed.")
