"""
Tests for app/deal_mismatch.py (the 7 discrepancy detectors + summary)
and app/deal_mismatch_export.py (PDF/Excel renderers), plus the
POST /portfolio/deal-mismatch-report[.pdf|.xlsx] routes.

Same conventions as test_investment_memo.py / test_ai_rent_roll_
validation.py: an isolated temp SQLite file per test, database.insert_
lease with FIELD_NAMES-shaped extracted_fields, filename extension
("...csv"/"...xlsx" vs "...pdf") deciding _is_rent_roll_import.
"""

import io
import os
import sys
import tempfile
from datetime import date

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from openpyxl import load_workbook
import pypdf

from app.api import app
from app import database
from app.portfolio import FIELD_NAMES
from app import deal_mismatch as dm
from app.deal_mismatch_export import generate_deal_mismatch_report_pdf, generate_deal_mismatch_report_excel


def _fresh_temp_db():
    tmp = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
    tmp.close()
    database.configure(tmp.name)
    database.init_db()
    return tmp.name


def _authed_client(role="analyst"):
    client = app.test_client()
    with client.session_transaction() as sess:
        sess["user_id"] = 1
        sess["email"] = f"test-{role}@example.com"
        sess["name"] = f"Test {role.capitalize()}"
        sess["role"] = role
    return client


def _fields(**overrides):
    result = {}
    for name in FIELD_NAMES:
        if name in overrides and overrides[name] is not None:
            value = overrides[name]
            result[name] = {"value": value, "source": {"page": 3, "quote": f"...{value}..."}, "confidence": "high"}
        else:
            result[name] = {"value": None, "source": None, "confidence": None}
    return result


def _insert_lease(filename="lease.pdf", **overrides):
    return database.insert_lease(filename, _fields(**overrides), display_name=overrides.get("tenant") or filename)


def _insert_rent_roll_row(filename="rentroll.csv", **overrides):
    return database.insert_lease(filename, _fields(**overrides), display_name=overrides.get("tenant") or filename)


def _rr_lease(**f):
    return {"id": 1, "filename": "rentroll.csv", "extracted_fields": _fields(**f)}


def _doc_lease(**f):
    return {"id": 2, "filename": "lease.pdf", "extracted_fields": _fields(**f)}


# ----------------------------------------------------------------------
# detect_rent_mismatch
# ----------------------------------------------------------------------

def test_rent_mismatch_flags_past_tolerance():
    rr = _rr_lease(property_address="1 Main St, Suite 100", tenant="Acme", rent_amount="$6,000.00")
    doc = _doc_lease(property_address="1 Main St, Suite 100", tenant="Acme", rent_amount="$6,250.00")
    rows = dm.detect_rent_mismatch([rr, doc])
    assert len(rows) == 1
    row = rows[0]
    assert row["discrepancy_type"] == "rent_mismatch"
    assert row["monthly_dollar_impact"] == 250.0
    assert row["annual_dollar_impact"] == 3000.0
    assert row["income_direction"] == "understate", "rent roll ($6,000) is below the lease ($6,250)"
    assert row["severity"] == "medium", "a $250 gap meets (>=) the $250/mo 'medium' threshold"
    assert row["source"] == {"page": 3, "quote": "...$6,250.00..."}
    print("✓ test_rent_mismatch_flags_past_tolerance: PASS")


def test_rent_mismatch_overstate_direction():
    rr = _rr_lease(property_address="1 Main St, Suite 100", tenant="Acme", rent_amount="$7,500.00")
    doc = _doc_lease(property_address="1 Main St, Suite 100", tenant="Acme", rent_amount="$6,000.00")
    rows = dm.detect_rent_mismatch([rr, doc])
    assert len(rows) == 1
    assert rows[0]["income_direction"] == "overstate", "rent roll ($7,500) is above the lease ($6,000)"
    assert rows[0]["severity"] == "high", "a $1,500/mo gap clears the $1,000/mo 'high' threshold"
    print("✓ test_rent_mismatch_overstate_direction: PASS")


def test_rent_mismatch_within_tolerance_not_flagged():
    rr = _rr_lease(property_address="1 Main St, Suite 100", tenant="Acme", rent_amount="$6,250.00")
    doc = _doc_lease(property_address="1 Main St, Suite 100", tenant="Acme", rent_amount="$6,251.00")
    rows = dm.detect_rent_mismatch([rr, doc])
    assert rows == [], "a $1 gap is within both the absolute and percentage tolerance"
    print("✓ test_rent_mismatch_within_tolerance_not_flagged: PASS")


# ----------------------------------------------------------------------
# detect_expired_but_occupied
# ----------------------------------------------------------------------

def test_expired_but_occupied_flags_past_end_date_with_real_tenant():
    rr = _rr_lease(property_address="1 Main St, Suite 100", tenant="Acme", rent_amount="$6,000.00")
    doc = _doc_lease(property_address="1 Main St, Suite 100", tenant="Acme", lease_end_date="January 1, 2024")
    rows = dm.detect_expired_but_occupied([rr, doc], today=date(2026, 1, 1))
    assert len(rows) == 1
    row = rows[0]
    assert row["income_direction"] == "overstate"
    assert row["monthly_dollar_impact"] == 6000.0
    assert row["annual_dollar_impact"] == 72000.0
    print("✓ test_expired_but_occupied_flags_past_end_date_with_real_tenant: PASS")


def test_expired_but_occupied_not_flagged_when_lease_still_active():
    rr = _rr_lease(property_address="1 Main St, Suite 100", tenant="Acme", rent_amount="$6,000.00")
    doc = _doc_lease(property_address="1 Main St, Suite 100", tenant="Acme", lease_end_date="January 1, 2030")
    rows = dm.detect_expired_but_occupied([rr, doc], today=date(2026, 1, 1))
    assert rows == []
    print("✓ test_expired_but_occupied_not_flagged_when_lease_still_active: PASS")


# ----------------------------------------------------------------------
# detect_unit_no_lease / detect_lease_no_unit
# ----------------------------------------------------------------------

def test_unit_no_lease_has_no_source_or_direction():
    rr = _rr_lease(property_address="2 Oak Ave", tenant="Beta LLC", rent_amount="$3,000.00")
    rows = dm.detect_unit_no_lease([rr])
    assert len(rows) == 1
    row = rows[0]
    assert row["discrepancy_type"] == "unit_no_lease"
    assert row["source"] is None, "no lease PDF exists to cite"
    assert row["income_direction"] is None, "unverified, not asserted as overstated"
    assert row["monthly_dollar_impact"] == 3000.0
    print("✓ test_unit_no_lease_has_no_source_or_direction: PASS")


def test_lease_no_unit_understates_income():
    doc = _doc_lease(property_address="3 Pine Rd", tenant="Gamma Inc", rent_amount="$4,500.00")
    rows = dm.detect_lease_no_unit([doc])
    assert len(rows) == 1
    row = rows[0]
    assert row["discrepancy_type"] == "lease_no_unit"
    assert row["income_direction"] == "understate"
    assert row["monthly_dollar_impact"] == 4500.0
    assert row["source"] is not None
    print("✓ test_lease_no_unit_understates_income: PASS")


def test_matched_unit_produces_neither_presence_finding():
    rr = _rr_lease(property_address="1 Main St", tenant="Acme", rent_amount="$6,000.00")
    doc = _doc_lease(property_address="1 Main St", tenant="Acme", rent_amount="$6,000.00")
    assert dm.detect_unit_no_lease([rr, doc]) == []
    assert dm.detect_lease_no_unit([rr, doc]) == []
    print("✓ test_matched_unit_produces_neither_presence_finding: PASS")


# ----------------------------------------------------------------------
# detect_concession_missing (Phase 1 stub)
# ----------------------------------------------------------------------

def test_concession_missing_is_stubbed_empty():
    rr = _rr_lease(property_address="1 Main St", tenant="Acme", rent_amount="$6,000.00")
    doc = _doc_lease(property_address="1 Main St", tenant="Acme", rent_amount="$6,000.00")
    assert dm.detect_concession_missing([rr, doc]) == []
    print("✓ test_concession_missing_is_stubbed_empty: PASS")


# ----------------------------------------------------------------------
# detect_dates_mismatch
# ----------------------------------------------------------------------

def test_dates_mismatch_flags_end_date_and_start_date_separately():
    rr = _rr_lease(
        property_address="1 Main St", tenant="Acme",
        lease_start_date="January 1, 2023", lease_end_date="December 31, 2027",
    )
    doc = _doc_lease(
        property_address="1 Main St", tenant="Acme",
        lease_start_date="June 1, 2023", lease_end_date="March 31, 2030",
    )
    rows = dm.detect_dates_mismatch([rr, doc])
    fields = {row["field"] for row in rows}
    assert fields == {"lease_start_date", "lease_end_date"}
    assert all(row["severity"] == "medium" for row in rows)
    assert all(row["monthly_dollar_impact"] is None for row in rows)
    print("✓ test_dates_mismatch_flags_end_date_and_start_date_separately: PASS")


def test_dates_mismatch_agrees_produces_nothing():
    rr = _rr_lease(property_address="1 Main St", tenant="Acme", lease_end_date="December 31, 2027")
    doc = _doc_lease(property_address="1 Main St", tenant="Acme", lease_end_date="December 31, 2027")
    assert dm.detect_dates_mismatch([rr, doc]) == []
    print("✓ test_dates_mismatch_agrees_produces_nothing: PASS")


# ----------------------------------------------------------------------
# detect_tenant_mismatch
# ----------------------------------------------------------------------

def test_tenant_mismatch_flags_disagreement_as_high_severity():
    rr = _rr_lease(property_address="1 Main St", tenant="Acme Corp")
    doc = _doc_lease(property_address="1 Main St", tenant="Beta LLC")
    rows = dm.detect_tenant_mismatch([rr, doc])
    assert len(rows) == 1
    assert rows[0]["severity"] == "high"
    assert rows[0]["monthly_dollar_impact"] is None
    print("✓ test_tenant_mismatch_flags_disagreement_as_high_severity: PASS")


def test_tenant_mismatch_tolerates_case_and_punctuation():
    rr = _rr_lease(property_address="1 Main St", tenant="Acme Corp.")
    doc = _doc_lease(property_address="1 Main St", tenant="acme corp")
    assert dm.detect_tenant_mismatch([rr, doc]) == []
    print("✓ test_tenant_mismatch_tolerates_case_and_punctuation: PASS")


# ----------------------------------------------------------------------
# build_deal_mismatch_report_data -- integration + summary math
# ----------------------------------------------------------------------

def test_empty_portfolio_honest_zero_state():
    db_path = _fresh_temp_db()
    try:
        data = dm.build_deal_mismatch_report_data()
        assert data["total_units_checked"] == 0
        assert data["total_discrepancies"] == 0
        assert data["annual_income_overstatement"] is None, "no data at all is 'unknown', not a guessed 0"
        assert data["discrepancies"] == []
    finally:
        os.unlink(db_path)
    print("✓ test_empty_portfolio_honest_zero_state: PASS")


def test_summary_sums_signed_impact_across_units():
    db_path = _fresh_temp_db()
    try:
        # Unit A: rent roll overstates by $500/mo -> +$6,000/yr
        _insert_lease("lease_a.pdf", tenant="Acme", rent_amount="$5,500.00", property_address="1 Main St, Suite A")
        _insert_rent_roll_row("rr_a.csv", tenant="Acme", rent_amount="$6,000.00", property_address="1 Main St, Suite A")
        # Unit B: lease on file, not on rent roll -> understates by full $4,000/mo -> -$48,000/yr
        _insert_lease("lease_b.pdf", tenant="Beta", rent_amount="$4,000.00", property_address="1 Main St, Suite B")

        data = dm.build_deal_mismatch_report_data()
        assert data["total_units_checked"] == 2
        assert data["total_discrepancies"] == 2
        assert data["annual_income_overstatement"] == 6000.0 - 48000.0
    finally:
        os.unlink(db_path)
    print("✓ test_summary_sums_signed_impact_across_units: PASS")


def test_property_address_scopes_report():
    db_path = _fresh_temp_db()
    try:
        _insert_lease("lease_a.pdf", tenant="Acme", rent_amount="$5,000.00", property_address="1 Main St, Suite A")
        _insert_lease("lease_b.pdf", tenant="Beta", rent_amount="$4,000.00", property_address="99 Oak Ave, Suite B")

        data = dm.build_deal_mismatch_report_data(property_address="1 Main St")
        assert data["total_units_checked"] == 1
    finally:
        os.unlink(db_path)
    print("✓ test_property_address_scopes_report: PASS")


# ----------------------------------------------------------------------
# Exporters
# ----------------------------------------------------------------------

def test_pdf_export_is_valid_and_readable():
    db_path = _fresh_temp_db()
    try:
        _insert_lease("lease_a.pdf", tenant="Acme", rent_amount="$5,500.00", property_address="1 Main St, Suite A")
        _insert_rent_roll_row("rr_a.csv", tenant="Acme", rent_amount="$6,000.00", property_address="1 Main St, Suite A")
        data = dm.build_deal_mismatch_report_data()
        pdf_bytes = generate_deal_mismatch_report_pdf(data)
        reader = pypdf.PdfReader(io.BytesIO(pdf_bytes))
        text = "\n".join(page.extract_text() for page in reader.pages)
        assert "Deal Mismatch Report" in text
        assert "overstates" in text.lower()
    finally:
        os.unlink(db_path)
    print("✓ test_pdf_export_is_valid_and_readable: PASS")


def test_excel_export_has_header_and_rows():
    db_path = _fresh_temp_db()
    try:
        _insert_lease("lease_a.pdf", tenant="Acme", rent_amount="$5,500.00", property_address="1 Main St, Suite A")
        _insert_rent_roll_row("rr_a.csv", tenant="Acme", rent_amount="$6,000.00", property_address="1 Main St, Suite A")
        data = dm.build_deal_mismatch_report_data()
        excel_bytes = generate_deal_mismatch_report_excel(data)
        workbook = load_workbook(io.BytesIO(excel_bytes))
        sheet = workbook.active
        assert sheet.cell(row=7, column=1).value == "Unit"
        assert sheet.cell(row=8, column=1).value == "1 Main St, Suite A"
    finally:
        os.unlink(db_path)
    print("✓ test_excel_export_has_header_and_rows: PASS")


# ----------------------------------------------------------------------
# Routes
# ----------------------------------------------------------------------

def test_json_route_requires_auth():
    db_path = _fresh_temp_db()
    try:
        client = app.test_client()
        resp = client.post('/portfolio/deal-mismatch-report')
        assert resp.status_code == 401
    finally:
        os.unlink(db_path)
    print("✓ test_json_route_requires_auth: PASS")


def test_routes_require_analyst_role():
    """
    Viewer role is below analyst on ROLE_RANK (app/auth.py), and every
    comparable portfolio export (investment-memo.pdf/xlsx, t12-
    reconciliation) is gated at 'analyst' -- these three routes must
    match that, not sit open to a bare logged-in viewer.
    """
    db_path = _fresh_temp_db()
    try:
        viewer = _authed_client(role="viewer")
        for path in (
            '/portfolio/deal-mismatch-report',
            '/portfolio/deal-mismatch-report.pdf',
            '/portfolio/deal-mismatch-report.xlsx',
        ):
            resp = viewer.post(path)
            assert resp.status_code == 403, f"{path} let a viewer through ({resp.status_code})"

        analyst = _authed_client(role="analyst")
        for path in (
            '/portfolio/deal-mismatch-report',
            '/portfolio/deal-mismatch-report.pdf',
            '/portfolio/deal-mismatch-report.xlsx',
        ):
            resp = analyst.post(path)
            assert resp.status_code == 200, f"{path} blocked an analyst ({resp.status_code})"
    finally:
        os.unlink(db_path)
    print("✓ test_routes_require_analyst_role: PASS")


def test_json_route_returns_report_and_persists_discrepancies():
    db_path = _fresh_temp_db()
    try:
        _insert_lease("lease_a.pdf", tenant="Acme", rent_amount="$5,500.00", property_address="1 Main St, Suite A")
        _insert_rent_roll_row("rr_a.csv", tenant="Acme", rent_amount="$6,000.00", property_address="1 Main St, Suite A")

        client = _authed_client()
        resp = client.post('/portfolio/deal-mismatch-report')
        assert resp.status_code == 200
        body = resp.get_json()
        assert body["total_discrepancies"] == 1
        assert body["discrepancies"][0]["discrepancy_id"] is not None, "route must persist via sync_deal_mismatch_report"
    finally:
        os.unlink(db_path)
    print("✓ test_json_route_returns_report_and_persists_discrepancies: PASS")


def test_pdf_route():
    db_path = _fresh_temp_db()
    try:
        client = _authed_client()
        resp = client.post('/portfolio/deal-mismatch-report.pdf')
        assert resp.status_code == 200
        assert resp.content_type == 'application/pdf'
        assert len(resp.data) > 500
    finally:
        os.unlink(db_path)
    print("✓ test_pdf_route: PASS")


def test_excel_route():
    db_path = _fresh_temp_db()
    try:
        client = _authed_client()
        resp = client.post('/portfolio/deal-mismatch-report.xlsx')
        assert resp.status_code == 200
        assert 'spreadsheetml' in resp.content_type
        assert len(resp.data) > 500
    finally:
        os.unlink(db_path)
    print("✓ test_excel_route: PASS")


if __name__ == "__main__":
    test_rent_mismatch_flags_past_tolerance()
    test_rent_mismatch_overstate_direction()
    test_rent_mismatch_within_tolerance_not_flagged()
    test_expired_but_occupied_flags_past_end_date_with_real_tenant()
    test_expired_but_occupied_not_flagged_when_lease_still_active()
    test_unit_no_lease_has_no_source_or_direction()
    test_lease_no_unit_understates_income()
    test_matched_unit_produces_neither_presence_finding()
    test_concession_missing_is_stubbed_empty()
    test_dates_mismatch_flags_end_date_and_start_date_separately()
    test_dates_mismatch_agrees_produces_nothing()
    test_tenant_mismatch_flags_disagreement_as_high_severity()
    test_tenant_mismatch_tolerates_case_and_punctuation()
    test_empty_portfolio_honest_zero_state()
    test_summary_sums_signed_impact_across_units()
    test_property_address_scopes_report()
    test_pdf_export_is_valid_and_readable()
    test_excel_export_has_header_and_rows()
    test_json_route_requires_auth()
    test_routes_require_analyst_role()
    test_json_route_returns_report_and_persists_discrepancies()
    test_pdf_route()
    test_excel_route()
    print("\nAll deal_mismatch tests passed.")
