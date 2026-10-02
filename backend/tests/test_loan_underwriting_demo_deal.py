"""
End-to-end test on the fictional Maple Ridge demo deal: real T-12 file ->
parsed line items -> underwritten NOI -> ratios and sizing -> stress
tests -> Word credit memo, through the actual HTTP routes.

Maple Ridge is entirely fabricated demo data (see
benchmark_data/demo_deal/README.md -- "No real property, owner, tenant,
or lender is represented anywhere in this package"). Nothing here touches
real customer data.

ON LOCATING THE DEMO FILES: benchmark_data/demo_deal/ is not committed on
this branch -- per TASKS.md it lands with fix/rent-roll-hardening. So this
test looks for the T-12 in a few places and SKIPS with a clear message if
it can't find one, rather than failing the suite for an absent fixture
that isn't this branch's to provide. Point it anywhere with
ABSTRACTLY_DEMO_DEAL_DIR.

The expected figures below were derived by hand from the T-12's own
printed lines (see the per-assertion comments) and independently matched
the engine's output, which is why they're asserted as exact values rather
than as "something was produced".
"""

import io
import os
import sys
import tempfile

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from docx import Document

from app.api import app
from app import database
from app.auth import hash_password
from app.t12_import import parse_t12_financials

_FLAG = "LOAN_UNDERWRITING_ENABLED"
_BACKEND = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))

# A realistic lender request against the demo property: 120 units,
# $18.5M purchase, 75% LTV sizing, 6.25% / 30-year amortization, 10-year
# term, 24 months interest-only.
_LOAN_REQUEST = {
    "deal_name": "Maple Ridge Apartments",
    "property_address": "4500 Maple Ridge Trail, Dallas, TX 75248",
    "loan_amount": 13_875_000,
    "annual_rate_pct": 6.25,
    "amortization_years": 30,
    "term_years": 10,
    "interest_only_months": 24,
    "purchase_price": 18_500_000,
    "appraised_value": 18_750_000,
    "unit_count": 120,
}


def _find_demo_t12():
    """
    The demo T-12, or None. Checks an explicit override first, then this
    branch, then the sibling worktrees this repo actually uses (the data
    is untracked and lives wherever it was generated).
    """
    candidates = []
    override = os.environ.get("ABSTRACTLY_DEMO_DEAL_DIR")
    if override:
        candidates.append(os.path.join(override, "t12", "maple_ridge_t12.xlsx"))

    candidates.append(os.path.join(_BACKEND, "benchmark_data", "demo_deal", "t12", "maple_ridge_t12.xlsx"))

    relative = os.path.join("backend", "benchmark_data", "demo_deal", "t12", "maple_ridge_t12.xlsx")
    projects = os.path.abspath(os.path.join(_BACKEND, "..", ".."))
    for sibling in ("lease-abstraction", "abstractly-rentroll-qa"):
        candidates.append(os.path.join(projects, sibling, relative))

    for path in candidates:
        if os.path.isfile(path):
            return path
    return None


def _fresh_temp_db():
    tmp = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
    tmp.close()
    database.configure(tmp.name)
    database.init_db()
    return tmp.name


def _analyst_client():
    created = database.create_user(
        "analyst@example.com", "Demo Analyst", hash_password("pw-not-used"), role="analyst"
    )
    client = app.test_client()
    with client.session_transaction() as sess:
        sess["user_id"] = created["id"]
        sess["email"] = "analyst@example.com"
        sess["name"] = "Demo Analyst"
        sess["role"] = "analyst"
        sess["is_owner"] = False
    return client


def _docx_text(docx_bytes):
    document = Document(io.BytesIO(docx_bytes))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _skip(name):
    print(f"- {name}: SKIPPED (demo deal T-12 not found; set ABSTRACTLY_DEMO_DEAL_DIR "
          f"to backend/benchmark_data/demo_deal to run it)")


# ----------------------------------------------------------------------

def test_demo_t12_parses_every_underwriting_line_with_citations():
    """
    Straight from the real demo workbook. Figures are the statement's own
    printed Total column values:
      Gross Potential Rent              2,177,280.00  (row 7)
      Vacancy Loss                       -107,685.39  (row 8)
      Loss to Lease                       -25,521.97  (row 9)
      Net Rental Income Billed          2,044,072.64  (row 10)
      Concessions                         -11,200.00  (row 11)
      Bad Debt / Collection Loss         -122,562.00  (row 12)
      Total Rent Collected              1,910,310.64  (row 13)
      Total Other Income                   86,340.27  (row 22)
      Management Fee                       59,899.54  (row 35)
      TOTAL OPERATING EXPENSES            766,762.83  (row 36)
      NET OPERATING INCOME (NOI)        1,229,888.08  (row 38)

    And the identity that proves the parse is coherent:
      2,177,280.00 - 107,685.39 - 25,521.97 = 2,044,072.64
    exactly the statement's own "Net Rental Income Billed" line.
    """
    path = _find_demo_t12()
    if path is None:
        return _skip("test_demo_t12_parses_every_underwriting_line_with_citations")

    with open(path, "rb") as handle:
        parsed = parse_t12_financials(handle.read(), "maple_ridge_t12.xlsx", unit_count=120)

    inputs = parsed["t12_inputs"]
    assert inputs["gross_potential_rent"] == 2_177_280.00, inputs
    assert inputs["loss_to_lease"] == 25_521.97, inputs
    assert inputs["concessions"] == 11_200.00, inputs
    assert inputs["bad_debt"] == 122_562.00, inputs
    assert inputs["other_income"] == 86_340.27, inputs
    assert inputs["historical_noi"] == 1_229_888.08, inputs
    assert inputs["total_rent_collected"] == 1_910_310.64, inputs
    assert inputs["net_rental_income_billed"] == 2_044_072.64, inputs

    # 766,762.83 total opex less the 59,899.54 management fee.
    assert inputs["operating_expenses_ex_management"] == 706_863.29, inputs
    assert parsed["management_fee_removed"] == 59_899.54, parsed

    # The statement's own internal identity.
    assert round(
        inputs["gross_potential_rent"] - 107_685.39 - inputs["loss_to_lease"], 2
    ) == inputs["net_rental_income_billed"]

    # No warnings: this statement has every line the build-up needs.
    assert parsed["warnings"] == [], parsed["warnings"]

    # Every figure cites its source row in the real file.
    assert parsed["sources"]["gross_potential_rent"]["row"] == 7, parsed["sources"]
    assert parsed["sources"]["gross_potential_rent"]["quote"] == "Gross Potential Rent"
    assert parsed["sources"]["historical_noi"]["row"] == 38, parsed["sources"]
    print("✓ test_demo_t12_parses_every_underwriting_line_with_citations: PASS")


def test_demo_deal_underwrites_end_to_end_through_the_routes():
    """
    Hand calculation at the default assumptions (5% vacancy, 3%
    management fee, $250/unit reserves, 120 units):

      Gross potential rent                   2,177,280.00
      less vacancy 5%                          108,864.00
      less loss to lease                        25,521.97
      less concessions                          11,200.00
      less bad debt                            122,562.00
      = effective rental income              1,909,132.03
      plus other income                         86,340.27
      = effective gross income               1,995,472.30
      less opex excluding mgmt fee             706,863.29
      less management fee 3% of EGI             59,864.17
      less reserves 120 x 250                   30,000.00
      = total operating expenses               796,727.46
      = UNDERWRITTEN NOI                     1,198,744.84

    Debt service on $13,875,000 at 6.25% / 30 years:
      monthly 85,430.76 -> annual 1,025,169.14
      interest-only monthly 72,265.62 -> annual 867,187.50

    Ratios:
      DSCR (amortizing) = 1,198,744.84 / 1,025,169.14 = 1.1693x
      DSCR (year one, fully IO) = /867,187.50          = 1.3823x
      LTV = 13,875,000 / 18,500,000 (price is the lesser) = 75.00%
      debt yield = 1,198,744.84 / 13,875,000            = 8.6396%
      breakeven = (796,727.46 + 1,025,169.14)
                  / (2,177,280.00 + 86,340.27)          = 80.486%

    Sizing:
      by DSCR       = (1,198,744.84/1.25) / 0.0738862  = 12,979,387.71
      by LTV        = 18,500,000 x 75%                 = 13,875,000.00
      by debt yield = 1,198,744.84 / 8%                = 14,984,310.50
      -> DSCR binds, and the requested $13,875,000 is ABOVE the
         $12,979,387.71 maximum. That's the interesting result for a demo:
         the deal as requested is oversized by $895,612.29.
    """
    path = _find_demo_t12()
    if path is None:
        return _skip("test_demo_deal_underwrites_end_to_end_through_the_routes")

    _fresh_temp_db()
    os.environ[_FLAG] = "true"
    try:
        client = _analyst_client()

        created = client.post("/loan-underwriting/requests", json=_LOAN_REQUEST)
        assert created.status_code == 201, created.data
        request_id = created.get_json()["id"]

        with open(path, "rb") as handle:
            upload = client.post(
                f"/loan-underwriting/requests/{request_id}/t12",
                data={"file": (io.BytesIO(handle.read()), "maple_ridge_t12.xlsx")},
                content_type="multipart/form-data",
            )
        assert upload.status_code == 201, upload.data

        result = client.get(f"/loan-underwriting/requests/{request_id}/underwriting").get_json()

        build_up = result["noi_build_up"]
        assert build_up["vacancy_amount"] == 108_864.00, build_up
        assert build_up["effective_rental_income"] == 1_909_132.03, build_up
        assert build_up["effective_gross_income"] == 1_995_472.30, build_up
        assert build_up["management_fee"] == 59_864.17, build_up
        assert build_up["replacement_reserves"] == 30_000.00, build_up
        assert build_up["total_operating_expenses"] == 796_727.46, build_up
        assert build_up["underwritten_noi"] == 1_198_744.84, build_up

        # Historical NOI is passed through untouched for comparison.
        assert result["historical_noi"] == 1_229_888.08, result["historical_noi"]

        debt_service = result["debt_service"]
        assert debt_service["monthly_payment_amortizing"] == 85_430.76, debt_service
        assert debt_service["annual_debt_service_amortizing"] == 1_025_169.14, debt_service
        assert debt_service["annual_debt_service_interest_only"] == 867_187.50, debt_service
        # 24 months IO means all of year one is interest-only.
        assert debt_service["annual_debt_service_year_one"] == 867_187.50, debt_service

        ratios = result["ratios"]
        assert ratios["dscr"] == 1.1693, ratios
        assert ratios["dscr_year_one"] == 1.3823, ratios
        assert ratios["ltv_pct"] == 75.0, ratios
        assert ratios["ltv_basis"] == "purchase_price", ratios
        assert ratios["debt_yield_pct"] == 8.6396, ratios
        assert ratios["breakeven_occupancy_pct"] == 80.486, ratios

        sizing = result["sizing"]
        assert sizing["max_loan_by_dscr"] == 12_979_387.71, sizing
        assert sizing["max_loan_by_ltv"] == 13_875_000.00, sizing
        assert sizing["max_loan_by_debt_yield"] == 14_984_310.50, sizing
        assert sizing["binding_constraint"] == "dscr", sizing
        assert sizing["maximum_loan"] == 12_979_387.71, sizing
        # The requested loan exceeds what the property supports.
        assert sizing["maximum_loan"] < _LOAN_REQUEST["loan_amount"], sizing
    finally:
        os.environ.pop(_FLAG, None)
    print("✓ test_demo_deal_underwrites_end_to_end_through_the_routes: PASS")


def test_demo_deal_stress_tests_break_coverage_at_plus_200bp():
    """
    Hand calculation: at 8.25% (6.25% + 200bp) the amortizing payment on
    $13,875,000 over 30 years rises enough to push DSCR below 1.00x --
    i.e. the property stops covering its debt. That is the single most
    useful thing a stress table can tell a credit officer, so it's
    asserted rather than left to inspection.
    """
    path = _find_demo_t12()
    if path is None:
        return _skip("test_demo_deal_stress_tests_break_coverage_at_plus_200bp")

    _fresh_temp_db()
    os.environ[_FLAG] = "true"
    try:
        client = _analyst_client()
        request_id = client.post("/loan-underwriting/requests", json=_LOAN_REQUEST).get_json()["id"]
        with open(path, "rb") as handle:
            client.post(
                f"/loan-underwriting/requests/{request_id}/t12",
                data={"file": (io.BytesIO(handle.read()), "maple_ridge_t12.xlsx")},
                content_type="multipart/form-data",
            )
        cases = {c["key"]: c for c in
                 client.get(f"/loan-underwriting/requests/{request_id}/underwriting").get_json()["stress_tests"]}

        assert len(cases) == 5, cases
        assert cases["rate_plus_100bp"]["stressed_annual_rate_pct"] == 7.25
        assert cases["rate_plus_200bp"]["stressed_annual_rate_pct"] == 8.25

        # Coverage degrades as the rate rises, and breaks below 1.00x
        # at +200bp.
        assert cases["rate_plus_100bp"]["dscr"] < 1.1693
        assert cases["rate_plus_200bp"]["dscr"] < 1.0, cases["rate_plus_200bp"]

        # A 10-point occupancy drop also breaks coverage.
        assert cases["occupancy_down_10pt"]["stressed_vacancy_pct"] == 15.0
        assert cases["occupancy_down_10pt"]["dscr"] < 1.0, cases["occupancy_down_10pt"]

        # The NOI stress is exactly 90% of the base NOI.
        assert cases["noi_down_10pct"]["underwritten_noi"] == round(1_198_744.84 * 0.90, 2), \
            cases["noi_down_10pct"]
    finally:
        os.environ.pop(_FLAG, None)
    print("✓ test_demo_deal_stress_tests_break_coverage_at_plus_200bp: PASS")


def test_demo_deal_credit_memo_exports_to_word_with_real_figures():
    """
    The whole pipeline's output: a .docx carrying the demo deal's computed
    figures, its source citations, the binding-constraint statement, and
    the recommendation placeholder. The Anthropic key on this project has
    no credit, so the narrative sections take the documented fallback
    path -- and the document must still be complete and correct.
    """
    path = _find_demo_t12()
    if path is None:
        return _skip("test_demo_deal_credit_memo_exports_to_word_with_real_figures")

    _fresh_temp_db()
    os.environ[_FLAG] = "true"
    saved_key = os.environ.pop("ANTHROPIC_API_KEY", None)
    try:
        client = _analyst_client()
        request_id = client.post("/loan-underwriting/requests", json=_LOAN_REQUEST).get_json()["id"]
        with open(path, "rb") as handle:
            client.post(
                f"/loan-underwriting/requests/{request_id}/t12",
                data={"file": (io.BytesIO(handle.read()), "maple_ridge_t12.xlsx")},
                content_type="multipart/form-data",
            )

        response = client.post(f"/loan-underwriting/requests/{request_id}/credit-memo.docx")
        assert response.status_code == 200, response.data
        assert response.mimetype == \
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document", response.mimetype
        assert "attachment" in response.headers["Content-Disposition"]
        assert ".docx" in response.headers["Content-Disposition"]

        docx_bytes = response.data
        assert docx_bytes[:2] == b"PK"
        text = _docx_text(docx_bytes)

        # The deal.
        assert "Maple Ridge Apartments" in text
        assert "4500 Maple Ridge Trail" in text

        # The computed figures, exactly as the engine produced them.
        assert "1,198,744.84" in text, "underwritten NOI"
        assert "1,229,888.08" in text, "historical T-12 NOI"
        assert "2,177,280.00" in text, "gross potential rent"
        assert "1.17x" in text, "DSCR rendered to 2dp"
        assert "8.64%" in text, "debt yield"
        assert "12,979,387.71" in text, "maximum supportable loan"

        # The binding constraint, named, and the shortfall called out.
        assert "Binding constraint:" in text
        assert "Target DSCR" in text
        assert "BELOW" in text, "the requested loan exceeds the maximum and the memo must say so"

        # Provenance: a real row from the real workbook.
        assert "maple_ridge_t12.xlsx" in text
        assert "row 7" in text

        # The management fee adjustment is disclosed.
        assert "59,899.54" in text

        # And the tool still does not recommend.
        assert "FOR CREDIT OFFICER DETERMINATION" in text
        assert "SPONSOR ANALYSIS TO BE COMPLETED" in text
    finally:
        os.environ.pop(_FLAG, None)
        if saved_key is not None:
            os.environ["ANTHROPIC_API_KEY"] = saved_key
    print("✓ test_demo_deal_credit_memo_exports_to_word_with_real_figures: PASS")


if __name__ == "__main__":
    if _find_demo_t12() is None:
        print("NOTE: benchmark_data/demo_deal/ is not present on this branch "
              "(TASKS.md: it lands with fix/rent-roll-hardening).")
        print("      Set ABSTRACTLY_DEMO_DEAL_DIR=<path to demo_deal> to run these tests.")
    test_demo_t12_parses_every_underwriting_line_with_citations()
    test_demo_deal_underwrites_end_to_end_through_the_routes()
    test_demo_deal_stress_tests_break_coverage_at_plus_200bp()
    test_demo_deal_credit_memo_exports_to_word_with_real_figures()
    print("\nDemo deal end-to-end tests complete.")
