"""
Credit memo tests: template structure, the AI/arithmetic boundary, and
the Word export.

NO REAL ANTHROPIC CALL IS EVER MADE. Every test either injects a fake
client or relies on the no-API-key fallback path, following
app/cre_qa.py's injectable-client pattern.

The two assertions that matter most are behavioural, not cosmetic:

  1. THE TOOL NEVER RECOMMENDS. The memo must always carry the
     recommendation PLACEHOLDER, and the recommendation line must not
     become a verdict -- even when the model supplies approval language.
     A prompt instruction is not a guarantee, so this is enforced on the
     rendered output.
  2. NUMBERS COME FROM THE ENGINE, NOT THE MODEL. A model response full
     of wrong figures must not change a single number in the document.
"""

import io
import os
import sys

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

import anthropic
from docx import Document

from app.credit_memo import (
    CREDIT_MEMO_SYSTEM_PROMPT,
    CreditMemoError,
    build_credit_memo_data,
    generate_narratives,
)
from app.credit_memo_export import generate_credit_memo_docx
from app.credit_memo_template import (
    RECOMMENDATION_PLACEHOLDER,
    SECTIONS,
    SPONSOR_PLACEHOLDER,
    narrative_sections,
)
from app.loan_underwriting import underwrite


# ----------------------------------------------------------------------
# Fakes
# ----------------------------------------------------------------------

class _FakeBlock:
    type = "text"

    def __init__(self, text):
        self.text = text


class _FakeResponse:
    def __init__(self, text):
        self.content = [_FakeBlock(text)]


class _FakeAPIError(anthropic.APIError):
    """A constructible anthropic.APIError -- the real one's __init__ wants a request object a unit test doesn't have."""

    def __init__(self):
        Exception.__init__(self, "simulated API failure")


class _FakeMessages:
    def __init__(self, text=None, error=False):
        self._text = text
        self._error = error
        self.calls = []

    def create(self, **kwargs):
        self.calls.append(kwargs)
        if self._error:
            raise _FakeAPIError()
        return _FakeResponse(self._text)


class _FakeClient:
    def __init__(self, text=None, error=False):
        self.messages = _FakeMessages(text, error)


def _good_response_text():
    """A well-formed model response: every requested section behind its sentinel."""
    return "\n\n".join(
        f"<<<SECTION:{key}>>>\nDrafted prose for {key}. It refers to figures without inventing any."
        for key, _heading, _brief in narrative_sections()
    )


_T12 = {
    "gross_potential_rent": 2_000_000.0,
    "loss_to_lease": 20_000.0,
    "concessions": 10_000.0,
    "bad_debt": 30_000.0,
    "other_income": 80_000.0,
    "operating_expenses_ex_management": 642_400.0,
    "unit_count": 100,
    "historical_noi": 1_220_000.0,
}
_TERMS = {
    "loan_amount": 10_000_000.0,
    "annual_rate_pct": 6.0,
    "amortization_years": 30,
    "term_years": 10,
    "interest_only_months": 24,
    "purchase_price": 15_000_000.0,
    "appraised_value": 15_500_000.0,
    "unit_count": 100,
}


def _underwriting():
    result = underwrite(
        _TERMS, _T12,
        sources={"gross_potential_rent": {"row": 5, "file": "t12.csv", "quote": "Gross Potential Rent"}},
    )
    result["loan_request"] = {
        "property_address": "4500 Maple Ridge Trail, Dallas, TX 75248",
        "deal_name": "Maple Ridge Apartments",
        "loan_amount": _TERMS["loan_amount"],
    }
    result["t12_snapshot"] = {
        "id": 1, "filename": "t12.csv", "uploaded_at": "2026-10-01T00:00:00+00:00",
        "management_fee_removed": 57_600.0, "warnings": [],
        "line_items": [{"label": "Gross Potential Rent", "annual_amount": 2_000_000.0,
                        "source": {"row": 5, "file": "t12.csv", "quote": "Gross Potential Rent"}}],
        "t12_inputs": {"total_rent_collected": 1_840_000.0, "net_rental_income_billed": 1_950_000.0},
    }
    result["assumption_rows"] = [
        {"key": "vacancy_pct", "label": "Vacancy", "value": 5.0, "unit": "percent",
         "kind": "assumption", "source": None, "is_default": True},
    ]
    result["input_rows"] = [
        {"key": "loan_amount", "label": "Loan amount", "value": 10_000_000.0,
         "origin": "user_entered", "source": None},
        {"key": "gross_potential_rent", "label": "Gross potential rent", "value": 2_000_000.0,
         "origin": "t12_document",
         "source": {"file": "t12.csv", "row": 5, "page": None, "quote": "Gross Potential Rent"}},
    ]
    result["rent_roll_analysis"] = {
        "has_lease_data": False, "total_units_checked": 0,
        "total_discrepancies": 0, "annual_income_overstatement": None, "discrepancies": [],
    }
    return result


def _docx_text(docx_bytes):
    """All text in the rendered document -- paragraphs and table cells."""
    document = Document(io.BytesIO(docx_bytes))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


# ----------------------------------------------------------------------
# Template structure
# ----------------------------------------------------------------------

def test_template_has_the_standard_bank_structure_in_order():
    assert [s["key"] for s in SECTIONS] == [
        "loan_summary", "property", "sponsor", "rent_roll_analysis",
        "financials", "ratios_and_stress", "risks", "conditions",
    ], [s["key"] for s in SECTIONS]
    print("✓ test_template_has_the_standard_bank_structure_in_order: PASS")


def test_only_four_sections_are_ai_narrated():
    """Financials, ratios and the two placeholder sections carry no model-written prose at all."""
    assert [key for key, _h, _b in narrative_sections()] == [
        "property", "rent_roll_analysis", "risks", "conditions",
    ]
    for key in ("loan_summary", "financials", "ratios_and_stress", "sponsor"):
        section = next(s for s in SECTIONS if s["key"] == key)
        assert not section.get("narrative"), key
        assert "narrative" not in section["blocks"], key
    print("✓ test_only_four_sections_are_ai_narrated: PASS")


def test_loan_summary_always_carries_the_recommendation_placeholder_block():
    """The one structural requirement a swapped-in lender template must preserve."""
    loan_summary = next(s for s in SECTIONS if s["key"] == "loan_summary")
    assert "recommendation_placeholder" in loan_summary["blocks"], loan_summary
    print("✓ test_loan_summary_always_carries_the_recommendation_placeholder_block: PASS")


def test_system_prompt_forbids_computing_and_recommending():
    """The prompt is not the guarantee, but it must still say both things."""
    prompt = CREDIT_MEMO_SYSTEM_PROMPT.lower()
    assert "never compute" in prompt, "prompt must forbid computing"
    assert "do not recommend" in prompt, "prompt must forbid recommending"
    print("✓ test_system_prompt_forbids_computing_and_recommending: PASS")


# ----------------------------------------------------------------------
# Narrative generation (mocked)
# ----------------------------------------------------------------------

def test_narratives_are_parsed_from_a_well_formed_response():
    client = _FakeClient(text=_good_response_text())
    result = generate_narratives(_underwriting(), client=client)

    assert result["generated"] is True, result
    assert result["reason"] is None, result
    for key, _h, _b in narrative_sections():
        assert key in result["narratives"], key
        assert "Drafted prose for" in result["narratives"][key]
    print("✓ test_narratives_are_parsed_from_a_well_formed_response: PASS")


def test_one_api_call_is_made_for_all_sections_with_a_cached_system_prompt():
    """Four sections share one figures context, so paying for it four times would be waste."""
    client = _FakeClient(text=_good_response_text())
    generate_narratives(_underwriting(), client=client)

    assert len(client.messages.calls) == 1, client.messages.calls
    call = client.messages.calls[0]
    assert call["system"][0]["cache_control"] == {"type": "ephemeral"}, call["system"]
    print("✓ test_one_api_call_is_made_for_all_sections_with_a_cached_system_prompt: PASS")


def test_the_model_is_given_precomputed_figures_not_asked_to_compute():
    """The prompt must hand over formatted figures -- that's what makes 'no AI for any math' structural."""
    client = _FakeClient(text=_good_response_text())
    generate_narratives(_underwriting(), client=client)
    user_message = client.messages.calls[0]["messages"][0]["content"]

    assert "Computed figures" in user_message
    assert "2,000,000" in user_message, "gross potential rent should be supplied pre-formatted"
    assert "never recompute" in user_message
    print("✓ test_the_model_is_given_precomputed_figures_not_asked_to_compute: PASS")


def test_api_failure_degrades_to_placeholders_without_raising():
    """A missing key or a failed call must not block a credit memo -- the tables are the part a lender needs."""
    result = generate_narratives(_underwriting(), client=_FakeClient(error=True))
    assert result["generated"] is False, result
    assert "failed" in result["reason"], result
    for key, _h, _b in narrative_sections():
        assert "NARRATIVE NOT GENERATED" in result["narratives"][key], key
    print("✓ test_api_failure_degrades_to_placeholders_without_raising: PASS")


def test_missing_api_key_degrades_without_attempting_a_call():
    saved = os.environ.pop("ANTHROPIC_API_KEY", None)
    try:
        result = generate_narratives(_underwriting(), client=None)
        assert result["generated"] is False, result
        assert "no Anthropic API key" in result["reason"], result
    finally:
        if saved is not None:
            os.environ["ANTHROPIC_API_KEY"] = saved
    print("✓ test_missing_api_key_degrades_without_attempting_a_call: PASS")


def test_unparseable_response_degrades_rather_than_emitting_garbage():
    """A response with no sentinels at all must not be dumped into the document as one blob."""
    result = generate_narratives(_underwriting(), client=_FakeClient(text="I'm afraid I can't do that."))
    assert result["generated"] is False, result
    assert "could not be parsed" in result["reason"], result
    assert "I'm afraid" not in result["narratives"]["risks"], result["narratives"]["risks"]
    print("✓ test_unparseable_response_degrades_rather_than_emitting_garbage: PASS")


def test_a_partially_complete_response_keeps_what_it_got():
    """Two good sections and two missing: keep the two, placeholder the rest, and say which were missing."""
    partial = ("<<<SECTION:property>>>\nReal property prose.\n\n"
               "<<<SECTION:risks>>>\nReal risk prose.")
    result = generate_narratives(_underwriting(), client=_FakeClient(text=partial))

    assert result["generated"] is True, result
    assert result["narratives"]["property"] == "Real property prose."
    assert result["narratives"]["risks"] == "Real risk prose."
    assert "NARRATIVE NOT GENERATED" in result["narratives"]["conditions"]
    assert "missing section" in result["reason"], result
    print("✓ test_a_partially_complete_response_keeps_what_it_got: PASS")


# ----------------------------------------------------------------------
# The AI / arithmetic boundary
# ----------------------------------------------------------------------

def test_model_output_cannot_change_a_single_number_in_the_memo():
    """
    The load-bearing test for "all numbers come from the engine". A model
    response stuffed with wrong figures must leave every computed value
    untouched, because the numbers never travel through the model -- they
    go from the engine straight to the renderer.
    """
    underwriting = _underwriting()
    engine_noi = underwriting["noi_build_up"]["underwritten_noi"]
    engine_dscr = underwriting["ratios"]["dscr"]

    liar = "\n\n".join(
        f"<<<SECTION:{key}>>>\nThe NOI is $99,999,999 and the DSCR is 42.00x."
        for key, _h, _b in narrative_sections()
    )
    memo = build_credit_memo_data(underwriting, client=_FakeClient(text=liar))

    assert memo["underwriting"]["noi_build_up"]["underwritten_noi"] == engine_noi
    assert memo["underwriting"]["ratios"]["dscr"] == engine_dscr

    text = _docx_text(generate_credit_memo_docx(memo))
    # The engine's figures are what the document's TABLES print.
    assert f"{engine_noi:,.2f}" in text, f"expected engine NOI {engine_noi:,.2f} in the document"
    assert f"{engine_dscr:.2f}x" in text, f"expected engine DSCR {engine_dscr:.2f}x in the document"
    print("✓ test_model_output_cannot_change_a_single_number_in_the_memo: PASS")


def test_memo_always_carries_the_recommendation_placeholder_and_no_verdict():
    """
    Enforced on the rendered output, not trusted to the prompt. Even when
    the model supplies approval language in its prose, the recommendation
    LINE is always the placeholder.
    """
    pushy = "\n\n".join(
        f"<<<SECTION:{key}>>>\nWe recommend approval of this loan." for key, _h, _b in narrative_sections()
    )
    memo = build_credit_memo_data(_underwriting(), client=_FakeClient(text=pushy))
    text = _docx_text(generate_credit_memo_docx(memo))

    assert RECOMMENDATION_PLACEHOLDER in text, "the recommendation placeholder must always be present"
    assert "FOR CREDIT OFFICER DETERMINATION" in text
    assert "does not recommend" in text
    print("✓ test_memo_always_carries_the_recommendation_placeholder_and_no_verdict: PASS")


def test_sponsor_section_is_always_a_placeholder():
    """This tool holds no sponsor data, so it must say so rather than let a model imagine a sponsor."""
    memo = build_credit_memo_data(_underwriting(), client=_FakeClient(text=_good_response_text()))
    text = _docx_text(generate_credit_memo_docx(memo))
    assert SPONSOR_PLACEHOLDER in text
    assert "SPONSOR ANALYSIS TO BE COMPLETED" in text
    print("✓ test_sponsor_section_is_always_a_placeholder: PASS")


def test_build_credit_memo_data_rejects_something_that_is_not_an_underwriting_result():
    try:
        build_credit_memo_data({"not": "an underwriting result"}, narratives={})
    except CreditMemoError as exc:
        assert "underwriting result" in str(exc)
    else:
        raise AssertionError("expected CreditMemoError")
    print("✓ test_build_credit_memo_data_rejects_something_that_is_not_an_underwriting_result: PASS")


def test_narratives_can_be_supplied_directly_skipping_the_model_entirely():
    memo = build_credit_memo_data(_underwriting(), narratives={"property": "Supplied prose."})
    assert memo["narratives"]["property"] == "Supplied prose."
    assert memo["narratives_generated"] is True
    print("✓ test_narratives_can_be_supplied_directly_skipping_the_model_entirely: PASS")


# ----------------------------------------------------------------------
# Word export
# ----------------------------------------------------------------------

def test_docx_is_a_valid_openable_word_document():
    memo = build_credit_memo_data(_underwriting(), narratives={})
    docx_bytes = generate_credit_memo_docx(memo)

    assert docx_bytes[:2] == b"PK", "a .docx is a zip container"
    assert len(docx_bytes) > 5_000, len(docx_bytes)
    document = Document(io.BytesIO(docx_bytes))  # raises if malformed
    assert len(document.tables) >= 8, f"expected the engine's tables, got {len(document.tables)}"
    print("✓ test_docx_is_a_valid_openable_word_document: PASS")


def test_docx_contains_every_template_heading():
    memo = build_credit_memo_data(_underwriting(), narratives={})
    text = _docx_text(generate_credit_memo_docx(memo))
    for section in SECTIONS:
        assert section["heading"] in text, section["heading"]
    print("✓ test_docx_contains_every_template_heading: PASS")


def test_docx_shows_the_noi_build_up_the_ratios_and_all_five_stress_cases():
    memo = build_credit_memo_data(_underwriting(), narratives={})
    text = _docx_text(generate_credit_memo_docx(memo))

    for label in ("Gross potential rent", "UNDERWRITTEN NOI", "Effective gross income",
                  "DSCR (fully amortizing)", "Breakeven occupancy", "Debt yield"):
        assert label in text, label
    for stress in ("Interest rate +100 bps", "Interest rate +200 bps",
                   "Occupancy -5 points", "Occupancy -10 points", "NOI -10%"):
        assert stress in text, stress
    print("✓ test_docx_shows_the_noi_build_up_the_ratios_and_all_five_stress_cases: PASS")


def test_docx_cites_source_rows_for_t12_figures_and_marks_analyst_entered_ones():
    memo = build_credit_memo_data(_underwriting(), narratives={})
    text = _docx_text(generate_credit_memo_docx(memo))

    assert "row 5" in text, "a T-12 figure must cite its source row"
    assert "t12.csv" in text
    assert "entered by analyst" in text, "analyst-entered inputs must be marked, not given a fake citation"
    print("✓ test_docx_cites_source_rows_for_t12_figures_and_marks_analyst_entered_ones: PASS")


def test_docx_names_the_binding_constraint():
    memo = build_credit_memo_data(_underwriting(), narratives={})
    text = _docx_text(generate_credit_memo_docx(memo))
    assert "Binding constraint:" in text
    assert "YES" in text, "the binding constraint row must be marked in the sizing table"
    print("✓ test_docx_names_the_binding_constraint: PASS")


def test_docx_discloses_the_management_fee_adjustment():
    """The double-count guard has to be visible to a reader, not just correct internally."""
    memo = build_credit_memo_data(_underwriting(), narratives={})
    text = _docx_text(generate_credit_memo_docx(memo))
    assert "was removed from the operating expense total" in text
    print("✓ test_docx_discloses_the_management_fee_adjustment: PASS")


def test_docx_says_plainly_when_no_leases_are_on_file():
    """
    "0 discrepancies" and "nothing was checked" are completely different
    statements to a credit officer. The empty case must read as a gap.
    """
    memo = build_credit_memo_data(_underwriting(), narratives={})
    text = _docx_text(generate_credit_memo_docx(memo))
    assert "has NOT been verified against signed leases" in text
    assert "data gap" in text
    print("✓ test_docx_says_plainly_when_no_leases_are_on_file: PASS")


def test_docx_warns_when_narratives_were_not_generated():
    memo = build_credit_memo_data(_underwriting(), client=_FakeClient(error=True))
    text = _docx_text(generate_credit_memo_docx(memo))
    assert "narrative sections were not generated" in text
    # ...but the figures are still all there.
    assert "UNDERWRITTEN NOI" in text
    print("✓ test_docx_warns_when_narratives_were_not_generated: PASS")


def test_docx_includes_the_assumptions_and_sources_appendices():
    memo = build_credit_memo_data(_underwriting(), narratives={})
    text = _docx_text(generate_credit_memo_docx(memo))
    assert "Appendix A" in text and "Assumptions and Constraints Used" in text
    assert "Appendix B" in text and "Input Sources" in text
    print("✓ test_docx_includes_the_assumptions_and_sources_appendices: PASS")


if __name__ == "__main__":
    test_template_has_the_standard_bank_structure_in_order()
    test_only_four_sections_are_ai_narrated()
    test_loan_summary_always_carries_the_recommendation_placeholder_block()
    test_system_prompt_forbids_computing_and_recommending()

    test_narratives_are_parsed_from_a_well_formed_response()
    test_one_api_call_is_made_for_all_sections_with_a_cached_system_prompt()
    test_the_model_is_given_precomputed_figures_not_asked_to_compute()
    test_api_failure_degrades_to_placeholders_without_raising()
    test_missing_api_key_degrades_without_attempting_a_call()
    test_unparseable_response_degrades_rather_than_emitting_garbage()
    test_a_partially_complete_response_keeps_what_it_got()

    test_model_output_cannot_change_a_single_number_in_the_memo()
    test_memo_always_carries_the_recommendation_placeholder_and_no_verdict()
    test_sponsor_section_is_always_a_placeholder()
    test_build_credit_memo_data_rejects_something_that_is_not_an_underwriting_result()
    test_narratives_can_be_supplied_directly_skipping_the_model_entirely()

    test_docx_is_a_valid_openable_word_document()
    test_docx_contains_every_template_heading()
    test_docx_shows_the_noi_build_up_the_ratios_and_all_five_stress_cases()
    test_docx_cites_source_rows_for_t12_figures_and_marks_analyst_entered_ones()
    test_docx_names_the_binding_constraint()
    test_docx_discloses_the_management_fee_adjustment()
    test_docx_says_plainly_when_no_leases_are_on_file()
    test_docx_warns_when_narratives_were_not_generated()
    test_docx_includes_the_assumptions_and_sources_appendices()

    print("\nAll credit memo tests passed.")
