"""
Self-check for the synthetic Section 8 fixture set in fixtures/section8/.

These files are test DATA for future Section 8 checks (rent-roll split,
HAP reconciliation, recert and inspection tracking). Nothing in the app
reads them yet, so this test guards the fixtures themselves:

  - the generator is the source of truth: regenerating gives
    byte-identical files, so nobody can hand-edit a PDF and leave
    ANSWER_KEY.md wrong;
  - every piece of evidence the answer key cites really is in the named
    file, row and page (a planted error that isn't actually planted would
    silently make a future checker look better than it is);
  - the clean control units really are clean;
  - every PDF is readable by the app's own PDFExtractor (read-only use);
  - nothing looks like real personal data (no SSNs, only 555-01xx phone
    numbers and .test email addresses).

No network, no Anthropic API.
"""
import csv
import importlib.util
import io
import json
import os
import re
import sys
import tempfile

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from openpyxl import load_workbook

from app.pdf_extractor import PDFExtractor

FIX = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures", "section8")
GEN_PATH = os.path.join(FIX, "generate_section8_fixtures.py")

_spec = importlib.util.spec_from_file_location("generate_section8_fixtures", GEN_PATH)
gen = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(gen)

with open(os.path.join(FIX, "answer_key.json"), encoding="utf-8") as _f:
    KEY = json.load(_f)

_pdf_cache = {}


def pdf_pages(rel):
    if rel not in _pdf_cache:
        path = os.path.join(FIX, rel)
        with open(path, "rb") as f:
            pages = PDFExtractor().extract_text(f, pdf_path=path)
        _pdf_cache[rel] = [" ".join((p["text"] or "").split()) for p in pages]
    return _pdf_cache[rel]


def table_rows(rel):
    """(header, rows) of a rent roll or register, skipping any title block."""
    path = os.path.join(FIX, rel)
    if rel.endswith(".xlsx"):
        ws = load_workbook(path, data_only=True)["Rent Roll"]
        raw = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        with open(path, encoding="utf-8") as f:
            raw = list(csv.reader(f))
    for i, r in enumerate(raw):
        cells = [str(c).strip() if c is not None else "" for c in r]
        if "Unit" in cells or "Bldg/Unit" in cells:
            header = cells
            return header, [dict(zip(header, row)) for row in raw[i + 1:] if any(c not in (None, "") for c in row)]
    raise AssertionError(f"{rel}: no header row found")


def find_row(rel, unit, prop=None):
    header, rows = table_rows(rel)
    unit_col = "Unit" if "Unit" in header else "Bldg/Unit"
    hits = [r for r in rows if str(r.get(unit_col) or "").strip() == unit
            and (prop is None or "Property" not in header or r.get("Property") == prop)]
    return hits


def same(cell, expected):
    if expected is None:
        return cell in (None, "")
    if isinstance(expected, (int, float)):
        try:
            return abs(float(str(cell).replace(",", "")) - expected) < 0.005
        except (TypeError, ValueError):
            return False
    return str(cell if cell is not None else "").strip() == str(expected)


def test_regenerates_byte_identical():
    with tempfile.TemporaryDirectory() as tmp:
        gen.generate(tmp)
        for rel in gen.generated_files():
            with open(os.path.join(tmp, rel), "rb") as a, open(os.path.join(FIX, rel), "rb") as b:
                assert a.read() == b.read(), f"{rel} differs from a fresh regenerate - re-run the generator and commit"
    committed = sorted(
        os.path.relpath(os.path.join(dp, f), FIX)
        for dp, _, fs in os.walk(FIX) for f in fs
        if not f.startswith(".") and "__pycache__" not in dp)
    expected = sorted(gen.generated_files() + ["generate_section8_fixtures.py", "README.md"])
    assert committed == expected, f"stray or missing files: {set(committed) ^ set(expected)}"
    print("✓ test_regenerates_byte_identical: PASS")


def test_file_counts_match_the_brief():
    f = KEY["files"]
    assert len(f["leases"]) == 10
    assert len(f["hap_contracts"]) == 5
    assert len(f["rent_rolls"]) == 5
    assert sum(r.endswith(".csv") for r in f["rent_rolls"]) == 3
    assert sum(r.endswith(".xlsx") for r in f["rent_rolls"]) == 2
    assert len([r for r in f["hap_registers"] if r.endswith(".csv")]) == 3
    assert len(f["inspections"]) == 3
    kinds = {v["kind"] for v in gen.HAP_CONTRACTS.values()}
    assert {"HCV", "PBV", "PBRA"} <= kinds
    # Every lease carries a tenancy addendum.
    for rel in f["leases"]:
        text = " ".join(pdf_pages(rel))
        assert "TENANCY ADDENDUM" in text, rel
    print("✓ test_file_counts_match_the_brief: PASS")


def test_every_evidence_item_is_really_there():
    checked = 0
    for err in KEY["errors"]:
        for ev in err["evidence"]:
            where = f"{err['id']}: {ev}"
            if ev["type"] == "pdf_text":
                pages = pdf_pages(ev["file"])
                assert ev["page"] <= len(pages), where
                assert ev["contains"] in pages[ev["page"] - 1], where
            elif ev["type"] == "pdf_absent":
                pages = pdf_pages(ev["file"])
                assert ev["page"] <= len(pages), where
                assert ev["text"] not in pages[ev["page"] - 1], where
            elif ev["type"] in ("rent_roll_cell", "register_cell"):
                hits = find_row(ev["file"], ev["unit"], ev.get("prop"))
                assert len(hits) == 1, f"{where}: {len(hits)} rows"
                assert ev["column"] in hits[0], f"{where}: no such column"
                assert same(hits[0][ev["column"]], ev["value"]), f"{where}: found {hits[0][ev['column']]!r}"
            elif ev["type"] == "register_absent":
                assert find_row(ev["file"], ev["unit"]) == [], where
            else:
                raise AssertionError(f"unknown evidence type {ev['type']}")
            checked += 1
    assert checked >= 60, checked
    print(f"✓ test_every_evidence_item_is_really_there: PASS ({checked} items)")


def test_answer_key_covers_required_error_kinds():
    errors = KEY["errors"]
    assert 20 <= len(errors) <= 25, len(errors)
    ids = [e["id"] for e in errors]
    assert len(set(ids)) == len(ids)
    flags = {e["flag"] for e in errors} | {f for e in errors for f in e.get("also_flags", [])}
    for required in ("s8_tenant_share_as_contract_rent", "s8_hap_mismatch_register", "s8_hap_contract_expired",
                     "s8_recert_overdue", "s8_hap_shortfall_unexplained", "s8_hap_abated_but_booked",
                     "s8_reinspection_overdue"):
        assert required in flags, required
    with open(os.path.join(FIX, "ANSWER_KEY.md"), encoding="utf-8") as f:
        md = f.read()
    for e in errors:
        assert f"### {e['id']} " in md and e["flag"] in md, e["id"]
        assert e["evidence"], e["id"]
        assert e["severity"] in ("high", "medium", "low"), e["id"]
    print("✓ test_answer_key_covers_required_error_kinds: PASS")


def test_clean_controls_are_clean():
    for c in KEY["clean_controls"]:
        prop_key = next(k for k, v in gen.PROPS.items() if v["name"] == c["property"])
        unit = c["unit"]
        for rr, spec in gen.RENT_ROLLS.items():
            if spec["prop"] not in (None, prop_key):
                continue
            hits = find_row(spec["file"], unit, c["property"])
            assert len(hits) == 1, (rr, unit)
            row = hits[0]
            contract, tenant, hap = (float(str(row[gen.col(rr, k)] or 0)) for k in ("contract", "tenant", "hap"))
            assert abs(contract - tenant - hap) < 0.005, (rr, unit, contract, tenant, hap)
            # Matching month's register pays exactly the HAP this roll books.
            month = spec["as_of"].strftime("%Y-%m")
            reg = find_row(gen.REGISTERS[month]["file"] + ".csv", unit)
            assert len(reg) == 1 and abs(float(reg[0]["Net Payment"]) - hap) < 0.005, (rr, unit)
        lease = next((s for s in gen.LEASES if s["prop"] == prop_key and s["unit"] == unit), None)
        if lease:
            u = gen.U[(prop_key, unit)]
            assert f"Monthly Rent: {gen.money(u['contract'])}" in pdf_pages(lease["file"])[0], unit
    # No planted error names a clean control unit.
    control_names = {f"{c['property']} {c['unit']}" for c in KEY["clean_controls"]}
    assert len(control_names) >= 3
    for e in KEY["errors"]:
        assert not control_names & set(e["units"]), e["id"]
    print("✓ test_clean_controls_are_clean: PASS")


def test_every_pdf_is_readable_by_the_app_extractor():
    scanned = {(s["file"], s["page"]) for s in KEY["scanned_image_pages"]}
    assert len(scanned) == 2
    pdfs = [r for group in KEY["files"].values() for r in group if r.endswith(".pdf")]
    assert len(pdfs) == 10 + 5 + 3 + 3
    for rel in pdfs:
        pages = pdf_pages(rel)
        assert sum(len(p) for p in pages) > 500, rel
        for i, text in enumerate(pages, start=1):
            if (rel, i) in scanned:
                # Only the synthetic-data footer/header is real text; the body is a raster image.
                body = text.replace(gen.FOOTER, "")
                assert "SIGNATURES" not in body and "/s/" not in body and len(body) < 120, \
                    f"{rel} p{i} should be image-only, got {body!r}"
            else:
                assert len(text) > 80, f"{rel} p{i} has no text layer"
    print(f"✓ test_every_pdf_is_readable_by_the_app_extractor: PASS ({len(pdfs)} PDFs)")


def test_no_real_looking_personal_data():
    texts = []
    for dp, _, fs in os.walk(FIX):
        for f in fs:
            path = os.path.join(dp, f)
            rel = os.path.relpath(path, FIX)
            if f.endswith((".csv", ".md", ".json")):
                with open(path, encoding="utf-8") as fh:
                    texts.append((rel, fh.read()))
            elif f.endswith(".pdf"):
                texts.append((rel, " ".join(pdf_pages(rel))))
            elif f.endswith(".xlsx"):
                wb = load_workbook(path, data_only=True)
                buf = io.StringIO()
                for ws in wb.worksheets:
                    for r in ws.iter_rows(values_only=True):
                        buf.write(" ".join("" if c is None else str(c) for c in r) + "\n")
                texts.append((rel, buf.getvalue()))
    for rel, t in texts:
        assert not re.search(r"\b\d{3}-\d{2}-\d{4}\b", t), f"{rel}: SSN-shaped number"
        for m in re.finditer(r"\(?\b\d{3}\)?[ .-]\d{3}-\d{4}\b", t):
            assert re.search(r"555-01\d\d$", m.group()), f"{rel}: non-fictional phone {m.group()}"
        for m in re.finditer(r"[\w.+-]+@([\w-]+\.)+([a-z]+)", t):
            assert m.group(2) == "test", f"{rel}: non-.test email {m.group()}"
    print(f"✓ test_no_real_looking_personal_data: PASS ({len(texts)} files)")


if __name__ == "__main__":
    test_regenerates_byte_identical()
    test_file_counts_match_the_brief()
    test_every_evidence_item_is_really_there()
    test_answer_key_covers_required_error_kinds()
    test_clean_controls_are_clean()
    test_every_pdf_is_readable_by_the_app_extractor()
    test_no_real_looking_personal_data()
    print("\nAll Section 8 fixture checks passed.")
