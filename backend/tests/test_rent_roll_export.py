"""
Tests for the rent roll CSV/Excel export.

Fixtures are hand-built lease records matching the database layer's
get_all_effective_leases() shape rather than PDFs run through the extractor:
the export's job is to render a given record faithfully, and coupling these
assertions to live extraction would make them fail for reasons that have
nothing to do with this module.

The fixtures deliberately include a lease with several fields missing, a lease
with no end date, and a CAM value carrying a qualifier ("per sq ft annually"),
since those are the three cases where a rent roll most easily says something
false — a "None" cell, a crash, or a number that drops its unit.
"""

import csv
import io
import os
import sys
from datetime import date

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from openpyxl import load_workbook

from app.rent_roll_export import COLUMNS, generate_rent_roll_csv, generate_rent_roll_excel


# Fixed so the Months Until Expiration assertions don't rot as time passes.
TODAY = date(2026, 8, 12)


def _field(value, page=1, confidence="high"):
    if value is None:
        return {"value": None, "source": None, "confidence": None}
    return {"value": value, "source": {"page": page, "quote": f"...{value}..."}, "confidence": confidence}


def _lease(lease_id, filename, **field_values):
    """Builds a full 15-field record, defaulting anything unspecified to not-found."""
    field_names = [
        "tenant", "landlord", "rent_amount", "lease_start_date", "lease_end_date",
        "property_address", "security_deposit", "cam_charges", "rent_escalation",
        "renewal_options", "permitted_use", "exclusivity_clause",
        "insurance_requirements", "default_cure_period", "square_footage",
    ]
    return {
        "id": lease_id,
        "filename": filename,
        "uploaded_at": "2026-08-01T10:00:00Z",
        "document_type": "lease",
        "base_lease_id": None,
        "amendment_count": 0,
        "extracted_fields": {
            name: _field(field_values.get(name)) for name in field_names
        },
    }


def make_leases():
    return [
        _lease(
            1, "blue_sky_coffee.pdf",
            tenant="Blue Sky Coffee Roasters, Inc.",
            landlord="Harborview Properties LLC",
            rent_amount="$6,250.00",
            lease_start_date="April 1, 2025",
            lease_end_date="March 31, 2028",
            property_address="412 Harbor Street, Suite 200, Portland, OR 97204",
            security_deposit="$12,500.00",
            cam_charges="$850.00",
            rent_escalation="3% annually",
            renewal_options="2 option(s) of 5 year(s) each; 180 days notice; renewal rent based on then-prevailing fair market rate",
            permitted_use="Retail coffee shop and cafe, including on-site roasting",
            exclusivity_clause="Landlord shall not lease other space in the shopping center to a coffee shop or cafe",
            insurance_requirements="$1,000,000 commercial general liability per occurrence",
            default_cure_period="10 days after written notice",
            square_footage="2,400 sq ft",
        ),
        # Sparse record: no square footage, deposit, CAM, escalation, or renewal.
        _lease(
            2, "northgate_dental.pdf",
            tenant="Northgate Dental Group, P.C.",
            landlord="Northgate Plaza Associates",
            rent_amount="$4,100.00",
            lease_start_date="June 1, 2024",
            lease_end_date="May 31, 2029",
            property_address="88 Northgate Plaza, Unit C",
        ),
        # No end date at all — Months Until Expiration must be blank, not a crash.
        _lease(
            3, "vertex_consulting.pdf",
            tenant="Vertex Consulting LLC",
            rent_amount="$3,000.00",
            square_footage="1,500 sq ft",
        ),
        # CAM carries a qualifier the bare number would drop.
        _lease(
            4, "anchor_hardware.pdf",
            tenant="Anchor Hardware Supply Co.",
            landlord="Ridgeline Retail Trust",
            rent_amount="$9,800.00",
            lease_start_date="January 6, 2022",
            lease_end_date="January 5, 2027",
            cam_charges="$4.50 per sq ft annually",
            security_deposit="$19,600.00",
            square_footage="7,000 sq ft",
        ),
    ]


def _read_csv(text):
    return list(csv.reader(io.StringIO(text)))


def test_csv_headers_and_row_count():
    rows = _read_csv(generate_rent_roll_csv(make_leases(), today=TODAY))

    assert rows[0] == COLUMNS, f"unexpected header row: {rows[0]}"
    assert rows[0][0] == "Filename" and rows[0][-1] == "Default/Cure Period"
    assert len(rows) == 5, f"expected 1 header + 4 lease rows, got {len(rows)}"
    assert all(len(row) == 18 for row in rows), "every row must have 18 columns"
    print("✓ test_csv_headers_and_row_count: PASS")


def test_csv_cell_values():
    rows = _read_csv(generate_rent_roll_csv(make_leases(), today=TODAY))
    header, first = rows[0], rows[1]
    cell = dict(zip(header, first))

    assert cell["Filename"] == "blue_sky_coffee.pdf"
    assert cell["Tenant"] == "Blue Sky Coffee Roasters, Inc."
    assert cell["Landlord"] == "Harborview Properties LLC"
    # Display strings pass through verbatim, not reformatted.
    assert cell["Monthly Rent"] == "$6,250.00"
    assert cell["Square Footage"] == "2,400 sq ft"
    assert cell["Lease Start"] == "April 1, 2025"
    assert cell["Rent Escalation"] == "3% annually"
    assert cell["Permitted Use"] == "Retail coffee shop and cafe, including on-site roasting"
    assert cell["Exclusivity Clause"] == "Landlord shall not lease other space in the shopping center to a coffee shop or cafe"
    assert cell["Insurance Requirements"] == "$1,000,000 commercial general liability per occurrence"
    assert cell["Default/Cure Period"] == "10 days after written notice"
    print("✓ test_csv_cell_values: PASS")


def test_csv_computed_columns():
    rows = _read_csv(generate_rent_roll_csv(make_leases(), today=TODAY))
    header = rows[0]
    by_file = {row[0]: dict(zip(header, row)) for row in rows[1:]}

    # $6,250.00 / 2,400 sq ft = $2.604... -> "$2.60"
    assert by_file["blue_sky_coffee.pdf"]["Rent/SqFt"] == "$2.60"
    # $9,800.00 / 7,000 sq ft = $1.40
    assert by_file["anchor_hardware.pdf"]["Rent/SqFt"] == "$1.40"
    # No square footage extracted -> blank, not a guessed number.
    assert by_file["northgate_dental.pdf"]["Rent/SqFt"] == ""

    # 2026-08-12 -> 2028-03-31 is 19 whole calendar months.
    assert by_file["blue_sky_coffee.pdf"]["Months Until Expiration"] == "19"
    # 2026-08-12 -> 2027-01-05 is 4 whole months (the 5th falls before the 12th).
    assert by_file["anchor_hardware.pdf"]["Months Until Expiration"] == "4"
    # No end date -> blank.
    assert by_file["vertex_consulting.pdf"]["Months Until Expiration"] == ""
    print("✓ test_csv_computed_columns: PASS")


def test_csv_missing_fields_render_empty_not_none():
    text = generate_rent_roll_csv(make_leases(), today=TODAY)
    rows = _read_csv(text)

    assert "None" not in text, "a missing field leaked the string 'None' into the CSV"
    sparse = dict(zip(rows[0], rows[2]))
    assert sparse["Filename"] == "northgate_dental.pdf"
    for column in ("Square Footage", "Security Deposit", "CAM Charges",
                   "Rent Escalation", "Renewal Options", "Rent/SqFt",
                   "Permitted Use", "Exclusivity Clause",
                   "Insurance Requirements", "Default/Cure Period"):
        assert sparse[column] == "", f"{column} should be an empty cell, got {sparse[column]!r}"
    print("✓ test_csv_missing_fields_render_empty_not_none: PASS")


def test_csv_empty_portfolio_still_has_headers():
    rows = _read_csv(generate_rent_roll_csv([], today=TODAY))
    assert len(rows) == 1 and rows[0] == COLUMNS
    print("✓ test_csv_empty_portfolio_still_has_headers: PASS")


def _load_sheet(leases):
    data = generate_rent_roll_excel(leases, today=TODAY)
    assert isinstance(data, bytes) and data[:2] == b"PK", "xlsx bytes should be a zip container"
    return load_workbook(io.BytesIO(data)).active


def test_excel_structure_and_headers():
    sheet = _load_sheet(make_leases())

    assert sheet.max_row == 5, f"expected 1 header + 4 lease rows, got {sheet.max_row}"
    assert sheet.max_column == 18, f"expected 18 columns, got {sheet.max_column}"
    headers = [sheet.cell(row=1, column=i).value for i in range(1, 19)]
    assert headers == COLUMNS, f"unexpected header row: {headers}"
    assert sheet.freeze_panes == "A2", "header row should be frozen"
    assert sheet.column_dimensions["A"].width and sheet.column_dimensions["R"].width, \
        "column widths should be set explicitly (openpyxl has no auto-fit)"
    print("✓ test_excel_structure_and_headers: PASS")


def test_excel_header_is_bold():
    sheet = _load_sheet(make_leases())
    assert all(sheet.cell(row=1, column=i).font.bold for i in range(1, 19)), \
        "every header cell should be bold"
    assert not sheet.cell(row=2, column=1).font.bold, "data cells should not be bold"
    print("✓ test_excel_header_is_bold: PASS")


def test_excel_currency_typing_and_number_format():
    sheet = _load_sheet(make_leases())
    column_index = {name: i for i, name in enumerate(COLUMNS, start=1)}

    rent = sheet.cell(row=2, column=column_index["Monthly Rent"])
    assert rent.value == 6250.0, f"rent should be a real number for summing, got {rent.value!r}"
    assert rent.number_format == "$#,##0.00", f"got {rent.number_format!r}"

    deposit = sheet.cell(row=2, column=column_index["Security Deposit"])
    assert deposit.value == 12500.0 and deposit.number_format == "$#,##0.00"

    per_sqft = sheet.cell(row=2, column=column_index["Rent/SqFt"])
    assert round(per_sqft.value, 4) == 2.6042 and per_sqft.number_format == "$#,##0.00"

    sqft = sheet.cell(row=2, column=column_index["Square Footage"])
    assert sqft.value == 2400.0 and sqft.number_format == "#,##0"

    months = sheet.cell(row=2, column=column_index["Months Until Expiration"])
    assert months.value == 19

    # A qualified CAM value stays text: "$4.50" alone would misstate it.
    cam = sheet.cell(row=5, column=column_index["CAM Charges"])
    assert cam.value == "$4.50 per sq ft annually", f"got {cam.value!r}"
    print("✓ test_excel_currency_typing_and_number_format: PASS")


def test_excel_missing_fields_are_blank_not_none_text():
    sheet = _load_sheet(make_leases())
    column_index = {name: i for i, name in enumerate(COLUMNS, start=1)}

    for column in ("Square Footage", "Security Deposit", "CAM Charges", "Renewal Options"):
        cell = sheet.cell(row=3, column=column_index[column])
        assert cell.value is None, f"{column} should be an empty cell, got {cell.value!r}"

    all_text = [
        sheet.cell(row=r, column=c).value
        for r in range(1, sheet.max_row + 1) for c in range(1, sheet.max_column + 1)
    ]
    assert "None" not in [v for v in all_text if isinstance(v, str)]
    print("✓ test_excel_missing_fields_are_blank_not_none_text: PASS")


def test_round_trip_integrity_full_row():
    """One complete lease survives both formats with every column intact."""
    leases = make_leases()
    csv_rows = _read_csv(generate_rent_roll_csv(leases, today=TODAY))
    csv_row = dict(zip(csv_rows[0], csv_rows[1]))

    expected_csv = {
        "Filename": "blue_sky_coffee.pdf",
        "Tenant": "Blue Sky Coffee Roasters, Inc.",
        "Landlord": "Harborview Properties LLC",
        "Property Address": "412 Harbor Street, Suite 200, Portland, OR 97204",
        "Square Footage": "2,400 sq ft",
        "Monthly Rent": "$6,250.00",
        "Rent/SqFt": "$2.60",
        "Security Deposit": "$12,500.00",
        "CAM Charges": "$850.00",
        "Rent Escalation": "3% annually",
        "Lease Start": "April 1, 2025",
        "Lease End": "March 31, 2028",
        "Renewal Options": (
            "2 option(s) of 5 year(s) each; 180 days notice; "
            "renewal rent based on then-prevailing fair market rate"
        ),
        "Months Until Expiration": "19",
        "Permitted Use": "Retail coffee shop and cafe, including on-site roasting",
        "Exclusivity Clause": "Landlord shall not lease other space in the shopping center to a coffee shop or cafe",
        "Insurance Requirements": "$1,000,000 commercial general liability per occurrence",
        "Default/Cure Period": "10 days after written notice",
    }
    assert csv_row == expected_csv, f"CSV round-trip mismatch: {csv_row}"

    sheet = _load_sheet(leases)
    xlsx_row = {
        COLUMNS[i - 1]: sheet.cell(row=2, column=i).value for i in range(1, 19)
    }
    # Same row, with the unambiguous numeric columns typed as numbers.
    for column, value in expected_csv.items():
        if column in ("Square Footage", "Monthly Rent", "Rent/SqFt",
                      "Security Deposit", "CAM Charges", "Months Until Expiration"):
            continue
        assert xlsx_row[column] == value, f"{column}: {xlsx_row[column]!r} != {value!r}"
    assert xlsx_row["Monthly Rent"] == 6250.0
    assert xlsx_row["Security Deposit"] == 12500.0
    assert xlsx_row["CAM Charges"] == 850.0
    assert xlsx_row["Square Footage"] == 2400.0
    assert xlsx_row["Months Until Expiration"] == 19
    print("✓ test_round_trip_integrity_full_row: PASS")


def test_excel_empty_portfolio():
    sheet = _load_sheet([])
    assert sheet.max_row == 1, "an empty portfolio should still produce the header row"
    assert [sheet.cell(row=1, column=i).value for i in range(1, 19)] == COLUMNS
    print("✓ test_excel_empty_portfolio: PASS")


if __name__ == "__main__":
    test_csv_headers_and_row_count()
    test_csv_cell_values()
    test_csv_computed_columns()
    test_csv_missing_fields_render_empty_not_none()
    test_csv_empty_portfolio_still_has_headers()
    test_excel_structure_and_headers()
    test_excel_header_is_bold()
    test_excel_currency_typing_and_number_format()
    test_excel_missing_fields_are_blank_not_none_text()
    test_round_trip_integrity_full_row()
    test_excel_empty_portfolio()
    print("\nAll rent roll export tests passed.")
