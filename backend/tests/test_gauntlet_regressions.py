"""
Regression tests for bugs found by the overnight reliability gauntlet
(backend/tools/gauntlet/, branch qa/overnight-gauntlet). One test (or a
small group) per fix; each was confirmed failing before its fix landed.

Pure-function tests where possible (the detectors in deal_mismatch.py take
plain lease dicts), so these run in milliseconds and need no fixtures.
"""
import os
import sys

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from app.portfolio import FIELD_NAMES, _normalize_address, _normalize_building_address
from app import deal_mismatch as dm
from datetime import date

TODAY = date(2026, 10, 5)


def _fields(**overrides):
    out = {}
    for name in FIELD_NAMES:
        v = overrides.get(name)
        out[name] = ({"value": v, "source": {"page": 1, "quote": str(v)}, "confidence": "high"} if v is not None
                     else {"value": None, "source": None, "confidence": None})
    return out


_ids = iter(range(1, 10_000))


def _rr(**f):
    return {"id": next(_ids), "filename": "rent_roll.csv", "extracted_fields": _fields(**f)}


def _doc(**f):
    return {"id": next(_ids), "filename": "lease.pdf", "extracted_fields": _fields(**f)}


def _client_with_fresh_db(team_id=1):
    """Fresh temp DB + an analyst session on a real team (real routes, real scoping)."""
    import tempfile
    from app.api import app
    from app import database, usage_limits
    from _session_users import sync_session_user
    tmp = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
    tmp.close()
    database.configure(tmp.name)
    database.init_db()
    usage_limits._reset_extraction_rate_limit_for_tests()
    client = app.test_client()
    with client.session_transaction() as s:
        s["user_id"] = 7
        s["team_id"] = team_id
        s["email"] = "gauntlet-test@abstractly.test"
        s["name"] = "Gauntlet Test"
        s["role"] = "analyst"
        sync_session_user(s)
    return client, tmp.name


def _post(client, url, filename, data, form=None):
    import io
    payload = dict(form or {})
    payload["file"] = (io.BytesIO(data), filename)
    return client.post(url, data=payload, content_type="multipart/form-data")


def _types(rows):
    return sorted(r["discrepancy_type"] for r in rows)


# ---------------------------------------------------------------- unit matching
def test_same_unit_matches_across_designator_wording():
    """
    The rent-roll importer turns a bare "101" into "<typed address>, Suite 101";
    a multifamily lease says "Apt 101" (or "Unit 101", "#101") with the city
    after it. Same unit -- before the fix every such pair was reported as a
    false unit_no_lease + lease_no_unit and every real check on it was lost.
    """
    rr = "1450 Cedar Bend Dr, Austin, TX 78745, Suite 101"
    for lease in ("1450 Cedar Bend Dr, Apt 101, Austin, TX 78745",
                  "1450 Cedar Bend Dr, Unit 101, Austin, TX 78745",
                  "1450 Cedar Bend Dr #101, Austin, TX 78745",
                  "1450 Cedar Bend Dr, Apartment No. 101, Austin, TX 78745",
                  "1450 Cedar Bend Dr, Apt #101, Austin, TX 78745",
                  "1450 Cedar Bend Drive, Austin, TX 78745, Suite 101",
                  "1450 Cedar Bend Dr, Austin, TX 78745, Suite 0101"):
        assert _normalize_address(rr) == _normalize_address(lease), (lease, _normalize_address(rr), _normalize_address(lease))
    print("✓ test_same_unit_matches_across_designator_wording: PASS")


def test_different_units_and_buildings_still_do_not_match():
    a = "1450 Cedar Bend Dr, Austin, TX 78745, Suite 101"
    for other in ("1450 Cedar Bend Dr, Austin, TX 78745, Suite 102",
                  "1450 Cedar Bend Dr, Austin, TX 78745, Suite 1010",
                  "1452 Cedar Bend Dr, Austin, TX 78745, Suite 101",
                  "1450 Cedar Bend Dr, Austin, TX 78745",
                  "1450 Cedar Bend Ct, Austin, TX 78745, Suite 101"):
        assert _normalize_address(a) != _normalize_address(other), other
    # "Unity" is a street name, not "Unit y".
    assert _normalize_building_address("12 Unity Ave, Suite 3") == _normalize_building_address("12 Unity Ave, Suite 4")
    assert _normalize_building_address("12 Unity Ave") != _normalize_building_address("12 Ave")
    print("✓ test_different_units_and_buildings_still_do_not_match: PASS")


def test_report_pairs_apt_lease_with_suite_rent_roll_row():
    rr = _rr(tenant="Greta Fairbanks", rent_amount="$1,600.00", property_address="1450 Cedar Bend Dr, Austin, TX 78745, Suite 101",
             lease_start_date="04/01/2026", lease_end_date="03/31/2027")
    doc = _doc(tenant="Greta Fairbanks", rent_amount="$1,450.00", property_address="1450 Cedar Bend Dr, Apt 101, Austin, TX 78745",
               lease_start_date="April 1, 2026", lease_end_date="March 31, 2027")
    rows = []
    for det in dm._DETECTORS:
        rows += det([rr, doc], TODAY) if det in dm._DATE_AWARE_DETECTORS else det([rr, doc])
    assert _types(rows) == ["rent_mismatch"], _types(rows)
    print("✓ test_report_pairs_apt_lease_with_suite_rent_roll_row: PASS")


def test_typed_long_street_suffix_matches_abbreviated_lease_building():
    """Uploader typed "Drive", the leases say "Dr": same building for report scoping."""
    assert _normalize_building_address("1450 Cedar Bend Drive, Austin, TX 78745") == \
        _normalize_building_address("1450 Cedar Bend Dr., Apt 101, Austin, TX 78745")
    print("✓ test_typed_long_street_suffix_matches_abbreviated_lease_building: PASS")


# ---------------------------------------------------------------- tenant names
def test_tenant_name_order_and_co_residents_are_not_a_mismatch():
    """
    Yardi/AppFolio export "SORENSEN, CELESTE"; the lease says "Celeste
    Sorensen". A rent roll lists "A & B" where the lease names "A and B" (or
    only one of them). None of these is a different tenant -- before the fix
    each was a high-severity tenant_mismatch.
    """
    same = [("SORENSEN, CELESTE", "Celeste Sorensen"),
            ("Noor Bellweather & Lucia Xenakis", "Noor Bellweather and Lucia Xenakis"),
            ("Noor Bellweather & Lucia Xenakis", "Lucia Xenakis"),
            ("Smith, John Q.", "John Q. Smith"),
            ("John Smith", "John Q. Smith")]
    for rr_name, lease_name in same:
        rr = _rr(tenant=rr_name, rent_amount="$1,000.00", property_address="1 A St, Suite 1")
        doc = _doc(tenant=lease_name, rent_amount="$1,000.00", property_address="1 A St, Apt 1")
        assert dm.detect_tenant_mismatch([rr, doc]) == [], (rr_name, lease_name)
    different = [("Mateo Nakamura", "Malik Nakamura"), ("NAKAMURA, MATEO", "Mateo Ingram"),
                 ("Avery Fairbanks", "Greta Okafor"), ("Jordan Lee & Sam Lee", "Pat Lee")]
    for rr_name, lease_name in different:
        rr = _rr(tenant=rr_name, rent_amount="$1,000.00", property_address="1 A St, Suite 1")
        doc = _doc(tenant=lease_name, rent_amount="$1,000.00", property_address="1 A St, Apt 1")
        assert _types(dm.detect_tenant_mismatch([rr, doc])) == ["tenant_mismatch"], (rr_name, lease_name)
    print("✓ test_tenant_name_order_and_co_residents_are_not_a_mismatch: PASS")


# ---------------------------------------------------------------- CSV decoding
def test_csv_rent_roll_decoding_cp1252_utf16_and_binary():
    """
    Excel on Windows saves CSV as Windows-1252 ("José Peña" became
    "Jos\ufffd Pe\ufffda" and then a false tenant_mismatch); Excel's
    "Unicode Text" is UTF-16 with a BOM; and a PNG renamed .csv crashed the
    import route with a 500 (_csv.Error: line contains NUL).
    """
    from app.rent_roll_import import parse_rent_roll_file, RentRollImportError
    rows = "Unit,Tenant,Rent\r\n101,José Peña,\"$1,200.00\"\r\n102,Zoë Lefèvre,1300\r\n"
    got = parse_rent_roll_file(rows.encode("cp1252"), "rr.csv", "1 A St")
    assert [l["extracted_fields"]["tenant"]["value"] for l in got["leases"]] == ["José Peña", "Zoë Lefèvre"], got["leases"]
    got = parse_rent_roll_file(rows.encode("utf-8"), "rr.csv", "1 A St")
    assert [l["extracted_fields"]["tenant"]["value"] for l in got["leases"]] == ["José Peña", "Zoë Lefèvre"]
    got = parse_rent_roll_file(rows.encode("utf-16"), "rr.csv", "1 A St")
    assert len(got["leases"]) == 2 and got["leases"][0]["extracted_fields"]["tenant"]["value"] == "José Peña"
    png = b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00(\x00\x00\x00(\x08\x02\x00\x00\x00"
    try:
        parse_rent_roll_file(png, "rr.csv", "1 A St")
        raise AssertionError("binary file accepted")
    except RentRollImportError as e:
        assert "text" in str(e).lower(), str(e)
    print("✓ test_csv_rent_roll_decoding_cp1252_utf16_and_binary: PASS")


# ---------------------------------------------------------------- rent-roll rows vs lease documents
def test_rent_roll_rows_are_rent_roll_rows_whatever_the_file_format():
    """
    Whether a record is a rent-roll row or a lease document was decided by
    the file EXTENSION (only .csv/.xlsx counted as rent rolls). A rent roll
    imported from .tsv/.txt/.xls/.docx/.pdf was therefore treated as a pile
    of LEASES by the Deal Mismatch Report: every unit came back as
    lease_no_unit and nothing was compared. The importer now marks its rows.
    """
    from app.portfolio import _is_rent_roll_import
    client, db = _client_with_fresh_db()
    try:
        tsv = "Unit\tTenant\tRent\n101\tAvery Ashgrove\t1200\n102\tJordan Bellweather\t1300\n".encode()
        r = _post(client, "/leases/import-rent-roll", "rent_roll.tsv", tsv, {"property_address": "9 Elm St, Austin, TX 78701"})
        assert r.status_code == 201, r.get_json()
        r = client.post("/portfolio/deal-mismatch-report", data={"property_address": "9 Elm St, Austin, TX 78701"},
                        content_type="multipart/form-data")
        types = sorted(row["discrepancy_type"] for row in r.get_json()["discrepancies"])
        assert types == ["unit_no_lease", "unit_no_lease"], types
    finally:
        os.unlink(db)
    assert _is_rent_roll_import({"filename": "rent_roll.pdf", "source_kind": "rent_roll"})
    assert not _is_rent_roll_import({"filename": "lease_abstract.xlsx", "source_kind": "document"})
    assert _is_rent_roll_import({"filename": "legacy_rent_roll.csv"})  # rows from before the marker existed
    print("✓ test_rent_roll_rows_are_rent_roll_rows_whatever_the_file_format: PASS")


def test_excel_unicode_text_txt_rent_roll_imports():
    """Excel's "Unicode Text" save: .txt, UTF-16, TAB-delimited. The route refused .txt outright."""
    client, db = _client_with_fresh_db()
    try:
        txt = "Unit\tTenant\tRent\r\n101\tAvery Ashgrove\t1,200.00\r\n102\tJosé Peña\t1,300.00\r\n".encode("utf-16")
        r = _post(client, "/leases/import-rent-roll", "rent_roll.txt", txt, {"property_address": "9 Elm St"})
        assert r.status_code == 201, r.get_json()
        body = r.get_json()
        assert body["imported_count"] == 2, body
        names = sorted(l.get("display_name", "") for l in body["leases"])
        assert any("José Peña" in n for n in names), names
    finally:
        os.unlink(db)
    print("✓ test_excel_unicode_text_txt_rent_roll_imports: PASS")


# ---------------------------------------------------------------- T-12 parsing
_T12_MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def _t12_csv(rows, months=None):
    import csv, io
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["Account"] + (months or _T12_MONTHS) + ["Total"])
    for label, monthly in rows:
        if monthly is None:
            w.writerow([label] + [""] * 13)
        else:
            w.writerow([label] + [f"{v:.2f}" for v in monthly] + [f"{sum(monthly):.2f}"])
    return buf.getvalue().encode()


def test_t12_section_header_rows_do_not_hide_the_real_line():
    """
    A T-12's "Rental Income" / "Other Income" SECTION HEADERS (label, no
    numbers) claimed those categories first, so the real "Net Rental
    Income" / "Total Other Income" lines below were never read and the
    T-12 cross-check had no collections figure at all.
    """
    from app.t12_statement import parse_csv_t12_statement
    data = _t12_csv([("INCOME", None), ("Rental Income", None), ("Gross Potential Rent", [10000.0] * 12),
                     ("Vacancy Loss", [-500.0] * 12), ("Net Rental Income", [9500.0] * 12),
                     ("Other Income", None), ("Utility Reimbursement", [300.0] * 12), ("Total Other Income", [300.0] * 12)])
    got = parse_csv_t12_statement(data, "t12.csv")
    assert got["rental_income_collected"] and got["rental_income_collected"]["annual"] == 114000.0, got["rental_income_collected"]
    assert got["other_income"] and got["other_income"]["annual"] == 3600.0, got["other_income"]
    assert got["gross_potential_rent"]["annual"] == 120000.0
    print("✓ test_t12_section_header_rows_do_not_hide_the_real_line: PASS")


def test_t12_month_headers_with_years_and_numeric_months():
    """
    Real T-12 exports label months "Oct 2025", "Oct-25", "10/2025",
    "October 2025", or put real Excel dates in the header row. Only bare
    "Oct"/"October" were recognized, so every one of these failed with
    "Couldn't find any month columns".
    """
    from datetime import datetime
    from io import BytesIO
    from openpyxl import Workbook
    from app.t12_statement import parse_csv_t12_statement, parse_xlsx_t12_statement
    from app.t12_import import parse_csv_t12
    period = [(2025, 10), (2025, 11), (2025, 12)] + [(2026, m) for m in range(1, 10)]
    names = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    longs = ["January", "February", "March", "April", "May", "June", "July", "August", "September",
             "October", "November", "December"]
    styles = {
        "mon_year": [f"{names[m - 1]} {y}" for y, m in period],
        "mon_dash_yy": [f"{names[m - 1]}-{str(y)[2:]}" for y, m in period],
        "mm_yyyy": [f"{m:02d}/{y}" for y, m in period],
        "long_year": [f"{longs[m - 1]} {y}" for y, m in period],
        "yyyy_mm": [f"{y}-{m:02d}" for y, m in period],
    }
    rows = [("Gross Potential Rent", [1000.0] * 12), ("Net Rental Income", [900.0] * 12)]
    for style, months in styles.items():
        got = parse_csv_t12_statement(_t12_csv(rows, months), "t12.csv")
        assert got["gross_potential_rent"] and got["gross_potential_rent"]["annual"] == 12000.0, (style, got["gross_potential_rent"])
        assert sum(v is not None for v in got["gross_potential_rent"]["monthly"].values()) == 12, style
        assert parse_csv_t12(_t12_csv(rows, months), "t12.csv")["annual_rental_income"] == 10800.0, style
    wb = Workbook()
    ws = wb.active
    ws.append(["Account"] + [datetime(y, m, 1) for y, m in period] + ["Total"])
    ws.append(["Gross Potential Rent"] + [1000.0] * 12 + [12000.0])
    buf = BytesIO()
    wb.save(buf)
    got = parse_xlsx_t12_statement(buf.getvalue(), "t12.xlsx")
    assert sum(v is not None for v in got["gross_potential_rent"]["monthly"].values()) == 12, got["gross_potential_rent"]
    # Words that merely contain a month name are still not month columns.
    from app.t12_import import _match_t12_columns
    assert _match_t12_columns(["Account", "Mayfield Rd", "Decor", "Marketing", "Total"])["months"] == {}
    print("✓ test_t12_month_headers_with_years_and_numeric_months: PASS")


def test_t12_blank_row_is_not_a_zero_total():
    """
    With no Total column, the annual figure is the sum of the 12 months --
    and a blank section-header row summed to $0.00. The T-12 cross-check
    route (t12_import) then took "Rental Income" (the header) as $0 of
    collections instead of reading "Net Rental Income" below it.
    """
    import csv, io
    from app.t12_import import parse_csv_t12, _row_annual_total, _match_t12_columns
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["Account"] + _T12_MONTHS)
    w.writerow(["Rental Income"] + [""] * 12)
    w.writerow(["Net Rental Income"] + ["900.00"] * 12)
    assert parse_csv_t12(buf.getvalue().encode(), "t12.csv")["annual_rental_income"] == 10800.0
    mapping = _match_t12_columns(["Account"] + _T12_MONTHS)
    assert _row_annual_total(["Rental Income"] + [""] * 12, mapping) is None
    print("✓ test_t12_blank_row_is_not_a_zero_total: PASS")


# ---------------------------------------------------------------- T-12 detectors
def _t12_cat(annual, monthly=None):
    return {"annual": annual, "monthly": monthly, "source": {"row": 5, "file": "t12.csv", "quote": "x"}}


def test_t12_income_gap_compares_annual_to_annual():
    """
    The rent roll's MONTHLY rent total was compared to the T-12's ANNUAL
    collections, so a rent roll overstating income by 17% was never
    flagged (tester-pack OVERNIGHT_REPORT bug #2).
    """
    leases = [_rr(tenant=f"T{i}", rent_amount="$1,000.00", property_address=f"9 Elm St, Austin, TX 78701, Suite {i}") for i in range(10)]
    t12 = {"rental_income_collected": _t12_cat(100000.0)}
    rows = dm.detect_t12_income_gap(leases, t12, 3.0)
    assert [r["annual_dollar_impact"] for r in rows] == [20000.0], rows
    assert rows[0]["income_direction"] == "overstate"
    # Collections within 3% of the rent roll: no finding.
    assert dm.detect_t12_income_gap(leases, {"rental_income_collected": _t12_cat(118500.0)}, 3.0) == []
    print("✓ test_t12_income_gap_compares_annual_to_annual: PASS")


def test_t12_loss_lines_are_read_as_magnitudes():
    """
    T-12s print vacancy loss / bad debt as NEGATIVE lines ("(17,100.00)").
    The occupancy check computed 1 - (-17,100 / 171,000) = 110% occupancy,
    and the bad-debt trend compared negative averages (never "rising").
    """
    leases = [_rr(tenant=f"T{i}", rent_amount="$1,000.00", property_address=f"9 Elm St, Suite {i}") for i in range(10)]
    rows = dm.detect_t12_occupancy_mismatch(leases, {"gross_potential_rent": _t12_cat(100000.0),
                                                     "vacancy_loss": _t12_cat(-20000.0)}, 3.0)
    assert len(rows) == 1 and rows[0]["lease_value"].startswith("80.0%"), rows
    months = ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"]
    rising = dict(zip(months, [-100.0] * 9 + [-900.0, -1000.0, -1100.0]))
    rows = dm.detect_t12_bad_debt_trend(leases, {"bad_debt": _t12_cat(sum(rising.values()), rising)}, 3.0)
    assert [r["discrepancy_type"] for r in rows] == ["t12_bad_debt_trend"], rows
    print("✓ test_t12_loss_lines_are_read_as_magnitudes: PASS")


def test_t12_zero_concessions_is_not_a_finding():
    """A T-12 whose Concessions line is $0.00 produced a "$0.00 (T12)" concession-gap row."""
    leases = [_rr(tenant="T", rent_amount="$1,000.00", property_address="9 Elm St, Suite 1")]
    assert dm.detect_t12_concession_gap(leases, {"concessions": _t12_cat(0.0)}, 3.0) == []
    rows = dm.detect_t12_concession_gap(leases, {"concessions": _t12_cat(-1800.0)}, 3.0)
    assert len(rows) == 1 and rows[0]["lease_value"].startswith("$1,800.00"), rows
    print("✓ test_t12_zero_concessions_is_not_a_finding: PASS")


def test_t12_bad_debt_trend_uses_the_statements_own_month_order():
    """
    A trailing T-12 runs Oct 2025 .. Sep 2026. "Last 3 months" were taken
    as Oct/Nov/Dec by calendar name -- the OLDEST months of that statement
    -- so a real spike in Jul-Sep was missed and an old spike that had
    since cleared was flagged as "rising".
    """
    from app.t12_statement import parse_csv_t12_statement
    period = ["Oct 2025", "Nov 2025", "Dec 2025"] + [f"{m} 2026" for m in _T12_MONTHS[:9]]
    leases = [_rr(tenant=f"T{i}", rent_amount="$1,000.00", property_address=f"9 Elm St, Suite {i}") for i in range(10)]
    recent_spike = [-100.0] * 9 + [-900.0, -1000.0, -1100.0]
    old_spike = [-1100.0, -1000.0, -900.0] + [-100.0] * 9
    t12 = parse_csv_t12_statement(_t12_csv([("Bad Debt", recent_spike)], period), "t12.csv")
    assert _types(dm.detect_t12_bad_debt_trend(leases, t12, 3.0)) == ["t12_bad_debt_trend"]
    t12 = parse_csv_t12_statement(_t12_csv([("Bad Debt", old_spike)], period), "t12.csv")
    assert dm.detect_t12_bad_debt_trend(leases, t12, 3.0) == []
    print("✓ test_t12_bad_debt_trend_uses_the_statements_own_month_order: PASS")


_GAUNTLET_FIX = os.path.join(os.path.dirname(__file__), "..", "tools", "gauntlet", "fixtures")


def test_t12_workbooks_any_sheet_legacy_xls_and_password():
    """
    T-12 workbooks: only the ACTIVE sheet was read (a "Summary" tab first
    -> "Couldn't find any month columns"); .xls went to openpyxl ("File is
    not a zip file"); a password-protected file got the same cryptic zip
    error instead of being told it's encrypted.
    """
    from io import BytesIO
    from openpyxl import Workbook
    from app.t12_statement import parse_xlsx_t12_statement
    from app.t12_import import parse_xlsx_t12, T12ImportError
    wb = Workbook()
    wb.active.title = "Summary"
    wb.active.append(["NOI", 123456])
    ws = wb.create_sheet("T12 Detail")
    ws.append(["Account"] + _T12_MONTHS + ["Total"])
    ws.append(["Gross Potential Rent"] + [1000.0] * 12 + [12000.0])
    ws.append(["Net Rental Income"] + [900.0] * 12 + [10800.0])
    buf = BytesIO()
    wb.save(buf)
    assert parse_xlsx_t12_statement(buf.getvalue(), "t12.xlsx")["rental_income_collected"]["annual"] == 10800.0
    assert parse_xlsx_t12(buf.getvalue(), "t12.xlsx")["annual_rental_income"] == 10800.0
    with open(os.path.join(_GAUNTLET_FIX, "t12", "p04__xls__clean.xls"), "rb") as f:
        xls = f.read()
    got = parse_xlsx_t12_statement(xls, "t12.xls")
    assert got["gross_potential_rent"] and got["gross_potential_rent"]["annual"] > 0, got
    with open(os.path.join(_GAUNTLET_FIX, "t12", "bad__password.xlsx"), "rb") as f:
        locked = f.read()
    for fn in (parse_xlsx_t12_statement, parse_xlsx_t12):
        try:
            fn(locked, "t12.xlsx")
            raise AssertionError("encrypted workbook parsed")
        except T12ImportError as e:
            assert "password" in str(e).lower(), str(e)
    print("✓ test_t12_workbooks_any_sheet_legacy_xls_and_password: PASS")


def test_password_protected_rent_roll_workbook_says_so():
    """An encrypted rent-roll .xlsx got "Couldn't read this file as an Excel workbook: File is not a zip file"."""
    from app.rent_roll_import import parse_rent_roll_file, RentRollImportError
    with open(os.path.join(_GAUNTLET_FIX, "rent_rolls", "bad__password.xlsx"), "rb") as f:
        locked = f.read()
    for name in ("rr.xlsx", "rr.xls"):
        try:
            parse_rent_roll_file(locked, name, "1 A St")
            raise AssertionError("encrypted workbook parsed")
        except RentRollImportError as e:
            assert "password" in str(e).lower(), str(e)
    print("✓ test_password_protected_rent_roll_workbook_says_so: PASS")


def test_t12_reconciliation_counts_rent_roll_rows_not_lease_pdfs_too():
    """
    compute_t12_reconciliation summed rent from EVERY record at the
    building -- rent-roll rows AND the lease PDFs for the same units -- so
    once leases were uploaded the "rent roll annual rent" doubled.
    """
    from app.portfolio import compute_t12_reconciliation
    rr = [_rr(tenant=f"T{i}", rent_amount="$1,000.00", property_address=f"9 Elm St, Austin, TX 78701, Suite {i}") for i in range(1, 4)]
    docs = [_doc(tenant=f"T{i}", rent_amount="$1,000.00", property_address=f"9 Elm St, Apt {i}, Austin, TX 78701") for i in range(1, 4)]
    got = compute_t12_reconciliation(rr + docs, "9 Elm St, Austin, TX 78701", 36000.0)
    assert got["rent_roll_annual_rent"] == 36000.0, got
    assert got["matched_lease_count"] == 3, got
    print("✓ test_t12_reconciliation_counts_rent_roll_rows_not_lease_pdfs_too: PASS")


# ---------------------------------------------------------------- occupancy (vacant/down units)
def _rr_with_vacancies_csv():
    rows = ["Unit,Tenant,Status,Rent"]
    for i in range(1, 9):
        rows.append(f"10{i},Resident {chr(64 + i)} Person,Current,1000.00")
    rows.append("109,VACANT,Vacant-Unrented,")
    rows.append("110,,Down,")
    rows.append("Total,,,8000.00")
    return ("\r\n".join(rows) + "\r\n").encode()


def test_vacant_and_down_units_count_toward_rent_roll_occupancy():
    """
    Vacant rows were never imported, so the rent roll always looked 100%
    occupied: the T-12 occupancy check false-alarmed on every property with
    a vacancy and the bad-debt check's "full occupancy" condition was always
    true (tester-pack OVERNIGHT_REPORT bug #4). The import now records each
    building's vacant/down units (team-scoped) and the T-12 checks use them.
    """
    import io
    from app.rent_roll_import import parse_rent_roll_file
    parsed = parse_rent_roll_file(_rr_with_vacancies_csv(), "rr.csv", "9 Elm St, Austin, TX 78701")
    summary = parsed["unit_summary"]["9 Elm St, Austin, TX 78701"]
    assert (summary["occupied_units"], sorted(summary["vacant_units"]), sorted(summary["down_units"])) == (8, ["109"], ["110"]), summary
    # Placeholder tenant cells are down units, not tenants; a real tenant
    # whose name merely contains "Model" is still a tenant.
    csv = b"Unit,Tenant,Rent\r\n101,Ann Lee,1000\r\n102,DOWN UNIT,\r\n103,Down - renovation,\r\n104,MODEL,\r\n105,Model Unit Corp,1500\r\n"
    p2 = parse_rent_roll_file(csv, "r.csv", "9 Elm St")
    assert sorted(p2["unit_summary"]["9 Elm St"]["down_units"]) == ["102", "103", "104"], p2["unit_summary"]
    assert [l["extracted_fields"]["tenant"]["value"] for l in p2["leases"]] == ["Ann Lee", "Model Unit Corp"]

    def t12_csv(vacancy_monthly):
        months = ",".join(_T12_MONTHS)
        lines = [f"Account,{months},Total",
                 "Gross Potential Rent," + ",".join(["10000"] * 12) + ",120000",
                 "Vacancy Loss," + ",".join([str(-vacancy_monthly)] * 12) + f",{-vacancy_monthly * 12}",
                 "Net Rental Income," + ",".join(["7900"] * 12) + ",94800"]
        return ("\r\n".join(lines) + "\r\n").encode()

    client, db = _client_with_fresh_db(team_id=1)
    try:
        r = _post(client, "/leases/import-rent-roll", "rr.csv", _rr_with_vacancies_csv(), {"property_address": "9 Elm St, Austin, TX 78701"})
        assert r.status_code == 201, r.get_json()
        for vacancy, want_rows in ((2000, []), (200, ["t12_occupancy_mismatch"])):  # T-12 at 80% (agrees) / 98% (disagrees)
            data = {"property_address": "9 Elm St, Austin, TX 78701", "t12_file": (io.BytesIO(t12_csv(vacancy)), "t12.csv")}
            rep = client.post("/portfolio/deal-mismatch-report", data=data, content_type="multipart/form-data").get_json()
            got = sorted(row["discrepancy_type"] for row in rep["rent_roll_vs_actual_collections"]
                         if row["discrepancy_type"] == "t12_occupancy_mismatch")
            assert got == want_rows, (vacancy, rep["rent_roll_vs_actual_collections"])
            if got:
                assert rep["rent_roll_vs_actual_collections"][0]["rent_roll_value"].startswith("80.0%"), rep
        # Another team's import of the same address never feeds this team's numbers.
        from app.api import app
        from _session_users import sync_session_user
        other = app.test_client()
        with other.session_transaction() as s:
            s.update(user_id=8, team_id=2, email="other@abstractly.test", name="Other", role="analyst")
            sync_session_user(s)
        full = "Unit,Tenant,Rent\r\n" + "".join(f"10{i},Other {chr(64 + i)} Resident,1000\r\n" for i in range(1, 9))
        _post(other, "/leases/import-rent-roll", "rr2.csv", full.encode(), {"property_address": "9 Elm St, Austin, TX 78701"})
        data = {"property_address": "9 Elm St, Austin, TX 78701", "t12_file": (io.BytesIO(t12_csv(2000)), "t12.csv")}
        rep = client.post("/portfolio/deal-mismatch-report", data=data, content_type="multipart/form-data").get_json()
        assert not [x for x in rep["rent_roll_vs_actual_collections"] if x["discrepancy_type"] == "t12_occupancy_mismatch"], rep
    finally:
        os.unlink(db)
    print("✓ test_vacant_and_down_units_count_toward_rent_roll_occupancy: PASS")


# ---------------------------------------------------------------- rent-roll headers & rows
def _parse_csv_text(text, base="9 Elm St"):
    from app.rent_roll_import import parse_rent_roll_file
    return parse_rent_roll_file(text.replace("\n", "\r\n").encode(), "rr.csv", base)


def _units_of(parsed):
    from app.portfolio import _normalize_address
    return [(l["extracted_fields"]["property_address"]["value"], l["extracted_fields"]["tenant"]["value"]) for l in parsed["leases"]]


def test_unit_column_headers_from_pms_exports():
    """Entrata's "Bldg-Unit" (and "Bldg/Unit", "Apt #") weren't recognized, so every row lost its unit number."""
    for header in ("Bldg-Unit", "Bldg/Unit", "Building/Unit", "Apt #", "Apartment", "Unit ID"):
        got = _parse_csv_text(f"{header},Resident,Scheduled Rent\nA-101,Ann Lee,1000\n")
        assert _units_of(got) == [("9 Elm St, Suite A-101", "Ann Lee")], (header, _units_of(got))
    print("✓ test_unit_column_headers_from_pms_exports: PASS")


def test_grand_total_row_is_not_a_tenant():
    """A broker sheet's "Grand Total" row (in the tenant column) was imported as a tenant paying the whole building's rent."""
    got = _parse_csv_text("Unit,Tenant,Rent\n101,Ann Lee,1000\n,Grand Total,1000\n,Grand Totals:,1000\n")
    assert [t for _, t in _units_of(got)] == ["Ann Lee"], _units_of(got)
    print("✓ test_grand_total_row_is_not_a_tenant: PASS")


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_") and callable(fn):
            fn()
    print("\nAll gauntlet regression tests passed.")
