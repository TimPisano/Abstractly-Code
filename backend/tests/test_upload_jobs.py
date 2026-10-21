"""
End-to-end tests for background large-file uploads (app/upload_jobs.py +
POST /uploads, GET /uploads/<id>, GET /uploads/config), with
LARGE_FILE_UPLOADS on. Uses the full-size fixtures from
create_large_fixtures.py. The Anthropic API is mocked (a fake client that
only reports what is in the text it was sent).

Most tests run the job inline (start_job patched to call run_upload_job
directly) so they are deterministic; test_thread_runner_* exercises the
real background thread, and test_rq_runner_* the Redis queue hand-off.
"""
import _test_env  # noqa: F401  -- first: temp DB, no .env, no real network (see _test_env.py)
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import time
import tracemalloc
from datetime import datetime, timedelta, timezone
from unittest import mock

# Must be set before app.api is imported: it sizes Flask's upload limit at import.
os.environ["LARGE_FILE_UPLOADS"] = "true"
os.environ.pop("REDIS_URL", None)
_UPLOAD_DIR = tempfile.mkdtemp(prefix="abstractly-test-uploads-")
os.environ["LARGE_FILE_UPLOAD_DIR"] = _UPLOAD_DIR

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from _session_users import sync_session_user  # noqa: E402
from create_large_fixtures import (  # noqa: E402
    RENT_ROLL_CSV, RENT_ROLL_XLSX, SCANNED_LEASE, TEXT_LEASE, ensure_large_fixtures, load_expected, path_of,
)
from test_ai_chunked_extraction import FaithfulReader  # noqa: E402

from app import ai_extraction, database, upload_jobs, usage_limits  # noqa: E402
from app.api import app  # noqa: E402
from app.pdf_extractor import PDFExtractor  # noqa: E402

TESTS_DIR = os.path.dirname(os.path.abspath(__file__))


def _fresh_db():
    tmp = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
    tmp.close()
    database.configure(tmp.name)
    database.init_db()
    usage_limits._reset_extraction_rate_limit_for_tests()
    return tmp.name


def _client(user_id=1, team_id=1, role="analyst"):
    c = app.test_client()
    with c.session_transaction() as s:
        s.update({"user_id": user_id, "email": f"u{user_id}@example.com", "name": f"U{user_id}", "role": role,
                  "team_id": team_id})
        sync_session_user(s)
    return c


def _inline(job_id):
    upload_jobs.run_upload_job(job_id)


def _post(client, path=None, filename=None, body=None, kind="lease", **form):
    if body is None:
        with open(path, "rb") as f:
            body = f.read()
    data = {"file": (io.BytesIO(body), filename or os.path.basename(path)), "kind": kind, **form}
    return client.post("/uploads", data=data, content_type="multipart/form-data")


def _upload_and_wait(client, *args, **kwargs):
    """POST, then run the queued job after the 202 has been returned -- the same order the real runners give."""
    queued = []
    with mock.patch.object(upload_jobs, "start_job", side_effect=queued.append):
        resp = _post(client, *args, **kwargs)
    assert resp.status_code == 202, (resp.status_code, resp.get_json())
    job_id = resp.get_json()["upload"]["id"]
    assert queued == [job_id]
    _inline(job_id)
    final = client.get(f"/uploads/{job_id}")
    assert final.status_code == 200
    return resp.get_json()["upload"], final.get_json()["upload"]


def _stored_files():
    return [n for n in os.listdir(_UPLOAD_DIR) if n.endswith(".upload")]


class _StageRecorder:
    def __init__(self):
        self.updates = []
        self._real = database.update_upload_job

    def __call__(self, job_id, **fields):
        self.updates.append(fields)
        return self._real(job_id, **fields)


# ---------------------------------------------------------------------------


def test_flag_off_means_routes_404_and_old_limits_unchanged():
    script = (
        "import os, sys, io; sys.path.insert(0, %r); os.environ.pop('LARGE_FILE_UPLOADS', None)\n"
        "from app.api import app, LEGACY_MAX_CONTENT_LENGTH\n"
        "assert app.config['MAX_CONTENT_LENGTH'] == 16 * 1024 * 1024 == LEGACY_MAX_CONTENT_LENGTH\n"
        "from app import upload_jobs; assert not upload_jobs.enabled()\n"
        "print('ok')\n" % os.path.abspath(os.path.join(TESTS_DIR, ".."))
    )
    env = {k: v for k, v in os.environ.items() if k != "LARGE_FILE_UPLOADS"}
    out = subprocess.run([sys.executable, "-c", script], capture_output=True, text=True, env=env, cwd=TESTS_DIR)
    assert out.returncode == 0 and "ok" in out.stdout, out.stderr[-2000:]

    db = _fresh_db()
    try:
        with mock.patch.object(upload_jobs, "enabled", return_value=False):
            c = _client()
            assert _post(c, filename="a.txt", body=b"hello").status_code == 404
            assert c.get("/uploads/1").status_code == 404
            assert c.get("/uploads/config").get_json()["large_file_uploads"] is False
    finally:
        os.unlink(db)
    print("✓ test_flag_off_means_routes_404_and_old_limits_unchanged: PASS")


def test_routes_require_login_and_role():
    db = _fresh_db()
    try:
        anon = app.test_client()
        assert anon.post("/uploads", data={}).status_code == 401
        assert anon.get("/uploads/1").status_code == 401
        assert anon.get("/uploads/config").status_code == 401
        viewer = _client(user_id=5, role="viewer")
        assert _post(viewer, filename="a.txt", body=b"hello").status_code == 403
        assert viewer.get("/uploads/config").status_code == 200
    finally:
        os.unlink(db)
    print("✓ test_routes_require_login_and_role: PASS")


def test_500_page_text_lease_regex_engine_end_to_end_with_progress():
    ensure_large_fixtures(only=[TEXT_LEASE])
    expected = load_expected()["text_lease"]
    db = _fresh_db()
    try:
        recorder = _StageRecorder()
        with mock.patch.object(database, "update_upload_job", recorder):
            accepted, done = _upload_and_wait(_client(), path_of(TEXT_LEASE))
        assert accepted["status"] == "queued" and accepted["detected_type"] == "PDF"
        assert done["status"] == "complete", done
        result = done["result"]
        assert result["page_count"] == 500 and result["split_count"] == 1
        lease = result["leases"][0]
        fields = lease["extracted_fields"]
        assert fields["tenant"]["value"] == expected["tenant"]
        assert fields["security_deposit"]["value"] == "$56,250.00"
        assert fields["security_deposit"]["source"]["page"] == 262
        assert lease["source_page_start"] == 1 and lease["source_page_end"] == 500

        reading = [u for u in recorder.updates if u.get("stage") == "reading" and u.get("progress_total")]
        assert reading[-1]["progress_current"] == 500 and reading[-1]["progress_total"] == 500
        assert len(reading) >= 2, "progress is reported while pages are read, not just at the end"
        stages = [u.get("stage") for u in recorder.updates if u.get("stage")]
        assert stages.index("reading") < stages.index("extracting") < stages.index("complete")

        assert _stored_files() == [], "stored upload deleted after the job"
        row = database.get_lease(lease["id"], team_id=1)
        assert row["content_hash"] and row["processing_status"] == "complete"
    finally:
        os.unlink(db)
    print("✓ test_500_page_text_lease_regex_engine_end_to_end_with_progress: PASS")


def test_500_page_lease_ai_engine_quota_checked_before_any_ai_call():
    ensure_large_fixtures(only=[TEXT_LEASE])
    db = _fresh_db()
    try:
        order = []
        reader = FaithfulReader()
        real_quota = usage_limits.check_team_quota

        def quota(team_id):
            order.append("quota")
            return real_quota(team_id)

        real_create = reader.create

        def create(**kwargs):
            order.append("ai")
            return real_create(**kwargs)

        reader.create = create
        with mock.patch.object(ai_extraction, "resolve_engine", return_value="ai"), \
             mock.patch.object(ai_extraction, "_make_client", return_value=reader), \
             mock.patch.object(usage_limits, "check_team_quota", side_effect=quota):
            _, done = _upload_and_wait(_client(), path_of(TEXT_LEASE))
        assert done["status"] == "complete", done
        chunks = len(ai_extraction.plan_chunks([{"page": n, "text": "x" * 2850} for n in range(1, 501)]))
        assert order.count("quota") == 2, order  # at upload time and again right before extraction
        assert order[1] == "quota" and order[2] == "ai", order
        assert order.count("ai") >= chunks - 1
        fields = done["result"]["leases"][0]["extracted_fields"]
        assert fields["renewal_options"]["source"]["page"] == 431
        assert fields["exclusivity_clause"]["source"]["page"] == 478
        assert database.get_connection().execute(
            "SELECT COUNT(*) FROM ai_extraction_runs WHERE lease_id = ?", (done["result"]["lease_ids"][0],)
        ).fetchone()[0] == 1, "one telemetry row for the whole chunked extraction"
    finally:
        os.unlink(db)
    print("✓ test_500_page_lease_ai_engine_quota_checked_before_any_ai_call: PASS")


def test_over_quota_fails_the_job_without_calling_the_model():
    ensure_large_fixtures(only=[TEXT_LEASE])
    db = _fresh_db()
    try:
        reader = FaithfulReader()
        answers = iter([None, "Monthly document quota reached (10 documents). Contact your admin to raise it."])
        with mock.patch.object(ai_extraction, "resolve_engine", return_value="ai"), \
             mock.patch.object(ai_extraction, "_make_client", return_value=reader), \
             mock.patch.object(usage_limits, "check_team_quota", side_effect=lambda team_id: next(answers)):
            _, done = _upload_and_wait(_client(), path_of(TEXT_LEASE))
        assert done["status"] == "failed" and "quota" in done["error"].lower(), done
        assert reader.prompts == [], "no AI call once the quota is used up"
        assert database.get_all_effective_leases(team_id=1) == []
        assert _stored_files() == []

        with mock.patch.object(usage_limits, "check_team_quota", return_value="Monthly document quota reached."):
            resp = _post(_client(), path_of(TEXT_LEASE))
        assert resp.status_code == 403 and _stored_files() == [], "rejected before the file is even stored"
    finally:
        os.unlink(db)
    print("✓ test_over_quota_fails_the_job_without_calling_the_model: PASS")


def test_ai_failure_shows_a_friendly_message_and_logs_the_details_with_the_team():
    """Review follow-up: the AI error's own text (billing state, configuration) went to the user as the job error."""
    ensure_large_fixtures(only=[TEXT_LEASE])
    db = _fresh_db()
    try:
        raw = ("AI processing is unavailable: the API credit balance is too low. "
               "Add credit (Plans & Billing) and try again.")
        cause = RuntimeError("request_id=req_test123 status=400")
        error = ai_extraction.AIExtractionError(raw)
        error.__cause__ = cause
        with mock.patch.object(ai_extraction, "resolve_engine", return_value="ai"), \
             mock.patch.object(ai_extraction, "extract_lease_fields_chunked", side_effect=error), \
             mock.patch.object(upload_jobs.logger, "warning") as warning:
            _, done = _upload_and_wait(_client(team_id=1), path_of(TEXT_LEASE))
        assert done["status"] == "failed" and done["error"] == upload_jobs.AI_FAILURE, done
        poll = json.dumps(done)
        assert "credit" not in poll and "req_test123" not in poll, "nothing of the AI error reaches the browser"
        logged = [call.args[0] % call.args[1:] for call in warning.call_args_list]
        assert any("team 1" in line and raw in line and "req_test123" in line for line in logged), logged
        assert database.get_all_effective_leases(team_id=1) == [] and _stored_files() == []
    finally:
        os.unlink(db)
    print("✓ test_ai_failure_shows_a_friendly_message_and_logs_the_details_with_the_team: PASS")


def test_120mb_scanned_pdf_is_streamed_not_held_in_memory():
    """The repro's scenario A (413) and B (OCR in-request, all pages at once). OCR itself is stubbed: real OCR of 300 pages is ~10 min (see run_large_file_scale_check.py)."""
    ensure_large_fixtures(only=[SCANNED_LEASE])
    size = os.path.getsize(path_of(SCANNED_LEASE))
    assert size > 100 * 1024 * 1024
    db = _fresh_db()
    try:
        with open(path_of(SCANNED_LEASE), "rb") as f:
            body = f.read()
        client = _client()
        jobs_seen = []
        tracemalloc.start()
        with mock.patch.object(upload_jobs, "start_job", side_effect=jobs_seen.append):
            resp = client.post("/uploads", data={"file": (io.BytesIO(body), "scan.pdf"), "kind": "lease"},
                               content_type="multipart/form-data")
        _, peak = tracemalloc.get_traced_memory()
        tracemalloc.stop()
        assert resp.status_code == 202, resp.get_json()
        # The test client holds `body` (traced before start), so the request's own
        # allocations are what's measured: a few MB of buffers, not the file.
        assert peak < 20 * 1024 * 1024, f"request allocated {peak / 1e6:.0f} MB for a {size / 1e6:.0f} MB file"
        assert len(_stored_files()) == 1 and os.path.getsize(os.path.join(_UPLOAD_DIR, _stored_files()[0])) == size
        del body

        ocr_calls = []

        def fake_ocr(self, pdf_path, page_num, dpi, timeout_s, rasterize_timeout_s):
            ocr_calls.append(page_num)
            text = ("This Commercial Lease is entered into by and between Harborline Commons LLC (\"Landlord\") and "
                    "Northwind Outfitters, Inc. (\"Tenant\"). Base Rent of $18,750.00 per month."
                    if page_num == 1 else f"Section {page_num}. General provisions continue on this page. " * 3)
            return text, 88.0

        with mock.patch.object(PDFExtractor, "_ocr_single_pdf_page", fake_ocr):
            upload_jobs.run_upload_job(jobs_seen[0])
        done = client.get(f"/uploads/{jobs_seen[0]}").get_json()["upload"]
        assert done["status"] == "complete", done
        assert ocr_calls == list(range(1, 301)), "every image-only page OCR'd, one at a time, in order"
        assert done["result"]["page_count"] == 300 and done["result"]["ocr_page_count"] == 300
        assert done["result"]["leases"][0]["extracted_fields"]["tenant"]["value"] == "Northwind Outfitters, Inc."
        assert _stored_files() == []
    finally:
        os.unlink(db)
    print("✓ test_120mb_scanned_pdf_is_streamed_not_held_in_memory: PASS")


def test_size_limits_are_friendly_and_leave_nothing_behind():
    db = _fresh_db()
    try:
        client = _client()
        # Werkzeug's per-request cap on POST /uploads (normally MAX_UPLOAD_MB + 1 MB of form overhead).
        with mock.patch.object(upload_jobs, "MAX_REQUEST_BYTES", 2 * 1024 * 1024):
            resp = _post(client, filename="big.pdf", body=b"%PDF-1.7\n" + b"0" * (3 * 1024 * 1024))
        assert resp.status_code == 413 and "The limit is 150 MB per file" in resp.get_json()["error"], resp.get_json()
        # save_upload's own guard, if a body slips past the header check (e.g. no Content-Length).
        with mock.patch.object(upload_jobs, "MAX_UPLOAD_BYTES", 1024 * 1024):
            resp = _post(client, filename="big.pdf", body=b"%PDF-1.7\n" + b"0" * (2 * 1024 * 1024))
        assert resp.status_code == 413 and "MB" in resp.get_json()["error"]
        # Formats parsed in memory have a lower cap; those read whole with no compression, lower still.
        with mock.patch.object(upload_jobs, "MAX_IN_MEMORY_DOCUMENT_MB", 1):
            resp = _post(client, filename="lease.docx", body=_docx_over_1mb())
        assert resp.status_code == 413 and "for a Word document" in resp.get_json()["error"], resp.get_json()
        with mock.patch.object(upload_jobs, "MAX_FLAT_DOCUMENT_MB", 1):
            resp = _post(client, filename="roll.csv", body=b"Unit,Tenant,Rent\n" + b"101,Acme,1200\n" * 100_000,
                         kind="rent_roll")
        assert resp.status_code == 413 and "for a CSV file" in resp.get_json()["error"], resp.get_json()
        # Review follow-up, at the real default: a 50 MB .txt took 135 MB just to read.
        resp = _post(client, filename="lease.txt", body=b"The Tenant shall pay rent monthly.\n" * 320_000)
        assert resp.status_code == 413 and "The limit is 10 MB for a text file" in resp.get_json()["error"], resp.get_json()
        assert _stored_files() == []

        # Every OTHER route keeps exactly 16 MB, even with the flag on.
        big = b"%PDF-1.7\n" + b"0" * (17 * 1024 * 1024)
        resp = client.post("/leases", data={"file": (io.BytesIO(big), "big.pdf")}, content_type="multipart/form-data")
        assert resp.status_code == 413 and resp.get_json()["error"] == "File too large. Maximum size is 16MB."
        resp = client.post("/leases/import-rent-roll", data={"file": (io.BytesIO(big), "big.csv")},
                           content_type="multipart/form-data")
        assert resp.status_code == 413
    finally:
        os.unlink(db)
    print("✓ test_size_limits_are_friendly_and_leave_nothing_behind: PASS")


def _docx_over_1mb():
    """A real .docx that's over 1 MB on disk (an embedded binary part that won't compress)."""
    import zipfile
    from docx import Document
    buf = io.BytesIO()
    Document().save(buf)
    with zipfile.ZipFile(buf, "a") as z:
        z.writestr("word/embeddings/oleObject1.bin", b"\xd0\xcf\x11\xe0" + os.urandom(1536 * 1024))
    return buf.getvalue()


def _rent_roll_docx(rows):
    from docx import Document
    buf = io.BytesIO()
    doc = Document()
    table = doc.add_table(rows=0, cols=len(rows[0]))
    for row in rows:
        for cell, value in zip(table.add_row().cells, row):
            cell.text = str(value)
    doc.save(buf)
    return buf.getvalue()


def test_tables_over_the_row_limit_are_refused_before_parsing():
    """
    Review follow-up: every parsed rent-roll row costs ~5.5 KB whatever the
    format (a 3 MB CSV of 50,000 units took 290 MB), so rows are counted --
    without building the table -- and a file over MAX_TABLE_ROWS is refused
    before the parser runs.
    """
    import openpyxl
    import xlrd
    db = _fresh_db()
    try:
        client = _client()
        rows = [["Unit", "Tenant", "Monthly Rent"]] + [[f"{100 + i}", f"Tenant {i}", 1200 + i] for i in range(60)]
        wb = openpyxl.Workbook()
        for row in rows:
            wb.active.append(row)
        for r in range(100, 400):  # formatted but empty rows aren't rows of data
            wb.active.cell(row=r, column=1).number_format = "0.00"
        xlsx = io.BytesIO()
        wb.save(xlsx)
        files = {
            "roll.csv": "".join(",".join(map(str, row)) + "\n" for row in rows).encode(),
            "roll.xlsx": xlsx.getvalue(),
            "roll.docx": _rent_roll_docx(rows),
        }
        for name, body in files.items():
            path = os.path.join(tempfile.mkdtemp(), name)
            with open(path, "wb") as f:
                f.write(body)
            assert upload_jobs._table_row_count(path, name.rsplit(".", 1)[1]) == 61, name
            shutil.rmtree(os.path.dirname(path))
        xls = os.path.join(TESTS_DIR, "..", "tools", "gauntlet", "fixtures", "rent_rolls", "p05__xls.xls")
        sheet = xlrd.open_workbook(xls).sheet_by_index(0)
        expected = sum(1 for r in range(sheet.nrows) if any(str(v).strip() for v in sheet.row_values(r)))
        assert upload_jobs._table_row_count(xls, "xls") == expected > 0

        with mock.patch.object(upload_jobs, "MAX_TABLE_ROWS", 50), \
             mock.patch("app.rent_roll_import.parse_rent_roll_file") as parse:
            for name, body in files.items():
                _, done = _upload_and_wait(client, filename=name, body=body, kind="rent_roll")
                assert done["status"] == "failed", (name, done)
                assert "61 rows. The limit is 50 rows per file" in done["error"], (name, done["error"])
            parse.assert_not_called()
            _, done = _upload_and_wait(client, filename="leases.csv", body=files["roll.csv"])
            assert done["status"] == "failed" and "61 rows" in done["error"], "a spreadsheet uploaded as a lease too"
        with mock.patch.object(upload_jobs, "MAX_TABLE_ROWS", 61):
            for name, body in files.items():
                _, done = _upload_and_wait(client, filename=name, body=body, kind="rent_roll")
                assert done["status"] == "complete" and done["result"]["imported_count"] == 60, (name, done)
        assert _stored_files() == []
    finally:
        os.unlink(db)
    print("✓ test_tables_over_the_row_limit_are_refused_before_parsing: PASS")


def test_the_largest_office_files_accepted_convert_well_under_512mb():
    """
    Review follow-up, measured: the densest Word document the upload
    accepts, and the largest .docx rent roll it accepts (the
    costliest format per row; its XML cap binds before the row cap), each
    converted in a fresh process. Peak memory above the process's starting
    point must stay under 200 MB.
    """
    script = r"""
import io, os, resource, subprocess, sys, tempfile
sys.path.insert(0, %(backend)r)
from app import document_extractor, file_types, rent_roll_import, upload_jobs

def current_mb():
    if sys.platform.startswith("linux"):
        with open("/proc/self/statm") as f:
            return int(f.read().split()[1]) * os.sysconf("SC_PAGE_SIZE") / 2**20
    return int(subprocess.check_output(["ps", "-o", "rss=", "-p", str(os.getpid())])) / 1024

def peak_mb():
    peak = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
    return peak / 2**20 if sys.platform == "darwin" else peak / 1024

import zipfile
from docx import Document
blank = io.BytesIO()
Document().save(blank)
with zipfile.ZipFile(blank) as z:  # the template's styles etc. count toward the cap too
    room = file_types.MAX_ZIP_XML_BYTES["docx"] - 8192 - sum(
        p.file_size for p in file_types.xml_parts(z) if p.filename != "word/document.xml")

def docx(body):
    out = io.BytesIO()
    with zipfile.ZipFile(blank) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            dst.writestr(item, body if item.filename == "word/document.xml" else src.read(item.filename))
    return out.getvalue()

W = b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
# The densest lease accepted: empty paragraphs up to the element cap (the template has ~19,000 elements).
lease = docx(W + b"<w:p/>" * (file_types.MAX_DOCX_ELEMENTS - 25000)
             + b"<w:p><w:r><w:t>The Tenant shall pay rent.</w:t></w:r></w:p></w:body></w:document>")
cell = lambda text: b"<w:tc><w:p><w:r><w:t>" + text.encode() + b"</w:t></w:r></w:p></w:tc>"
header = ["Unit", "Tenant", "Lease Start", "Lease End", "Monthly Rent", "Status", "Sq Ft"]
row = lambda i: b"<w:tr>" + b"".join(cell(v) for v in (str(100000 + i), "Tenant Name %%d" %% (100000 + i), "01/01/2025",
                                                       "12/31/2025", str(1200 + i %% 500), "Occupied", "850")) + b"</w:tr>"
fits = room // len(row(0))
units = min(upload_jobs.MAX_TABLE_ROWS - 1, fits)
rows = [b"<w:tr>" + b"".join(cell(h) for h in header) + b"</w:tr>"] + [row(i) for i in range(units)]
roll = docx(W + b"<w:tbl>" + b"".join(rows) + b"</w:tbl></w:body></w:document>")
del rows
print("expected", units)

tmp = tempfile.mkdtemp()
for name, data, kind in (("lease.docx", lease, "lease"), ("roll.docx", roll, "rent_roll")):
    path = os.path.join(tmp, name)
    with open(path, "wb") as f:
        f.write(data)
    assert file_types.detect_file_type(path, name, kind).extension == "docx", name
    assert upload_jobs._table_row_count(path, "docx") <= upload_jobs.MAX_TABLE_ROWS
del lease, roll

base = current_mb()
document_extractor.extract_pages_from_path(os.path.join(tmp, "lease.docx"), "docx")
print("lease", round(peak_mb() - base))
with open(os.path.join(tmp, "roll.docx"), "rb") as f:
    units = len(rent_roll_import.parse_rent_roll_file(f.read(), "roll.docx", None)["leases"])
print("roll", round(peak_mb() - base), units)
""" % {"backend": os.path.abspath(os.path.join(TESTS_DIR, ".."))}
    out = subprocess.run([sys.executable, "-c", script], capture_output=True, text=True, env=dict(os.environ),
                         cwd=TESTS_DIR, timeout=600)
    assert out.returncode == 0, out.stderr[-3000:]
    lines = dict(line.split(" ", 1) for line in out.stdout.strip().splitlines())
    lease_mb, (roll_mb, units) = int(lines["lease"]), map(int, lines["roll"].split())
    assert units == int(lines["expected"]) > 5000, (units, lines)
    assert lease_mb < 200 and roll_mb < 200, f"peak above start: lease {lease_mb} MB, rent roll {roll_mb} MB"
    print(f"✓ test_the_largest_office_files_accepted_convert_well_under_512mb: PASS "
          f"(lease .docx +{lease_mb} MB, {units:,}-unit .docx rent roll +{roll_mb} MB)")


def test_page_limit_is_checked_before_any_ocr():
    import pypdf
    db = _fresh_db()
    try:
        writer = pypdf.PdfWriter()
        for _ in range(upload_jobs.MAX_PAGES + 1):
            writer.add_blank_page(612, 792)
        buf = io.BytesIO()
        writer.write(buf)
        with mock.patch.object(PDFExtractor, "_ocr_single_pdf_page", side_effect=AssertionError("OCR must not start")):
            _, done = _upload_and_wait(_client(), filename="huge.pdf", body=buf.getvalue())
        assert done["status"] == "failed" and "1,001 pages" in done["error"] and "1,000" in done["error"], done
    finally:
        os.unlink(db)
    print("✓ test_page_limit_is_checked_before_any_ocr: PASS")


def test_wrong_content_is_rejected_and_mislabeled_files_are_read_as_what_they_are():
    db = _fresh_db()
    try:
        client = _client()
        from test_file_types import _pe_bytes
        resp = _post(client, filename="tool.pdf", body=_pe_bytes())
        assert resp.status_code == 415 and "a program" in resp.get_json()["error"], resp.get_json()
        with open(os.path.join(TESTS_DIR, "multiformat_lease.doc"), "rb") as f:
            resp = _post(client, filename="roll.doc", body=f.read(), kind="rent_roll")
        assert resp.status_code == 415 and ".docx" in resp.get_json()["error"]
        assert _stored_files() == []

        with open(os.path.join(TESTS_DIR, "multiformat_lease.png"), "rb") as f:
            png = f.read()
        _, done = _upload_and_wait(client, filename="lease.pdf", body=png)
        assert done["status"] == "complete", done
        assert done["detected_type"] == "PNG image"
        assert any("actually a PNG image" in w for w in done["result"]["warnings"]), done["result"]["warnings"]
        assert done["result"]["leases"][0]["filename"] == "lease.pdf", "the user's own filename is kept for display"
    finally:
        os.unlink(db)
    print("✓ test_wrong_content_is_rejected_and_mislabeled_files_are_read_as_what_they_are: PASS")


def test_every_accepted_format_goes_through_the_one_pipeline():
    import pillow_heif
    from PIL import Image
    pillow_heif.register_heif_opener()
    db = _fresh_db()
    tmpdir = tempfile.mkdtemp()
    try:
        heic = os.path.join(tmpdir, "lease.heic")
        Image.open(os.path.join(TESTS_DIR, "multiformat_lease.png")).convert("RGB").save(heic, "HEIF")
        files = [os.path.join(TESTS_DIR, f"multiformat_lease.{ext}")
                 for ext in ("docx", "doc", "xlsx", "xls", "csv", "txt", "jpg", "png", "tiff")]
        files += [heic, os.path.join(TESTS_DIR, "sample_lease.pdf")]
        client = _client()
        for path in files:
            usage_limits._reset_extraction_rate_limit_for_tests()  # 11 uploads in a row; the 10/min limit is real
            _, done = _upload_and_wait(client, path)
            assert done["status"] == "complete", (os.path.basename(path), done)
            fields = done["result"]["leases"][0]["extracted_fields"]
            assert any((fields.get(k) or {}).get("value") for k in ("tenant", "landlord", "rent_amount")), \
                (os.path.basename(path), {k: (v or {}).get("value") for k, v in fields.items()})
    finally:
        shutil.rmtree(tmpdir)
        os.unlink(db)
    print("✓ test_every_accepted_format_goes_through_the_one_pipeline: PASS")


def test_5000_row_messy_rent_rolls_xlsx_and_csv():
    ensure_large_fixtures(only=[RENT_ROLL_XLSX, RENT_ROLL_CSV])
    expected = load_expected()["rent_roll"]
    db = _fresh_db()
    try:
        client = _client()
        timings = {}
        for name in (RENT_ROLL_XLSX, RENT_ROLL_CSV):
            started = time.perf_counter()
            accepted, done = _upload_and_wait(client, path_of(name), kind="rent_roll",
                                              property_address="1000 Juniper Row")
            timings[name] = time.perf_counter() - started
            assert done["status"] == "complete", done
            result = done["result"]
            assert result["imported_count"] == expected["occupied_rows"], (name, result["imported_count"])
            assert len(result["leases"]) == expected["occupied_rows"]
            assert not any("subtotal" in (l["display_name"] or "").lower() or "total" == (l["display_name"] or "").lower()
                           for l in result["leases"]), "subtotal/total rows never become tenants"
        # The CSV is the same roll, so the second import replaces the first.
        assert result["replaced_previous_rows"] == expected["occupied_rows"], result["replaced_previous_rows"]
        assert len(database.get_all_effective_leases(team_id=1)) == expected["occupied_rows"]
        # The old route took 56 s (XLSX) and 72 s (CSV) on this machine (docs/plans/large-files-repro.txt).
        assert max(timings.values()) < 45, timings
        print(f"    rent roll timings: { {k: round(v, 1) for k, v in timings.items()} }")
    finally:
        os.unlink(db)
    print("✓ test_5000_row_messy_rent_rolls_xlsx_and_csv: PASS")


def test_bulk_insert_writes_what_insert_lease_writes():
    db = _fresh_db()
    try:
        fields = {"tenant": {"value": "Acme", "source": None, "confidence": "high"}}
        one = database.insert_lease("roll.csv", fields, team_id=1, display_name="Acme - 101", source_kind="rent_roll")
        [bulk] = database.insert_leases_bulk([{"filename": "roll.csv", "extracted_fields": fields,
                                               "display_name": "Acme - 101", "source_kind": "rent_roll"}], team_id=1)
        a, b = database.get_lease(one, team_id=1), database.get_lease(bulk, team_id=1)
        for column in set(a) - {"id", "uploaded_at"}:
            assert a[column] == b[column], (column, a[column], b[column])
        try:
            database.insert_leases_bulk([{"filename": "x", "extracted_fields": fields},
                                         {"filename": None, "extracted_fields": fields}], team_id=1)
        except Exception:
            pass
        assert len(database.get_all_leases(team_id=1)) == 2, "a failed bulk insert writes nothing"
    finally:
        os.unlink(db)
    print("✓ test_bulk_insert_writes_what_insert_lease_writes: PASS")


def test_another_team_cannot_see_an_upload_job():
    db = _fresh_db()
    try:
        _, done = _upload_and_wait(_client(user_id=1, team_id=1), filename="lease.txt",
                                   body=open(os.path.join(TESTS_DIR, "multiformat_lease.txt"), "rb").read())
        job_id = done["id"]
        other = _client(user_id=2, team_id=2)
        assert other.get(f"/uploads/{job_id}").status_code == 404
        assert database.get_upload_job(job_id, team_id=2) is None
        assert database.get_all_effective_leases(team_id=2) == []
        assert _client(user_id=1, team_id=1).get(f"/uploads/{job_id}").status_code == 200
        assert _client().get(f"/uploads/{2 ** 70}").status_code == 404
    finally:
        os.unlink(db)
    print("✓ test_another_team_cannot_see_an_upload_job: PASS")


def test_stale_job_is_failed_on_poll_and_never_resumed():
    db = _fresh_db()
    try:
        client = _client()
        with mock.patch.object(upload_jobs, "start_job", lambda job_id: None):
            resp = _post(client, filename="lease.txt", body=b"This lease is between A and B for a store in town.\n" * 5)
        job_id = resp.get_json()["upload"]["id"]
        old = (datetime.now(timezone.utc) - timedelta(hours=4)).isoformat()
        conn = database.get_connection()
        conn.execute("UPDATE upload_jobs SET updated_at = ? WHERE id = ?", (old, job_id))
        conn.commit()
        conn.close()
        assert len(_stored_files()) == 1
        job = client.get(f"/uploads/{job_id}").get_json()["upload"]
        assert job["status"] == "failed" and job["error"] == upload_jobs.STALE_MESSAGE
        assert _stored_files() == [], "a stale job's file is deleted with it (review finding)"
        assert database.get_upload_job(job_id, team_id=1)["stored_path"] is None
        upload_jobs.run_upload_job(job_id)  # a late worker pick-up must not resurrect it
        assert client.get(f"/uploads/{job_id}").get_json()["upload"]["status"] == "failed"
        assert database.get_all_effective_leases(team_id=1) == []
    finally:
        for name in _stored_files():
            os.unlink(os.path.join(_UPLOAD_DIR, name))
        os.unlink(db)
    print("✓ test_stale_job_is_failed_on_poll_and_never_resumed: PASS")


def test_unexpected_crash_fails_cleanly_and_deletes_the_file():
    db = _fresh_db()
    try:
        with mock.patch("app.document_extractor.extract_pages_from_path", side_effect=KeyError("boom")):
            _, done = _upload_and_wait(_client(), path_of(TEXT_LEASE))
        assert done["status"] == "failed" and done["error"] == upload_jobs.GENERIC_FAILURE
        assert "boom" not in done["error"] and _stored_files() == []
    finally:
        os.unlink(db)
    print("✓ test_unexpected_crash_fails_cleanly_and_deletes_the_file: PASS")


def test_unreadable_pages_policy():
    import pypdf
    from create_large_fixtures import TEXT_LEASE as _T  # noqa: F401
    db = _fresh_db()
    tmpdir = tempfile.mkdtemp()
    try:
        reader = pypdf.PdfReader(path_of(TEXT_LEASE))
        # A real lease with one blank separator page: extracted, with a warning.
        writer = pypdf.PdfWriter()
        for i in range(6):
            writer.add_page(reader.pages[i])
        writer.add_blank_page(612, 792)
        mostly_text = os.path.join(tmpdir, "with_blank.pdf")
        writer.write(mostly_text)
        _, done = _upload_and_wait(_client(), mostly_text)
        assert done["status"] == "complete", done
        assert done["result"]["leases"][0]["processing_status"] == "complete"
        assert any("page 7" in w or ": 7." in w for w in done["result"]["warnings"]), done["result"]["warnings"]

        # Mostly blank: saved as "needs OCR", fields not guessed.
        writer = pypdf.PdfWriter()
        writer.add_page(reader.pages[0])
        for _ in range(3):
            writer.add_blank_page(612, 792)
        mostly_blank = os.path.join(tmpdir, "mostly_blank.pdf")
        writer.write(mostly_blank)
        _, done = _upload_and_wait(_client(), mostly_blank)
        lease = done["result"]["leases"][0]
        assert lease["processing_status"] == "ocr_needed" and "3 of 4" in lease["processing_error"], lease
        assert lease["extracted_fields"]["tenant"]["value"] is None
    finally:
        shutil.rmtree(tmpdir)
        os.unlink(db)
    print("✓ test_unreadable_pages_policy: PASS")


def test_rent_roll_uploaded_as_a_lease_is_redirected():
    db = _fresh_db()
    try:
        with open(os.path.join(TESTS_DIR, "synthetic_rent_roll_report.pdf"), "rb") as f:
            body = f.read()
        _, done = _upload_and_wait(_client(), filename="report.pdf", body=body)
        assert done["status"] == "failed" and "Rent Roll section" in done["error"], done
    finally:
        os.unlink(db)
    print("✓ test_rent_roll_uploaded_as_a_lease_is_redirected: PASS")


def test_duplicate_upload_returns_the_existing_lease():
    db = _fresh_db()
    try:
        client = _client()
        body = open(os.path.join(TESTS_DIR, "multiformat_lease.txt"), "rb").read()
        _, first = _upload_and_wait(client, filename="lease.txt", body=body)
        resp = _post(client, filename="copy.txt", body=body)
        assert resp.status_code == 200 and resp.get_json()["reused_existing_upload"] is True
        assert resp.get_json()["leases"][0]["id"] == first["result"]["lease_ids"][0]
        assert _stored_files() == []
    finally:
        os.unlink(db)
    print("✓ test_duplicate_upload_returns_the_existing_lease: PASS")


def test_queue_unavailable_is_a_clear_503():
    db = _fresh_db()
    try:
        with mock.patch.object(upload_jobs, "start_job", side_effect=ConnectionError("redis down")):
            resp = _post(_client(), filename="lease.txt", body=b"Lease between A and B.\n" * 10)
        assert resp.status_code == 503 and "try again" in resp.get_json()["error"].lower()
        assert _stored_files() == []
        job = database.get_connection().execute("SELECT status, stored_path FROM upload_jobs").fetchone()
        assert job["status"] == "failed" and job["stored_path"] is None
    finally:
        os.unlink(db)
    print("✓ test_queue_unavailable_is_a_clear_503: PASS")


def test_thread_runner_processes_the_job_in_the_background():
    db = _fresh_db()
    try:
        assert upload_jobs.runner_name() == "thread", "no REDIS_URL -> in-process thread (local dev)"
        client = _client()
        body = open(os.path.join(TESTS_DIR, "multiformat_lease.txt"), "rb").read()
        resp = _post(client, filename="lease.txt", body=body)
        assert resp.status_code == 202
        job_id = resp.get_json()["upload"]["id"]
        deadline = time.time() + 30
        status = None
        while time.time() < deadline:
            status = client.get(f"/uploads/{job_id}").get_json()["upload"]["status"]
            if status in ("complete", "failed"):
                break
            time.sleep(0.2)
        assert status == "complete", status
    finally:
        os.unlink(db)
    print("✓ test_thread_runner_processes_the_job_in_the_background: PASS")


def test_rq_runner_enqueues_the_job_on_the_existing_queue():
    from app import jobs
    with mock.patch.dict(os.environ, {"REDIS_URL": "redis://example.invalid:6379/0"}):
        assert upload_jobs.runner_name() == "rq"
        with mock.patch.object(jobs.extraction_queue, "enqueue") as enqueue:
            upload_jobs.start_job(42)
        enqueue.assert_called_once_with(jobs.run_upload_job, 42, job_timeout=upload_jobs.JOB_TIMEOUT_SECONDS)
    with mock.patch.dict(os.environ, {"UPLOAD_JOB_RUNNER": "thread", "REDIS_URL": "redis://x"}):
        assert upload_jobs.runner_name() == "thread"
    # The RQ entry point delegates to the same job body.
    with mock.patch.object(upload_jobs, "run_upload_job") as body:
        jobs.run_upload_job(7)
    body.assert_called_once_with(7)
    print("✓ test_rq_runner_enqueues_the_job_on_the_existing_queue: PASS")


def test_config_endpoint_describes_limits_and_types():
    db = _fresh_db()
    try:
        cfg = _client().get("/uploads/config").get_json()
        assert cfg["large_file_uploads"] is True and cfg["max_upload_mb"] == 150 and cfg["max_pages"] == 1000
        lease_types = cfg["accepted"]["lease"]["extensions"]
        for ext in (".pdf", ".docx", ".doc", ".xlsx", ".xls", ".csv", ".txt", ".jpg", ".png", ".heic", ".tiff"):
            assert ext in lease_types, ext
    finally:
        os.unlink(db)
    print("✓ test_config_endpoint_describes_limits_and_types: PASS")


def test_existing_bulk_routes_keep_working():
    """
    Regression: commit 89d9654 appended a second database.get_leases_by_ids
    (returning a list) that silently replaced the existing one (a dict by
    id), breaking POST /leases/bulk-delete (500) and /leases/bulk-tag
    (every lease "not found"). Neither route had an offline test.
    """
    db = _fresh_db()
    try:
        fields = {"tenant": {"value": "Acme", "source": None, "confidence": "high"}}
        ids = [database.insert_lease(f"l{n}.pdf", fields, team_id=1) for n in range(3)]
        other_team = database.insert_lease("theirs.pdf", fields, team_id=2)
        client = _client()
        resp = client.post("/leases/bulk-tag", json={"ids": ids[:2] + [other_team], "tag": "Q3"})
        assert resp.status_code == 200, resp.get_json()
        body = resp.get_json()
        assert sorted(body.get("tagged") or []) == sorted(ids[:2]), body
        assert other_team in (body.get("not_found") or []), "another team's lease is never touched"
        resp = client.post("/leases/bulk-delete", json={"ids": [ids[0], other_team]})
        assert resp.status_code == 200, resp.get_json()
        assert database.get_lease(ids[0], team_id=1) is None
        assert database.get_lease(other_team, team_id=2) is not None
    finally:
        os.unlink(db)
    print("✓ test_existing_bulk_routes_keep_working: PASS")


def test_no_module_defines_the_same_function_twice():
    """Guard for the bug above: a second top-level def with the same name silently replaces the first."""
    import ast
    import collections
    import glob
    app_dir = os.path.abspath(os.path.join(TESTS_DIR, "..", "app"))
    duplicates = {}
    for path in sorted(glob.glob(os.path.join(app_dir, "*.py"))):
        with open(path) as f:
            tree = ast.parse(f.read())
        counts = collections.Counter(n.name for n in tree.body
                                     if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)))
        dupes = sorted(name for name, c in counts.items() if c > 1)
        if dupes:
            duplicates[os.path.basename(path)] = dupes
    assert not duplicates, f"defined more than once: {duplicates}"
    print("✓ test_no_module_defines_the_same_function_twice: PASS")


def test_a_busy_database_does_not_fail_the_job_or_the_poll():
    """
    Found while screenshotting: with ~5,000 leases, POST /alerts/generate holds
    SQLite's write lock for 5-9 s; jobs and polls that waited the default 5 s
    failed with "database is locked". Here another connection holds the write
    lock for 6.5 s while a job runs and the browser polls.
    """
    import sqlite3
    import threading
    db = _fresh_db()
    try:
        client = _client()
        queued = []
        body = open(os.path.join(TESTS_DIR, "multiformat_lease.txt"), "rb").read()
        with mock.patch.object(upload_jobs, "start_job", side_effect=queued.append):
            resp = _post(client, filename="lease.txt", body=body)
        job_id = queued[0]

        locked, release = threading.Event(), threading.Event()

        def hold_write_lock():
            conn = sqlite3.connect(database.get_db_path())
            conn.execute("BEGIN IMMEDIATE")
            locked.set()
            release.wait(6.5)
            conn.rollback()
            conn.close()

        holder = threading.Thread(target=hold_write_lock)
        holder.start()
        locked.wait(5)
        polls = []
        poller = threading.Thread(target=lambda: polls.append(client.get(f"/uploads/{job_id}").status_code))
        poller.start()
        upload_jobs.run_upload_job(job_id)
        poller.join()
        holder.join()
        assert polls == [200], polls
        done = client.get(f"/uploads/{job_id}").get_json()["upload"]
        assert done["status"] == "complete", done
        assert len(done["result"]["lease_ids"]) == 1
    finally:
        os.unlink(db)
    print("✓ test_a_busy_database_does_not_fail_the_job_or_the_poll: PASS")


def _multipart(filename, size_bytes, boundary="XyZbound"):
    head = (f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"{filename}\"\r\n"
            "Content-Type: application/pdf\r\n\r\n").encode()
    body = b"%PDF-1.7\n" + b"0" * size_bytes
    tail = (f"\r\n--{boundary}\r\nContent-Disposition: form-data; name=\"kind\"\r\n\r\nlease\r\n"
            f"--{boundary}--\r\n").encode()
    return head + body + tail, f"multipart/form-data; boundary={boundary}"


def test_a_chunked_body_cannot_skip_the_old_16mb_limit():
    """
    Security review: with the flag on, the old routes' 16 MB cap was checked
    against Content-Length only, so a chunked body (no Content-Length; gunicorn
    sets wsgi.input_terminated) streamed 17 MB into POST /leases, which reads
    the whole file into memory. Fails on 8fbcfce.
    """
    db = _fresh_db()
    try:
        client = _client()
        seen = []
        import app.api as api_module
        real = api_module._validate_upload

        def spy():
            from flask import request
            f = request.files.get("file")
            seen.append(len(f.read()) if f else None)
            if f:
                f.seek(0)
            return real()

        payload, content_type = _multipart("big.pdf", 17 * 1024 * 1024)
        chunked = {"headers": {"Transfer-Encoding": "chunked"}, "environ_overrides": {"wsgi.input_terminated": True}}
        with mock.patch.object(api_module, "_validate_upload", side_effect=spy):
            resp = client.post("/leases", input_stream=io.BytesIO(payload), content_type=content_type, **chunked)
        assert resp.status_code == 413, (resp.status_code, resp.get_json(), seen)
        assert seen == [], "the old route never sees the file"

        # POST /uploads still takes a big chunked body (it streams to disk).
        queued = []
        with mock.patch.object(upload_jobs, "start_job", side_effect=queued.append):
            resp = client.post("/uploads", input_stream=io.BytesIO(payload), content_type=content_type, **chunked)
        assert resp.status_code == 202, resp.get_json()
        # ...but not past its own limit.
        with mock.patch.object(upload_jobs, "MAX_REQUEST_BYTES", 5 * 1024 * 1024):
            resp = client.post("/uploads", input_stream=io.BytesIO(payload), content_type=content_type, **chunked)
        assert resp.status_code == 413
    finally:
        for name in _stored_files():
            os.unlink(os.path.join(_UPLOAD_DIR, name))
        os.unlink(db)
    print("✓ test_a_chunked_body_cannot_skip_the_old_16mb_limit: PASS")


def test_a_team_can_have_only_a_few_uploads_in_flight():
    """Security review: one analyst (or script) could queue unlimited 150 MB uploads and fill the disk."""
    db = _fresh_db()
    try:
        client = _client()
        with mock.patch.object(upload_jobs, "start_job", lambda job_id: None):
            for n in range(upload_jobs.MAX_ACTIVE_JOBS_PER_TEAM):
                assert _post(client, filename=f"r{n}.csv", body=f"Unit,Tenant,Rent\n{n},A,1\n".encode(),
                             kind="rent_roll").status_code == 202
            resp = _post(client, filename="one-more.csv", body=b"Unit,Tenant,Rent\n9,A,1\n", kind="rent_roll")
            assert resp.status_code == 429 and "Wait for one to finish" in resp.get_json()["error"]
            assert len(_stored_files()) == upload_jobs.MAX_ACTIVE_JOBS_PER_TEAM, "the refused file was never stored"
            # Another team isn't affected by this team's queue.
            assert _post(_client(user_id=2, team_id=2), filename="t2.csv", body=b"Unit,Tenant,Rent\n1,B,1\n",
                         kind="rent_roll").status_code == 202
        with mock.patch.object(upload_jobs, "has_room_for_an_upload", return_value=False):
            resp = _post(_client(user_id=3, team_id=3), filename="x.csv", body=b"Unit,Tenant,Rent\n1,C,1\n",
                         kind="rent_roll")
        assert resp.status_code == 503
    finally:
        for name in _stored_files():
            os.unlink(os.path.join(_UPLOAD_DIR, name))
        os.unlink(db)
    print("✓ test_a_team_can_have_only_a_few_uploads_in_flight: PASS")


def test_orphaned_files_are_swept():
    old = os.path.join(_UPLOAD_DIR, "deadbeef.upload")
    fresh = os.path.join(_UPLOAD_DIR, "cafebabe.upload")
    for path in (old, fresh):
        with open(path, "wb") as f:
            f.write(b"x")
    long_ago = time.time() - upload_jobs.ORPHAN_AFTER_SECONDS - 60
    os.utime(old, (long_ago, long_ago))
    try:
        assert upload_jobs.sweep_orphaned_files() == 1
        assert not os.path.exists(old) and os.path.exists(fresh), "only files too old to belong to a live job"
    finally:
        upload_jobs.remove_file(fresh)
    print("✓ test_orphaned_files_are_swept: PASS")


def test_a_job_failed_as_stale_is_not_brought_back_by_its_runner():
    """Review finding: a heartbeat used to set status='processing' unconditionally, reviving a job the user was told had failed."""
    db = _fresh_db()
    try:
        client = _client()
        queued = []
        with mock.patch.object(upload_jobs, "start_job", side_effect=queued.append):
            _post(client, path_of(TEXT_LEASE))
        job_id = queued[0]
        from app import document_extractor
        real = document_extractor.extract_pages_from_path

        def mark_stale_midway(path, extension, progress=None, max_pages=None):
            def on_page(done, total):
                if done == 10:  # meanwhile, a poll decided this job was stale
                    database.update_upload_job(job_id, status="failed", stage="failed", error=upload_jobs.STALE_MESSAGE)
                progress(done, total)
            return real(path, extension, progress=on_page, max_pages=max_pages)

        with mock.patch.object(document_extractor, "extract_pages_from_path", side_effect=mark_stale_midway):
            upload_jobs.run_upload_job(job_id)
        job = client.get(f"/uploads/{job_id}").get_json()["upload"]
        assert job["status"] == "failed" and job["error"] == upload_jobs.STALE_MESSAGE, job
        assert database.get_all_effective_leases(team_id=1) == [], "the runner stopped instead of saving"
        assert _stored_files() == []
    finally:
        os.unlink(db)
    print("✓ test_a_job_failed_as_stale_is_not_brought_back_by_its_runner: PASS")


if __name__ == "__main__":
    try:
        test_flag_off_means_routes_404_and_old_limits_unchanged()
        test_routes_require_login_and_role()
        test_500_page_text_lease_regex_engine_end_to_end_with_progress()
        test_500_page_lease_ai_engine_quota_checked_before_any_ai_call()
        test_over_quota_fails_the_job_without_calling_the_model()
        test_ai_failure_shows_a_friendly_message_and_logs_the_details_with_the_team()
        test_120mb_scanned_pdf_is_streamed_not_held_in_memory()
        test_size_limits_are_friendly_and_leave_nothing_behind()
        test_tables_over_the_row_limit_are_refused_before_parsing()
        test_the_largest_office_files_accepted_convert_well_under_512mb()
        test_page_limit_is_checked_before_any_ocr()
        test_wrong_content_is_rejected_and_mislabeled_files_are_read_as_what_they_are()
        test_every_accepted_format_goes_through_the_one_pipeline()
        test_5000_row_messy_rent_rolls_xlsx_and_csv()
        test_bulk_insert_writes_what_insert_lease_writes()
        test_another_team_cannot_see_an_upload_job()
        test_stale_job_is_failed_on_poll_and_never_resumed()
        test_unexpected_crash_fails_cleanly_and_deletes_the_file()
        test_unreadable_pages_policy()
        test_rent_roll_uploaded_as_a_lease_is_redirected()
        test_duplicate_upload_returns_the_existing_lease()
        test_queue_unavailable_is_a_clear_503()
        test_thread_runner_processes_the_job_in_the_background()
        test_rq_runner_enqueues_the_job_on_the_existing_queue()
        test_config_endpoint_describes_limits_and_types()
        test_existing_bulk_routes_keep_working()
        test_no_module_defines_the_same_function_twice()
        test_a_busy_database_does_not_fail_the_job_or_the_poll()
        test_a_chunked_body_cannot_skip_the_old_16mb_limit()
        test_a_team_can_have_only_a_few_uploads_in_flight()
        test_orphaned_files_are_swept()
        test_a_job_failed_as_stale_is_not_brought_back_by_its_runner()
    finally:
        shutil.rmtree(_UPLOAD_DIR, ignore_errors=True)
    print("\nAll upload job tests passed!")
