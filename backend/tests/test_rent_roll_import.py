"""
Tests for app/rent_roll_import.py -- the broker Excel/CSV rent roll
importer. No fixed format is assumed (that's the whole point), so
these tests deliberately exercise header-naming diversity, vacant/
total rows, currency formatting variety, and missing columns, not just
one clean example.
"""

import csv
import io
import os
import sys

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from openpyxl import Workbook

from app.rent_roll_import import (
    RentRollImportError,
    parse_csv_rent_roll,
    parse_rent_roll_rows,
    parse_xlsx_rent_roll,
)


def _csv_bytes(rows):
    """
    Proper CSV encoding, not a naive ",".join() -- several test fixtures
    below intentionally use values like "$4,500.00" that themselves
    contain a comma, which a naive join would incorrectly split into
    two fields. Using csv.writer catches exactly that class of bug
    (caught for real during initial test authoring here, not
    hypothetically -- see DECISIONS.md).
    """
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerows(rows)
    return buf.getvalue().encode("utf-8")


def _xlsx_bytes(rows):
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_csv_happy_path_standard_headers():
    csv_bytes = _csv_bytes([
        ["Tenant", "Rent", "Square Footage", "Lease Start", "Lease End"],
        ["Acme Corp", "$4,500.00", "2,000 sq ft", "01/01/2024", "12/31/2028"],
        ["Beta LLC", "$3,200.00", "1,500 sq ft", "06/01/2023", "05/31/2027"],
    ])
    result = parse_csv_rent_roll(csv_bytes, "rentroll.csv", base_property_address="100 Main St")

    assert result["skipped_rows"] == []
    assert len(result["leases"]) == 2
    acme = result["leases"][0]["extracted_fields"]
    assert acme["tenant"]["value"] == "Acme Corp"
    assert acme["rent_amount"]["value"] == "$4,500.00"
    assert acme["square_footage"]["value"] == "2,000 sq ft"
    assert acme["property_address"]["value"] == "100 Main St"
    assert acme["lease_start_date"]["source"]["row"] == 2
    assert acme["lease_start_date"]["source"]["file"] == "rentroll.csv"
    print("✓ test_csv_happy_path_standard_headers: PASS")


def test_xlsx_happy_path_with_real_numeric_and_date_cell_types():
    """xlsx cells commonly come back as real Python int/float/date objects, not strings -- confirms those are handled, not just string cells."""
    from datetime import date
    xlsx_bytes = _xlsx_bytes([
        ["Tenant Name", "Monthly Rent", "Sq Ft", "Commencement Date"],
        ["Acme Corp", 4500, 2000, date(2024, 1, 1)],  # real numbers/date, not strings
    ])
    result = parse_xlsx_rent_roll(xlsx_bytes, "rentroll.xlsx")

    assert len(result["leases"]) == 1
    fields = result["leases"][0]["extracted_fields"]
    assert fields["rent_amount"]["value"] == "$4,500.00"
    assert fields["square_footage"]["value"] == "2,000 sq ft"
    assert fields["lease_start_date"]["value"] == "January 01, 2024"
    print("✓ test_xlsx_happy_path_with_real_numeric_and_date_cell_types: PASS")


def test_header_alias_diversity():
    """Different brokers name columns differently -- confirms several realistic header variants all map correctly, not just one exact spelling."""
    csv_bytes = _csv_bytes([
        ["Lessee", "Base Rent", "RSF", "Suite #"],
        ["Gamma Inc", "$2,000.00", "1,000", "204"],
    ])
    result = parse_csv_rent_roll(csv_bytes, "rr.csv", base_property_address="500 Plaza Dr")
    fields = result["leases"][0]["extracted_fields"]
    assert fields["tenant"]["value"] == "Gamma Inc"
    assert fields["rent_amount"]["value"] == "$2,000.00"
    assert fields["square_footage"]["value"] == "1,000 sq ft"
    assert fields["property_address"]["value"] == "500 Plaza Dr, Suite 204"
    print("✓ test_header_alias_diversity: PASS")


def test_unit_column_already_spelled_out_does_not_double_prefix():
    """
    Regression test: found via live messy-data testing, not written
    speculatively. Many real rent rolls' Unit/Suite column already
    contains the designator itself (e.g. "Suite 101"), not a bare
    number -- unconditionally prepending "Suite " produced "Suite Suite
    101", which then silently broke compute_rent_roll_reconciliation's
    exact-address matching against the same unit's lease PDF (whose
    property_address reads plain "Suite 101"), since the two strings
    normalize differently. Covers "Suite", "Ste.", "Unit", "Apt", and
    "#"-style cells, plus confirms a bare identifier still gets "Suite "
    prepended as before.
    """
    for unit_cell, expected_address in [
        ("Suite 101", "500 Commerce Blvd, Suite 101"),
        ("Ste. 101", "500 Commerce Blvd, Ste. 101"),
        ("Unit 5", "500 Commerce Blvd, Unit 5"),
        ("Apt 2B", "500 Commerce Blvd, Apt 2B"),
        ("#12", "500 Commerce Blvd, #12"),
        ("101", "500 Commerce Blvd, Suite 101"),
    ]:
        csv_bytes = _csv_bytes([
            ["Tenant", "Unit", "Rent"],
            ["Acme Corp", unit_cell, "$4,950.00"],
        ])
        result = parse_csv_rent_roll(csv_bytes, "rr.csv", base_property_address="500 Commerce Blvd")
        address = result["leases"][0]["extracted_fields"]["property_address"]["value"]
        assert address == expected_address, f"unit cell {unit_cell!r} produced {address!r}, expected {expected_address!r}"
    print("✓ test_unit_column_already_spelled_out_does_not_double_prefix: PASS")


def test_bare_numeric_rent_with_no_dollar_sign():
    """A rent column stored as a plain number string (no "$") -- common in CSV exports -- must still parse."""
    csv_bytes = _csv_bytes([
        ["Tenant", "Rent"],
        ["Delta Co", "3200.50"],
    ])
    result = parse_csv_rent_roll(csv_bytes, "rr.csv")
    assert result["leases"][0]["extracted_fields"]["rent_amount"]["value"] == "$3,200.50"
    print("✓ test_bare_numeric_rent_with_no_dollar_sign: PASS")


def test_vacant_row_skipped():
    csv_bytes = _csv_bytes([
        ["Tenant", "Rent"],
        ["Real Tenant Co", "$1,000.00"],
        ["VACANT", "$0.00"],
        ["Vacant", ""],
    ])
    result = parse_csv_rent_roll(csv_bytes, "rr.csv")
    assert len(result["leases"]) == 1
    assert len(result["skipped_rows"]) == 2
    assert result["skipped_rows"][0]["row"] == 3
    print("✓ test_vacant_row_skipped: PASS")


def test_total_row_skipped():
    csv_bytes = _csv_bytes([
        ["Tenant", "Rent"],
        ["Real Tenant Co", "$1,000.00"],
        ["Total", "$1,000.00"],
        ["Subtotal", "$1,000.00"],
    ])
    result = parse_csv_rent_roll(csv_bytes, "rr.csv")
    assert len(result["leases"]) == 1
    assert len(result["skipped_rows"]) == 2
    print("✓ test_total_row_skipped: PASS")


def test_subtotal_row_with_extra_text_is_skipped():
    """
    Real bug found via live stress-testing: a subtotal/total row's
    tenant cell is often not the bare keyword alone (e.g. "Subtotal
    Floor 1", "Total (12 units)") -- the old exact-match-only check let
    these through as fake tenants. Also confirms a real tenant whose
    name merely starts with a similar-looking word ("Totally Awesome
    Tenant") is still correctly imported, not caught by the fix.
    """
    csv_bytes = _csv_bytes([
        ["Tenant", "Rent"],
        ["Real Tenant Co", "$1,000.00"],
        ["Subtotal Floor 1", "$1,000.00"],
        ["Total (2 units)", "$1,000.00"],
        ["Totally Awesome Tenant", "$900.00"],
    ])
    result = parse_csv_rent_roll(csv_bytes, "rr.csv")
    tenants = {l["extracted_fields"]["tenant"]["value"] for l in result["leases"]}
    assert tenants == {"Real Tenant Co", "Totally Awesome Tenant"}
    assert len(result["skipped_rows"]) == 2
    print("✓ test_subtotal_row_with_extra_text_is_skipped: PASS")


def test_tenant_literal_zero_is_skipped():
    """
    Real bug found via live stress-testing: some rent rolls mark a
    vacant unit's tenant cell with a literal "0" instead of blank/
    VACANT/Vacant -- this used to import a fake tenant literally named
    "0" with $0.00 rent.
    """
    csv_bytes = _csv_bytes([
        ["Tenant", "Rent"],
        ["0", "$0.00"],
        ["Real Tenant After Zero", "$1,800.00"],
    ])
    result = parse_csv_rent_roll(csv_bytes, "rr.csv")
    assert len(result["leases"]) == 1
    assert result["leases"][0]["extracted_fields"]["tenant"]["value"] == "Real Tenant After Zero"
    assert len(result["skipped_rows"]) == 1
    print("✓ test_tenant_literal_zero_is_skipped: PASS")


def test_blank_tenant_row_skipped():
    csv_bytes = _csv_bytes([
        ["Tenant", "Rent"],
        ["Real Tenant Co", "$1,000.00"],
        ["", "$500.00"],  # blank tenant, e.g. a formatting/spacer row
    ])
    result = parse_csv_rent_roll(csv_bytes, "rr.csv")
    assert len(result["leases"]) == 1
    assert len(result["skipped_rows"]) == 1
    print("✓ test_blank_tenant_row_skipped: PASS")


def test_missing_tenant_and_rent_columns_raises():
    csv_bytes = _csv_bytes([
        ["Property Name", "Notes"],
        ["Building A", "some note"],
    ])
    try:
        parse_csv_rent_roll(csv_bytes, "rr.csv")
        assert False, "should have raised RentRollImportError"
    except RentRollImportError:
        pass
    print("✓ test_missing_tenant_and_rent_columns_raises: PASS")


def test_missing_optional_column_does_not_block_import():
    """No square-footage column at all -- rows must still import, with square_footage honestly 'not found', not a crash."""
    csv_bytes = _csv_bytes([
        ["Tenant", "Rent"],
        ["Epsilon LLC", "$2,500.00"],
    ])
    result = parse_csv_rent_roll(csv_bytes, "rr.csv")
    assert len(result["leases"]) == 1
    fields = result["leases"][0]["extracted_fields"]
    assert fields["square_footage"]["value"] is None
    assert fields["square_footage"]["source"] is None
    print("✓ test_missing_optional_column_does_not_block_import: PASS")


def test_unrecognized_extra_columns_ignored():
    csv_bytes = _csv_bytes([
        ["Tenant", "Rent", "Parking Spaces", "Notes"],
        ["Zeta Inc", "$1,800.00", "4", "corner unit"],
    ])
    result = parse_csv_rent_roll(csv_bytes, "rr.csv")
    assert len(result["leases"]) == 1  # extra columns don't break anything
    print("✓ test_unrecognized_extra_columns_ignored: PASS")


def test_blank_rows_in_file_are_skipped_entirely():
    csv_bytes = _csv_bytes([
        ["Tenant", "Rent"],
        ["Eta Co", "$1,200.00"],
        ["", ""],  # a fully blank spacer row
        ["Theta Co", "$1,400.00"],
    ])
    result = parse_csv_rent_roll(csv_bytes, "rr.csv")
    assert len(result["leases"]) == 2
    assert result["skipped_rows"] == []  # fully-blank rows are dropped before parsing, not reported as skipped
    print("✓ test_blank_rows_in_file_are_skipped_entirely: PASS")


def test_no_base_address_and_no_unit_column_leaves_address_unset():
    csv_bytes = _csv_bytes([
        ["Tenant", "Rent"],
        ["Iota Inc", "$900.00"],
    ])
    result = parse_csv_rent_roll(csv_bytes, "rr.csv")  # no base_property_address passed
    fields = result["leases"][0]["extracted_fields"]
    assert fields["property_address"]["value"] is None
    print("✓ test_no_base_address_and_no_unit_column_leaves_address_unset: PASS")


def test_empty_csv_raises():
    try:
        parse_csv_rent_roll(b"", "empty.csv")
        assert False, "should have raised RentRollImportError"
    except RentRollImportError:
        pass
    print("✓ test_empty_csv_raises: PASS")


def test_empty_xlsx_raises():
    empty_bytes = _xlsx_bytes([])
    try:
        parse_xlsx_rent_roll(empty_bytes, "empty.xlsx")
        assert False, "should have raised RentRollImportError"
    except RentRollImportError:
        pass
    print("✓ test_empty_xlsx_raises: PASS")


def test_xlsx_checks_every_sheet_not_just_the_active_one():
    """
    Real bug found via live stress-testing: a workbook with a blank/
    decorative "Summary" sheet first and the real per-unit data on a
    second "Detail" sheet used to silently fail to import anything --
    only workbook.active (whichever sheet was selected when the file
    was last saved) was ever read. Now every sheet is checked, using
    the first one with a recognizable tenant+rent header.
    """
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(["This tab is a portfolio summary, not per-unit data"])
    detail = wb.create_sheet("Detail")
    detail.append(["Tenant", "Unit", "Rent"])
    detail.append(["Sheet2 Tenant A", "D1", "$1,500.00"])
    buf = io.BytesIO()
    wb.save(buf)

    result = parse_xlsx_rent_roll(buf.getvalue(), "multi_sheet.xlsx")
    assert len(result["leases"]) == 1
    assert result["leases"][0]["extracted_fields"]["tenant"]["value"] == "Sheet2 Tenant A"
    print("✓ test_xlsx_checks_every_sheet_not_just_the_active_one: PASS")


def test_row_numbers_reflect_actual_file_position():
    """Row numbering must match what a human looking at the actual spreadsheet would count (1-indexed, header is row 1)."""
    csv_bytes = _csv_bytes([
        ["Tenant", "Rent"],   # row 1
        ["First Co", "$1,000.00"],  # row 2
        ["Second Co", "$2,000.00"],  # row 3
    ])
    result = parse_csv_rent_roll(csv_bytes, "rr.csv")
    assert result["leases"][0]["extracted_fields"]["tenant"]["source"]["row"] == 2
    assert result["leases"][1]["extracted_fields"]["tenant"]["source"]["row"] == 3
    print("✓ test_row_numbers_reflect_actual_file_position: PASS")


def test_specific_date_column_not_misattributed_to_rent_amount():
    """
    Regression test for a real bug: "Rent Commencement"/"Rent
    Expiration" are standard commercial lease terms (rent can start
    later than the lease itself during a free-rent period) -- these
    must map to lease_start_date/lease_end_date, not get grabbed by
    rent_amount's bare "rent" alias just because that field happens to
    be checked first. With no unambiguous rent-amount column present at
    all, rent_amount should end up genuinely unmapped (honest failure)
    rather than incorrectly pointed at a date column.
    """
    csv_bytes = _csv_bytes([
        ["Tenant", "Monthly Payment", "Rent Commencement", "Rent Expiration"],
        ["Acme Corp", "$4,000.00", "01/01/2024", "12/31/2028"],
    ])
    try:
        parse_csv_rent_roll(csv_bytes, "rr.csv")
        assert False, "should have raised -- 'Monthly Payment' isn't a recognized rent alias, so rent_amount is genuinely unmapped"
    except RentRollImportError:
        pass

    # But WITH an unambiguous "Rent" column present, the date columns
    # must still land correctly, not collide with it.
    csv_bytes2 = _csv_bytes([
        ["Tenant", "Rent", "Rent Commencement", "Rent Expiration"],
        ["Acme Corp", "$4,000.00", "01/01/2024", "12/31/2028"],
    ])
    result = parse_csv_rent_roll(csv_bytes2, "rr.csv")
    fields = result["leases"][0]["extracted_fields"]
    assert fields["rent_amount"]["value"] == "$4,000.00"
    assert fields["lease_start_date"]["value"] == "01/01/2024"
    assert fields["lease_end_date"]["value"] == "12/31/2028"
    print("✓ test_specific_date_column_not_misattributed_to_rent_amount: PASS")


def test_header_row_auto_detection_skips_decorative_rows():
    """
    Real PMS canned reports (Yardi, AppFolio, ...) routinely open with a
    few decorative rows -- property name, report title, an "As of" date
    -- before the actual column-header row, unlike a broker's own
    hand-built spreadsheet where row 1 almost always already is the
    header. Includes a genuinely blank spacer row in the decorative
    block too, since that's a real pattern and specifically exercises
    the true-row-number tracking (not just index arithmetic) that fixes
    citation drift when blank rows get dropped before the header.
    """
    csv_bytes = _csv_bytes([
        ["Riverside Commons"],
        ["Rent Roll"],
        ["As of 08/01/2026"],
        [],
        ["Tenant", "Unit", "Rent"],
        ["Acme Corp", "101", "2000.00"],
    ])
    result = parse_csv_rent_roll(csv_bytes, "yardi_style.csv", base_property_address="1 Main St")
    assert result["column_mapping"] == {"tenant": 0, "unit": 1, "rent_amount": 2}, result["column_mapping"]
    fields = result["leases"][0]["extracted_fields"]
    assert fields["tenant"]["value"] == "Acme Corp"
    # True file row 6, not 5 -- the dropped blank spacer row (file row 4)
    # must not shift this citation off by one.
    assert fields["tenant"]["source"]["row"] == 6, fields["tenant"]["source"]
    print("✓ test_header_row_auto_detection_skips_decorative_rows: PASS")


def test_header_row_already_at_top_is_unaffected():
    """The common broker-spreadsheet case (header already row 1) must behave identically to before header detection existed -- no regression."""
    csv_bytes = _csv_bytes([
        ["Tenant", "Rent"],
        ["Acme Corp", "$2,000.00"],
    ])
    result = parse_csv_rent_roll(csv_bytes, "rr.csv")
    assert result["leases"][0]["extracted_fields"]["tenant"]["source"]["row"] == 2
    print("✓ test_header_row_already_at_top_is_unaffected: PASS")


def test_no_header_found_within_scan_window_gives_clear_error():
    """
    A file with no recognizable header row anywhere near the top (not a
    rent roll at all, or a genuinely malformed export) must still fail
    with the same clear, honest error as before -- not silently
    misinterpret some deep decorative or data row as if it were the
    header.
    """
    csv_bytes = _csv_bytes([["Notes"], ["This file has nothing recognizable in it"], ["Parking Spaces: 12"]])
    try:
        parse_csv_rent_roll(csv_bytes, "not_a_rent_roll.csv")
        assert False, "should have raised RentRollImportError"
    except RentRollImportError as e:
        assert "recognizable" in str(e).lower()
    print("✓ test_no_header_found_within_scan_window_gives_clear_error: PASS")


def test_decorative_title_row_with_a_stray_recognizable_word_not_mistaken_for_header():
    """
    A decorative title row containing a word that partially overlaps a
    real alias ("Rent Roll Report" contains "Rent") must NOT be mistaken
    for the real header -- header detection requires BOTH a tenant AND
    a rent match, and a title row has neither a real tenant column nor
    a real rent column, just a word that happens to appear in one.
    """
    csv_bytes = _csv_bytes([
        ["Rent Roll Report"],
        ["Tenant", "Rent"],
        ["Acme Corp", "$2,000.00"],
    ])
    result = parse_csv_rent_roll(csv_bytes, "rr.csv")
    assert result["column_mapping"] == {"tenant": 0, "rent_amount": 1}
    assert result["leases"][0]["extracted_fields"]["tenant"]["value"] == "Acme Corp"
    print("✓ test_decorative_title_row_with_a_stray_recognizable_word_not_mistaken_for_header: PASS")


def test_header_row_auto_detection_works_for_xlsx_too():
    """Same decorative-header scenario, but through the .xlsx path -- both readers share _find_header_row, but that's exactly the kind of 'should be covered' assumption worth actually checking."""
    xlsx_bytes = _xlsx_bytes([
        ["Sunset Plaza"],
        ["Rent Roll"],
        ["Tenant", "Unit", "Rent"],
        ["Beta LLC", "Suite 200", 3000],
    ])
    result = parse_xlsx_rent_roll(xlsx_bytes, "yardi_style.xlsx", base_property_address="2 Oak Ave")
    fields = result["leases"][0]["extracted_fields"]
    assert fields["tenant"]["value"] == "Beta LLC"
    assert fields["property_address"]["value"] == "2 Oak Ave, Suite 200"
    assert fields["tenant"]["source"]["row"] == 4
    print("✓ test_header_row_auto_detection_works_for_xlsx_too: PASS")


def test_per_row_property_column_overrides_base_address_for_multi_property_exports():
    """
    A portfolio-wide PMS export can cover several DIFFERENT properties
    in one file, one row per unit -- unlike a single-property rent roll
    where one uploader-typed base_property_address is correct for every
    row. When the file has its own Property column, each row's real
    building must win, not be flattened into whatever address the
    uploader happened to type (which would be actively wrong for every
    row except whichever property they were thinking of).
    """
    csv_bytes = _csv_bytes([
        ["Property", "Tenant", "Unit", "Rent"],
        ["100 Alpha St", "Acme Corp", "101", "2000.00"],
        ["200 Beta Ave", "Beta LLC", "5", "3000.00"],
    ])
    result = parse_csv_rent_roll(csv_bytes, "multi_property.csv", base_property_address="Should Not Be Used")
    addresses = {l["extracted_fields"]["tenant"]["value"]: l["extracted_fields"]["property_address"]["value"] for l in result["leases"]}
    assert addresses["Acme Corp"] == "100 Alpha St, Suite 101", addresses
    assert addresses["Beta LLC"] == "200 Beta Ave, Suite 5", addresses
    print("✓ test_per_row_property_column_overrides_base_address_for_multi_property_exports: PASS")


def test_property_column_falls_back_to_base_address_when_a_row_is_blank():
    """A mostly-complete Property column shouldn't lose the uploader's fallback for the rows it's actually missing on."""
    csv_bytes = _csv_bytes([
        ["Property", "Tenant", "Unit", "Rent"],
        ["100 Alpha St", "Acme Corp", "101", "2000.00"],
        ["", "Beta LLC", "5", "3000.00"],
    ])
    result = parse_csv_rent_roll(csv_bytes, "multi_property.csv", base_property_address="200 Beta Ave")
    addresses = {l["extracted_fields"]["tenant"]["value"]: l["extracted_fields"]["property_address"]["value"] for l in result["leases"]}
    assert addresses["Acme Corp"] == "100 Alpha St, Suite 101", addresses
    assert addresses["Beta LLC"] == "200 Beta Ave, Suite 5", addresses
    print("✓ test_property_column_falls_back_to_base_address_when_a_row_is_blank: PASS")


def test_market_rent_is_never_used_as_rent_amount():
    """
    Regression/design test for a real risk caught during design, not
    found broken in production: "Market Rent" is a theoretical
    achievable-at-100%-occupancy figure, not what a tenant is actually,
    contractually paying. The bare "rent" alias would otherwise match
    it via substring containment. A file with ONLY a Market Rent column
    (no real rent column) must fail the same honest way as a file with
    no rent column at all -- silently treating market rent as if it
    were actual rent would be systematically wrong, not just imprecise.
    """
    csv_bytes = _csv_bytes([["Tenant", "Market Rent"], ["Acme Corp", "2500.00"]])
    try:
        parse_csv_rent_roll(csv_bytes, "market_only.csv")
        assert False, "should have raised -- Market Rent must never be treated as rent_amount"
    except RentRollImportError:
        pass

    # With BOTH a market rent column and a real one, the real one must
    # win -- market rent must not even be considered a lower-priority
    # fallback candidate.
    csv_bytes2 = _csv_bytes([
        ["Tenant", "Market Rent", "Current Rent"],
        ["Acme Corp", "2500.00", "2000.00"],
    ])
    result = parse_csv_rent_roll(csv_bytes2, "market_and_real.csv")
    assert result["leases"][0]["extracted_fields"]["rent_amount"]["value"] == "$2,000.00", \
        result["leases"][0]["extracted_fields"]["rent_amount"]
    print("✓ test_market_rent_is_never_used_as_rent_amount: PASS")


def test_yardi_appfolio_terminology_aliases():
    """PMS-specific column names (Resident, Unit SF, Scheduled Rent, Lease From/To) map to the same fields as their broker-spreadsheet equivalents."""
    csv_bytes = _csv_bytes([
        ["Resident", "Unit SF", "Scheduled Rent", "Lease From", "Lease To"],
        ["Gamma Inc", "1200", "3200.00", "01/01/2024", "12/31/2026"],
    ])
    result = parse_csv_rent_roll(csv_bytes, "pms_terms.csv")
    fields = result["leases"][0]["extracted_fields"]
    assert fields["tenant"]["value"] == "Gamma Inc"
    assert fields["square_footage"]["value"] == "1,200 sq ft"
    assert fields["rent_amount"]["value"] == "$3,200.00"
    assert fields["lease_start_date"]["value"] == "01/01/2024"
    assert fields["lease_end_date"]["value"] == "12/31/2026"
    print("✓ test_yardi_appfolio_terminology_aliases: PASS")


def test_realpage_mri_terminology_aliases():
    """
    RealPage/MRI-style column names: "Actual Rent" alongside "Market
    Rent" (RealPage), and a BARE "Commence" (no "date"/"lease" suffix,
    common in MRI-style exports) -- must correctly map to rent_amount
    and lease_start_date respectively, not fall through unmapped just
    because the header is shorter than the existing full-phrase
    aliases.
    """
    csv_bytes = _csv_bytes([
        ["Suite", "Occupant", "SF", "Commence", "Expire", "Market Rent", "Actual Rent"],
        ["101", "Acme Corp", "1200", "01/01/2024", "12/31/2029", "3600.00", "3200.00"],
    ])
    result = parse_csv_rent_roll(csv_bytes, "realpage_style.csv")
    fields = result["leases"][0]["extracted_fields"]
    assert fields["tenant"]["value"] == "Acme Corp"
    assert fields["rent_amount"]["value"] == "$3,200.00"  # Actual Rent, not Market Rent's $3,600.00
    assert fields["lease_start_date"]["value"] == "01/01/2024"
    assert fields["lease_end_date"]["value"] == "12/31/2029"
    print("✓ test_realpage_mri_terminology_aliases: PASS")


def test_rent_psf_rate_column_never_used_as_rent_amount():
    """
    Regression/design test for a real bug caught during self-review,
    not found broken in production: "Rent PSF" (rent per square foot --
    a common RealPage/MRI-style column) is a RATE, not the tenant's
    total dollar rent. With no better rent column present, the bare
    "rent" alias would otherwise match it and silently treat e.g.
    "2.75" ($2.75/sqft/month) as if it were $2.75/month total rent --
    wrong by roughly the unit's entire square footage. Must fail
    honestly instead, same as a file with no rent column at all.
    """
    csv_bytes = _csv_bytes([["Tenant", "Rent PSF"], ["Acme Corp", "2.75"]])
    try:
        parse_csv_rent_roll(csv_bytes, "psf_only.csv")
        assert False, "should have raised -- Rent PSF must never be treated as rent_amount"
    except RentRollImportError:
        pass

    # With BOTH a PSF rate column and a real one, the real one must win.
    csv_bytes2 = _csv_bytes([
        ["Tenant", "Rent PSF", "Actual Rent"],
        ["Acme Corp", "2.75", "4400.00"],
    ])
    result = parse_csv_rent_roll(csv_bytes2, "psf_and_real.csv")
    assert result["leases"][0]["extracted_fields"]["rent_amount"]["value"] == "$4,400.00", \
        result["leases"][0]["extracted_fields"]["rent_amount"]
    print("✓ test_rent_psf_rate_column_never_used_as_rent_amount: PASS")


def test_property_manager_column_not_mistaken_for_property_address():
    """
    Regression test caught during self-review, not found broken in
    production: "Property Manager" (a real, common rent-roll column --
    the on-site PM's name) contains the bare "property" alias as a
    whole word, which pass 2's substring matching would otherwise grab.
    A file with BOTH a real "Property" column and a "Property Manager"
    column must use the real one for the address, not the person's name.
    """
    csv_bytes = _csv_bytes([
        ["Property", "Property Manager", "Tenant", "Unit", "Rent"],
        ["100 Alpha St", "Jane Smith", "Acme Corp", "101", "2000.00"],
    ])
    result = parse_csv_rent_roll(csv_bytes, "pm_test.csv")
    assert result["column_mapping"]["property"] == 0, result["column_mapping"]  # the real "Property" column, not "Property Manager"
    address = result["leases"][0]["extracted_fields"]["property_address"]["value"]
    assert address == "100 Alpha St, Suite 101", address
    print("✓ test_property_manager_column_not_mistaken_for_property_address: PASS")


def test_move_in_move_out_not_mistaken_for_lease_dates():
    """
    Move-in/move-out are occupancy dates (when a tenant physically took/
    vacated possession) -- a different real-world fact from the lease's
    own contractual start/end. Deliberately not aliased to lease_start_
    date/lease_end_date, so a file with ONLY these (no real lease-date
    columns) must leave those fields honestly unset, not silently
    conflate the two.
    """
    csv_bytes = _csv_bytes([
        ["Tenant", "Rent", "Move-in", "Move-out"],
        ["Acme Corp", "$2,000.00", "01/15/2024", "12/20/2026"],
    ])
    result = parse_csv_rent_roll(csv_bytes, "moveinout.csv")
    fields = result["leases"][0]["extracted_fields"]
    assert fields["lease_start_date"]["value"] is None, fields["lease_start_date"]
    assert fields["lease_end_date"]["value"] is None, fields["lease_end_date"]
    print("✓ test_move_in_move_out_not_mistaken_for_lease_dates: PASS")


if __name__ == "__main__":
    test_csv_happy_path_standard_headers()
    test_xlsx_happy_path_with_real_numeric_and_date_cell_types()
    test_header_alias_diversity()
    test_unit_column_already_spelled_out_does_not_double_prefix()
    test_bare_numeric_rent_with_no_dollar_sign()
    test_vacant_row_skipped()
    test_total_row_skipped()
    test_subtotal_row_with_extra_text_is_skipped()
    test_tenant_literal_zero_is_skipped()
    test_blank_tenant_row_skipped()
    test_missing_tenant_and_rent_columns_raises()
    test_missing_optional_column_does_not_block_import()
    test_unrecognized_extra_columns_ignored()
    test_blank_rows_in_file_are_skipped_entirely()
    test_no_base_address_and_no_unit_column_leaves_address_unset()
    test_empty_csv_raises()
    test_empty_xlsx_raises()
    test_xlsx_checks_every_sheet_not_just_the_active_one()
    test_row_numbers_reflect_actual_file_position()
    test_specific_date_column_not_misattributed_to_rent_amount()
    test_header_row_auto_detection_skips_decorative_rows()
    test_header_row_already_at_top_is_unaffected()
    test_no_header_found_within_scan_window_gives_clear_error()
    test_decorative_title_row_with_a_stray_recognizable_word_not_mistaken_for_header()
    test_header_row_auto_detection_works_for_xlsx_too()
    test_per_row_property_column_overrides_base_address_for_multi_property_exports()
    test_property_column_falls_back_to_base_address_when_a_row_is_blank()
    test_market_rent_is_never_used_as_rent_amount()
    test_yardi_appfolio_terminology_aliases()
    test_realpage_mri_terminology_aliases()
    test_rent_psf_rate_column_never_used_as_rent_amount()
    test_property_manager_column_not_mistaken_for_property_address()
    test_move_in_move_out_not_mistaken_for_lease_dates()
    print("\nAll rent roll import tests passed.")
