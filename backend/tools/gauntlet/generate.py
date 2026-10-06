"""
Overnight gauntlet fixture generator. Deterministic: same code -> same files.

    cd backend && venv/bin/python tools/gauntlet/generate.py

Writes backend/tools/gauntlet/fixtures/:
    leases/<prop>/<unit>.pdf      one signed lease per leased unit
    rent_rolls/<prop>__<fmt>.<ext> messy rent roll exports (+ bad files)
    t12/<prop>__<fmt>.<ext>        trailing-12 operating statements (+ bad files)
    manifest.json                  the correct answer for every file

All data is FICTIONAL (see truth.py). Nothing here touches a database.

Optional generation-only libraries (password-protected xlsx, legacy .xls)
are imported if available on PYTHONPATH; without them those fixtures are
kept from the previous run (they're committed) or skipped.
"""
import csv
import io
import json
import os
import random
import re
import shutil
import sys
import zipfile
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab import rl_config
rl_config.invariant = 1  # byte-identical PDFs on every run (no timestamps / random ids)
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import truth  # noqa: E402
from truth import AS_OF  # noqa: E402

FIX = os.path.join(HERE, "fixtures")
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
LONG_MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August",
               "September", "October", "November", "December"]
# Trailing-12 period ending the month before AS_OF: Oct 2025 .. Sep 2026.
PERIOD = [(2025, 10), (2025, 11), (2025, 12)] + [(2026, m) for m in range(1, 10)]

try:
    import msoffcrypto  # noqa: F401
    HAVE_CRYPTO = True
except Exception:
    HAVE_CRYPTO = False
try:
    import xlwt  # noqa: F401
    HAVE_XLWT = True
except Exception:
    HAVE_XLWT = False


# ---------------------------------------------------------------- helpers
def d(s):
    return date.fromisoformat(s) if s else None


def money(v, style="dollar"):
    if v is None:
        return ""
    if style == "plain":
        return f"{v:.2f}"
    if style == "int":
        return f"{v:.0f}"
    if style == "comma":
        return f"{v:,.2f}"
    if style == "paren":
        return f"({abs(v):,.2f})" if v < 0 else f"{v:,.2f}"
    if style == "dollar_paren":
        return f"(${abs(v):,.2f})" if v < 0 else f"${v:,.2f}"
    return f"-${abs(v):,.2f}" if v < 0 else f"${v:,.2f}"


def fdate(x, style="mdy"):
    if not x:
        return ""
    x = d(x) if isinstance(x, str) else x
    if style == "mdy":
        return x.strftime("%m/%d/%Y")
    if style == "mdy_short":
        return f"{x.month}/{x.day}/{x.strftime('%y')}"
    if style == "iso":
        return x.isoformat()
    if style == "dmon":
        return x.strftime("%d-%b-%Y")
    if style == "dmon2":
        return x.strftime("%d-%b-%y")
    if style == "long":
        return f"{LONG_MONTHS[x.month - 1]} {x.day}, {x.year}"
    raise ValueError(style)


def long_date(x):
    return fdate(x, "long")


def write(path, data):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    mode = "wb" if isinstance(data, (bytes, bytearray)) else "w"
    with open(path, mode) as f:
        f.write(data)


def csv_bytes(rows, encoding="utf-8", delimiter=","):
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=delimiter, lineterminator="\r\n")
    for r in rows:
        w.writerow(["" if c is None else c for c in r])
    text = buf.getvalue()
    if encoding == "utf-16":
        return text.encode("utf-16")  # with BOM, like Excel "Unicode Text"
    return text.encode(encoding)


_FIXED_TS = __import__("datetime").datetime(2026, 10, 5, 0, 0, 0)


def stable_zip(raw):
    """Rewrite an OOXML zip with fixed entry timestamps so regenerating doesn't churn the repo."""
    src = zipfile.ZipFile(io.BytesIO(raw))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:)", rb"\g<1>2026-10-05T00:00:00Z\g<2>", data)
            fixed = zipfile.ZipInfo(info.filename, date_time=(2026, 10, 5, 0, 0, 0))
            fixed.compress_type = zipfile.ZIP_DEFLATED
            fixed.external_attr = info.external_attr
            dst.writestr(fixed, data)
    return out.getvalue()


def xlsx_bytes(wb):
    wb.properties.created = _FIXED_TS
    wb.properties.modified = _FIXED_TS
    buf = io.BytesIO()
    wb.save(buf)
    return stable_zip(buf.getvalue())


def unit_cell(u, style="bare"):
    return u["uid"] if style == "bare" else f"{style} {u['uid']}"


def bed_bath(u):
    return f"{u['beds']}/{u['baths']}"


def rr_rows_units(prop):
    return truth.rr_units(prop)


# ---------------------------------------------------------------- leases
def lease_address(prop, u):
    st = prop["lease_address_style"]
    if st == "apt_inline":
        return f"{prop['street']}, Apt {u['uid']}, {prop['city']}, {prop['state']} {prop['zip']}"
    if st == "unit_inline":
        return f"{prop['street']}, Unit {u['uid']}, {prop['city']}, {prop['state']} {prop['zip']}"
    return f"{prop['base_address']}, Suite {u['uid']}"


def concession_text(u):
    c = u["concession"]
    s = d(u["start"])
    if c["kind"] == "free_month":
        return (f"One (1) month of Base Rent (the first full calendar month, {LONG_MONTHS[s.month - 1]} "
                f"{s.year}) is abated as a move-in incentive.")
    if c["months"] == 12:
        return (f"Base Rent is reduced by {money(c['amount'])} per month for the full twelve (12) month "
                "Lease Term as a retention incentive.")
    e = truth._month_add(s, c["months"] - 1)
    return (f"Base Rent is reduced by {money(c['amount'])} per month for the first six (6) months of the "
            f"Lease Term ({LONG_MONTHS[s.month - 1]} {s.year} through {LONG_MONTHS[e.month - 1]} {e.year}) "
            "as a leasing incentive.")


def lease_paragraphs(prop, u):
    owner = f"{prop['name']} Owner LLC"
    mgmt = "Fictional Residential Management Co."
    addr = lease_address(prop, u)
    beds = {1: "a one-bedroom apartment", 2: "a two-bedroom apartment", 3: "a three-bedroom apartment"}[u["beds"]]
    tm = truth.term_months(d(u["start"]), d(u["end"]))
    paras = [
        (mgmt, "title"), (f"{prop['name']}  |  {prop['base_address']}", "subtitle"), "", ("", "rule"),
        "RESIDENTIAL APARTMENT LEASE AGREEMENT", "",
        f"This Lease Agreement (\"Lease\") is entered into by and between {owner} (\"Owner\"), acting "
        f"through its agent {mgmt} (\"Management\"), and {u['tenant']} (\"Resident\"), for the apartment "
        "home described below.", "",
        f"1. PREMISES. Owner leases to Resident the apartment located at {addr} (the \"Premises\"), "
        f"consisting of {beds}, approximately {u['sqft']:,} square feet.", "",
        f"2. LEASE TERM. The Lease Term begins on {long_date(u['start'])} and ends on {long_date(u['end'])} "
        f"({tm} months), unless renewed or terminated as provided in this Lease.", "",
        f"3. RENT. Resident shall pay Owner Base Rent of {money(u['lease_rent'])} per month, due in advance "
        "on or before the 1st day of each month.",
    ]
    if u["s8"]:
        paras += ["", "3B. HOUSING ASSISTANCE. This tenancy is assisted under the Housing Choice Voucher "
                  f"program. Of the monthly amount in Section 3, the Public Housing Agency pays "
                  f"{money(u['s8']['hap'])} each month as a Housing Assistance Payment and Resident pays "
                  f"{money(u['s8']['tenant_portion'])}."]
    if u["concession"]:
        paras += ["", "3A. RENT CONCESSION. As a leasing incentive, and notwithstanding Section 3 above: "
                  f"{concession_text(u)} All other terms of this Lease remain unchanged."]
    paras += [
        "", "4. SECURITY DEPOSIT. Resident has paid a security deposit of $500.00, refundable as provided by law.",
        "", "5. LATE CHARGES. Rent not received by the 3rd day of the month incurs a late charge of $50.00.",
        "", "6. UTILITIES. Resident is responsible for electricity and internet service.",
        "", "7. USE. The Premises shall be used solely as a private residence.",
        "", f"IN WITNESS WHEREOF, the parties have executed this Lease as of {long_date(u['start'])}.",
        "", ("", "rule"), "", "OWNER:", owner, f"By: {mgmt}, as authorized agent", "",
        "Signature: ______________________        Date: ____________", "", "RESIDENT:", u["tenant"], "",
        "Signature: ______________________        Date: ____________",
    ]
    return paras


_STY = {"title": ("Helvetica-Bold", 15), "subtitle": ("Helvetica", 10), "small": ("Helvetica-Oblique", 8)}


def _wrap(c, text, font, size, max_w):
    words, lines, cur = text.split(" "), [], ""
    for w in words:
        t = f"{cur} {w}".strip()
        if c.stringWidth(t, font, size) <= max_w:
            cur = t
        else:
            lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines


def render_paragraph_pdf(paragraphs, footer):
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    c.setTitle("Lease")
    W, H = letter
    left, right = inch, W - inch
    y = H - inch
    page = 1

    def foot():
        c.setFont("Helvetica", 8)
        c.drawString(left, 0.6 * inch, footer)
        c.drawRightString(right, 0.6 * inch, f"Page {page}")

    for item in paragraphs:
        text, style = item if isinstance(item, tuple) else (item, None)
        if style == "rule":
            c.line(left, y, right, y)
            y -= 0.2 * inch
            continue
        if text == "":
            y -= 0.14 * inch
            continue
        font, size = _STY.get(style, ("Helvetica-Bold", 11) if text.isupper() else ("Helvetica", 10.5))
        for line in _wrap(c, text, font, size, right - left):
            if y < 1.1 * inch:
                foot()
                c.showPage()
                page += 1
                y = H - inch
            c.setFont(font, size)
            c.drawString(left, y, line)
            y -= size * 1.35
    foot()
    c.save()
    return buf.getvalue()


def write_leases(prop):
    out = []
    for u in truth.lease_units(prop):
        pdf = render_paragraph_pdf(lease_paragraphs(prop, u), f"{prop['name']} -- Unit {u['uid']} -- FICTIONAL TEST DATA")
        rel = f"leases/{prop['id']}/{u['uid']}.pdf"
        write(os.path.join(FIX, rel), pdf)
        out.append({"file": rel, "unit": u["uid"], "tenant": u["tenant"], "rent": u["lease_rent"],
                    "start": u["start"], "end": u["end"], "address": lease_address(prop, u),
                    "concession": u["concession"]})
    return out


# ---------------------------------------------------------------- rent roll writers
# Each returns (ext, bytes, carries:set, note:str). The manifest's expected
# answer for the file comes from truth.expected_rent_roll/expected_findings
# with the same `carries`.

def rr_appfolio_csv(prop, rng):
    rows = [["Rent Roll"], [f"Properties: {prop['name']} - {prop['street']} {prop['city']}, {prop['state']} {prop['zip']}"],
            [f"As of: {fdate(AS_OF)}"], [],
            ["Unit", "BD/BA", "Tenant", "Status", "Sqft", "Market Rent", "Rent", "Deposit", "Lease From", "Lease To",
             "Move-in", "Move-out", "Past Due"]]
    tot = 0.0
    for u in rr_rows_units(prop):
        if u["status"] == "occupied":
            tot += u["rr_rent"]
            rows.append([u["uid"], bed_bath(u), u["rr_tenant"], "Current", u["sqft"], money(u["market_rent"]),
                         money(u["rr_rent"]), "$500.00", fdate(u["rr_start"]), fdate(u["rr_end"]),
                         fdate(u["rr_start"]), "", "$0.00"])
        else:
            rows.append([u["uid"], bed_bath(u), "VACANT" if u["status"] == "vacant" else "",
                         "Vacant-Unrented" if u["status"] == "vacant" else "Down", u["sqft"],
                         money(u["market_rent"]), "", "", "", "", "", "", ""])
    rows.append(["Total", "", "", "", "", "", money(tot), "", "", "", "", "", ""])
    return "csv", csv_bytes(rows), {"start", "end"}, "AppFolio CSV: title block, VACANT/Down rows, Total row"


def _appfolio_wb(prop, dup=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Rent Roll"
    ws.append(["Rent Roll"])
    ws.append([f"Properties: {prop['name']}"])
    ws.append([f"As of: {fdate(AS_OF)}"])
    ws.append([])
    ws.append(["Unit", "BD/BA", "Tenant", "Status", "Sqft", "Market Rent", "Rent", "Deposit", "Lease From",
               "Lease To", "Move-in", "Move-out", "Past Due"])
    units = rr_rows_units(prop)
    for u in units:
        if u["status"] == "occupied":
            r = [u["uid"], bed_bath(u), u["rr_tenant"], "Current", u["sqft"], u["market_rent"], u["rr_rent"], 500.0,
                 d(u["rr_start"]), d(u["rr_end"]), d(u["rr_start"]), None, 0.0]
            ws.append(r)
            if dup and u is [x for x in units if x["status"] == "occupied"][1]:
                ws.append(r)  # export glitch: exact duplicate row
        else:
            ws.append([u["uid"], bed_bath(u), "VACANT" if u["status"] == "vacant" else None,
                       "Vacant-Unrented" if u["status"] == "vacant" else "Down", u["sqft"], u["market_rent"],
                       None, None, None, None, None, None, None])
    for row in ws.iter_rows(min_row=6):
        for c in row:
            if c.column in (6, 7, 8, 13):
                c.number_format = '"$"#,##0.00'
            if c.column in (9, 10, 11):
                c.number_format = "mm/dd/yyyy"
    ws.append([])
    ws.append(["Total", None, None, None, None, None,
               sum(u["rr_rent"] for u in truth.rr_occupied(prop))])
    return wb


def rr_appfolio_xlsx(prop, rng):
    return "xlsx", xlsx_bytes(_appfolio_wb(prop)), {"start", "end"}, "AppFolio xlsx: real date cells, $ formats"


def rr_appfolio_xlsx_dup(prop, rng):
    return "xlsx", xlsx_bytes(_appfolio_wb(prop, dup=True)), {"start", "end"}, "AppFolio xlsx with one exact duplicate unit row (export glitch) -- must count once"


def rr_yardi_xlsx(prop, rng):
    """Yardi 'Rent Roll': merged title block, TWO-ROW header, Last, First names, no lease start column."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Report1"
    ws["A1"] = prop["name"]
    ws.merge_cells("A1:L1")
    ws["A2"] = "Rent Roll"
    ws.merge_cells("A2:L2")
    ws["A3"] = f"As Of = {fdate(AS_OF)}"
    ws["A4"] = f"Month Year = {AS_OF.month:02d}/{AS_OF.year}"
    for cell in ("A1", "A2"):
        ws[cell].alignment = Alignment(horizontal="center")
        ws[cell].font = Font(bold=True)
    h1 = ["Unit", "Unit", "Unit", "Resident", "Name", "Market", "Actual", "Resident", "Other", "Move In", "Lease", "Move Out", "Balance"]
    h2 = [None, "Type", "Sq Ft", None, None, "Rent", "Rent", "Deposit", "Deposit", None, "Expiration", None, None]
    ws.append([])
    ws.append(h1)
    ws.append(h2)
    for u in rr_rows_units(prop):
        if u["status"] == "occupied":
            nm = u["rr_tenant"]
            if u["trap"] != "last_first_name" and "," not in nm:
                parts = nm.split(" ")
                if len(parts) == 2:
                    nm = f"{parts[1]}, {parts[0]}"
            ws.append([u["uid"], f"{u['beds']}x{u['baths']}", u["sqft"], f"t{rng.randint(100000, 999999):07d}", nm,
                       u["market_rent"], u["rr_rent"], 500.0, 0.0, d(u["rr_start"]), d(u["rr_end"]), None, 0.0])
        else:
            ws.append([u["uid"], f"{u['beds']}x{u['baths']}", u["sqft"], "VACANT", "VACANT" if u["status"] == "vacant" else "DOWN",
                       u["market_rent"], 0.0, 0.0, 0.0, None, None, None, 0.0])
    ws.append([])
    ws.append(["Summary Groups", None, None, None, None, "Square Footage", "Market Rent", "Leased Units"])
    ws.append(["Current/Notice/Vacant Residents", None, sum(u["sqft"] for u in rr_rows_units(prop)), None, None, None,
               sum(u["market_rent"] for u in rr_rows_units(prop)), len(truth.rr_occupied(prop))])
    for row in ws.iter_rows(min_row=8):
        for c in row:
            if c.column in (10, 11):
                c.number_format = "mm/dd/yyyy"
    return "xlsx", xlsx_bytes(wb), {"end"}, "Yardi xlsx: merged title, two-row header, 'Last, First' names (name order must not be a tenant mismatch)"


def rr_yardi_charges_xlsx(prop, rng):
    """Yardi 'Rent Roll with Lease Charges': one row per charge code, unit info only on the first."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Rent Roll with Lease Charges"
    ws.append([prop["name"]])
    ws.append(["Rent Roll with Lease Charges"])
    ws.append([f"As Of = {fdate(AS_OF)}"])
    ws.append(["Unit", "Unit Type", "Unit Sq Ft", "Resident", "Name", "Market Rent", "Charge Code", "Amount",
               "Resident Deposit", "Move In", "Lease Expiration", "Move Out", "Balance"])
    for u in rr_rows_units(prop):
        if u["status"] == "occupied":
            ws.append([u["uid"], f"{u['beds']}x{u['baths']}", u["sqft"], f"t{rng.randint(100000, 999999):07d}", u["rr_tenant"],
                       u["market_rent"], "rent", u["rr_rent"], 500.0, d(u["rr_start"]), d(u["rr_end"]), None, 0.0])
            ws.append([None, None, None, None, None, None, "trash", 15.0])
            if u["rr_concession"]:
                ws.append([None, None, None, None, None, None, "conc", -u["rr_concession"]])
            ws.append([None, None, None, None, None, None, "Total", u["rr_rent"] + 15.0 - (u["rr_concession"] or 0)])
        else:
            ws.append([u["uid"], f"{u['beds']}x{u['baths']}", u["sqft"], "VACANT", "VACANT", u["market_rent"], None, None,
                       None, None, None, None, None])
    return "xlsx", xlsx_bytes(wb), {"end", "concession"}, "Yardi 'with Lease Charges': one row per charge code (rent/trash/conc/Total); rent = the 'rent' code only"


def _realpage_rows(prop, rng, money_style):
    hdr = ["Bldg/Unit", "Floorplan", "SQFT", "Unit/Lease Status", "Name", "Phone Number", "Move-In", "Move-Out",
           "Lease Start", "Lease End", "Market + Addl.", "Lease Rent", "Other Charges/Credits", "Total Billing",
           "Dep On Hand", "Balance"]
    rows = [["RealPage OneSite"], [f"{prop['name']} - Rent Roll Detail"], [f"As of {fdate(AS_OF)}"], [], hdr]
    for u in rr_rows_units(prop):
        if u["status"] == "occupied":
            other = -u["rr_concession"] if u["rr_concession"] else 25.0
            rows.append([u["uid"], f"{u['beds']}BR", u["sqft"], "Occupied", u["rr_tenant"], f"555-01{rng.randint(10, 99)}",
                         fdate(u["rr_start"]), "", fdate(u["rr_start"]), fdate(u["rr_end"]),
                         money(u["market_rent"], money_style), money(u["rr_rent"], money_style), money(other, money_style),
                         money(u["rr_rent"] + other, money_style), money(500.0, money_style), money(0.0, money_style)])
        else:
            rows.append([u["uid"], f"{u['beds']}BR", u["sqft"], "Vacant-Unrented" if u["status"] == "vacant" else "Down Unit",
                         "VACANT" if u["status"] == "vacant" else "", "", "", "", "", "", money(u["market_rent"], money_style),
                         "", "", "", "", ""])
    return rows


def rr_realpage_xlsx(prop, rng):
    wb = Workbook()
    ws = wb.active
    ws.title = "Rent Roll Detail"
    for r in _realpage_rows(prop, rng, "dollar"):
        ws.append(r)
    return "xlsx", xlsx_bytes(wb), {"start", "end"}, "RealPage OneSite: 'Name' header (no 'Tenant'), 'Lease Rent' vs 'Market + Addl.', text money"


def rr_realpage_csv(prop, rng):
    return "csv", csv_bytes(_realpage_rows(prop, rng, "dollar_paren")), {"start", "end"}, "RealPage CSV: parentheses negatives in Other Charges/Credits"


def rr_entrata_csv(prop, rng):
    rows = [[f"{prop['name']}"], ["Rent Roll"], [f"Report Date: {fdate(AS_OF, 'iso')}"],
            ["Bldg-Unit", "Floor Plan", "Resident", "Status", "Lease Start", "Lease End", "Market Rent",
             "Scheduled Rent", "Concession", "Balance"]]
    styles = ["mdy", "iso", "mdy_short", "dmon"]
    for i, u in enumerate(rr_rows_units(prop)):
        ds = styles[i % len(styles)]
        if u["status"] == "occupied":
            rows.append([u["uid"], f"{u['beds']}x{u['baths']}", u["rr_tenant"], "Current", fdate(u["rr_start"], ds),
                         fdate(u["rr_end"], ds), money(u["market_rent"], "comma"), money(u["rr_rent"], "comma"),
                         money(-u["rr_concession"], "paren") if u["rr_concession"] else "0.00", "0.00"])
        else:
            rows.append([u["uid"], f"{u['beds']}x{u['baths']}", "", "Vacant" if u["status"] == "vacant" else "Down/Model",
                         "", "", money(u["market_rent"], "comma"), "", "", ""])
    rows.append([])
    rows.append(["Totals:", "", "", "", "", "", "", money(sum(u["rr_rent"] for u in truth.rr_occupied(prop)), "comma"), "", ""])
    return "csv", csv_bytes(rows), {"start", "end", "concession"}, "Entrata CSV: mixed date formats per row, concession as (75.00), Totals row"


def rr_broker_xlsx(prop, rng):
    """Broker workbook: Summary sheet first, merged title, floor group/subtotal rows, footnotes, extra sheet."""
    wb = Workbook()
    s = wb.active
    s.title = "Summary"
    s.append(["Property", prop["name"]])
    s.append(["Units", len(rr_rows_units(prop))])
    s.append(["Occupied", len(truth.rr_occupied(prop))])
    s.append(["Monthly Rent", sum(u["rr_rent"] for u in truth.rr_occupied(prop))])
    ws = wb.create_sheet("Rent Roll")
    ws["A1"] = f"{prop['name']} -- Rent Roll (Broker Prepared)"
    ws.merge_cells("A1:H1")
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Prepared {fdate(AS_OF)} -- information deemed reliable but not guaranteed"
    ws.merge_cells("A2:H2")
    ws.append([])
    ws.append(["Unit #", "Tenant Name", "Unit Type", "SF", "Monthly Rent", "Lease Start", "Lease End", "Notes"])
    hdr_row = ws.max_row
    for c in ws[hdr_row]:
        c.font = Font(bold=True)
        c.fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
    units = rr_rows_units(prop)
    by_floor = {}
    for u in units:
        fl = [ch for ch in u["uid"] if ch.isdigit()]
        key = fl[-3] if len(fl) >= 3 else "1"
        by_floor.setdefault(key, []).append(u)
    for fl, us in sorted(by_floor.items()):
        ws.append([f"Floor {fl}"])
        ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=8)
        for u in us:
            if u["status"] == "occupied":
                note = "Section 8" if u["s8"] else ("MTM" if d(u["rr_end"]) < AS_OF else None)
                ws.append([u["uid"], u["rr_tenant"], f"{u['beds']}BR/{u['baths']}BA", u["sqft"], u["rr_rent"],
                           d(u["rr_start"]), d(u["rr_end"]), note])
            else:
                ws.append([u["uid"], "Vacant" if u["status"] == "vacant" else "Down - renovation",
                           f"{u['beds']}BR/{u['baths']}BA", u["sqft"], None, None, None, None])
        ws.append([None, f"Subtotal Floor {fl}", None, sum(u["sqft"] for u in us), sum(u["rr_rent"] or 0 for u in us)])
        ws.append([])
    ws.append([None, "Grand Total", None, sum(u["sqft"] for u in units), sum(u["rr_rent"] or 0 for u in units)])
    ws.append([])
    ws.append(["* Rents shown are contract rents before concessions."])
    ws.append(["Source: owner records, unaudited."])
    for row in ws.iter_rows(min_row=hdr_row + 1):
        for c in row:
            if c.column == 5:
                c.number_format = '"$"#,##0.00'
            if c.column in (6, 7):
                c.number_format = "m/d/yyyy"
    a = wb.create_sheet("Assumptions")
    a.append(["Rent growth", 0.03])
    a.append(["Vacancy", 0.05])
    return "xlsx", xlsx_bytes(wb), {"start", "end"}, "Broker workbook: Summary sheet first, merged title + floor rows, subtotals, footnotes, Assumptions sheet"


def rr_broker_csv_messy(prop, rng):
    rows = [[f"{prop['name'].upper()} RENT ROLL"], [], ["Unit No", "Tenant Nmae", "Montly Rent", "Lease Strat",
                                                          "Lease Expiraton", "Sq. Ft."]]
    dstyles = ["mdy_short", "iso", "dmon2", "long", "mdy"]
    mstyles = ["dollar", "int", "comma", "plain"]
    for i, u in enumerate(rr_rows_units(prop)):
        if i % 4 == 2:
            rows.append([])
        if u["status"] == "occupied":
            rows.append([u["uid"], u["rr_tenant"], " " + money(u["rr_rent"], mstyles[i % 4]) + " ",
                         fdate(u["rr_start"], dstyles[i % 5]), fdate(u["rr_end"], dstyles[(i + 2) % 5]), f"{u['sqft']:,}"])
        else:
            rows.append([u["uid"], "VACANT" if u["status"] == "vacant" else "DOWN UNIT", "", "", "", f"{u['sqft']:,}"])
    rows.append([])
    rows.append(["TOTAL", "", money(sum(u["rr_rent"] for u in truth.rr_occupied(prop))), "", "", ""])
    rows.append(["Note: unit rents per owner; verify against leases."])
    return "csv", csv_bytes(rows), {"start", "end"}, "Broker CSV: misspelled headers (Tenant Nmae / Montly Rent / Lease Strat / Lease Expiraton), mixed money + date formats, blank rows, TOTAL + note rows"


def rr_s8_xlsx(prop, rng):
    wb = Workbook()
    ws = wb.active
    ws.title = "Affordable Rent Roll"
    ws.append([f"{prop['name']} - Affordable Rent Roll"])
    ws.append(["Unit", "Resident", "Contract Rent", "Tenant Rent", "HAP Amount", "Utility Allowance", "Lease Start", "Lease End", "Cert Date"])
    for u in rr_rows_units(prop):
        if u["status"] == "occupied":
            tp = u["s8"]["tenant_portion"] if u["s8"] else u["rr_rent"]
            hap = u["s8"]["hap"] if u["s8"] else 0.0
            # Overstated contract rents keep HAP fixed: the difference lands in the tenant portion.
            if u["s8"]:
                tp = u["rr_rent"] - hap
            ws.append([u["uid"], u["rr_tenant"], u["rr_rent"], tp, hap, 45.0 if u["s8"] else 0.0, d(u["rr_start"]), d(u["rr_end"]),
                       d(u["rr_start"])])
        else:
            ws.append([u["uid"], "VACANT", None, None, None, None, None, None, None])
    return "xlsx", xlsx_bytes(wb), {"start", "end"}, "Section 8 rent roll: Contract Rent + Tenant Rent + HAP columns; rent = contract rent"


def rr_s8_split_csv(prop, rng):
    rows = [["Unit", "Resident Name", "Tenant Portion", "HAP", "Lease Begin", "Lease End"]]
    for u in rr_rows_units(prop):
        if u["status"] == "occupied":
            hap = u["s8"]["hap"] if u["s8"] else 0.0
            rows.append([u["uid"], u["rr_tenant"], money(u["rr_rent"] - hap, "plain"), money(hap, "plain"),
                         fdate(u["rr_start"]), fdate(u["rr_end"])])
        else:
            rows.append([u["uid"], "VACANT", "", "", "", ""])
    return "csv", csv_bytes(rows), {"start", "end"}, "Section 8 split with NO total column: Tenant Portion + HAP must be added (contract rent = sum)"


def _simple_rows(prop, money_style="dollar", date_style="mdy"):
    rows = [["Unit", "Tenant", "Rent", "Lease Start", "Lease End"]]
    for u in rr_rows_units(prop):
        if u["status"] == "occupied":
            rows.append([u["uid"], u["rr_tenant"], money(u["rr_rent"], money_style), fdate(u["rr_start"], date_style),
                         fdate(u["rr_end"], date_style)])
        else:
            rows.append([u["uid"], "VACANT", "", "", ""])
    return rows


def rr_clean_csv(prop, rng):
    return "csv", csv_bytes(_simple_rows(prop)), {"start", "end"}, "Clean simple CSV (baseline)"


def rr_cp1252_csv(prop, rng):
    return "csv", csv_bytes(_simple_rows(prop), encoding="cp1252"), {"start", "end"}, "Windows-1252 encoded CSV (Excel 'CSV' on Windows) -- accented names must survive"


def rr_utf16_txt(prop, rng):
    return "txt", csv_bytes(_simple_rows(prop, "comma"), encoding="utf-16", delimiter="\t"), {"start", "end"}, "Excel 'Unicode Text' export: UTF-16 with BOM, tab-delimited, .txt"


def _rr_table_pdf(prop, title_rows=True):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(letter), leftMargin=0.5 * inch, rightMargin=0.5 * inch)
    styles = getSampleStyleSheet()
    data = [["Unit", "Tenant", "Status", "Sq Ft", "Market Rent", "Rent", "Lease Start", "Lease End"]]
    for u in rr_rows_units(prop):
        if u["status"] == "occupied":
            data.append([u["uid"], u["rr_tenant"], "Occupied", f"{u['sqft']:,}", money(u["market_rent"]), money(u["rr_rent"]),
                         fdate(u["rr_start"]), fdate(u["rr_end"])])
        else:
            data.append([u["uid"], "VACANT", "Vacant" if u["status"] == "vacant" else "Down", f"{u['sqft']:,}",
                         money(u["market_rent"]), "", "", ""])
    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, 0), "Helvetica-Bold"), ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                           ("FONTSIZE", (0, 0), (-1, -1), 10)]))
    story = []
    if title_rows:
        story += [Paragraph(f"{prop['name']} Rent Roll", styles["Title"]), Paragraph(f"As of {fdate(AS_OF)}", styles["Normal"]), Spacer(1, 12)]
    story.append(t)
    doc.build(story)
    return buf.getvalue()


def rr_pdf_text(prop, rng):
    return "pdf", _rr_table_pdf(prop), {"start", "end"}, "Digital PDF rent roll (real text layer, gridded table)"


def _scan(pdf_bytes, rotate=0, skew=0.0, dpi=200, noise=True):
    from pdf2image import convert_from_bytes
    from PIL import Image, ImageFilter
    pages = convert_from_bytes(pdf_bytes, dpi=dpi)
    out = []
    r = random.Random(7)
    for p in pages:
        p = p.convert("L")
        if skew:
            p = p.rotate(skew, expand=True, fillcolor=255)
        if rotate:
            p = p.rotate(rotate, expand=True, fillcolor=255)
        if noise:
            px = p.load()
            w, h = p.size
            for _ in range(w * h // 400):
                x, y = r.randrange(w), r.randrange(h)
                px[x, y] = r.choice([0, 90, 160])
            p = p.filter(ImageFilter.GaussianBlur(0.4))
        out.append(p.convert("RGB"))
    buf = io.BytesIO()
    out[0].save(buf, format="PDF", save_all=True, append_images=out[1:], resolution=dpi)
    # PIL stamps the current time; pin it (same byte length keeps the xref valid).
    return re.sub(rb"D:\d{14}", b"D:20261005000000", buf.getvalue())


def rr_pdf_scanned(prop, rng):
    return "pdf", _scan(_rr_table_pdf(prop)), {"start", "end"}, "Scanned PDF (image only, 200dpi, speckle noise) -- OCR path"


def rr_pdf_scanned_rotated(prop, rng):
    return "pdf", _scan(_rr_table_pdf(prop), rotate=90, skew=1.2), {"start", "end"}, "Scanned PDF fed sideways (rotated 90 deg + 1.2 deg skew) -- OCR must detect orientation"


def rr_xls(prop, rng):
    if not HAVE_XLWT:
        return None
    import xlwt
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Rent Roll")
    datefmt = xlwt.easyxf(num_format_str="MM/DD/YYYY")
    rows = [["Unit", "Tenant", "Rent", "Lease Start", "Lease End"]]
    ws.write(0, 0, f"{prop['name']} rent roll")
    for ci, h in enumerate(rows[0]):
        ws.write(2, ci, h)
    r = 3
    for u in rr_rows_units(prop):
        if u["status"] == "occupied":
            ws.write(r, 0, u["uid"])
            ws.write(r, 1, u["rr_tenant"])
            ws.write(r, 2, u["rr_rent"])
            ws.write(r, 3, d(u["rr_start"]), datefmt)
            ws.write(r, 4, d(u["rr_end"]), datefmt)
        else:
            ws.write(r, 0, u["uid"])
            ws.write(r, 1, "VACANT")
        r += 1
    buf = io.BytesIO()
    wb.save(buf)
    return "xls", buf.getvalue(), {"start", "end"}, "Legacy Excel 97-2003 .xls with real date cells"


def rr_docx(prop, rng):
    from docx import Document
    doc = Document()
    doc.add_heading(f"{prop['name']} Rent Roll", 1)
    doc.add_paragraph(f"As of {fdate(AS_OF)}. Prepared by ownership.")
    rows = _simple_rows(prop)
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            t.cell(i, j).text = str(v)
    doc.core_properties.created = _FIXED_TS
    doc.core_properties.modified = _FIXED_TS
    buf = io.BytesIO()
    doc.save(buf)
    return "docx", stable_zip(buf.getvalue()), {"start", "end"}, "Word document with a rent roll table"


def rr_merged_roommates_xlsx(prop, rng):
    """Unit + rent cells merged vertically over a roommate's own row (two residents, one lease)."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Unit", "Tenant", "Rent", "Lease Start", "Lease End"])
    first_occ = [u for u in rr_rows_units(prop) if u["status"] == "occupied"][0]
    for u in rr_rows_units(prop):
        if u["status"] == "occupied":
            ws.append([u["uid"], u["rr_tenant"], u["rr_rent"], d(u["rr_start"]), d(u["rr_end"])])
            if u is first_occ:
                ws.append([None, "Roommate: Pat Fernsby", None, None, None])
                r = ws.max_row
                ws.merge_cells(start_row=r - 1, start_column=1, end_row=r, end_column=1)
                ws.merge_cells(start_row=r - 1, start_column=3, end_row=r, end_column=3)
        else:
            ws.append([u["uid"], "VACANT", None, None, None])
    return "xlsx", xlsx_bytes(wb), {"start", "end"}, "Merged Unit/Rent cells spanning a roommate row -- roommate is not a second lease"


RR_WRITERS = {
    "appfolio_csv": rr_appfolio_csv, "appfolio_xlsx": rr_appfolio_xlsx, "appfolio_xlsx_dup": rr_appfolio_xlsx_dup,
    "yardi_xlsx": rr_yardi_xlsx, "yardi_charges_xlsx": rr_yardi_charges_xlsx,
    "realpage_xlsx": rr_realpage_xlsx, "realpage_csv": rr_realpage_csv, "entrata_csv": rr_entrata_csv,
    "broker_xlsx": rr_broker_xlsx, "broker_csv_messy": rr_broker_csv_messy,
    "s8_xlsx": rr_s8_xlsx, "s8_split_csv": rr_s8_split_csv,
    "clean_csv": rr_clean_csv, "cp1252_csv": rr_cp1252_csv, "utf16_txt": rr_utf16_txt,
    "pdf_text": rr_pdf_text, "pdf_scanned": rr_pdf_scanned, "pdf_scanned_rotated": rr_pdf_scanned_rotated,
    "xls": rr_xls, "docx": rr_docx, "merged_roommates_xlsx": rr_merged_roommates_xlsx,
}

# Which formats each property is rendered in. Every property also gets
# clean_csv (the canonical file the T-12 cases import).
RR_PLAN = [
    ["appfolio_csv", "yardi_xlsx", "broker_csv_messy", "pdf_text"],
    ["appfolio_xlsx", "realpage_xlsx", "yardi_xlsx", "entrata_csv"],
    ["s8_xlsx", "s8_split_csv", "realpage_csv", "broker_xlsx"],
    ["entrata_csv", "appfolio_xlsx_dup", "docx", "utf16_txt"],
    ["cp1252_csv", "broker_xlsx", "yardi_charges_xlsx", "pdf_scanned"],
    ["entrata_csv", "realpage_xlsx", "appfolio_csv", "xls"],
    ["s8_xlsx", "yardi_xlsx", "merged_roommates_xlsx", "pdf_text"],
    ["appfolio_csv", "broker_csv_messy", "realpage_csv", "yardi_charges_xlsx"],
    ["cp1252_csv", "entrata_csv", "broker_xlsx", "pdf_scanned_rotated"],
    ["appfolio_xlsx", "yardi_xlsx", "utf16_txt", "docx"],
    ["realpage_xlsx", "appfolio_xlsx_dup", "broker_csv_messy", "xls"],
    ["s8_split_csv", "s8_xlsx", "entrata_csv", "pdf_text"],
    ["appfolio_csv", "realpage_csv", "merged_roommates_xlsx", "yardi_charges_xlsx"],
    ["broker_xlsx", "appfolio_xlsx", "entrata_csv", "pdf_scanned"],
]


# ---------------------------------------------------------------- T-12
def t12_truth(prop, scenario, rng):
    """
    Monthly line items for Oct 2025..Sep 2026 (rounded to cents), plus the
    expected T-12 rows of the Deal Mismatch Report under `scenario`:
      clean     collections within 1% of the rent roll, occupancy consistent
      gap       collections 14% below the rent roll -> t12_income_gap
      occ       T-12 vacancy implies ~98% occupancy vs a rent roll with vacancies
      baddebt   bad debt rising over the LAST 3 months, full occupancy
      baddebt_old  bad debt high in the FIRST 3 months, falling -- no trend
    """
    units = truth.rr_units(prop)
    occ = truth.rr_occupied(prop)
    total_units = len(units)
    rr_monthly = sum(u["rr_rent"] for u in occ)
    rr_annual = round(rr_monthly * 12, 2)
    gpr_m = sum(u["market_rent"] for u in units)
    rr_occ = len(occ) / total_units
    conc_m = 0.0
    for u in occ:
        if u["concession"]:
            tot, tm = truth.concession_value(u)
            conc_m += tot / tm
    net_target = rr_annual * (0.86 if scenario == "gap" else 0.992)
    vac_frac = (1 - rr_occ) if scenario != "occ" else 0.02
    lines = {k: [] for k in ["gpr", "ltl", "vacancy", "concessions", "bad_debt", "net", "utility", "late", "pet",
                             "other_total", "egi", "payroll", "rm", "taxes", "insurance", "admin", "opex", "noi"]}
    for i in range(12):
        gpr = round(gpr_m, 2)
        vac = round(gpr_m * vac_frac, 2)
        con = round(conc_m, 2)
        if scenario == "baddebt":
            bd = round([150, 140, 160, 150, 145, 155, 150, 160, 150, 900, 1100, 1250][i] * total_units / 10, 2)
        elif scenario == "baddebt_old":
            bd = round([1250, 1100, 900, 150, 140, 160, 150, 145, 155, 150, 160, 150][i] * total_units / 10, 2)
        else:
            bd = round(rng.uniform(100, 200) * total_units / 10, 2)
        net = round(net_target / 12 + rng.uniform(-40, 40), 2)
        ltl = round(gpr - vac - con - bd - net, 2)
        util, late, pet = round(32.0 * total_units, 2), round(rng.uniform(80, 160), 2), round(18.0 * total_units, 2)
        other = round(util + late + pet, 2)
        egi = round(net + other, 2)
        payroll, rm, taxes, ins, adm = (round(95.0 * total_units, 2), round(rng.uniform(60, 90) * total_units, 2),
                                        round(160.0 * total_units, 2), round(45.0 * total_units, 2), round(25.0 * total_units, 2))
        opex = round(payroll + rm + taxes + ins + adm, 2)
        for k, v in [("gpr", gpr), ("ltl", -ltl), ("vacancy", -vac), ("concessions", -con), ("bad_debt", -bd), ("net", net),
                     ("utility", util), ("late", late), ("pet", pet), ("other_total", other), ("egi", egi),
                     ("payroll", payroll), ("rm", rm), ("taxes", taxes), ("insurance", ins), ("admin", adm),
                     ("opex", opex), ("noi", round(egi - opex, 2))]:
            lines[k].append(v)
    ann = {k: round(sum(v), 2) for k, v in lines.items()}

    expected_rows = []
    gap = rr_annual - ann["net"]
    if gap > 0 and gap / max(rr_annual, ann["net"]) * 100 > 3.0:
        expected_rows.append({"type": "t12_income_gap", "annual": round(gap, 2)})
    t12_occ = (1 - abs(ann["vacancy"]) / ann["gpr"]) * 100
    if abs(rr_occ * 100 - t12_occ) > 3.0:
        expected_rows.append({"type": "t12_occupancy_mismatch", "annual": None})
    if round(abs(ann["concessions"]), 2) != 0:
        expected_rows.append({"type": "t12_concession_gap", "annual": None})
    bd = [abs(x) for x in lines["bad_debt"]]
    avg_all, avg_last3 = sum(bd) / 12, sum(bd[-3:]) / 3
    if avg_all > 0 and (avg_last3 - avg_all) / avg_all * 100 > 25 and len(occ) == total_units:
        expected_rows.append({"type": "t12_bad_debt_trend", "annual": None})
    parsed = {
        "gross_potential_rent": ann["gpr"], "vacancy_loss": abs(ann["vacancy"]), "concessions": abs(ann["concessions"]),
        "bad_debt": abs(ann["bad_debt"]), "rental_income_collected": ann["net"], "other_income": ann["other_total"],
    }
    return lines, ann, parsed, expected_rows, rr_annual


T12_LABELS = [
    ("INCOME", None, "section"),
    ("Rental Income", None, "section"),
    ("Gross Potential Rent", "gpr", None),
    ("Loss to Lease", "ltl", None),
    ("Vacancy Loss", "vacancy", None),
    ("Concessions", "concessions", None),
    ("Bad Debt", "bad_debt", None),
    ("Net Rental Income", "net", "total"),
    ("Other Income", None, "section"),
    ("Utility Reimbursement", "utility", None),
    ("Late Fees", "late", None),
    ("Pet Fees", "pet", None),
    ("Total Other Income", "other_total", "total"),
    ("EFFECTIVE GROSS INCOME", "egi", "total"),
    ("OPERATING EXPENSES", None, "section"),
    ("Payroll", "payroll", None),
    ("Repairs & Maintenance", "rm", None),
    ("Real Estate Taxes", "taxes", None),
    ("Insurance", "insurance", None),
    ("General & Administrative", "admin", None),
    ("Total Operating Expenses", "opex", "total"),
    ("NET OPERATING INCOME", "noi", "total"),
]


def month_headers(style):
    out = []
    for y, m in PERIOD:
        if style == "mon_year":
            out.append(f"{MONTHS[m - 1]} {y}")
        elif style == "mon":
            out.append(MONTHS[m - 1])
        elif style == "mon_dash_yy":
            out.append(f"{MONTHS[m - 1]}-{str(y)[2:]}")
        elif style == "mm_yyyy":
            out.append(f"{m:02d}/{y}")
        elif style == "long_year":
            out.append(f"{LONG_MONTHS[m - 1]} {y}")
    return out


def _t12_table(prop, lines, mstyle, money_style, total_col=True, label_header="Account"):
    hdr = [label_header] + month_headers(mstyle) + (["Total"] if total_col else [])
    rows = [hdr]
    for label, key, kind in T12_LABELS:
        if key is None:
            rows.append([label] + [None] * (len(hdr) - 1))
            continue
        vals = lines[key]
        cells = [money(v, money_style) if money_style else v for v in vals]
        if total_col:
            t = round(sum(vals), 2)
            cells.append(money(t, money_style) if money_style else t)
        rows.append([label] + cells)
    return rows


def t12_xlsx(prop, lines, mstyle="mon_year", extra_sheets=False, total_col=True):
    wb = Workbook()
    if extra_sheets:
        s = wb.active
        s.title = "Summary"
        s.append(["Property", prop["name"]])
        s.append(["NOI", round(sum(lines["noi"]), 2)])
        ws = wb.create_sheet("T12 Detail")
    else:
        ws = wb.active
        ws.title = "12 Month Statement"
    ws.append([prop["name"]])
    ws.append(["12 Month Operating Statement"])
    ws.append([f"Period = {month_headers('mon_year')[0]} - {month_headers('mon_year')[-1]}"])
    ws.append(["Book = Accrual"])
    for r in _t12_table(prop, lines, mstyle, None, total_col):
        ws.append(r)
    for row in ws.iter_rows(min_row=6):
        for c in row:
            if c.column > 1:
                c.number_format = "#,##0.00;(#,##0.00)"
    return xlsx_bytes(wb)


def t12_csv(prop, lines, mstyle, money_style, total_col=True, encoding="utf-8"):
    rows = [[prop["name"]], ["Trailing 12 Month Income Statement"], []]
    rows += _t12_table(prop, lines, mstyle, money_style, total_col)
    return csv_bytes(rows, encoding=encoding)


def t12_pdf(prop, lines):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(letter), leftMargin=0.3 * inch, rightMargin=0.3 * inch)
    data = _t12_table(prop, lines, "mon_dash_yy", "paren")
    data = [[("" if c is None else c) for c in r] for r in data]
    t = Table(data)
    t.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 6), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey)]))
    doc.build([Paragraph(f"{prop['name']} T-12", getSampleStyleSheet()["Title"]), t])
    return buf.getvalue()


def t12_xls(prop, lines):
    if not HAVE_XLWT:
        return None
    import xlwt
    wb = xlwt.Workbook()
    ws = wb.add_sheet("T12")
    ws.write(0, 0, prop["name"])
    for ri, r in enumerate(_t12_table(prop, lines, "mon_year", None), start=2):
        for ci, v in enumerate(r):
            if v is not None:
                ws.write(ri, ci, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# (format, scenario) per property -- 2 each = 28, + bad files below.
T12_PLAN = [
    [("xlsx_monyear", "clean"), ("csv_dollar_mondash", "gap")],
    [("xlsx_mon", "occ"), ("csv_mm_yyyy_nototal", "clean")],
    [("xlsx_extra_sheets", "gap"), ("pdf_text", "clean")],
    [("xlsx_monyear", "baddebt"), ("csv_long_year", "baddebt_old")],
    [("xls", "clean"), ("csv_dollar_mondash", "occ")],
    [("xlsx_mon", "gap"), ("xlsx_extra_sheets", "clean")],
    [("csv_mm_yyyy_nototal", "gap"), ("xlsx_monyear", "occ")],
    [("pdf_text", "baddebt"), ("xlsx_mon", "baddebt_old")],
    [("csv_long_year", "clean"), ("xls", "gap")],
    [("xlsx_monyear", "baddebt_old"), ("csv_dollar_mondash", "baddebt")],
    [("xlsx_extra_sheets", "occ"), ("csv_mm_yyyy_nototal", "gap")],
    [("pdf_text", "gap"), ("xlsx_monyear", "clean")],
    [("csv_long_year", "occ"), ("xlsx_mon", "clean")],
    [("xlsx_monyear", "gap"), ("csv_dollar_mondash", "baddebt")],
]


def render_t12(prop, fmt, lines):
    if fmt == "xlsx_monyear":
        return "xlsx", t12_xlsx(prop, lines, "mon_year"), "xlsx, 'Oct 2025' month headers, section header rows, accounting negatives"
    if fmt == "xlsx_mon":
        return "xlsx", t12_xlsx(prop, lines, "mon"), "xlsx, bare 'Oct'..'Sep' headers (trailing, NOT calendar order)"
    if fmt == "xlsx_extra_sheets":
        return "xlsx", t12_xlsx(prop, lines, "mon_year", extra_sheets=True), "xlsx with a Summary sheet first; detail on sheet 2"
    if fmt == "csv_dollar_mondash":
        return "csv", t12_csv(prop, lines, "mon_dash_yy", "dollar_paren"), "CSV, 'Oct-25' headers, ($1,234.56) negatives"
    if fmt == "csv_mm_yyyy_nototal":
        return "csv", t12_csv(prop, lines, "mm_yyyy", "comma", total_col=False), "CSV, '10/2025' headers, no Total column (annual = sum of months), -1,234.56 negatives"
    if fmt == "csv_long_year":
        return "csv", t12_csv(prop, lines, "long_year", "paren"), "CSV, 'October 2025' headers, (1,234.56) negatives"
    if fmt == "pdf_text":
        return "pdf", t12_pdf(prop, lines), "Digital PDF T-12 table, 'Oct-25' headers, tiny font"
    if fmt == "xls":
        b = t12_xls(prop, lines)
        return ("xls", b, "Legacy .xls T-12") if b else None
    raise ValueError(fmt)


# ---------------------------------------------------------------- bad files
_PREVIOUS_ENCRYPTED = {}  # rel path -> bytes from the last run (encryption salts are random)


def _encrypt_xlsx(raw, password="gauntlet", reuse=None):
    if reuse and reuse in _PREVIOUS_ENCRYPTED:
        return _PREVIOUS_ENCRYPTED[reuse]
    import msoffcrypto
    from msoffcrypto.format.ooxml import OOXMLFile
    out = io.BytesIO()
    f = OOXMLFile(io.BytesIO(raw))
    f.encrypt(password, out)
    return out.getvalue()


def _encrypt_pdf(raw, password="gauntlet"):
    from pypdf import PdfReader, PdfWriter
    w = PdfWriter()
    for p in PdfReader(io.BytesIO(raw)).pages:
        w.add_page(p)
    w.encrypt(password)
    out = io.BytesIO()
    w.write(out)
    return out.getvalue()


def _png_bytes():
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (40, 40), (200, 30, 30)).save(buf, format="PNG")
    return buf.getvalue()


def bad_rent_rolls(props):
    p0 = props[0]
    clean_xlsx = xlsx_bytes(_appfolio_wb(p0))
    empty_wb = Workbook()
    out = [
        ("empty.csv", b"", ["empty"], "0-byte CSV"),
        ("empty_sheet.xlsx", xlsx_bytes(empty_wb), ["empty"], "xlsx with one blank sheet"),
        ("header_only.csv", b"Unit,Tenant,Rent,Lease Start,Lease End\r\n", ["no ", "rows|units|tenants"], "header row, zero data rows -- importing 0 units silently is not OK"),
        ("only_vacant.csv", csv_bytes([["Unit", "Tenant", "Rent"], ["101", "VACANT", ""], ["102", "VACANT", ""]]), ["vacant|no occupied|no "], "all units vacant -- say so, don't 'import 0' silently"),
        ("image_named.csv", _png_bytes(), ["not a|doesn't look|binary|image|text"], "PNG bytes with a .csv name"),
        ("zip_named.xlsx", _zip_bytes(), ["excel|workbook|read"], "random zip with an .xlsx name"),
        ("truncated.xlsx", clean_xlsx[: len(clean_xlsx) // 2], ["excel|workbook|read|corrupt|damaged"], "xlsx cut off halfway (partial download)"),
        ("prose.docx", _prose_docx(), ["table"], "Word doc with no table at all"),
        ("rentroll.numbers", b"PK\x03\x04fake", ["file type|supported"], "Apple Numbers file"),
        ("blank_page.pdf", _blank_pdf(), ["blank|no text|no readable|empty|unreadable|table|scanned"], "PDF with a single blank page"),
        ("random_bytes.pdf", bytes(random.Random(3).randrange(256) for _ in range(4000)), ["pdf|read|corrupt|damaged"], "not actually a PDF"),
        ("lease_instead_of_rent_roll.pdf", render_paragraph_pdf(lease_paragraphs(p0, truth.lease_units(p0)[0]), "x"),
         ["rent roll|table|tenant"], "a single lease PDF uploaded into the rent-roll slot"),
    ]
    out.append(("password.pdf", _encrypt_pdf(_rr_table_pdf(p0)), ["password|encrypted|protected"], "password-protected PDF rent roll"))
    if HAVE_CRYPTO:
        out.append(("password.xlsx", _encrypt_xlsx(clean_xlsx, reuse="rent_rolls/bad__password.xlsx"), ["password|encrypted|protected"], "password-protected xlsx (Office encryption)"))
    return out


def _zip_bytes():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr(zipfile.ZipInfo("readme.txt", date_time=(2026, 10, 5, 0, 0, 0)), "not a workbook")
    return buf.getvalue()


def _prose_docx():
    from docx import Document
    doc = Document()
    doc.add_paragraph("Dear investor, attached please find the rent roll. Let me know if you have questions.")
    doc.core_properties.created = _FIXED_TS
    doc.core_properties.modified = _FIXED_TS
    buf = io.BytesIO()
    doc.save(buf)
    return stable_zip(buf.getvalue())


def _blank_pdf():
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    c.showPage()
    c.save()
    return buf.getvalue()


def bad_t12s(props, t12_lines_p0):
    p0 = props[0]
    out = [
        ("empty.csv", b"", ["empty"], "0-byte T-12"),
        ("image_named.xlsx", _png_bytes(), ["excel|workbook|read"], "PNG named .xlsx"),
        ("rent_roll_instead.xlsx", xlsx_bytes(_appfolio_wb(p0)), ["month|t-12|t12|income"], "a rent roll uploaded into the T-12 slot"),
        ("budget_only.csv", csv_bytes([["Account", "Budget 2026"], ["Gross Potential Rent", "1,000,000"]]), ["month|t-12|t12|total"], "a budget, not a T-12"),
    ]
    if HAVE_CRYPTO:
        out.append(("password.xlsx", _encrypt_xlsx(t12_xlsx(p0, t12_lines_p0), reuse="t12/bad__password.xlsx"), ["password|encrypted|protected"], "password-protected T-12"))
    return out


# ---------------------------------------------------------------- main
def main():
    props = truth.all_properties()
    for rel in ("rent_rolls/bad__password.xlsx", "t12/bad__password.xlsx"):
        if os.path.exists(os.path.join(FIX, rel)):
            with open(os.path.join(FIX, rel), "rb") as f:
                _PREVIOUS_ENCRYPTED[rel] = f.read()
    for sub in ("leases", "rent_rolls", "t12"):
        shutil.rmtree(os.path.join(FIX, sub), ignore_errors=True)
    manifest = {"as_of": AS_OF.isoformat(), "generated_by": "backend/tools/gauntlet/generate.py",
                "properties": [], "rent_rolls": [], "t12s": [], "bad_rent_rolls": [], "bad_t12s": []}
    t12_lines_p0 = None
    for prop, fmts, t12s in zip(props, RR_PLAN, T12_PLAN):
        rng = random.Random(prop["id"])
        leases = write_leases(prop)
        manifest["properties"].append({"id": prop["id"], "name": prop["name"], "typed_address": prop["typed_address"],
                                       "base_address": prop["base_address"], "leases": leases, "units": prop["units"]})
        for fmt in ["clean_csv"] + fmts:
            res = RR_WRITERS[fmt](prop, rng)
            if res is None:
                print(f"  skip {prop['id']} {fmt} (generation lib unavailable)")
                continue
            ext, data, carries, note = res
            rel = f"rent_rolls/{prop['id']}__{fmt}.{ext}"
            write(os.path.join(FIX, rel), data)
            manifest["rent_rolls"].append({
                "file": rel, "property": prop["id"], "format": fmt, "ext": ext, "note": note,
                "canonical": fmt == "clean_csv", "carries": sorted(carries),
                "expected_rent_roll": truth.expected_rent_roll(prop, carries),
                "expected_findings": truth.expected_findings(prop, carries),
            })
        for fmt, scenario in t12s:
            lines, ann, parsed, rows, rr_annual = t12_truth(prop, scenario, random.Random(f"{prop['id']}{fmt}{scenario}"))
            if t12_lines_p0 is None:
                t12_lines_p0 = lines
            res = render_t12(prop, fmt, lines)
            if res is None:
                print(f"  skip {prop['id']} t12 {fmt}")
                continue
            ext, data, note = res
            rel = f"t12/{prop['id']}__{fmt}__{scenario}.{ext}"
            write(os.path.join(FIX, rel), data)
            manifest["t12s"].append({"file": rel, "property": prop["id"], "format": fmt, "scenario": scenario, "ext": ext,
                                     "note": note, "expected_parse": parsed, "expected_t12_rows": rows,
                                     "rent_roll_annual": rr_annual})
    for name, data, kw, note in bad_rent_rolls(props):
        rel = f"rent_rolls/bad__{name}"
        write(os.path.join(FIX, rel), data)
        manifest["bad_rent_rolls"].append({"file": rel, "expect_status": 400, "expect_keywords": kw, "note": note})
    for name, data, kw, note in bad_t12s(props, t12_lines_p0):
        rel = f"t12/bad__{name}"
        write(os.path.join(FIX, rel), data)
        manifest["bad_t12s"].append({"file": rel, "expect_status": 400, "expect_keywords": kw, "note": note,
                                     "property": props[0]["id"]})
    with open(os.path.join(FIX, "manifest.json"), "w") as f:
        json.dump(manifest, f, indent=1, default=str)
    n_rr = len([r for r in manifest["rent_rolls"] if not r["canonical"]]) + len(manifest["bad_rent_rolls"])
    n_t12 = len(manifest["t12s"]) + len(manifest["bad_t12s"])
    n_leases = sum(len(p["leases"]) for p in manifest["properties"])
    print(f"properties={len(props)} rent_rolls={n_rr} (+{len(props)} canonical) t12s={n_t12} leases={n_leases}")


if __name__ == "__main__":
    main()
